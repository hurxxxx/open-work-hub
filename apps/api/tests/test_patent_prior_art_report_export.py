from __future__ import annotations

from io import BytesIO
from typing import cast
import zipfile

from docx import Document
import pytest

from ai_do_api.domains.patent_prior_art.report_export import render_report_export
from ai_do_api.domains.patent_prior_art.reporting import (
    build_detailed_report,
    render_markdown_report,
)
from ai_do_api.domains.patent_prior_art.schemas import PatentPriorArtReportFormat


_PRIVATE_RANKING_MARKER = "DO_NOT_EXPORT_92731"


def _payload() -> dict[str, object]:
    return {
        "schema_version": "1.0",
        "report_type": "patent_prior_art_research",
        "title": "<script>열관리 제어 조사</script>",
        "scope": {
            "jurisdictions": ["KR", "US"],
            "category_ids": ["vehicle"],
            "invention_text": "센서 신호에 따라 열교환기와 송풍기를 통합 제어하는 기술",
            "invention_text_truncated": False,
            "technology_summary": "센서 기반 열관리 장치의 에너지 절감 제어",
        },
        "search_plan": {
            "category_ids": {"values": ["vehicle"], "source": "user"},
            "keywords_ko": {"values": ["열교환기", "송풍기"], "source": "input_derived"},
            "keywords_en": {
                "values": ["heat exchanger", "blower control"],
                "source": "input_derived",
            },
            "ipc_codes": {"values": ["B60H"], "source": "input_derived"},
            "cpc_codes": {"values": ["B60H1/00"], "source": "input_derived"},
            "applicants": {"values": [], "source": "user"},
            "excluded_terms": {"values": ["가정용"], "source": "input_derived"},
            "display_query": "B60H AND (열교환기 OR 송풍기)",
        },
        "executed_queries": [
            {
                "source_id": "source-1",
                "source_label": "공개 특허 데이터",
                "jurisdiction": "KR",
                "query_text": "B60H AND 열교환기",
                "result_count": 18,
            }
        ],
        "selection": {
            "candidate_count": 3,
            "ranking_profile": {
                "input_candidate_count": 18,
                "ranked_candidate_count": 12,
                "returned_candidate_count": 3,
                "methods": ["bm25", "semantic", "rrf", "cross_encoder"],
                "raw_ranking_value": _PRIVATE_RANKING_MARKER,
            },
        },
        "candidates": [
            {
                "rank": 1,
                "publication_number": "KR-2024-000001",
                "title": "통합 열관리 제어 장치",
                "assignees": ["출원인 A"],
                "jurisdiction": "KR",
                "filing_date": "2022-01-02",
                "publication_date": "2024-02-03",
                "classification_codes": ["B60H1/00", "B60H3/06"],
                "abstract": "복수 센서와 열교환기 제어를 포함합니다.",
                "summary": "센서 입력과 송풍기 제어 구성이 입력 기술과 대응합니다.",
                "relevance_band": "high",
                "match_reasons": ["센서 입력", "송풍기 제어"],
                "external_url": "https://patents.example/KR-2024-000001",
                "raw_ranking_value": _PRIVATE_RANKING_MARKER,
            },
            {
                "rank": 2,
                "publication_number": "US-2021-000002",
                "title": "Thermal management controller",
                "assignees": ["연구기관 B"],
                "jurisdiction": "US",
                "filing_date": "2019-03-04",
                "publication_date": "2021-04-05",
                "classification_codes": ["B60H1/00", "F24F11/00"],
                "abstract": "A controller coordinates a heat exchanger and a blower.",
                "summary": "열교환기와 송풍기의 협조 제어를 개시합니다.",
                "relevance_band": "medium",
                "match_reasons": ["열교환기", "협조 제어"],
                "external_url": "https://patents.example/US-2021-000002",
            },
            {
                "rank": 3,
                "publication_number": "KR-2018-000003",
                "title": "송풍기 운전 방법",
                "assignees": ["출원인 A"],
                "jurisdiction": "KR",
                "filing_date": "2016-05-06",
                "publication_date": "2018-06-07",
                "classification_codes": ["F24F11/00"],
                "abstract": "송풍기 운전 조건을 조절합니다.",
                "summary": "송풍기 제어 요소 일부가 대응합니다.",
                "relevance_band": "low",
                "match_reasons": ["송풍기 운전"],
                "external_url": "https://patents.example/KR-2018-000003",
            },
        ],
    }


def _docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(values)


def test_detailed_projection_builds_neutral_counts_and_classification_clusters() -> None:
    report = build_detailed_report(_payload())

    assert [(item.label, item.count) for item in report.relevance_distribution] == [
        ("높음", 1),
        ("보통", 1),
        ("낮음", 1),
        ("미평가", 0),
    ]
    assert [(item.label, item.count) for item in report.applicant_distribution] == [
        ("출원인 A", 2),
        ("연구기관 B", 1),
    ]
    assert [(item.label, item.count) for item in report.year_distribution] == [
        ("2018", 1),
        ("2021", 1),
        ("2024", 1),
    ]
    assert [(item.code, item.candidate_count) for item in report.technology_clusters] == [
        ("B60H", 2),
        ("F24F", 2),
    ]
    assert _PRIVATE_RANKING_MARKER not in "\n".join(report.selection_process)


def test_markdown_has_nine_detailed_sections_without_private_ranking_value() -> None:
    markdown = render_markdown_report(_payload())

    headings = [line for line in markdown.splitlines() if line.startswith("## ")]
    assert headings == [
        "## 1. 과제 개요",
        "## 2. 검색식 도출 근거",
        "## 3. 선별 프로세스",
        "## 4. 핵심 특허 표",
        "## 5. 후보별 요지시트",
        "## 6. 정량 분석",
        "## 7. 분류코드 기반 기술 클러스터",
        "## 8. 종합 기술 검토",
        "## 9. 출처·한계 및 고지",
    ]
    assert _PRIVATE_RANKING_MARKER not in markdown


@pytest.mark.parametrize(
    ("report_format", "extension", "media_type", "magic"),
    [
        ("html", "html", "text/html; charset=utf-8", b"<!doctype html>"),
        ("pdf", "pdf", "application/pdf", b"%PDF-"),
        ("docx", "docx", "application/vnd.openxmlformats", b"PK"),
        ("summary_pdf", "pdf", "application/pdf", b"%PDF-"),
        ("summary_docx", "docx", "application/vnd.openxmlformats", b"PK"),
    ],
)
def test_all_report_formats_render_with_typed_metadata(
    report_format: str,
    extension: str,
    media_type: str,
    magic: bytes,
) -> None:
    rendered = render_report_export(
        _payload(),
        report_format=cast(PatentPriorArtReportFormat, report_format),
        job_id="job/unsafe value",
    )

    assert rendered.content.startswith(magic)
    assert rendered.extension == extension
    assert rendered.media_type.startswith(media_type)
    assert rendered.filename == (
        f"patent-prior-art-{'summary' if report_format.startswith('summary_') else 'detailed'}-"
        f"job-unsafe-value.{extension}"
    )


def test_html_is_self_contained_escaped_and_contains_no_active_remote_assets() -> None:
    rendered = render_report_export(_payload(), report_format="html", job_id="job-1")
    html = rendered.content.decode("utf-8")

    assert "&lt;script&gt;열관리 제어 조사&lt;/script&gt;" in html
    assert "<script" not in html.lower()
    assert "<link" not in html.lower()
    assert "<img" not in html.lower()
    assert "connect-src 'none'" in html
    assert _PRIVATE_RANKING_MARKER not in html
    for title in (
        "과제 개요",
        "검색식 도출 근거",
        "선별 프로세스",
        "핵심 특허 표",
        "후보별 요지시트",
        "정량 분석",
        "분류코드 기반 기술 클러스터",
        "종합 기술 검토",
        "출처·한계 및 고지",
    ):
        assert title in html


def test_docx_has_full_sections_and_no_external_relationships() -> None:
    rendered = render_report_export(_payload(), report_format="docx", job_id="job-1")
    text = _docx_text(rendered.content)

    assert "1. 과제 개요" in text
    assert "9. 출처·한계 및 고지" in text
    assert "후보별 요지시트" in text
    assert _PRIVATE_RANKING_MARKER not in text
    with zipfile.ZipFile(BytesIO(rendered.content)) as archive:
        relationships = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith(".rels")
        )
    assert 'TargetMode="External"' not in relationships


def test_summary_docx_contains_only_summary_sections() -> None:
    rendered = render_report_export(
        _payload(),
        report_format="summary_docx",
        job_id="job-1",
    )
    text = _docx_text(rendered.content)

    for title in ("1. 조사 범위", "2. 후보 요약", "3. 종합 기술 검토", "4. 고지"):
        assert title in text
    for detailed_title in (
        "검색식 도출 근거",
        "선별 프로세스",
        "후보별 요지시트",
        "정량 분석",
        "출처·한계",
    ):
        assert detailed_title not in text


def test_pdf_contains_no_clickable_remote_uri_annotation() -> None:
    rendered = render_report_export(_payload(), report_format="pdf", job_id="job-1")

    assert rendered.content.startswith(b"%PDF-")
    assert b"/URI" not in rendered.content
