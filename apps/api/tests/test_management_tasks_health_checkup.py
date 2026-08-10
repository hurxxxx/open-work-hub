from __future__ import annotations

import asyncio
import zipfile
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor
from dataclasses import replace
from datetime import date, datetime
from io import BytesIO
from threading import Barrier, Event
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.routing import APIRoute
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from ai_do_api.core.db import Base, get_db_session
from ai_do_api.domains.auth.models import OrgUnit, User
from ai_do_api.domains.hr.master import (
    HR_MASTER_ERP_VIEW_SCHEMA_VERSION,
    HrMasterErpEmployee,
    HrMasterErpEmployeeSnapshot,
    HrMasterErpEmployeeSnapshotMetadata,
)
from ai_do_api.domains.management_tasks import router as health_router
from ai_do_api.domains.management_tasks import service
from ai_do_api.domains.management_tasks import xlsx as health_xlsx
from ai_do_api.domains.management_tasks.models import (
    ManagementHealthCheckupSettings,
    ManagementHealthCheckupSettingsHistory,
    ManagementHealthDecisionRow,
    ManagementHealthDecisionRun,
    ManagementHealthPriorRow,
    ManagementHealthPriorUpload,
)
from ai_do_api.domains.management_tasks.rules import (
    AGE_CALC_INTERNATIONAL,
    AGE_CALC_KOREAN,
    DeterminationSettings,
    determine,
)
from ai_do_api.domains.management_tasks.schemas import (
    HealthCheckupSettingsUpdateRequest,
)
from ai_do_api.domains.management_tasks.xlsx import (
    MAX_UPLOAD_BYTES,
    UnsafeXlsxError,
    build_employee_table_xlsx,
    parse_prior_exam_workbook,
    read_upload_limited,
)


DOMAIN_TABLES = [
    ManagementHealthCheckupSettings.__table__,
    ManagementHealthCheckupSettingsHistory.__table__,
    ManagementHealthPriorUpload.__table__,
    ManagementHealthPriorRow.__table__,
    ManagementHealthDecisionRun.__table__,
    ManagementHealthDecisionRow.__table__,
    # Homonym disambiguation reads groupware departments (auth.User → OrgUnit),
    # so these tables must exist for prior-exam upload tests.
    OrgUnit.__table__,
    User.__table__,
]


@pytest.fixture
def db() -> Session:
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine, tables=DOMAIN_TABLES)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    with session_factory() as session:
        yield session
    engine.dispose()


@pytest.fixture
def erp_snapshot() -> HrMasterErpEmployeeSnapshot:
    return HrMasterErpEmployeeSnapshot(
        run_id="master-run-1",
        erp_run_id="erp-run-1",
        schema_version=HR_MASTER_ERP_VIEW_SCHEMA_VERSION,
        captured_at=datetime(2026, 7, 20, 3, 0, 0),
        snapshot_hash="a" * 64,
        employees=(
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-1",
                employee_code="1001",
                name="직원일",
                department_code="D1",
                department_name="설계팀",
                position="책임",
                occupation="정규직",
                birth_date=datetime(1980, 6, 1).date(),
                hire_date=datetime(2010, 6, 1).date(),
            ),
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-2",
                employee_code="1002",
                name="직원이",
                department_code="D2",
                department_name="생산팀",
                position=None,
                occupation="정규직",
                birth_date=datetime(1960, 6, 1).date(),
                hire_date=datetime(2025, 1, 1).date(),
            ),
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-3",
                employee_code="1003",
                name="파견원",
                department_code="D3",
                department_name="지원팀",
                position=None,
                occupation="파견직",
                birth_date=datetime(1960, 1, 1).date(),
                hire_date=datetime(2000, 1, 1).date(),
            ),
        ),
    )


def _workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "수검명단"
    worksheet.append(["부서명", "성명", "관계 및 사번"])
    for row in rows:
        worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _replace_zip_entry(content: bytes, entry_name: str, replacement: bytes) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(content)) as source:
        with zipfile.ZipFile(output, "w") as target:
            for entry in source.infolist():
                target.writestr(
                    entry,
                    replacement if entry.filename == entry_name else source.read(entry),
                )
    return output.getvalue()


def _snapshot_metadata(
    snapshot: HrMasterErpEmployeeSnapshot,
) -> HrMasterErpEmployeeSnapshotMetadata:
    return HrMasterErpEmployeeSnapshotMetadata(
        run_id=snapshot.run_id,
        erp_run_id=snapshot.erp_run_id,
        schema_version=snapshot.schema_version,
        captured_at=snapshot.captured_at,
        snapshot_hash=snapshot.snapshot_hash,
        employee_count=len(snapshot.employees),
    )


def _patch_snapshot_readers(
    monkeypatch: pytest.MonkeyPatch,
    provider: Callable[[Session], HrMasterErpEmployeeSnapshot],
) -> None:
    snapshot_provider = provider
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot",
        snapshot_provider,
    )
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot_metadata",
        lambda db: _snapshot_metadata(snapshot_provider(db)),
    )


def _assert_stale_run_blocks_all_exports(db: Session, *, target_year: int) -> None:
    with pytest.raises(service.DecisionRunStaleError):
        service.list_determinations(db, target_year=target_year)

    for export_kind in ("employees", "targets", "roster"):
        with pytest.raises(service.DecisionNotPublishableError) as error:
            service.export_determinations(
                db,
                target_year=target_year,
                export_kind=export_kind,
            )
        assert error.value.blockers == [service.PUBLISH_BLOCKER_STALE_INPUTS]


def test_rules_apply_age_service_prior_exam_and_dispatch() -> None:
    korean = determine(
        birth_date=datetime(1970, 12, 31).date(),
        hire_date=datetime(2020, 1, 2).date(),
        occupation="정규직",
        prior_year_examined=False,
        settings=DeterminationSettings(age_calc_method=AGE_CALC_KOREAN),
        target_year=2026,
    )
    international = determine(
        birth_date=datetime(1969, 12, 31).date(),
        hire_date=datetime(2020, 1, 2).date(),
        occupation="정규직",
        prior_year_examined=False,
        settings=DeterminationSettings(age_calc_method=AGE_CALC_INTERNATIONAL),
        target_year=2026,
    )
    deferred = determine(
        birth_date=datetime(1980, 1, 1).date(),
        hire_date=datetime(2020, 1, 1).date(),
        occupation="정규직",
        prior_year_examined=True,
        settings=DeterminationSettings(),
        target_year=2026,
    )
    dispatch = determine(
        birth_date=datetime(1960, 1, 1).date(),
        hire_date=datetime(2000, 1, 1).date(),
        occupation="파견직",
        prior_year_examined=False,
        settings=DeterminationSettings(),
        target_year=2026,
    )

    assert korean.age == 57 and korean.reason == "senior" and korean.is_target
    assert international.age == 56 and international.reason == "adult_or_service"
    assert deferred.reason == "deferred" and deferred.is_target is False
    assert dispatch.reason == "dispatch" and dispatch.is_target is False


def test_parser_keeps_exact_code_but_does_not_treat_self_marker_as_code() -> None:
    parsed = parse_prior_exam_workbook(
        _workbook_bytes(
            [
                ["설계팀", "직원일", "1001"],
                ["설계팀", "직원일", "본인"],
                ["설계팀", "가족", "배우자"],
            ]
        )
    )

    assert parsed[0].provided_employee_code == "1001"
    assert parsed[1].provided_employee_code is None
    assert parsed[1].is_dependent is False
    assert parsed[2].is_dependent is True


def test_parser_allows_formula_only_in_unused_columns() -> None:
    # Benign computed columns (row numbers, cost sums) that we never read must not
    # abort the upload — the roster still parses from its plain identity columns.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["No", "부서명", "성명", "관계 및 사번", "비고"])
    worksheet.append(["=ROW()-1", "설계팀", "직원일", "1001", "=1+1"])
    output = BytesIO()
    workbook.save(output)

    parsed = parse_prior_exam_workbook(output.getvalue())

    assert [row.person_name for row in parsed] == ["직원일"]
    assert parsed[0].provided_employee_code == "1001"


def test_parser_rejects_formula_in_identity_column() -> None:
    # A formula in an identity/relation column cannot be trusted (its cached value
    # may be missing/stale and would misclassify the row), so the upload is
    # rejected fail-closed rather than silently reading a blank.
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["부서명", "성명"])
    worksheet.append(["설계팀", "직원일"])
    worksheet.append(["설계팀", '=HYPERLINK("https://example.test")'])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(UnsafeXlsxError, match="formula_in_identity_column"):
        parse_prior_exam_workbook(output.getvalue())


def test_parser_rejects_blank_department_on_self_row() -> None:
    with pytest.raises(UnsafeXlsxError, match="missing_department"):
        parse_prior_exam_workbook(
            _workbook_bytes(
                [
                    ["설계팀", "직원일", "1001"],
                    [None, "직원이", "1002"],
                ]
            )
        )


def test_parser_inherits_merged_department_for_dependent_row() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["부서명", "성명", "관계 및 사번"])
    worksheet.append(["설계팀", "직원일", "1001"])
    worksheet.append([None, "가족", "배우자"])
    worksheet.merge_cells("A2:A3")
    output = BytesIO()
    workbook.save(output)

    parsed = parse_prior_exam_workbook(output.getvalue())

    assert parsed[1].dept_name == "설계팀"
    assert parsed[1].is_dependent is True


def test_parser_rejects_unmerged_blank_department_on_dependent_row() -> None:
    with pytest.raises(UnsafeXlsxError, match="missing_department"):
        parse_prior_exam_workbook(
            _workbook_bytes(
                [
                    ["설계팀", "직원일", "1001"],
                    [None, "가족", "배우자"],
                ]
            )
        )


def test_parser_does_not_carry_department_across_blank_separator() -> None:
    with pytest.raises(UnsafeXlsxError, match="missing_department"):
        parse_prior_exam_workbook(
            _workbook_bytes(
                [
                    ["설계팀", "직원일", "1001"],
                    [None, None, None],
                    [None, "가족", "배우자"],
                ]
            )
        )


def test_parser_does_not_inherit_merged_department_across_blank_separator() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["부서명", "성명", "관계 및 사번"])
    worksheet.append(["설계팀", "직원일", "1001"])
    worksheet.append([None, None, None])
    worksheet.append([None, "가족", "배우자"])
    worksheet.merge_cells("A2:A4")
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(UnsafeXlsxError, match="missing_department"):
        parse_prior_exam_workbook(output.getvalue())


def test_parser_rejects_absurd_merged_range_before_expansion() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["부서명", "성명", "관계 및 사번"])
    worksheet.append(["설계팀", "직원일", "1001"])
    worksheet.append([None, "가족", "배우자"])
    worksheet.merge_cells("A2:A3")
    output = BytesIO()
    workbook.save(output)

    with zipfile.ZipFile(BytesIO(output.getvalue())) as archive:
        worksheet_xml = archive.read("xl/worksheets/sheet1.xml")
    malformed = _replace_zip_entry(
        output.getvalue(),
        "xl/worksheets/sheet1.xml",
        worksheet_xml.replace(b'ref="A2:A3"', b'ref="A2:A999999999"'),
    )

    with pytest.raises(UnsafeXlsxError, match="too_many_rows"):
        parse_prior_exam_workbook(malformed)


def test_counted_upload_stream_rejects_oversize() -> None:
    upload = UploadFile(
        filename="oversize.xlsx",
        file=BytesIO(b"x" * (MAX_UPLOAD_BYTES + 1)),
    )
    with pytest.raises(UnsafeXlsxError, match="file_too_large"):
        asyncio.run(read_upload_limited(upload))


@pytest.mark.parametrize(
    ("limit_name", "limit_value", "error_code"),
    [
        ("MAX_ZIP_ENTRIES", 1, "too_many_archive_entries"),
        ("MAX_DECOMPRESSED_BYTES", 1, "decompressed_size_exceeded"),
        ("MAX_SHEETS", 0, "too_many_sheets"),
        ("MAX_ROWS_PER_SHEET", 1, "too_many_rows"),
        ("MAX_CELLS_PER_ROW", 2, "too_many_cells"),
    ],
)
def test_parser_enforces_archive_and_workbook_limits(
    monkeypatch: pytest.MonkeyPatch,
    limit_name: str,
    limit_value: int,
    error_code: str,
) -> None:
    monkeypatch.setattr(health_xlsx, limit_name, limit_value)
    with pytest.raises(UnsafeXlsxError, match=error_code):
        parse_prior_exam_workbook(_workbook_bytes([["설계팀", "직원일", "1001"]]))


def test_parser_rejects_external_links() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["부서명", "성명", "관계 및 사번"])
    worksheet.append(["설계팀", "직원일", "1001"])
    worksheet["B2"].hyperlink = "https://example.test/person"
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(UnsafeXlsxError, match="active_or_external_content"):
        parse_prior_exam_workbook(output.getvalue())


def test_isolated_parser_enforces_killable_deadline() -> None:
    with pytest.raises(UnsafeXlsxError, match="parse_timeout"):
        health_xlsx.parse_prior_exam_workbook_isolated(
            _workbook_bytes([["설계팀", "직원일", "1001"]]),
            timeout_seconds=0.001,
        )


def test_export_escapes_spreadsheet_formula_leads() -> None:
    content = build_employee_table_xlsx(
        sheet_title="전직원",
        headers=["이름", "사번", "표식"],
        rows=[["=cmd", "+1001", "-"]],
    )
    workbook = load_workbook(BytesIO(content), data_only=False)
    try:
        worksheet = workbook.active
        assert worksheet["A2"].value == "'=cmd"
        assert worksheet["B2"].value == "'+1001"
        # A lone hyphen placeholder stays clean (no leading apostrophe).
        assert worksheet["C2"].value == "-"
    finally:
        workbook.close()


def test_employee_export_uses_marks_and_korean_reasons(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["설계팀", "직원일", "1001"]]),
        actor_user_id="admin",
    )
    service.refresh_determinations(db, target_year=2026, actor_user_id="admin")
    export = service.export_determinations(db, target_year=2026, export_kind="employees")
    workbook = load_workbook(BytesIO(export.content), data_only=False)
    try:
        worksheet = workbook.active
        header = [cell.value for cell in worksheet[1]]
        prior_col = header.index("전년도 수검")
        target_col = header.index("대상 여부")
        reason_col = header.index("판정 사유")
        by_name = {
            row[header.index("이름")]: row
            for row in worksheet.iter_rows(min_row=2, values_only=True)
        }
        # 직원일: examined last year -> deferred, not a target this year.
        assert by_name["직원일"][prior_col] == "O"
        assert by_name["직원일"][target_col] == "-"
        assert "유예" in by_name["직원일"][reason_col]
        # 직원이: senior threshold met -> target, reason states the age.
        assert by_name["직원이"][target_col] == "O"
        assert by_name["직원이"][reason_col] == "57세 이상"
        # Boolean columns are marks, never raw Python booleans / TRUE-FALSE.
        assert not any(isinstance(cell, bool) for row in by_name.values() for cell in row)
    finally:
        workbook.close()


def test_source_employee_limit_rejects_before_raw_snapshot_materialization(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(service, "MAX_SOURCE_EMPLOYEES", 2)
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot_metadata",
        lambda _db: _snapshot_metadata(erp_snapshot),
    )
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot",
        lambda _db: pytest.fail("oversized source rows must not be materialized"),
    )

    status = service.get_source_status(db)
    assert status.available is False
    assert status.reason == "source_employee_limit_exceeded"
    with pytest.raises(
        service.SourceUnavailableError,
        match="source_employee_limit_exceeded",
    ):
        service.refresh_determinations(
            db,
            target_year=2026,
            actor_user_id="admin",
        )
    assert db.scalar(select(func.count(ManagementHealthDecisionRun.id))) == 0


def test_source_status_reports_missing_integrated_hr_erp_basis(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot_metadata",
        lambda _db: None,
    )

    status = service.get_source_status(db)

    assert status.available is False
    assert status.basis == "erp"
    assert status.reason == "integrated_hr_erp_snapshot_not_found"
    assert status.run_id is None
    assert status.erp_run_id is None


def test_no_code_row_matches_by_unique_name_and_publishes(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["설계팀", "직원일", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 1
    assert uploaded.ambiguous_count == 0
    assert uploaded.unmatched_count == 0
    assert uploaded.publish_blockers == []
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.match_status == "matched"
    assert stored.matched_employee_code == "1001"
    assert stored.candidate_employee_codes == []

    preview = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )
    assert preview.publishable is True
    assert preview.status == "ready"
    assert preview.publish_blockers == []
    by_code = {item.employee_code: item for item in preview.items}
    assert by_code["1001"].reason == "deferred"
    assert service.export_determinations(
        db,
        target_year=2026,
        export_kind="targets",
    ).content.startswith(b"PK")


def test_unmatched_row_is_excluded_but_does_not_block_publish(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["설계팀", "직원일", "본인"], ["없는팀", "퇴사자", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 1
    assert uploaded.unmatched_count == 1
    assert uploaded.ambiguous_count == 0
    # A person absent from the current ERP roster (retired/left) is excluded and
    # must not block publishing.
    assert uploaded.publish_blockers == []
    preview = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )
    assert preview.publishable is True


def test_legacy_unresolved_count_does_not_block_unmatched_only_upload(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["없는팀", "퇴사자", "본인"]]),
        actor_user_id="admin",
    )
    stored_upload = db.get(ManagementHealthPriorUpload, uploaded.upload_id)
    assert stored_upload is not None
    # Historical uploads counted unmatched rows as unresolved. The persisted
    # aggregate is therefore not authoritative for the current blocker meaning.
    stored_upload.unresolved_count = 1
    db.flush()

    upload_result = service._prior_upload_out(db, stored_upload, target_year=2026)
    status = service.get_prior_exam_status(db, target_year=2026)
    preview = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )

    assert upload_result.unmatched_count == 1
    assert upload_result.ambiguous_count == 0
    assert upload_result.publish_blockers == []
    assert status.publish_blockers == []
    assert preview.publishable is True
    assert preview.publish_blockers == []


def _homonym_snapshot(*, departments: tuple[str, str]) -> HrMasterErpEmployeeSnapshot:
    return HrMasterErpEmployeeSnapshot(
        run_id="master-run-h",
        erp_run_id="erp-run-h",
        schema_version=HR_MASTER_ERP_VIEW_SCHEMA_VERSION,
        captured_at=datetime(2026, 7, 20, 3, 0, 0),
        snapshot_hash="b" * 64,
        employees=(
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-h1",
                employee_code="2001",
                name="정성문",
                department_code="D1",
                department_name=departments[0],
                position="책임",
                occupation="정규직",
                birth_date=datetime(1980, 6, 1).date(),
                hire_date=datetime(2010, 6, 1).date(),
            ),
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-h2",
                employee_code="2002",
                name="정성문",
                department_code="D2",
                department_name=departments[1],
                position=None,
                occupation="정규직",
                birth_date=datetime(1981, 6, 1).date(),
                hire_date=datetime(2011, 6, 1).date(),
            ),
        ),
    )


def test_actual_ambiguous_rows_block_even_when_persisted_count_is_zero(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = _homonym_snapshot(departments=("생산팀", "구매팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업팀", "정성문", "본인"]]),
        actor_user_id="admin",
    )
    stored_upload = db.get(ManagementHealthPriorUpload, uploaded.upload_id)
    assert stored_upload is not None
    stored_upload.unresolved_count = 0
    db.flush()

    upload_result = service._prior_upload_out(db, stored_upload, target_year=2026)
    status = service.get_prior_exam_status(db, target_year=2026)
    preview = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )

    assert upload_result.publish_blockers == ["unresolved_prior_exam_rows"]
    assert status.publish_blockers == ["unresolved_prior_exam_rows"]
    assert preview.publishable is False
    assert preview.publish_blockers == ["unresolved_prior_exam_rows"]


def test_homonym_resolved_by_department_after_stripping_suffix(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = _homonym_snapshot(departments=("기술연구소 시작평가팀", "품질본부 서비스"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes(
            [["기술연구소", "정성문A", "본인"], ["품질본부", "정성문B", "본인"]]
        ),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 2
    assert uploaded.ambiguous_count == 0
    rows = {row.person_name: row for row in db.scalars(select(ManagementHealthPriorRow)).all()}
    assert rows["정성문A"].matched_employee_code == "2001"
    assert rows["정성문B"].matched_employee_code == "2002"


def test_homonym_without_distinguishing_department_is_ambiguous(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Both candidates sit in the identical department, so the file department
    # ties and the row is left for manual confirmation.
    snapshot = _homonym_snapshot(departments=("생산팀", "생산팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["생산팀", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    assert uploaded.publish_blockers == ["unresolved_prior_exam_rows"]
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.match_status == "ambiguous"
    assert sorted(stored.candidate_employee_codes) == ["2001", "2002"]


def test_homonym_with_only_weak_department_overlap_stays_ambiguous(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # The file department shares only an incidental token ("팀") with each
    # candidate, so no candidate clears the confidence floor: the row must stay
    # ambiguous rather than silently confirm the wrong same-name employee.
    snapshot = _homonym_snapshot(departments=("생산팀", "구매팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업팀", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.match_status == "ambiguous"


def test_homonym_with_similar_sibling_department_stays_ambiguous(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = _homonym_snapshot(departments=("생산2팀", "영업팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)

    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["생산1팀", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.match_status == "ambiguous"


def test_homonym_with_department_name_as_unsafe_substring_stays_ambiguous(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = _homonym_snapshot(departments=("해외영업팀", "구매팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)

    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업팀", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1


def test_homonym_resolved_by_groupware_department_when_erp_differs(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Neither ERP department matches the file's "영업본부"; only employee 2002's
    # groupware org-unit chain does, so groupware breaks the tie.
    db.add(OrgUnit(id="ou-sales-hq", name="영업본부", slug="sales-hq"))
    db.add(
        OrgUnit(
            id="ou-sales-domestic",
            name="영업본부 국내영업팀",
            slug="sales-domestic",
            parent_id="ou-sales-hq",
        )
    )
    db.add(
        User(
            id="local-user-2001",
            login_id="local2001",
            email="local2001@example.test",
            full_name="정성문",
            password_hash="x",
            employee_code="2001",
            primary_org_unit_id="ou-sales-domestic",
            auth_provider="local",
            status="active",
        )
    )
    db.add(
        User(
            id="suspended-user-2001",
            login_id="suspended2001",
            email="suspended2001@example.test",
            full_name="정성문",
            password_hash="x",
            employee_code="2001",
            primary_org_unit_id="ou-sales-domestic",
            auth_provider="groupware",
            status="suspended",
            hr_source_system="groupware",
            hr_domain_num=1,
            hr_user_num=2001,
            hr_com_state=1,
        )
    )
    db.add(
        User(
            id="user-2002",
            login_id="jsm2002",
            email="jsm@example.test",
            full_name="정성문",
            password_hash="x",
            employee_code="2002",
            primary_org_unit_id="ou-sales-domestic",
            auth_provider="groupware",
            status="active",
            hr_source_system="groupware",
            hr_domain_num=1,
            hr_user_num=2002,
            hr_com_state=1,
        )
    )
    db.flush()
    snapshot = _homonym_snapshot(departments=("경영지원팀", "고객지원팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업본부", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 1
    assert uploaded.ambiguous_count == 0
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.matched_employee_code == "2002"


def test_duplicate_active_groupware_identities_do_not_resolve_homonym(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    db.add(OrgUnit(id="ou-sales-hq", name="영업본부", slug="sales-hq"))
    for index in (1, 2):
        db.add(
            User(
                id=f"duplicate-user-{index}",
                login_id=f"duplicate{index}",
                email=f"duplicate{index}@example.test",
                full_name="정성문",
                password_hash="x",
                employee_code="2002",
                primary_org_unit_id="ou-sales-hq",
                auth_provider="groupware",
                status="active",
                hr_source_system="groupware",
                hr_domain_num=1,
                hr_user_num=3000 + index,
                hr_com_state=1,
            )
        )
    db.flush()
    snapshot = _homonym_snapshot(departments=("경영지원팀", "고객지원팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)

    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업본부", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    assert "2002" not in service._groupware_department_chains(db, {"2002"})


def test_inactive_groupware_org_unit_does_not_resolve_homonym(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    db.add(
        OrgUnit(
            id="ou-inactive-sales",
            name="영업본부",
            slug="inactive-sales",
            active=False,
        )
    )
    db.add(
        User(
            id="user-inactive-org-2002",
            login_id="inactive-org-2002",
            email="inactive-org-2002@example.test",
            full_name="정성문",
            password_hash="x",
            employee_code="2002",
            primary_org_unit_id="ou-inactive-sales",
            auth_provider="groupware",
            status="active",
            hr_source_system="groupware",
            hr_domain_num=1,
            hr_user_num=2002,
            hr_com_state=1,
        )
    )
    db.flush()
    snapshot = _homonym_snapshot(departments=("경영지원팀", "고객지원팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)

    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업본부", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    assert "2002" not in service._groupware_department_chains(db, {"2002"})


def test_erp_suffixed_homonym_matches_exact_name_before_stripping(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # ERP distinguishes homonyms with its own A/B suffix and both sit under the
    # same parent department, so only an exact-name match (not stripped) is safe.
    snapshot = HrMasterErpEmployeeSnapshot(
        run_id="master-run-suffix",
        erp_run_id="erp-run-suffix",
        schema_version=HR_MASTER_ERP_VIEW_SCHEMA_VERSION,
        captured_at=datetime(2026, 7, 20, 3, 0, 0),
        snapshot_hash="c" * 64,
        employees=(
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-s1",
                employee_code="3001",
                name="정호철A",
                department_code="D1",
                department_name="제조팀 제조 RAD 가공 라인",
                position=None,
                occupation="정규직",
                birth_date=datetime(1980, 6, 1).date(),
                hire_date=datetime(2010, 6, 1).date(),
            ),
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-s2",
                employee_code="3002",
                name="정호철B",
                department_code="D2",
                department_name="제조팀 제조 TUBE MILL 가공 라인",
                position=None,
                occupation="정규직",
                birth_date=datetime(1981, 6, 1).date(),
                hire_date=datetime(2011, 6, 1).date(),
            ),
        ),
    )
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["제조팀", "정호철B", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 1
    assert uploaded.ambiguous_count == 0
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.matched_employee_code == "3002"


@pytest.mark.parametrize("file_name", ["직원일", "직원일B"])
def test_missing_or_mismatched_suffix_does_not_match_unique_suffixed_employee(
    file_name: str,
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    suffixed_employee = replace(erp_snapshot.employees[0], name="직원일A")
    snapshot = replace(erp_snapshot, employees=(suffixed_employee,))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)

    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["설계팀", file_name, "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    assert uploaded.publish_blockers == ["unresolved_prior_exam_rows"]
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.match_status == "ambiguous"
    assert stored.candidate_employee_codes == ["1001"]


def _mixed_suffix_snapshot() -> HrMasterErpEmployeeSnapshot:
    # ERP holds a plain 정호철 AND a suffixed 정호철A — two different people.
    return HrMasterErpEmployeeSnapshot(
        run_id="master-run-mixed",
        erp_run_id="erp-run-mixed",
        schema_version=HR_MASTER_ERP_VIEW_SCHEMA_VERSION,
        captured_at=datetime(2026, 7, 20, 3, 0, 0),
        snapshot_hash="d" * 64,
        employees=(
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-m1",
                employee_code="5001",
                name="정호철",
                department_code="D1",
                department_name="관리팀 관리",
                position="책임",
                occupation="정규직",
                birth_date=datetime(1980, 6, 1).date(),
                hire_date=datetime(2010, 6, 1).date(),
            ),
            HrMasterErpEmployee(
                snapshot_row_id="erp-row-m2",
                employee_code="5002",
                name="정호철A",
                department_code="D2",
                department_name="생산팀 라인",
                position=None,
                occupation="정규직",
                birth_date=datetime(1981, 6, 1).date(),
                hire_date=datetime(2011, 6, 1).date(),
            ),
        ),
    )


def test_bare_name_colliding_with_suffixed_homonym_is_not_blindly_confirmed(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # A file row of bare "정호철" exact-matches ERP's plain 정호철, but a suffixed
    # homonym 정호철A also exists; with a department that fits neither, the row must
    # stay ambiguous instead of silently confirming (and deferring) the wrong 정호철.
    _patch_snapshot_readers(monkeypatch, lambda _db: _mixed_suffix_snapshot())
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["영업팀", "정호철", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 0
    assert uploaded.ambiguous_count == 1
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.match_status == "ambiguous"
    assert sorted(stored.candidate_employee_codes) == ["5001", "5002"]


def test_bare_name_with_matching_department_resolves_the_right_homonym(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: _mixed_suffix_snapshot())
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([["관리팀", "정호철", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.matched_count == 1
    assert uploaded.ambiguous_count == 0
    stored = db.scalar(select(ManagementHealthPriorRow))
    assert stored is not None
    assert stored.matched_employee_code == "5001"


def test_match_key_strips_single_letter_homonym_suffix() -> None:
    assert service._match_key("정성문A") == "정성문"
    assert service._match_key("정성문 B") == "정성문"
    assert service._match_key("정성문") == "정성문"
    # A purely latin name is left intact (no Korean stem to key on).
    assert service._match_key("David") == "david"


def test_resolve_homonym_prefers_groupware_department_over_erp() -> None:
    first = HrMasterErpEmployee(
        snapshot_row_id="r1",
        employee_code="4001",
        name="정성문",
        department_code="D1",
        department_name="경영지원팀",
        position=None,
        occupation="정규직",
        birth_date=date(1980, 6, 1),
        hire_date=date(2010, 6, 1),
    )
    second = HrMasterErpEmployee(
        snapshot_row_id="r2",
        employee_code="4002",
        name="정성문",
        department_code="D2",
        department_name="고객지원팀",
        position=None,
        occupation="정규직",
        birth_date=date(1981, 6, 1),
        hire_date=date(2011, 6, 1),
    )
    matched = service._resolve_homonym(
        "영업본부",
        [first, second],
        {"4002": ["영업본부 국내영업팀", "영업본부"]},
    )
    assert matched is not None
    assert matched.employee_code == "4002"


def test_upload_details_prioritize_ambiguous_rows_over_earlier_matches(
    db: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    snapshot = _homonym_snapshot(departments=("생산팀", "구매팀"))
    _patch_snapshot_readers(monkeypatch, lambda _db: snapshot)
    matched_rows = [["생산팀", "정성문", "2001"]] * health_xlsx.MAX_DETAIL_ROWS

    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=_workbook_bytes([*matched_rows, ["영업팀", "정성문", "본인"]]),
        actor_user_id="admin",
    )

    assert uploaded.details_truncated is True
    assert len(uploaded.matched) + len(uploaded.unmatched) == health_xlsx.MAX_DETAIL_ROWS
    assert len(uploaded.unmatched) == 1
    assert uploaded.unmatched[0].match_status == "ambiguous"


def test_exact_upload_is_distinct_and_each_refresh_is_an_immutable_run(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    uploaded = service.upload_prior_exam(
        db,
        target_year=2026,
        filename="C:\\fakepath\\prior.xlsx",
        content=_workbook_bytes(
            [
                ["설계팀", "직원일", "1001"],
                ["설계팀", "직원일", "1001"],
                ["설계팀", "가족", "배우자"],
            ]
        ),
        actor_user_id="admin",
    )
    assert uploaded.source_filename == "prior.xlsx"
    assert uploaded.matched_count == 1
    assert uploaded.spouse_excluded_count == 1
    assert uploaded.publish_blockers == []

    first = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )
    second = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )

    assert first.run_id != second.run_id
    assert first.items == second.items
    assert first.publishable is True and first.status == "ready"
    assert first.source.basis == "erp"
    assert first.source.run_id == erp_snapshot.run_id
    assert first.source.erp_run_id == erp_snapshot.erp_run_id
    assert first.target_count == 1
    by_code = {item.employee_code: item for item in first.items}
    assert by_code["1001"].reason == "deferred"
    assert by_code["1002"].reason == "senior"
    assert by_code["1003"].reason == "dispatch"
    assert db.scalar(select(func.count(ManagementHealthDecisionRun.id))) == 2
    stored_upload = db.scalar(select(ManagementHealthPriorUpload))
    assert stored_upload is not None
    assert stored_upload.source_run_id == erp_snapshot.run_id
    assert stored_upload.source_erp_run_id == erp_snapshot.erp_run_id
    stored_runs = list(db.scalars(select(ManagementHealthDecisionRun)).all())
    assert {run.source_run_id for run in stored_runs} == {erp_snapshot.run_id}
    assert {run.source_erp_run_id for run in stored_runs} == {erp_snapshot.erp_run_id}
    stored_rows = list(db.scalars(select(ManagementHealthDecisionRow)).all())
    assert {row.source_snapshot_row_id for row in stored_rows} == {
        "erp-row-1",
        "erp-row-2",
        "erp-row-3",
    }
    roster_export = service.export_determinations(
        db,
        target_year=2026,
        export_kind="roster",
    )
    assert roster_export.run_id == second.run_id
    assert roster_export.exported_count == 1
    roster = load_workbook(BytesIO(roster_export.content), data_only=False)
    try:
        assert roster.active["A1"].value == "2026년 종합검진 대상자 명단"
        assert roster.active["D1"].value == ("한국식 나이 40세 이상 또는 근속 10년 이상: 2년 1회")
        assert roster.active["D2"].value == "한국식 나이 57세 이상: 매년"
        # The age-criteria cells carry no outline.
        assert roster.active["D1"].border.left.style is None
        assert roster.active["D2"].border.top.style is None
    finally:
        roster.close()


def test_determination_and_export_limits_fail_before_unbounded_output(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    refreshed = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )

    monkeypatch.setattr(service, "MAX_DETERMINATION_RESPONSE_ROWS", 2)
    with pytest.raises(
        service.SourceUnavailableError,
        match="determination_row_limit_exceeded",
    ):
        service.list_determinations(db, target_year=2026)

    monkeypatch.setattr(service, "MAX_EXPORT_ROWS", 2)
    monkeypatch.setattr(
        service,
        "build_employee_table_xlsx",
        lambda **_kwargs: pytest.fail("oversized export rows must not be rendered"),
    )
    with pytest.raises(
        service.SourceUnavailableError,
        match="determination_row_limit_exceeded",
    ):
        service.export_determinations(
            db,
            target_year=2026,
            export_kind="employees",
        )

    monkeypatch.setattr(service, "MAX_EXPORT_ROWS", 3)
    monkeypatch.setattr(service, "MAX_EXPORT_BYTES", 2)
    monkeypatch.setattr(service, "build_employee_table_xlsx", lambda **_kwargs: b"PKX")
    with pytest.raises(
        service.SourceUnavailableError,
        match="export_size_limit_exceeded",
    ):
        service.export_determinations(
            db,
            target_year=2026,
            export_kind="employees",
        )

    assert refreshed.total == 3


def test_missing_prior_upload_produces_preview_and_settings_history(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_snapshot_readers(monkeypatch, lambda _db: erp_snapshot)
    preview = service.refresh_determinations(
        db,
        target_year=2026,
        actor_user_id="admin",
    )
    assert preview.publishable is False
    assert preview.publish_blockers == ["missing_prior_exam_upload"]

    updated = service.update_settings(
        db,
        payload=HealthCheckupSettingsUpdateRequest(age_calc_method="international"),
        actor_user_id="admin",
    )
    history = db.scalar(select(ManagementHealthCheckupSettingsHistory))
    assert updated.age_calc_method == "international"
    assert history is not None
    assert history.field_key == "age_calc_method"
    assert history.old_value == "korean"
    assert history.new_value == "international"


def test_replayed_noop_settings_update_does_not_create_false_history(
    db: Session,
) -> None:
    result = service.update_settings(
        db,
        payload=HealthCheckupSettingsUpdateRequest(age_calc_method="korean"),
        actor_user_id="admin",
    )

    assert result.updated_at is not None
    stored = db.get(ManagementHealthCheckupSettings, "company")
    assert stored is not None
    assert stored.updated_by_user_id is None
    assert stored.age_calc_method == "korean"
    assert stored.senior_age == 57
    assert stored.adult_age == 40
    assert stored.service_years_threshold == 10
    assert db.scalar(select(func.count(ManagementHealthCheckupSettingsHistory.id))) == 0


def test_changed_settings_upload_and_erp_projection_invalidate_latest_run(
    db: Session,
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    current_snapshot = [erp_snapshot]
    _patch_snapshot_readers(monkeypatch, lambda _db: current_snapshot[0])
    workbook = _workbook_bytes([["설계팀", "직원일", "1001"]])
    service.upload_prior_exam(
        db,
        target_year=2026,
        filename="prior.xlsx",
        content=workbook,
        actor_user_id="admin",
    )
    service.refresh_determinations(db, target_year=2026, actor_user_id="admin")

    service.update_settings(
        db,
        payload=HealthCheckupSettingsUpdateRequest(service_years_threshold=12),
        actor_user_id="settings-admin",
    )
    _assert_stale_run_blocks_all_exports(db, target_year=2026)
    service.refresh_determinations(db, target_year=2026, actor_user_id="admin")

    service.upload_prior_exam(
        db,
        target_year=2026,
        filename="corrected-prior.xlsx",
        content=workbook,
        actor_user_id="upload-admin",
    )
    _assert_stale_run_blocks_all_exports(db, target_year=2026)
    service.refresh_determinations(db, target_year=2026, actor_user_id="admin")

    current_snapshot[0] = replace(
        erp_snapshot,
        run_id="master-run-2",
        erp_run_id="erp-run-2",
        captured_at=datetime(2026, 7, 21, 3, 0, 0),
    )
    # Provenance-only changes do not stale a determination when the ERP-basis
    # projection and its schema are identical.
    service.list_determinations(db, target_year=2026)

    current_snapshot[0] = replace(
        current_snapshot[0],
        snapshot_hash="e" * 64,
    )
    _assert_stale_run_blocks_all_exports(db, target_year=2026)


def test_concurrent_disjoint_settings_updates_preserve_both_changes_and_history(
    application_postgres_dsn: str,
) -> None:
    engine = create_engine(application_postgres_dsn)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    barrier = Barrier(2)

    def update(payload: HealthCheckupSettingsUpdateRequest, actor_user_id: str) -> None:
        with session_factory() as session:
            barrier.wait(timeout=10)
            service.update_settings(
                session,
                payload=payload,
                actor_user_id=actor_user_id,
            )
            session.commit()

    try:
        with ThreadPoolExecutor(max_workers=2) as executor:
            futures = [
                executor.submit(
                    update,
                    HealthCheckupSettingsUpdateRequest(senior_age=60),
                    "senior-admin",
                ),
                executor.submit(
                    update,
                    HealthCheckupSettingsUpdateRequest(service_years_threshold=12),
                    "service-admin",
                ),
            ]
            for future in futures:
                future.result(timeout=20)

        with session_factory() as session:
            stored = session.get(ManagementHealthCheckupSettings, "company")
            history = list(
                session.scalars(
                    select(ManagementHealthCheckupSettingsHistory).order_by(
                        ManagementHealthCheckupSettingsHistory.field_key
                    )
                ).all()
            )

        assert stored is not None
        assert stored.senior_age == 60
        assert stored.service_years_threshold == 12
        assert [
            (
                row.field_key,
                row.old_value,
                row.new_value,
                row.changed_by_user_id,
            )
            for row in history
        ] == [
            ("senior_age", "57", "60", "senior-admin"),
            ("service_years_threshold", "10", "12", "service-admin"),
        ]
    finally:
        engine.dispose()


def test_upload_rejects_oversized_or_unverifiable_body_before_parsing() -> None:
    route = next(
        route
        for route in health_router.router.routes
        if getattr(route, "endpoint", None) is health_router.upload_prior_exam
    )
    limit = health_router._BODY_LIMIT_BY_ENDPOINT[health_router.upload_prior_exam]

    for headers in (
        [(b"content-length", str(limit + 1).encode())],
        [],
        [(b"content-length", b"not-a-number")],
    ):
        messages: list[dict] = []
        scope = {
            "type": "http",
            "http_version": "1.1",
            "method": "POST",
            "scheme": "http",
            "path": "/management-tasks/health-checkup/prior-exams",
            "raw_path": b"/management-tasks/health-checkup/prior-exams",
            "query_string": b"",
            "headers": headers,
            "client": ("testclient", 50000),
            "server": ("testserver", 80),
        }

        async def receive() -> dict:
            raise AssertionError("invalid request body must not be read")

        async def send(message: dict) -> None:
            messages.append(message)

        try:
            asyncio.run(route.app(scope, receive, send))
        except HTTPException as error:
            assert error.status_code == 413
        else:
            assert any(
                message.get("type") == "http.response.start" and message.get("status") == 413
                for message in messages
            )


def test_upload_audit_unresolved_count_means_ambiguous_only(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    result = SimpleNamespace(
        upload_id="upload-1",
        target_year=2026,
        exam_year=2025,
        total_rows=4,
        matched_count=1,
        unmatched_count=2,
        ambiguous_count=1,
    )
    captured_payload: dict[str, object] = {}
    commits: list[bool] = []
    monkeypatch.setattr(
        service,
        "upload_prior_exam",
        lambda *_args, **_kwargs: result,
    )

    def capture_audit(*_args: object, **kwargs: object) -> None:
        captured_payload.update(kwargs["payload"])

    monkeypatch.setattr(health_router, "record_audit_log", capture_audit)

    returned = health_router._upload_prior_exam_and_audit(
        SimpleNamespace(commit=lambda: commits.append(True)),
        target_year=2026,
        filename="prior.xlsx",
        content=b"xlsx",
        actor_user_id="admin",
        workspace_id="workspace-1",
        workspace_key="management",
    )

    assert returned is result
    assert captured_payload["unresolved_row_count"] == 1
    assert commits == [True]


@pytest.mark.anyio
async def test_upload_parser_and_database_work_do_not_block_event_loop(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    started = Event()
    release = Event()
    sentinel = SimpleNamespace()

    def blocking_upload(*args: object, **kwargs: object) -> object:
        started.set()
        if not release.wait(timeout=1):
            raise AssertionError("event loop could not release the worker")
        return sentinel

    monkeypatch.setattr(health_router, "_upload_prior_exam_and_audit", blocking_upload)
    task = asyncio.create_task(
        health_router.upload_prior_exam(
            year=2026,
            file=UploadFile(file=BytesIO(b"PK upload"), filename="prior.xlsx"),
            db=SimpleNamespace(),
            context=SimpleNamespace(
                auth=SimpleNamespace(user=SimpleNamespace(id="admin")),
                workspace=SimpleNamespace(id="workspace-1", key="management"),
            ),
        )
    )

    for _ in range(100):
        if started.is_set():
            break
        await asyncio.sleep(0.01)
    assert started.is_set()
    assert not task.done()
    release.set()
    assert await asyncio.wait_for(task, timeout=1) is sentinel


def test_router_has_workspace_member_and_entitlement_gates_and_source_status_api(
    erp_snapshot: HrMasterErpEmployeeSnapshot,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    for route in health_router.router.routes:
        assert isinstance(route, APIRoute)
        dependency_calls = {dependency.call for dependency in route.dependant.dependencies}
        assert health_router.require_management_tasks_enabled in dependency_calls
        assert health_router.require_management_tasks_member in dependency_calls

    app = FastAPI()
    app.include_router(health_router.router, prefix="/api/v1")
    app.dependency_overrides[health_router.require_management_tasks_enabled] = lambda: None
    app.dependency_overrides[health_router.require_management_tasks_member] = lambda: (
        SimpleNamespace(auth=SimpleNamespace(user=SimpleNamespace(id="admin")))
    )
    app.dependency_overrides[get_db_session] = lambda: None
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot_metadata",
        lambda _db: _snapshot_metadata(erp_snapshot),
    )
    monkeypatch.setattr(
        service,
        "load_latest_hr_master_erp_employee_snapshot",
        lambda _db: pytest.fail("source status must not materialize integrated HR rows"),
    )

    with TestClient(app) as client:
        response = client.get("/api/v1/management-tasks/health-checkup/source/status")
    assert response.status_code == 200
    assert response.json() == {
        "available": True,
        "basis": "erp",
        "reason": None,
        "run_id": "master-run-1",
        "erp_run_id": "erp-run-1",
        "captured_at": "2026-07-20T03:00:00",
        "employee_count": 3,
        "schema_version": HR_MASTER_ERP_VIEW_SCHEMA_VERSION,
    }
