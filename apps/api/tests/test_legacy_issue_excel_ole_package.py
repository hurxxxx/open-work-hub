from __future__ import annotations

import struct
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import olefile
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from open_alm_api.domains.legacy_issues.excel_ole_package import (
    OleAttachmentPlacement,
    _compound_file_header,
    _plan_compound_file,
    embed_ole_attachments,
)


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
VML_NS = "urn:schemas-microsoft-com:vml"
EXCEL_NS = "urn:schemas-microsoft-com:office:excel"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OLE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"


class _GuardedStream(BytesIO):
    def __init__(
        self,
        content: bytes,
        requested_sizes: list[int],
        released_streams: list[BytesIO],
    ) -> None:
        super().__init__(content)
        self._requested_sizes = requested_sizes
        self._released_streams = released_streams

    def read(self, size: int = -1) -> bytes:
        assert size >= 0, "attachment payload must not be read without a bound"
        assert size <= 1024 * 1024
        self._requested_sizes.append(size)
        return super().read(size)

    def release_conn(self) -> None:
        self._released_streams.append(self)


def _save_workbook(path: Path, *, with_comment: bool = False) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Master"
    worksheet["C2"] = "근거📎.txt"
    worksheet.column_dimensions["C"].width = 15
    if with_comment:
        worksheet["A1"].comment = Comment("existing comment", "tester")
    workbook.save(path)


def _read_zero_terminated(stream: BytesIO) -> bytes:
    value = bytearray()
    while True:
        character = stream.read(1)
        if character == b"\x00":
            return bytes(value)
        assert character
        value.extend(character)


def _parse_ole10native(content: bytes) -> tuple[bytes, list[str]]:
    stream = BytesIO(content)
    total_size = struct.unpack("<I", stream.read(4))[0]
    assert total_size == len(content) - 4
    assert struct.unpack("<H", stream.read(2))[0] == 2
    _read_zero_terminated(stream)
    _read_zero_terminated(stream)
    stream.read(4)
    command_size = struct.unpack("<I", stream.read(4))[0]
    stream.read(command_size)
    payload_size = struct.unpack("<I", stream.read(4))[0]
    payload = stream.read(payload_size)
    unicode_values: list[str] = []
    for _ in range(3):
        character_count = struct.unpack("<I", stream.read(4))[0]
        encoded = stream.read(character_count * 2)
        assert encoded.endswith(b"\x00\x00")
        unicode_values.append(encoded[:-2].decode("utf-16le"))
    assert stream.read() == b""
    return payload, unicode_values


def test_embed_ole_attachments_streams_payloads_and_wraps_icons(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _save_workbook(source)
    payloads = [
        ("근거📎.txt", b"exact original bytes"),
        ("second.bin", b"\x00\x01\x02"),
        ("third.dat", b"x" * 5000),
        ("fourth.txt", b"four"),
        ("fifth.txt", b"five"),
    ]
    requested_sizes: list[int] = []
    released_streams: list[BytesIO] = []
    progress: list[tuple[int, int]] = []
    placements = [
        OleAttachmentPlacement(
            sheet_name="Master",
            cell_coordinate="c2",
            filename=filename,
            size_bytes=len(payload),
            stream_factory=lambda payload=payload: _GuardedStream(
                payload,
                requested_sizes,
                released_streams,
            ),
        )
        for filename, payload in payloads
    ]

    embed_ole_attachments(
        source,
        output,
        placements,
        lambda completed_count, completed_bytes: progress.append(
            (completed_count, completed_bytes)
        ),
    )

    assert requested_sizes
    assert len(released_streams) == len(payloads)
    assert all(stream.closed for stream in released_streams)
    assert progress == [
        (index, sum(len(payload) for _, payload in payloads[:index]))
        for index in range(1, len(payloads) + 1)
    ]
    with ZipFile(output) as package:
        assert package.testzip() is None
        embedding_parts = sorted(
            name for name in package.namelist() if name.startswith("xl/embeddings/oleObject")
        )
        assert len(embedding_parts) == 5
        assert all(package.getinfo(name).extract_version >= 45 for name in embedding_parts)
        assert len([name for name in package.namelist() if name.startswith("xl/media/image")]) == 1

        worksheet_xml = package.read("xl/worksheets/sheet1.xml")
        assert "근거📎.txt".encode() not in worksheet_xml
        worksheet = ET.fromstring(worksheet_xml)
        cell = worksheet.find(f".//{{{MAIN_NS}}}c[@r='C2']")
        assert cell is not None
        assert list(cell) == []

        ole_objects = worksheet.findall(f".//{{{MAIN_NS}}}oleObject")
        assert len(ole_objects) == 5
        shape_ids = [item.get("shapeId") for item in ole_objects]
        assert len(set(shape_ids)) == 5
        from_markers = [item.find(f".//{{{MAIN_NS}}}from") for item in ole_objects]
        row_offsets = [
            int(marker.findtext(f"{{{DRAWING_NS}}}rowOff", "-1"))
            for marker in from_markers
            if marker is not None
        ]
        column_offsets = [
            int(marker.findtext(f"{{{DRAWING_NS}}}colOff", "-1"))
            for marker in from_markers
            if marker is not None
        ]
        to_markers = [item.find(f".//{{{MAIN_NS}}}to") for item in ole_objects]
        assert all(
            to_marker is not None
            and from_marker is not None
            and int(to_marker.findtext(f"{{{DRAWING_NS}}}colOff", "-1"))
            - int(from_marker.findtext(f"{{{DRAWING_NS}}}colOff", "-1"))
            == 24 * 9525
            for from_marker, to_marker in zip(from_markers, to_markers, strict=True)
        )
        assert row_offsets[:4] == [row_offsets[0]] * 4
        assert column_offsets[:4] == sorted(column_offsets[:4])
        assert len(set(column_offsets[:4])) == 4
        assert row_offsets[4] > row_offsets[0]
        assert column_offsets[4] == column_offsets[0]
        assert {
            marker.findtext(f"{{{DRAWING_NS}}}col") for marker in from_markers if marker is not None
        } == {"2"}
        assert {
            marker.findtext(f"{{{DRAWING_NS}}}row") for marker in from_markers if marker is not None
        } == {"1"}

        relationships = ET.fromstring(package.read("xl/worksheets/_rels/sheet1.xml.rels"))
        relationship_ids = [item.get("Id") for item in relationships]
        assert len(relationship_ids) == len(set(relationship_ids))
        assert sum(item.get("Type") == OLE_REL_TYPE for item in relationships) == 5

        vml_part = next(name for name in package.namelist() if name.endswith(".vml"))
        vml = ET.fromstring(package.read(vml_part))
        vml_shapes = vml.findall(f".//{{{VML_NS}}}shape")
        assert all("width:18pt;height:18pt" in (shape.get("style") or "") for shape in vml_shapes)
        assert {shape.get("id") for shape in vml_shapes} == {
            f"_x0000_s{shape_id}" for shape_id in shape_ids
        }
        vml_anchors = vml.findall(f".//{{{EXCEL_NS}}}Anchor")
        assert all(anchor.text and anchor.text.strip().startswith("2,") for anchor in vml_anchors)

        content_types = package.read("[Content_Types].xml")
        assert content_types.count(b"application/vnd.openxmlformats-officedocument.oleObject") == 5

        for index, embedding_part in enumerate(embedding_parts):
            compound_file = package.read(embedding_part)
            with olefile.OleFileIO(BytesIO(compound_file)) as embedded:
                assert str(embedded.root.clsid) == "0003000C-0000-0000-C000-000000000046"
                native = embedded.openstream("\x01Ole10Native").read()
            payload, unicode_names = _parse_ole10native(native)
            assert payload == payloads[index][1]
            assert unicode_names == [payloads[index][0]] * 3

    loaded = load_workbook(output, read_only=True)
    assert loaded.sheetnames == ["Master"]
    loaded.close()


def test_embed_reuses_existing_vml_without_relationship_or_shape_id_collisions(
    tmp_path: Path,
) -> None:
    source = tmp_path / "comments.xlsx"
    output = tmp_path / "embedded.xlsx"
    _save_workbook(source, with_comment=True)

    embed_ole_attachments(
        source,
        output,
        [OleAttachmentPlacement("Master", "C2", "proof.bin", 3, lambda: BytesIO(b"abc"))],
    )

    with ZipFile(output) as package:
        worksheet = ET.fromstring(package.read("xl/worksheets/sheet1.xml"))
        assert len(worksheet.findall(f"{{{MAIN_NS}}}legacyDrawing")) == 1
        relationships = ET.fromstring(package.read("xl/worksheets/_rels/sheet1.xml.rels"))
        ids = [item.get("Id") for item in relationships]
        assert len(ids) == len(set(ids))
        legacy_relationship = next(
            item for item in relationships if item.get("Type", "").endswith("/vmlDrawing")
        )
        target = legacy_relationship.get("Target")
        assert target is not None
        vml_part = (
            target.lstrip("/") if target.startswith("/") else f"xl/{target.removeprefix('../')}"
        )
        vml = ET.fromstring(package.read(vml_part))
        shape_ids = [shape.get("id") for shape in vml.findall(f".//{{{VML_NS}}}shape")]
        assert len(shape_ids) == len(set(shape_ids))
        assert any(shape_id and shape_id.startswith("_x0000_s") for shape_id in shape_ids)


def test_embed_failure_does_not_replace_existing_workbook(tmp_path: Path) -> None:
    workbook = tmp_path / "same.xlsx"
    _save_workbook(workbook)
    original = workbook.read_bytes()

    with pytest.raises(ValueError, match="ended before size_bytes"):
        embed_ole_attachments(
            workbook,
            workbook,
            [
                OleAttachmentPlacement(
                    "Master",
                    "C2",
                    "short.bin",
                    5,
                    lambda: BytesIO(b"123"),
                )
            ],
        )

    assert workbook.read_bytes() == original
    assert list(tmp_path.glob(".same.xlsx.*.tmp")) == []


@pytest.mark.parametrize(
    ("placement", "message"),
    [
        (OleAttachmentPlacement("Missing", "A1", "x", 0, lambda: BytesIO()), "Worksheet not found"),
        (OleAttachmentPlacement("Master", "not-a-cell", "x", 0, lambda: BytesIO()), "Invalid cell"),
        (OleAttachmentPlacement("Master", "A1", "bad\x00name", 0, lambda: BytesIO()), "NUL"),
        (OleAttachmentPlacement("Master", "A1", "x", -1, lambda: BytesIO()), "negative"),
    ],
)
def test_embed_rejects_invalid_placements(
    tmp_path: Path,
    placement: OleAttachmentPlacement,
    message: str,
) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _save_workbook(source)

    with pytest.raises(ValueError, match=message):
        embed_ole_attachments(source, output, [placement])

    assert not output.exists()


def test_compound_file_plan_supports_more_fat_sectors_than_header_difat() -> None:
    plan = _plan_compound_file(80 * 1024 * 1024)
    header = _compound_file_header(plan)

    assert plan.fat_sector_count > 109
    assert plan.difat_sector_count > 0
    assert struct.unpack_from("<I", header, 44)[0] == plan.fat_sector_count
    assert struct.unpack_from("<I", header, 68)[0] == plan.difat_start
    assert struct.unpack_from("<I", header, 72)[0] == plan.difat_sector_count
