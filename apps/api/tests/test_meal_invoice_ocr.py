"""식당 명세서 OCR 도메인의 순수 함수 단위 테스트.

인프라(비전 엔드포인트/MinIO/DB)에 의존하지 않는 파싱·정규화·검증·내보내기 로직만 검증한다.
회귀 방지 대상은 코드 리뷰에서 확정된 결함들이다(원산지 연결표기, kg→k 중복판정,
_num 숫자 추출, 카탈로그 헤더 폴백 등).
"""

from __future__ import annotations

import asyncio
import io
import json
import zipfile
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from open_alm_api.domains.auth.workspace_apps import get_workspace_app_catalog_item
from open_alm_api.domains.meal_invoice_ocr import APP_ID
from open_alm_api.domains.meal_invoice_ocr import catalog as catalog_mod
from open_alm_api.domains.meal_invoice_ocr import service, vocab
from open_alm_api.domains.meal_invoice_ocr.export import (
    _format_date,
    _num,
    _quantity_value,
    _split_qty,
    export_documents_xlsx,
)
from open_alm_api.domains.meal_invoice_ocr.schemas import (
    _MAX_CORRECTION_ITEMS,
    _MAX_EXPORT_DOCS,
    _MAX_IMAGE_CHARS,
    MealInvoiceCorrectionItem,
    MealInvoiceCorrectionSaveRequest,
    MealInvoiceDocument,
    MealInvoiceExportDocument,
    MealInvoiceExportRequest,
    MealInvoiceExportRow,
    MealInvoiceRow,
)


# --------------------------------------------------------------------------- #
# 앱 등록
# --------------------------------------------------------------------------- #
def test_app_registered_without_hardcoded_category() -> None:
    item = get_workspace_app_catalog_item(APP_ID)
    assert item is not None
    assert item.app_id == APP_ID
    assert item.route_base.endswith("/meal-invoice-ocr")


# --------------------------------------------------------------------------- #
# 숫자 파싱
# --------------------------------------------------------------------------- #
def test_parse_number_strips_commas_units_and_uncertainty() -> None:
    assert service.parse_number("3,500") == 3500
    assert service.parse_number("3,500?") == 3500
    assert service.parse_number("12통") == 12
    assert service.parse_number("") is None
    assert service.parse_number("미정") is None
    assert service.parse_number("-") is None


# --------------------------------------------------------------------------- #
# 품명 정제
# --------------------------------------------------------------------------- #
def test_clean_item_name_removes_parens_brand_grade_and_units() -> None:
    assert service._clean_item_name("생삼겹(국비상)") == "생삼겹"
    assert service._clean_item_name("[비축]된장") == "된장"
    assert service._clean_item_name("햇살담은 진간장골드13L").endswith("진간장")
    assert service._clean_item_name("양조식초") in {"양조식초", "식초"}


# --------------------------------------------------------------------------- #
# 원산지 추출/정규화  (리뷰 #8 회귀: 연결표기 '중국산낙지')
# --------------------------------------------------------------------------- #
def test_extract_origin_from_concatenated_country() -> None:
    # '~산'이 명시되면 뒤에 한글이 붙어도 원산지를 뽑고 품명에서 제거한다.
    origin, name = service._extract_origin("중국산낙지")
    assert origin == "중국산"
    assert "중국" not in name
    assert name == "낙지"


def test_extract_origin_paren_and_comma_forms() -> None:
    assert service._extract_origin("낙지(중국산/냉동/절단)")[0] == "중국산"
    assert service._extract_origin("두부(미국)")[0] == "미국산"
    assert service._extract_origin("…,중국산,스카이푸드")[0] == "중국산"


def test_extract_origin_paren_abbreviation_markers() -> None:
    # 손글씨 축약 마커: '(중)'→중국산, '(베트)'→베트남산. 원산지만 옮기고 품명은 순수 품목명.
    assert service._extract_origin("고춧가루(중)") == ("중국산", "고춧가루")
    assert service._extract_origin("깐감자(중)") == ("중국산", "깐감자")
    assert service._extract_origin("디포리(베트)") == ("베트남산", "디포리")
    assert service._extract_origin("두부(미)") == ("미국산", "두부")


def test_extract_origin_ignores_domestic_grade_markers() -> None:
    # '국비상'(국내산 비상품)은 원산지가 아니다.
    origin, _ = service._extract_origin("생삼겹(국비상)")
    assert origin == ""


def test_normalize_origin_validates_against_known_countries() -> None:
    assert service._normalize_origin("중국") == "중국산"
    assert service._normalize_origin("중") == "중국산"
    assert service._normalize_origin("中") == "중국산"
    # 축약 마커도 나라 원산지로 인정.
    assert service._normalize_origin("베트") == "베트남산"
    assert service._normalize_origin("베트남산") == "베트남산"
    # 상호·브랜드는 원산지가 아니므로 국내산으로 강제.
    assert service._normalize_origin("태화") == "국내산"
    assert service._normalize_origin("") == "국내산"


def test_confirmed_origin_distinguishes_unknown_from_domestic() -> None:
    # 근거가 없는 값은 '국내산'을 만들지 않고 빈 문자열로 둔다(호출측이 다른 칸을 이어 볼 수 있게).
    assert service._confirmed_origin("") == ""
    assert service._confirmed_origin("스카이푸드") == ""
    # 명시적 국내산은 '근거 있는 값'이므로 빈 문자열이 아니다(뒤 칸 조회를 멈춰야 한다).
    assert service._confirmed_origin("국내산") == "국내산"
    # 나라·축약 마커 판정은 _normalize_origin 과 동일하다.
    assert service._confirmed_origin("중국") == "중국산"
    assert service._confirmed_origin("中") == "중국산"
    assert service._confirmed_origin("베트") == "베트남산"


# --------------------------------------------------------------------------- #
# 금액/합계 검증
# --------------------------------------------------------------------------- #
def test_annotate_row_amount_verdict() -> None:
    ok = service.annotate_row(MealInvoiceRow(수량="10", 단가="100", 금액="1000"))
    assert ok.금액검증 == "OK"
    bad = service.annotate_row(MealInvoiceRow(수량="10", 단가="100", 금액="999"))
    assert bad.금액검증.startswith("확인")
    unknown = service.annotate_row(MealInvoiceRow(수량="", 단가="100", 금액="1000"))
    assert unknown.금액검증 == ""


def test_annotate_document_total_matches_supply_or_total() -> None:
    doc = MealInvoiceDocument(
        합계금액="1100",
        공급가액="1000",
        품목=[
            MealInvoiceRow(금액="600"),
            MealInvoiceRow(금액="400"),
        ],
    )
    service.annotate_document(doc)
    # 행합 1000 == 공급가액 → 일치(부가세 분리 양식).
    assert doc.합계검증 == "일치"


def test_annotate_document_flags_mismatch() -> None:
    doc = MealInvoiceDocument(
        합계금액="5000",
        공급가액="",
        품목=[MealInvoiceRow(금액="600"), MealInvoiceRow(금액="400")],
    )
    service.annotate_document(doc)
    assert doc.합계검증.startswith("불일치")


# --------------------------------------------------------------------------- #
# 페이로드 → 문서  (리뷰 #4 회귀: kg→k 중복판정)
# --------------------------------------------------------------------------- #
def test_document_from_payload_qty_with_embedded_kg_not_duplicated() -> None:
    # 모델이 수량에 이미 단위를 붙이고(10kg) 단위 칸에도 kg 를 넣는 흔한 경우.
    payload = {
        "품목": [{"품명": "쌀", "수량": "10kg", "단위": "kg", "단가": "1000", "금액": "10000"}],
    }
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    row = doc.품목[0]
    # "10kg k" 같은 문자열 손상이 없어야 한다 → split 시 숫자 수량이 보존된다.
    num, unit = _split_qty(row.수량)
    assert num == "10"
    assert unit == "k"
    assert service.parse_number(row.수량) == 10


def test_document_from_payload_qty_and_separate_unit() -> None:
    payload = {"품목": [{"품명": "무", "수량": "12", "단위": "박스"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    num, unit = _split_qty(doc.품목[0].수량)
    assert num == "12"
    assert unit == "박스"


@pytest.mark.parametrize(
    ("quantity", "unit", "expected"),
    (
        ("미정", "kg", "미정 k"),
        ("1/2", "box", "1/2 box"),
        ("1-2", "kg", "1-2 k"),
        ("12 ?", "kg", "12 ? k"),
        ("+3", "kg", "+3 k"),
        ("10kg", "box", "10k box"),
        ("12kg?", "kg", "12k?"),
        ("10pack?", "k", "10pack? k"),
        ("10american?", "can", "10american? can"),
        ("10outbox?", "box", "10outbox? box"),
        ("10backpack?", "pack", "10backpack? pack"),
        ("10낱개?", "개", "10낱개? 개"),
        ("10낱개?", "낱개", "10낱개?"),
        ("1/2 box", "box", "1/2 box"),
        ("1/2box", "box", "1/2box"),
        ("1-2kg", "kg", "1-2k"),
        (".5kg", "kg", ".5k"),
        ("12~13kg", "kg", "12~13k"),
        ("미정 kg", "kg", "미정 k"),
        ("미정kg", "kg", "미정k"),
        ("12 ? kg", "kg", "12 ? k"),
    ),
)
def test_document_from_payload_preserves_noncanonical_quantity_with_unit(
    quantity: str,
    unit: str,
    expected: str,
) -> None:
    payload = {"품목": [{"품명": "무", "수량": quantity, "단위": unit}]}

    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)

    assert doc.품목[0].수량 == expected


def test_document_from_payload_preserves_noncanonical_quantity_with_promoted_spec_unit() -> None:
    payload = {"품목": [{"품명": "무", "수량": "1/2", "규격": "박스", "단위": ""}]}

    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)

    assert doc.품목[0].수량 == "1/2 박스"
    assert doc.품목[0].규격 == ""


def test_document_from_payload_origin_from_name() -> None:
    payload = {"품목": [{"품명": "중국산낙지", "수량": "1", "단위": "박스"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "중국산"
    assert doc.품목[0].품명 == "낙지"


def test_document_from_payload_origin_from_parenthesized_country() -> None:
    # 손글씨 '품목(미국산)'처럼 괄호 안 나라 원산지(중국뿐 아니라 미국 등)를 원산지로 옮기고
    # 품명은 괄호 앞만 남긴다. (회귀: 우전각(미국산)이 국내산으로 잘못 인식되던 문제.)
    payload = {"품목": [{"품명": "우전각(미국산)", "수량": "16", "단위": "봉"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "미국산"
    assert doc.품목[0].품명 == "우전각"


def test_document_from_payload_origin_from_field_normalized() -> None:
    # 품명엔 나라가 없고 별도 '원산지' 칸에만 나라가 온 경우도 '~산'으로 정규화해 채운다.
    payload = {"품목": [{"품명": "우전각", "원산지": "미국", "수량": "16", "단위": "봉"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "미국산"
    assert doc.품목[0].품명 == "우전각"


def test_document_from_payload_origin_from_규격_when_other_columns_lack_it() -> None:
    # 회귀: 인쇄된 '중국산'을 모델이 원산지 칸이 아니라 규격 칸에만 남기면 국내산으로 단정됐다.
    # (관측 사례: 품명 '낙지볶음소스' / 규격 '중국산,스카이푸드' → 국내산)
    payload = {
        "품목": [
            {"품명": "낙지볶음소스", "규격": "중국산,스카이푸드", "수량": "3", "단위": "ea"}
        ]
    }
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    row = doc.품목[0]
    assert row.원산지 == "중국산"
    # 규격 원문은 표시용으로 보존한다(OCR 원문 충실도).
    assert row.규격 == "중국산,스카이푸드"
    # 학습 키(원문)도 보정 없이 이 판정을 그대로 갖는다.
    assert row.원문.원산지 == "중국산"


def test_document_from_payload_origin_from_규격_when_원산지_칸이_브랜드() -> None:
    # 모델이 원산지 칸에 상호·브랜드를 넣어 근거가 없을 때도 규격 칸의 나라 표기를 살린다.
    payload = {
        "품목": [
            {
                "품명": "낙지볶음소스",
                "원산지": "스카이푸드",
                "규격": "중국산,2KG/EA",
                "수량": "3",
                "단위": "ea",
            }
        ]
    }
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "중국산"


def test_document_from_payload_origin_from_규격_bare_country_stem() -> None:
    # '냉동,중국'처럼 '산' 없이 적힌 규격도 나라 원산지로 인정한다(품명 규칙과 동일).
    payload = {"품목": [{"품명": "무청시래기", "규격": "냉동,중국", "수량": "5", "단위": "pk"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "중국산"


def test_document_from_payload_규격_whole_paren_group_is_still_origin() -> None:
    # 규격 전체가 괄호인 관측 형태('(냉동,중국)')는 원료 주석이 아니라 그 품목의 표기다.
    # 단어에 붙은 괄호만 원료 주석으로 떼어 내므로 이 경로는 계속 인식돼야 한다.
    for 규격 in ("(냉동,중국)", "[중국산]", "(중국산/절단)"):
        payload = {"품목": [{"품명": "낙지", "규격": 규격, "수량": "1"}]}
        doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
        assert doc.품목[0].원산지 == "중국산", 규격


def test_origin_from_규격_prefers_foreign_stem_regardless_of_order() -> None:
    # 규격에 국내·수입 표기가 함께 오면 작성 순서와 무관하게 비국내산을 택한다(_confirmed_origin 과 동일).
    assert service._origin_from_spec("국내산/중국산") == "중국산"
    assert service._origin_from_spec("중국,국내") == "중국산"
    assert service._origin_from_spec("냉동,국내산") == "국내산"
    assert service._origin_from_spec("국내산") == "국내산"


def test_document_from_payload_규격_origin_does_not_override_earlier_columns() -> None:
    # 앞선 칸에 근거가 있으면 규격은 보지 않는다(우선순위: 품명 > 원산지 칸 > 규격).
    from_name = service._document_from_payload(
        {"품목": [{"품명": "낙지(미국산)", "규격": "중국산,스카이푸드", "수량": "1"}]},
        filename="a.pdf",
        page_no=1,
    )
    assert from_name.품목[0].원산지 == "미국산"
    from_field = service._document_from_payload(
        {"품목": [{"품명": "낙지", "원산지": "미국", "규격": "중국산,스카이푸드", "수량": "1"}]},
        filename="a.pdf",
        page_no=1,
    )
    assert from_field.품목[0].원산지 == "미국산"
    # 반면 '원산지' 칸의 국내산은 모델 기본값("표시 없으면 국내산")이므로 체인을 멈추지 않는다.
    # 문서에 인쇄된 나라 표기가 뒤에 있으면 그것을 쓴다.
    model_default_domestic = service._document_from_payload(
        {"품목": [{"품명": "낙지", "원산지": "국내산", "규격": "중국산,스카이푸드", "수량": "1"}]},
        filename="a.pdf",
        page_no=1,
    )
    assert model_default_domestic.품목[0].원산지 == "중국산"


def test_foreign_origin_treats_model_default_domestic_as_unknown() -> None:
    # 판단 필드(원산지 칸)의 국내산은 근거가 아니라 기본값이므로 '미확인'으로 넘긴다.
    assert service._foreign_origin("국내산") == ""
    assert service._foreign_origin("") == ""
    assert service._foreign_origin("스카이푸드") == ""
    # 실제 나라는 그대로 근거로 인정한다.
    assert service._foreign_origin("중국") == "중국산"
    assert service._foreign_origin("미국산") == "미국산"


def test_rescan_promotes_origin_from_verbatim_name_text(monkeypatch) -> None:
    # 회귀(실측): pass-1 이 품명을 정제하며 ',중국산,스카이푸드'를 지우고 원산지 칸엔 기본값 국내산을
    # 넣어 인쇄된 '중국산'이 유실됐다. 재판독은 품명을 원문 그대로 읽으므로, 재판독 원산지 칸이
    # 국내산으로 와도 품명 텍스트의 나라를 근거로 승격해야 한다.
    payload = {
        "품목": [
            {
                "품명": "고추기나치볶음소스",
                "원산지": "국내산",
                "규격": "스카이푸드, 2Kg/EA",
                "수량": "3",
                "단위": "ea",
                "box": [400, 430],
            }
        ]
    }
    rescanned = {
        "행": [
            {"품명": "고추명가낙지볶음소스,중국산,스카이푸드", "원산지": "국내산", "단위": "ea"}
        ]
    }
    monkeypatch.setattr(service, "_render_band_png", lambda png, y0, y1: png)
    monkeypatch.setattr(
        service, "_invoke_vision", lambda png, **kwargs: json.dumps(rescanned, ensure_ascii=False)
    )
    assert service._rescan_band_into_payload(
        b"png", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["원산지"] == "중국산"
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "중국산"
    # 표시용 품명은 pass-1 값을 유지한다(재판독 원문을 화면 값으로 승격하지 않는다).
    assert doc.품목[0].품명 == "고추기나치볶음소스"


def test_rescan_count_mismatch_requires_same_item_similarity(monkeypatch) -> None:
    # 행 수가 어긋난 경로는 위치 근거가 없어 유사도가 유일한 판단이다. 서로 다른 실재 품목
    # ('양파'↔'양상추' 0.615, '깐감자'↔'깐강마늘' 0.667)에 원산지·단위가 이식되면 안 된다.
    payload = {
        "품목": [
            {"품명": "양파", "수량": "10", "단위": "k", "box": [100, 130]},
            {"품명": "깐감자", "수량": "25", "단위": "k", "box": [140, 170]},
        ]
    }
    rescanned = {
        "행": [{"품명": "양상추", "원산지": "중국산", "단위": "박스"}],
    }
    monkeypatch.setattr(service, "_render_band_png", lambda png, y0, y1: png)
    monkeypatch.setattr(
        service, "_invoke_vision", lambda png, **kwargs: json.dumps(rescanned, ensure_ascii=False)
    )
    service._rescan_band_into_payload(
        b"png", payload, db=None, workspace_id="w", actor_user_id=None
    )
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert [r.원산지 for r in doc.품목] == ["국내산", "국내산"]
    assert [r.수량 for r in doc.품목] == ["10 k", "25 k"]


def test_rescan_count_mismatch_still_matches_same_item(monkeypatch) -> None:
    # 같은 품목의 오독 변형('청양고추'↔'청양추' 0.889)은 계속 짝지어 원산지·단위를 보강해야 한다.
    payload = {"품목": [{"품명": "청양고추", "수량": "2", "단위": "k", "box": [100, 130]},
                       {"품명": "무", "수량": "1", "단위": "k", "box": [140, 170]}]}
    rescanned = {"행": [{"품명": "청양추", "원산지": "중국산", "단위": "박스"}]}
    monkeypatch.setattr(service, "_render_band_png", lambda png, y0, y1: png)
    monkeypatch.setattr(
        service, "_invoke_vision", lambda png, **kwargs: json.dumps(rescanned, ensure_ascii=False)
    )
    service._rescan_band_into_payload(
        b"png", payload, db=None, workspace_id="w", actor_user_id=None
    )
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "중국산"
    assert doc.품목[0].수량 == "2 박스"
    assert doc.품목[1].원산지 == "국내산"


def test_rescan_does_not_invent_origin_from_name_text(monkeypatch) -> None:
    # 재판독 품명 원문에 나라가 없으면 원산지를 만들지 않는다(오탐 방지).
    payload = {
        "품목": [
            {"품명": "냉면육수", "원산지": "국내산", "수량": "3", "단위": "ea", "box": [400, 430]}
        ]
    }
    rescanned = {
        "행": [{"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA", "원산지": "국내산", "단위": "ea"}]
    }
    monkeypatch.setattr(service, "_render_band_png", lambda png, y0, y1: png)
    monkeypatch.setattr(
        service, "_invoke_vision", lambda png, **kwargs: json.dumps(rescanned, ensure_ascii=False)
    )
    service._rescan_band_into_payload(
        b"png", payload, db=None, workspace_id="w", actor_user_id=None
    )
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "국내산"


@pytest.mark.parametrize(
    "rescanned_name",
    ["감자 중", "감자 中", "감자(중)", "감자 대"],
)
def test_rescan_does_not_read_size_grade_in_name_as_country(
    monkeypatch, rescanned_name: str
) -> None:
    # 재판독은 품명을 원문 그대로 읽으므로 크기 등급('중'·'中'·'(중)')이 함께 온다. 축약 마커를
    # 나라로 읽으면 pass-1 의 국내산 농산물이 중국산으로 덮인다(규격 칸과 같은 판단 기준).
    payload = {
        "품목": [{"품명": "감자", "원산지": "국내산", "수량": "10", "단위": "k", "box": [100, 130]}]
    }
    rescanned = {"행": [{"품명": rescanned_name, "원산지": "국내산", "단위": "k"}]}
    monkeypatch.setattr(service, "_render_band_png", lambda *a, **k: b"PNG")
    monkeypatch.setattr(
        service, "_invoke_vision", lambda png, **kwargs: json.dumps(rescanned, ensure_ascii=False)
    )
    service._rescan_band_into_payload(
        b"png", payload, db=None, workspace_id="w", actor_user_id=None
    )
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "국내산", rescanned_name
    assert doc.품목[0].원문.원산지 == "국내산", rescanned_name


@pytest.mark.parametrize(
    "규격",
    [
        # 농산물 명세표의 크기 등급. '중'을 손글씨 축약 마커로 읽으면 국내산 농산물이 중국산이 된다.
        "대",
        "중",
        "소",
        "특호",
        "1호",
        # 축약 한 글자는 규격 칸에서 크기·등급과 구분할 수 없어 원산지로 인정하지 않는다.
        "미",
        "호",
        "(중)",
        # 문장에 섞인 '원료' 원산지는 그 품목의 원산지가 아니다(가공식품 규격에 흔하다).
        "중국산 고춧가루 사용",
        "중국산 대두 혼합",
        # 같은 원료 원산지를 괄호로 적은 형태. 구분자로 나누면 조각 전체가 나라 표기가 되므로,
        # 단어에 붙은 괄호 주석을 떼어 내 걸러야 한다(국내산 가공식품이 중국산으로 표기되던 경로).
        "고춧가루(중국산)",
        "고춧가루 (중국산)",
        "원료:고춧가루(중국산),국내제조",
        "대두(미국산)/밀(호주산)",
        "정제소금(국내산)/고추양념(중국산)",
    ],
)
def test_document_from_payload_규격_grade_and_ingredient_notes_are_not_origin(규격: str) -> None:
    payload = {"품목": [{"품명": "무", "규격": 규격, "수량": "1"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "국내산", 규격


def test_document_from_payload_규격_ingredient_note_does_not_override_원산지_칸() -> None:
    # 회귀: 원산지 칸의 국내산이 fallback 을 막지 않게 바꾼 뒤, 규격의 원료 원산지 문구가
    # 품목 원산지를 덮는 부작용이 생겼다.
    payload = {
        "품목": [{"품명": "고추장", "원산지": "국내산", "규격": "중국산 고춧가루 사용", "수량": "1"}]
    }
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    assert doc.품목[0].원산지 == "국내산"


def test_document_from_payload_규격_without_country_stays_domestic() -> None:
    # 나라 표기가 없는 평범한 규격은 원산지에 영향을 주지 않는다(오탐 방지).
    for 규격 in ("2KG/EA", "71-90 1kg/EA", "DC,광동,10L", "청대구40%", "특호"):
        payload = {"품목": [{"품명": "품목", "규격": 규격, "수량": "1"}]}
        doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
        assert doc.품목[0].원산지 == "국내산", 규격


# --------------------------------------------------------------------------- #
# 원산지 재판독(pass-2): 데이터 행 band 확대 재판독으로 손글씨 원산지 보강
# (전체 페이지 다운샘플로 1차에서 놓친 '(미국산)' 등을 잡되, 국내산→나라 승격만 한다.)
# --------------------------------------------------------------------------- #
def _patch_rescan(monkeypatch: "pytest.MonkeyPatch", rescan_json: str) -> None:
    # PIL 렌더를 우회하고(가짜 PNG), 확대 band 재판독 응답을 주입한다.
    monkeypatch.setattr(service, "_render_band_png", lambda *a, **k: b"PNG")
    monkeypatch.setattr(service, "_call_vision", lambda *a, **k: rescan_json)


def test_rescan_upgrades_domestic_to_country(monkeypatch: "pytest.MonkeyPatch") -> None:
    payload = {
        "품목": [
            {"품명": "돈전지", "원산지": "국내산", "box": [380, 420]},
            {"품명": "우전각", "원산지": "국내산", "box": [460, 500]},
        ]
    }
    _patch_rescan(
        monkeypatch,
        '{"행":[{"품명":"돈전지","원산지":"국내산"},{"품명":"우전각","원산지":"미국산"}]}',
    )
    made = service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert made is True
    assert payload["품목"][0]["원산지"] == "국내산"  # 실제 국내산은 그대로.
    assert payload["품목"][1]["원산지"] == "미국산"  # 놓쳤던 미국산 보강.


def test_rescan_does_not_downgrade_existing_country(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    # 1차에서 이미 나라가 잡힌 행은 재판독이 국내산이라 해도 되돌리지 않는다.
    payload = {"품목": [{"품명": "낙지", "원산지": "중국산", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"낙지","원산지":"국내산"}]}')
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["원산지"] == "중국산"


def test_rescan_skips_when_no_boxes(monkeypatch: "pytest.MonkeyPatch") -> None:
    # box 가 없으면 band 를 만들 수 없으므로 재판독 호출 자체를 하지 않는다(예산 낭비 방지).
    calls = {"n": 0}

    def spy(*_a: object, **_k: object) -> str:
        calls["n"] += 1
        return "{}"

    monkeypatch.setattr(service, "_call_vision", spy)
    payload = {"품목": [{"품명": "돈전지", "원산지": "국내산"}]}
    made = service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert made is False
    assert calls["n"] == 0
    assert payload["품목"][0]["원산지"] == "국내산"


def test_rescan_uses_dedicated_workload(monkeypatch: "pytest.MonkeyPatch") -> None:
    from open_alm_api.domains.meal_invoice_ocr.task_kinds import (
        MEAL_INVOICE_OCR_RESCAN_WORKLOAD_ID,
    )

    captured: dict[str, object] = {}

    monkeypatch.setattr(service, "_render_band_png", lambda *a, **k: b"PNG")

    def spy(*_args: object, **kwargs: object) -> str:
        captured.update(kwargs)
        return '{"행":[]}'

    monkeypatch.setattr(service, "_call_vision", spy)
    made = service._rescan_band_into_payload(
        b"PAGE",
        {"품목": [{"품명": "양파", "box": [100, 140]}]},
        db=None,
        workspace_id="w",
        actor_user_id=None,
    )

    assert made is True
    assert captured["workload_id"] == MEAL_INVOICE_OCR_RESCAN_WORKLOAD_ID


def test_rescan_ignores_unmatched_names_on_count_mismatch(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    # 재판독 행 수가 1차와 다르면(순서 신뢰 불가) 품명 유사도로 매칭한다. 전혀 다른 이름들만 있으면
    # 매칭 실패 → 원산지를 함부로 옮기지 않는다. (행 수가 같을 땐 위치 정렬로 보강함 — 별도 테스트.)
    payload = {"품목": [{"품명": "돈전지", "원산지": "국내산", "box": [380, 420]}]}
    _patch_rescan(
        monkeypatch,
        '{"행":[{"품명":"전혀다른A","원산지":"미국산"},{"품명":"전혀다른B","원산지":"중국산"}]}',
    )
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["원산지"] == "국내산"


def test_rescan_fills_unit_from_zoomed_read(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 전체 페이지가 단위를 놓쳐 비었더라도, 확대 재판독이 읽은 단위를 채운다(숫자는 보존).
    payload = {"품목": [{"품명": "양파", "수량": "6", "단위": "", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"양파","원산지":"국내산","단위":"박스"}]}')
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["단위"] == "박스"
    assert payload["품목"][0]["수량"] == "6"


def test_rescan_unit_overrides_full_page_unit(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 확대 재판독 단위가 전체 페이지 단위와 다르면 재판독값을 우선한다(고해상도 신뢰). 숫자는 보존.
    payload = {"품목": [{"품명": "양상추", "수량": "10", "단위": "k", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"양상추","원산지":"국내산","단위":"박스"}]}')
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["단위"] == "박스"
    assert payload["품목"][0]["수량"] == "10"


def test_rescan_rejects_unit_outside_server_owned_options(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    payload = {"품목": [{"품명": "양상추", "수량": "10", "단위": "송이", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"양상추","원산지":"국내산","단위":"통나무"}]}')

    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )

    assert payload["품목"][0]["단위"] == "송이"


def test_rescan_keeps_unit_when_zoomed_read_empty(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 확대 재판독이 단위를 못 읽으면(빈값) 기존 단위를 덮지 않는다(억지 지우기 방지).
    payload = {"품목": [{"품명": "양파", "수량": "6", "단위": "박스", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"양파","원산지":"국내산","단위":""}]}')
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["단위"] == "박스"


def _name_catalog(*items: str) -> dict[str, "Any"]:
    return {
        it: {"단위": "k", "단위들": ["k"], "단위카운트": {"k": 3}, "단가": "1000", "count": 3}
        for it in items
    }


def test_rescan_suggests_name_from_catalog_match(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 확대 재판독이 다른 카탈로그 품명을 읽어도 원값은 유지하고 검수 후보로만 제시한다.
    payload = {
        "품목": [
            {"품명": "후추", "수량": "5", "box": [100, 140]},
            {"품명": "양파", "수량": "6", "box": [150, 190]},
        ]
    }
    _patch_rescan(
        monkeypatch,
        '{"행":[{"품명":"당근","원산지":"국내산","단위":""},{"품명":"양파","원산지":"국내산","단위":""}]}',
    )
    service._rescan_band_into_payload(
        b"PAGE",
        payload,
        db=None,
        workspace_id="w",
        actor_user_id=None,
        catalog=_name_catalog("당근", "후추", "양파"),
    )
    assert payload["품목"][0]["품명"] == "후추"
    assert payload["품목"][0]["_rescan_name_candidate"] == "당근"
    assert payload["품목"][1]["품명"] == "양파"  # 동일하면 변화 없음.


def test_rescan_never_overwrites_one_exact_catalog_name_with_another(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    payload = {"품목": [{"품명": "양파", "수량": "5", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"양상추","원산지":"국내산","단위":""}]}')

    service._rescan_band_into_payload(
        b"PAGE",
        payload,
        db=None,
        workspace_id="w",
        actor_user_id=None,
        catalog=_name_catalog("양파", "양상추"),
    )

    assert payload["품목"][0]["품명"] == "양파"
    assert payload["품목"][0]["_rescan_name_candidate"] == "양상추"
    doc = service._document_from_payload(payload, filename="x.jpg", page_no=1)
    assert doc.품목[0].품명 == "양파"
    assert doc.품목[0].사전후보 == "양상추"


def test_rescan_keeps_name_when_rescan_not_in_catalog(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 확대 재판독 이름이 대장 실제 품목과 안 맞으면(지어낸 값 가능성) 1차 품명을 유지한다.
    payload = {"품목": [{"품명": "후추", "수량": "5", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"없는품목xyz","원산지":"국내산","단위":""}]}')
    service._rescan_band_into_payload(
        b"PAGE",
        payload,
        db=None,
        workspace_id="w",
        actor_user_id=None,
        catalog=_name_catalog("후추"),
    )
    assert payload["품목"][0]["품명"] == "후추"


def test_rescan_no_name_correction_without_catalog(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 대장이 없으면 품명 교정을 하지 않는다(원산지·단위만 보강).
    payload = {"품목": [{"품명": "후추", "수량": "5", "box": [100, 140]}]}
    _patch_rescan(monkeypatch, '{"행":[{"품명":"당근","원산지":"국내산","단위":""}]}')
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["품명"] == "후추"


def test_rescan_skips_origin_on_misaligned_rows(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 재판독 행 수는 1차와 같지만(위치 정렬 후보) 이름이 전혀 달라 '같은 행'이라 믿을 수 없다.
    # 이런 쌍에는 원산지/단위를 옮기지 않는다 — 개수만 우연히 맞고 행이 어긋난 경우(모델 병합·분리)
    # 엉뚱한 품목에 원산지가 이식되는 것을 막기 위함(대장 교정 확정 쌍은 예외).
    payload = {
        "품목": [
            {"품명": "돈전지", "원산지": "국내산", "box": [100, 140]},
            {"품명": "우전각", "원산지": "국내산", "box": [150, 190]},
        ]
    }
    _patch_rescan(
        monkeypatch,
        '{"행":[{"품명":"전혀다른A","원산지":"미국산"},{"품명":"전혀다른B","원산지":"중국산"}]}',
    )
    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )
    assert payload["품목"][0]["원산지"] == "국내산"
    assert payload["품목"][1]["원산지"] == "국내산"


def test_rescan_count_mismatch_consumes_each_candidate_once(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    payload = {
        "품목": [
            {"품명": "양파", "원산지": "국내산", "단위": "k", "box": [100, 140]},
            {"품명": "양파", "원산지": "국내산", "단위": "k", "box": [150, 190]},
        ]
    }
    _patch_rescan(
        monkeypatch,
        '{"행":[{"품명":"양파","원산지":"미국산","단위":"박스"},'
        '{"품명":"전혀다른품목","원산지":"중국산","단위":"통"},'
        '{"품명":"또다른품목","원산지":"호주산","단위":"병"}]}',
    )

    service._rescan_band_into_payload(
        b"PAGE", payload, db=None, workspace_id="w", actor_user_id=None
    )

    changed = [item for item in payload["품목"] if item["원산지"] == "미국산"]
    assert len(changed) == 1
    assert [item["단위"] for item in payload["품목"]].count("박스") == 1


@pytest.mark.parametrize(
    ("size", "pixel_limit", "expected_relation"),
    (
        ((100, 20), 10_000, "doubled"),
        ((100, 100), 10_000, "unchanged"),
        ((200, 100), 10_000, "downscaled"),
    ),
)
def test_render_band_png_respects_output_pixel_cap(
    monkeypatch: "pytest.MonkeyPatch",
    size: tuple[int, int],
    pixel_limit: int,
    expected_relation: str,
) -> None:
    from PIL import Image

    source = io.BytesIO()
    Image.new("RGB", size, "white").save(source, format="PNG")
    monkeypatch.setattr(service, "MAX_RENDER_PIXELS", pixel_limit)

    rendered = service._render_band_png(source.getvalue(), 0.0, 1.0)

    with Image.open(io.BytesIO(rendered)) as result:
        assert result.width * result.height <= pixel_limit
        if expected_relation == "doubled":
            assert result.size == (size[0] * 2, size[1] * 2)
        elif expected_relation == "unchanged":
            assert result.size == size
        else:
            assert result.width < size[0]
            assert result.height < size[1]


# --------------------------------------------------------------------------- #
# 다중 페이지/파일 순차 처리 — 순서·trace context 보존 + 실패 격리
# --------------------------------------------------------------------------- #
class _FakeDoc:
    def __init__(self, name: str, pages: int) -> None:
        self._name = name
        self.page_count = pages

    def __getitem__(self, i: int) -> str:
        return f"{self._name}#p{i}"

    def close(self) -> None:  # noqa: D401
        pass


def _patch_extract_pipeline(
    monkeypatch: "pytest.MonkeyPatch", pages_by_file: dict[str, int], fail_on: set[str]
) -> None:
    import json as _json

    monkeypatch.setattr(
        service,
        "_open_document",
        lambda upload: _FakeDoc(upload.filename, pages_by_file[upload.filename]),
    )
    # page 마커("a.jpg#p0")를 그대로 png 바이트로. 페이지마다 고유.
    monkeypatch.setattr(service, "_render_page_png", lambda page: page.encode())
    monkeypatch.setattr(service, "_render_page_display_data_url", lambda page: "")
    monkeypatch.setattr(service, "_rescan_band_into_payload", lambda *a, **k: False)
    # 각 작업은 자체 세션을 연다 → get_session_factory 를 더미로.
    import open_alm_api.core.db as _db

    monkeypatch.setattr(
        _db, "get_session_factory", lambda: (lambda: SimpleNamespace(close=lambda: None))
    )

    def fake_vision(png: bytes, **_k: object) -> str:
        marker = png.decode()
        if marker in fail_on:
            raise RuntimeError("boom")
        # 거래처에 페이지 마커를 실어 순서 검증에 쓴다.
        return _json.dumps({"거래처": marker, "품목": []})

    monkeypatch.setattr(service, "_call_vision", fake_vision)


def test_extract_documents_sequential_preserves_order(monkeypatch: "pytest.MonkeyPatch") -> None:
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 1, "b.jpg": 2}, fail_on=set())
    uploads = [
        service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x"),
        service.MealInvoiceUpload("b.jpg", "image/jpeg", b"x"),
    ]
    resp = service.extract_documents(uploads, workspace_id="w")
    # 입력(렌더) 순서대로 즉시 처리하므로 문서 순서가 보존돼야 한다.
    assert [d.거래처 for d in resp.documents] == ["a.jpg#p0", "b.jpg#p0", "b.jpg#p1"]
    assert resp.warnings == []


def _patch_extract_with_rescan_counts(
    monkeypatch: "pytest.MonkeyPatch", counts: tuple[int, int]
) -> None:
    """재판독이 (pass-1, pass-2) 행 수를 보고하도록 파이프라인을 대체한다."""
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 1}, fail_on=set())

    def fake_rescan(page_png, payload, *, row_counts=None, **_kwargs):
        if row_counts is not None:
            row_counts.append(counts)
        return True

    monkeypatch.setattr(service, "_rescan_band_into_payload", fake_rescan)


def test_extract_documents_warns_when_rescan_reads_more_rows(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    # 재판독이 더 많은 행을 읽었다 = 1차가 인접 두 행을 합쳐 결과에서 데이터가 빠졌을 수 있다 → 경고.
    _patch_extract_with_rescan_counts(monkeypatch, (10, 12))
    resp = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )
    assert len(resp.warnings) == 1
    # 어느 회차가 틀렸다고 단정하지 않고, 검수자가 실제로 확인할 수 있는 것(원본 줄 수)만 요청한다.
    assert "추출 10행" in resp.warnings[0]
    assert "재판독 12행" in resp.warnings[0]


def test_extract_documents_silent_when_rescan_reads_fewer_rows(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    # 손글씨 사진에서는 재판독이 중간에 끊겨 더 적게 읽는 일이 흔하다(실측 14행 → 10행). 이때 1차가
    # 더 완전한 판독이라 결과에서 빠진 행이 없으므로 경고하지 않는다 — 매번 뜨면 진짜 경고가 묻힌다.
    _patch_extract_with_rescan_counts(monkeypatch, (14, 10))
    resp = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )
    assert resp.warnings == []


def test_extract_documents_preserves_context_and_processes_each_page_incrementally(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    from contextvars import ContextVar

    _patch_extract_pipeline(monkeypatch, {"a.jpg": 2}, fail_on=set())
    trace_marker: ContextVar[str] = ContextVar("ocr_test_trace", default="missing")
    trace_marker.set("request-trace")
    events: list[str] = []

    def render(page: str) -> bytes:
        events.append(f"render:{page}")
        return page.encode()

    def vision(png: bytes, **_kwargs: object) -> str:
        marker = png.decode()
        assert trace_marker.get() == "request-trace"
        events.append(f"vision:{marker}")
        return '{"품목":[]}'

    monkeypatch.setattr(service, "_render_page_png", render)
    monkeypatch.setattr(service, "_call_vision", vision)

    service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )

    assert events == [
        "render:a.jpg#p0",
        "vision:a.jpg#p0",
        "render:a.jpg#p1",
        "vision:a.jpg#p1",
    ]


def test_extract_documents_sequential_isolates_failures(monkeypatch: "pytest.MonkeyPatch") -> None:
    # 한 페이지의 OCR 실패가 다른 페이지를 막지 않고, 순서도 유지된다.
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 3}, fail_on={"a.jpg#p1"})
    resp = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )
    assert [d.거래처 for d in resp.documents] == ["a.jpg#p0", "a.jpg#p2"]
    assert any("p2" in w or "OCR 호출 실패" in w for w in resp.warnings)


def test_extract_documents_isolates_session_checkout_failure(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 2}, fail_on=set())
    import open_alm_api.core.db as db_module

    calls = 0

    def flaky_factory() -> SimpleNamespace:
        nonlocal calls
        calls += 1
        if calls == 1:
            raise RuntimeError("checkout")
        return SimpleNamespace(close=lambda: None)

    monkeypatch.setattr(db_module, "get_session_factory", lambda: flaky_factory)

    response = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )

    assert [doc.거래처 for doc in response.documents] == ["a.jpg#p1"]
    assert any("DB 세션 실패" in warning for warning in response.warnings)


def test_extract_documents_isolates_session_close_failure(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 2}, fail_on=set())
    import open_alm_api.core.db as db_module

    calls = 0

    def session_factory() -> SimpleNamespace:
        nonlocal calls
        calls += 1

        def close() -> None:
            if calls == 1:
                raise RuntimeError("close")

        return SimpleNamespace(close=close)

    monkeypatch.setattr(db_module, "get_session_factory", lambda: session_factory)

    response = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )

    assert [doc.거래처 for doc in response.documents] == ["a.jpg#p0", "a.jpg#p1"]
    assert any("DB 세션 종료 실패" in warning for warning in response.warnings)


def test_extract_documents_preserves_40_page_one_pass_capacity(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    pages_by_file = {f"{index}.jpg": 5 for index in range(8)}
    _patch_extract_pipeline(monkeypatch, pages_by_file, fail_on=set())
    uploads = [
        service.MealInvoiceUpload(filename, "image/jpeg", b"x") for filename in pages_by_file
    ]

    response = service.extract_documents(uploads, workspace_id="w")

    assert len(response.documents) == 40
    assert response.warnings == []


def test_extract_documents_caps_actual_pass1_and_rescan_calls_at_40(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    pages_by_file = {f"{index}.jpg": 5 for index in range(8)}
    _patch_extract_pipeline(monkeypatch, pages_by_file, fail_on=set())
    calls = 0
    original_fake_vision = service._call_vision

    def counting_vision(*args: object, **kwargs: object) -> str:
        nonlocal calls
        calls += 1
        return original_fake_vision(*args, **kwargs)

    def rescan_with_budget(
        page_png: bytes,
        _payload: object,
        *,
        budget: service._VisionCallBudget,
        db: object,
        workspace_id: str,
        actor_user_id: str | None,
        **_kwargs: object,
    ) -> bool:
        service._invoke_vision(
            page_png,
            budget=budget,
            db=db,
            workspace_id=workspace_id,
            actor_user_id=actor_user_id,
        )
        return True

    monkeypatch.setattr(service, "_call_vision", counting_vision)
    monkeypatch.setattr(service, "_rescan_band_into_payload", rescan_with_budget)
    uploads = [
        service.MealInvoiceUpload(filename, "image/jpeg", b"x") for filename in pages_by_file
    ]

    response = service.extract_documents(uploads, workspace_id="w")

    assert calls == 40
    assert len(response.documents) == 20
    assert any("호출 40" in warning for warning in response.warnings)


def test_extract_documents_warns_when_last_call_leaves_rescan_unavailable(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 1}, fail_on=set())
    monkeypatch.setattr(service, "MAX_VISION_CALLS_PER_REQUEST", 1)
    monkeypatch.setattr(service, "MAX_TOTAL_OCR_PAGES", 1)
    monkeypatch.setattr(
        service,
        "_call_vision",
        lambda *_args, **_kwargs: (
            '{"거래처":"a.jpg#p0","품목":[{"품명":"양파","수량":"1","box":[100,140]}]}'
        ),
    )
    rescan_calls = 0

    def rescan(*_args: object, **_kwargs: object) -> bool:
        nonlocal rescan_calls
        rescan_calls += 1
        return True

    monkeypatch.setattr(service, "_rescan_band_into_payload", rescan)

    response = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )

    assert len(response.documents) == 1
    assert rescan_calls == 0
    assert any(
        "재판독 생략" in warning and "호출 상한 1" in warning for warning in response.warnings
    )


def test_extract_documents_reports_rescan_failure(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 1}, fail_on=set())

    def fail_rescan(*_args: object, **_kwargs: object) -> bool:
        raise RuntimeError("rescan")

    monkeypatch.setattr(service, "_rescan_band_into_payload", fail_rescan)

    response = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )

    assert [doc.거래처 for doc in response.documents] == ["a.jpg#p0"]
    assert any("재판독 실패" in warning for warning in response.warnings)


def test_extract_documents_isolates_assembly_failure(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    # 한 페이지의 '조립'(payload→문서/주석) 예외가 executor.map 재던짐으로 배치 전체를 날리지 않고
    # 그 페이지만 경고로 격리되며, 나머지 페이지와 순서는 유지된다.
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 2}, fail_on=set())
    real_build = service._document_from_payload

    def flaky_build(payload: object, **kwargs: object) -> object:
        if kwargs.get("page_no") == 1:
            raise ValueError("assembly boom")
        return real_build(payload, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(service, "_document_from_payload", flaky_build)
    resp = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")], workspace_id="w"
    )
    assert [d.거래처 for d in resp.documents] == ["a.jpg#p1"]
    assert any("조립 실패" in w for w in resp.warnings)


# --------------------------------------------------------------------------- #
# 카탈로그(영양사 대장) 단위 제안 — OCR 이 읽은 단위는 덮어쓰지 않는다(카탈로그는 제안만).
# --------------------------------------------------------------------------- #
def _unit_catalog(item: str, counts: dict[str, int]) -> dict[str, Any]:
    ordered = [u for u, _ in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))]
    return {
        item: {
            "단위": ordered[0],
            "단위들": ordered,
            "단위카운트": counts,
            "단가": "9000",
            "count": sum(counts.values()),
        }
    }


def test_annotate_catalog_does_not_override_read_unit() -> None:
    # OCR 이 '박스'로 읽었으면 카탈로그 이력이 'k'뿐이어도 덮어쓰지 않는다(정상 단위를 뭉개지 않음).
    # 카탈로그 최빈 단위는 단위후보 칩으로만 제안한다.
    doc = MealInvoiceDocument(품목=[MealInvoiceRow(품명="양파", 수량="6 박스")])
    service.annotate_catalog(doc, _unit_catalog("양파", {"k": 9}))
    assert doc.품목[0].수량 == "6 박스"  # OCR 단위 유지.
    assert doc.품목[0].단위후보 == "k"  # 카탈로그 단위는 제안 칩으로만.


def test_annotate_catalog_preserves_varied_units_across_document() -> None:
    # 한 명세표에서 품목마다 단위가 달라도(박스/송이/k) OCR 이 읽은 대로 보존한다(전파 없음).
    doc = MealInvoiceDocument(
        품목=[
            MealInvoiceRow(품명="양파", 수량="6 박스"),
            MealInvoiceRow(품명="양상추", 수량="10 송이"),
            MealInvoiceRow(품명="청양고추", 수량="2 k"),
        ]
    )
    catalog = {
        **_unit_catalog("양파", {"k": 9}),
        **_unit_catalog("양상추", {"k": 5}),
        **_unit_catalog("청양고추", {"k": 7}),
    }
    service.annotate_catalog(doc, catalog)
    assert [r.수량 for r in doc.품목] == ["6 박스", "10 송이", "2 k"]


def test_annotate_catalog_fills_missing_unit_with_catalog() -> None:
    # OCR 이 단위를 못 뽑은 순수 숫자 수량은 카탈로그 최빈 단위로 채운다(빈 단위 보완).
    doc = MealInvoiceDocument(품목=[MealInvoiceRow(품명="포기김치", 수량="2")])
    service.annotate_catalog(doc, _unit_catalog("포기김치", {"k": 3}))
    assert doc.품목[0].수량 == "2 k"


def test_annotate_catalog_keeps_matching_unit() -> None:
    # OCR 단위가 카탈로그 최빈 단위와 같으면 그대로 둔다(불필요한 재작성 없음).
    doc = MealInvoiceDocument(품목=[MealInvoiceRow(품명="돈전지", 수량="12 k")])
    service.annotate_catalog(doc, _unit_catalog("돈전지", {"k": 5}))
    assert doc.품목[0].수량 == "12 k"


def test_document_from_payload_promotes_unit_written_in_규격() -> None:
    # 전용 '단위' 칸이 없고 손글씨로 '규격' 칸에 단위(낱)를 적은 명세표.
    payload = {"품목": [{"품명": "돈전지", "규격": "낱", "수량": "12", "단위": ""}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    row = doc.품목[0]
    num, unit = _split_qty(row.수량)
    assert (num, unit) == ("12", "낱")
    # 단위로 승격한 규격 원문은 중복 표시를 막기 위해 비운다.
    assert row.규격 == ""


def test_document_from_payload_규격_unit_normalized_like_units() -> None:
    # 영문 단위는 소문자, kg 는 k 로 통일(수량·단위와 동일 규칙).
    for 규격, expected in (("EA", "ea"), ("KG", "k"), ("박스", "박스")):
        payload = {"품목": [{"품명": "쌀", "규격": 규격, "수량": "3", "단위": ""}]}
        doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
        _, unit = _split_qty(doc.품목[0].수량)
        assert unit == expected


def test_document_from_payload_real_규격_not_promoted() -> None:
    # 크기·호수·등급 등 진짜 규격은 단위로 승격하지 않고 규격에 남고 단위는 빈칸.
    payload = {"품목": [{"품명": "고등어", "규격": "특호", "수량": "12", "단위": ""}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    row = doc.품목[0]
    _, unit = _split_qty(row.수량)
    assert unit == ""
    assert row.규격 == "특호"


@pytest.mark.parametrize(
    ("규격", "expected_규격", "expected_quantity"),
    (("L", "L", "12"), ("l", "", "12 l")),
)
def test_document_from_payload_preserves_ambiguous_uppercase_l_specification(
    규격: str,
    expected_규격: str,
    expected_quantity: str,
) -> None:
    payload = {"품목": [{"품명": "고등어", "규격": 규격, "수량": "12", "단위": ""}]}

    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)

    assert doc.품목[0].규격 == expected_규격
    assert doc.품목[0].수량 == expected_quantity


def test_document_from_payload_explicit_unit_column_wins_over_규격() -> None:
    # 전용 '단위' 칸이 있으면 규격은 승격하지 않고 단위 칸 값을 쓴다.
    payload = {"품목": [{"품명": "무", "규격": "낱", "수량": "12", "단위": "통"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    _, unit = _split_qty(doc.품목[0].수량)
    assert unit == "통"


def test_unit_correction_never_enters_the_ocr_prompt() -> None:
    """단위·수량·금액 교정은 그 행에서만 참이라 프롬프트에 넣지 않는다(후처리 alias 전용).

    "'1 박스'→'1 10k'" 를 예시로 주면 비전 모델이 다음 명세표의 멀쩡한 단위/숫자까지 바꿔 읽는다.
    """
    from open_alm_api.domains.meal_invoice_ocr import corrections as corr

    records = [
        {
            "거래처": "김치상회",
            "원본": {"품명": "포기김치", "수량": "1 박스"},
            "교정": {"품명": "포기김치", "수량": "1 10k"},
        }
    ]

    assert corr.build_fewshot(records) == ""

    doc = service._document_from_payload(
        {
            "거래처": "김치상회",
            "품목": [
                {"품명": "포기김치10KG", "수량": "2 박스"},
                {"품명": "포기김치5KG", "수량": "2 박스"},
                {"품명": "포기김치(10KG)", "수량": "2 박스"},
                {"품명": "포기김치,5KG", "수량": "2 박스"},
                {"품명": "포기김치", "수량": "2 박스"},
            ],
        },
        filename="a.pdf",
        page_no=1,
    )

    assert [row.품명 for row in doc.품목] == ["포기김치"] * 5
    assert [row.수량 for row in doc.품목] == ["2 박스"] * 5


def test_split_qty_preserves_uncertainty_marker() -> None:
    # OCR 불확실 표시 '?'는 단위가 아니라 숫자에 붙는다(프롬프트 계약). 음수 부호도 숫자에 보존.
    assert service.split_qty("12?") == ("12?", "")
    assert service.split_qty("3,500?") == ("3,500?", "")
    assert service.split_qty("-3") == ("-3", "")
    # 불확실 수량 + 단위칸: '?'가 유실되지 않고 단위와 공존한다.
    payload = {"품목": [{"품명": "쌀", "수량": "12?", "단위": "kg"}]}
    doc = service._document_from_payload(payload, filename="a.pdf", page_no=1)
    num, unit = _split_qty(doc.품목[0].수량)
    assert (num, unit) == ("12?", "k")


# --------------------------------------------------------------------------- #
# 내보내기 헬퍼  (리뷰 #5 회귀: _num 숫자 추출)
# --------------------------------------------------------------------------- #
def test_num_extracts_numeric_from_uncertain_value() -> None:
    assert _num("3,500?") == 3500
    assert _num("3,500") == 3500
    assert _num("1234.5") == 1234.5
    assert _num("") == ""
    # 숫자가 전혀 없으면 원문 보존.
    assert _num("미정") == "미정"


def test_quantity_value_only_coerces_lossless_numbers() -> None:
    assert _quantity_value("3,500?") == 3500
    assert _quantity_value("-3") == -3
    assert _quantity_value("+3") == "+3"
    assert _quantity_value("1/2") == "1/2"
    assert _quantity_value("12~13") == "12~13"
    assert _quantity_value("--") == "--"


def test_split_qty_space_and_legacy_forms() -> None:
    assert _split_qty("418 10k") == ("418", "10k")
    assert _split_qty("12박스") == ("12", "박스")
    assert _split_qty("30") == ("30", "")
    assert _split_qty("1/2 box") == ("1/2", "box")
    assert _split_qty("12 ? k") == ("12 ?", "k")
    assert _split_qty("+3k") == ("+3", "k")
    assert _split_qty("1/2box") == ("1/2", "box")
    assert _split_qty("12~13k") == ("12~13", "k")
    assert _split_qty("--") == ("--", "")


def test_format_date_normalizes_to_yy_mm_dd() -> None:
    assert _format_date("2026-07-08") == "26.07.08"
    assert _format_date("2026.7.8") == "26.07.08"
    assert _format_date("없음") == "없음"


def test_export_documents_xlsx_roundtrip() -> None:
    doc = MealInvoiceDocument(
        원본파일="a.pdf",
        거래일="2026-07-08",
        품목=[
            MealInvoiceRow(품명="쌀", 수량="10 k", 단가="1,000?", 금액="10000", 원산지="국내산"),
        ],
    )
    body = export_documents_xlsx([doc])
    workbook = load_workbook(io.BytesIO(body))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert header[:7] == ["납품일자", "원산지", "품목", "단위", "수량", "단가", "금액"]
    first = [cell.value for cell in sheet[2]]
    assert first[0] == "26.07.08"  # 첫 줄에만 납품일자.
    assert first[2] == "쌀"
    assert first[3] == "k"  # 단위 분리.
    assert first[4] == 10  # 수량 숫자.
    assert first[5] == 1000  # 불확실 단가도 숫자로.
    assert first[6] == 10000


@pytest.mark.parametrize(
    ("quantity", "unit", "expected_quantity", "expected_unit"),
    (
        ("+3kg", "kg", "'+3", "k"),
        ("1/2", "box", "1/2", "box"),
        ("12~13kg", "kg", "12~13", "k"),
        ("--", "", "--", None),
    ),
)
def test_noncanonical_quantity_survives_payload_to_xlsx_roundtrip(
    quantity: str,
    unit: str,
    expected_quantity: object,
    expected_unit: object,
) -> None:
    parsed = service._document_from_payload(
        {"품목": [{"품명": "무", "수량": quantity, "단위": unit}]},
        filename="a.pdf",
        page_no=1,
    )
    export_doc = MealInvoiceExportDocument(
        품목=[MealInvoiceExportRow(품명="무", 수량=parsed.품목[0].수량)]
    )

    body = export_documents_xlsx([export_doc])
    sheet = load_workbook(io.BytesIO(body)).active

    assert sheet["D2"].value == expected_unit
    assert sheet["E2"].value == expected_quantity
    if isinstance(expected_quantity, str):
        assert sheet["E2"].data_type == "s"


# --------------------------------------------------------------------------- #
# 카탈로그 파싱  (리뷰 #6 회귀: 헤더 없는 시트 폴백)
# --------------------------------------------------------------------------- #
def _catalog_xlsx(rows: list[list[object]], *, header: list[object] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if header is not None:
        sheet.append(header)
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_parse_catalog_with_header() -> None:
    content = _catalog_xlsx(
        [
            ["26.07.01", "국내산", "포기김치", "통", 2, 5000, 10000],
            ["26.07.01", "중국산", "대파", "단", 3, 1000, 3000],
        ],
        header=["납품일자", "원산지", "품목", "단위", "수량", "단가", "금액"],
    )
    catalog, units = catalog_mod.parse_catalog_xlsx(content)
    assert "포기김치" in catalog
    assert catalog["포기김치"]["단위"] == "통"
    assert catalog["포기김치"]["단가"] == "5000"
    assert "통" in units


def test_parse_catalog_without_header_uses_fixed_columns() -> None:
    # 리뷰 #6: 헤더 라벨이 없는 시트도 고정 열 순서(원산지·품목·단위·_·단가)로 적재돼야 한다.
    content = _catalog_xlsx(
        [
            ["26.07.01", "국내산", "포기김치", "통", 2, 5000, 10000],
            ["26.07.01", "중국산", "대파", "단", 3, 1000, 3000],
        ],
        header=None,
    )
    catalog, _units = catalog_mod.parse_catalog_xlsx(content)
    assert "포기김치" in catalog
    assert "대파" in catalog
    assert catalog["대파"]["원산지"] == "중국산"


def test_parse_catalog_reads_later_sheets_past_a_padded_first_sheet(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    """앞 시트의 빈 꼬리 행이 전역 행 예산을 먹어 뒤 주차 시트를 통째로 놓치면 안 된다.

    실제 영양사 대장은 주차별 시트 24장인데 서식 때문에 시트마다 max_row 가 100만 행이 넘는다.
    예전에는 남은 예산만큼 잘라 list() 로 물질화해서, 첫 시트의 빈 행이 예산을 다 쓰고 나머지
    23주차가 적재되지 않았다(771종 중 66종만 들어옴).
    """
    header = ["납품일자", "원산지", "품목", "단위", "수량", "단가", "금액"]
    workbook = Workbook()
    first = workbook.active
    first.title = "6월 5주차"
    first.append(header)
    first.append(["26.06.29", "국내산", "포기김치", "10k", 4, 52000, 208000])
    for _ in range(100):  # 서식만 남은 빈 꼬리.
        first.append([None] * len(header))
    second = workbook.create_sheet("6월 4주차")
    second.append(header)
    second.append(["26.06.22", "국내산", "깻잎", "박스", 1, 23000, 23000])
    stream = io.BytesIO()
    workbook.save(stream)

    # 빈 꼬리(100행)보다 작은 예산으로 좁혀, 예산이 빈 행에 소모되는지를 결정적으로 검증한다.
    monkeypatch.setattr(catalog_mod, "_MAX_TOTAL_ROWS", 60)
    catalog, _units = catalog_mod.parse_catalog_xlsx(stream.getvalue())

    assert "포기김치" in catalog
    assert "깻잎" in catalog


def test_parse_catalog_finds_header_below_long_decorative_preamble() -> None:
    """제목·결재란처럼 빈 행이 길게 이어져도 헤더를 찾아야 한다.

    헤더 탐색을 고정 행 수로 자르면 장식이 그보다 깊은 시트는 조용히 고정 열 순서로 오파싱된다
    (품목·단위·단가·원산지가 전부 엉뚱한 열에서 읽힌다). 빈 행은 탐색 예산으로 세지 않는다.
    """
    header = ["납품일자", "원산지", "품목", "단위", "수량", "단가", "금액"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "7월 1주차"
    for _ in range(catalog_mod._MAX_HEADER_SCAN_ROWS + 50):  # 병합 셀 장식(빈 행).
        sheet.append([None] * len(header))
    sheet.append(header)
    sheet.append(["26.07.01", "국내산", "포기김치", "통", 2, 5000, 10000])
    stream = io.BytesIO()
    workbook.save(stream)

    catalog, _units = catalog_mod.parse_catalog_xlsx(stream.getvalue())

    assert catalog["포기김치"]["단위"] == "통"
    assert catalog["포기김치"]["원산지"] == "국내산"


def test_build_item_hint_prefers_frequent_items_and_reports_dropped() -> None:
    """후보 목록은 등장 횟수 내림차순 상위만 넣고, 잘려나간 수를 함께 돌려준다.

    목록에서 빠진 품목은 목록 안의 다른 품목으로 끌려가므로(실측), 몇 종이 빠졌는지 호출부가 알고
    사용자에게 알릴 수 있어야 한다.
    """
    catalog = {
        "당근": {"count": 100},
        "통마늘": {"count": 99},
        "생삼겹": {"count": 1},
    }

    hint, dropped = catalog_mod.build_item_hint(catalog, limit=2)

    assert "당근" in hint and "통마늘" in hint
    assert "생삼겹" not in hint
    assert dropped == 1
    # 강제가 아니라 우선순위다. 대장에 없는 신규 품목은 읽은 대로 적게 둔다.
    assert "읽은 대로" in hint


def test_build_item_hint_is_empty_without_catalog() -> None:
    assert catalog_mod.build_item_hint({}) == ("", 0)


def test_parse_catalog_rejects_expanding_zip_before_openpyxl(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", "0" * 200_000)
    content = stream.getvalue()
    called = False

    def fail_if_called(*_args: object, **_kwargs: object) -> None:
        nonlocal called
        called = True
        raise AssertionError("load_workbook must not run before ZIP safety checks")

    monkeypatch.setattr(catalog_mod, "_MAX_XLSX_ARCHIVE_EXPANSION_RATIO", 1)
    monkeypatch.setattr(catalog_mod, "load_workbook", fail_if_called)

    with pytest.raises(catalog_mod.CatalogArchiveTooLarge):
        catalog_mod.parse_catalog_xlsx(content)

    assert called is False


# --------------------------------------------------------------------------- #
# 사전 유사 매칭
# --------------------------------------------------------------------------- #
def test_vocab_best_match_finds_closest() -> None:
    best, score = vocab.best_match("진간장", ["진간장", "국간장", "대파"])
    assert best == "진간장"
    assert score == 1.0
    empty, zero = vocab.best_match("", ["진간장"])
    assert empty == ""
    assert zero == 0.0


def test_confusable_alt_surfaces_same_first_syllable_lookalike() -> None:
    # 읽은 품명이 사전에 그대로 있어도(양상추=정상 품명), 첫 음절이 같은 다른 품목(양파)이 사전에 있으면
    # 대안 후보로 제안한다(양파↔양상추 오독 원클릭 교정용).
    alt, score = vocab.confusable_alt("양상추", ["양파", "양상추", "당근"])
    assert alt == "양파"
    assert score >= 0.5


def test_confusable_alt_skips_when_read_name_not_in_vocab() -> None:
    # 읽은 품명이 사전에 없으면 일반 best_match 제안이 담당하므로 대안 칩은 내지 않는다.
    alt, score = vocab.confusable_alt("양상추", ["양파", "당근"])
    assert alt == ""
    assert score == 0.0


def test_confusable_alt_skips_when_no_shared_first_syllable() -> None:
    # 첫 글자가 다른 품목은 헷갈리는 대안이 아니다.
    alt, _ = vocab.confusable_alt("당근", ["당근", "양파", "감자"])
    assert alt == ""


def test_annotate_vocab_suggests_confusable_alt_for_valid_read() -> None:
    # 읽은 품명이 사전에 있는 정상 품명이라 기본 사전후보는 안 뜨지만, 첫 음절이 같은 대안(양파)을
    # 사전후보 칩으로 제안해 검수자가 한 번에 뒤집게 한다(값은 안 바꿈).
    doc = MealInvoiceDocument(품목=[MealInvoiceRow(품명="양상추")])
    service.annotate_vocab(doc, ["양파", "양상추", "당근"])
    assert doc.품목[0].품명 == "양상추"
    assert doc.품목[0].사전후보 == "양파"


def test_annotate_vocab_no_confusable_chip_when_unique() -> None:
    # 첫 음절이 같은 헷갈리는 이웃이 없으면 대안 칩을 띄우지 않는다.
    doc = MealInvoiceDocument(품목=[MealInvoiceRow(품명="당근")])
    service.annotate_vocab(doc, ["당근", "양파", "감자"])
    assert doc.품목[0].사전후보 == ""


def test_extract_documents_keeps_confusable_duplicates_as_suggestions_only(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    _patch_extract_pipeline(monkeypatch, {"a.jpg": 1}, fail_on=set())
    monkeypatch.setattr(service, "_rescan_band_into_payload", lambda *a, **k: False)
    monkeypatch.setattr(service, "_render_band_png", lambda *a, **k: b"ROW")
    calls = 0

    def vision(*_args: object, **_kwargs: object) -> str:
        nonlocal calls
        calls += 1
        if calls == 1:
            return (
                '{"거래처":"상선유통","품목":['
                '{"품명":"양파","box":[100,140]},'
                '{"품명":"양파","box":[300,340]}]}'
            )
        return '{"품명":"양상추"}'

    monkeypatch.setattr(service, "_call_vision", vision)

    response = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")],
        workspace_id="w",
        vocab=["양파", "양상추", "당근"],
    )

    assert calls == 1
    assert [row.품명 for row in response.documents[0].품목] == ["양파", "양파"]
    assert [row.사전후보 for row in response.documents[0].품목] == ["양상추", "양상추"]


# --------------------------------------------------------------------------- #
# 관리자 권한(공유 학습/카탈로그 데이터 변경 보호)
# --------------------------------------------------------------------------- #
def _route_admin_map() -> dict[tuple[str, str], bool]:
    from open_alm_api.domains.meal_invoice_ocr.router import (
        require_meal_invoice_ocr_admin,
        router,
    )

    result: dict[tuple[str, str], bool] = {}
    for route in router.routes:
        deps = [getattr(d, "dependency", None) for d in getattr(route, "dependencies", [])]
        is_admin = require_meal_invoice_ocr_admin in deps
        for method in sorted(getattr(route, "methods", []) or []):
            result[(method, route.path)] = is_admin
    return result


def test_shared_data_mutations_require_admin() -> None:
    """공유 학습/카탈로그 상태를 바꾸는 작업은 관리자 전용이어야 한다.

    교정 저장(POST /corrections)도 포함한다 — 저장된 교정이 few-shot/vocab 로 워크스페이스 전체
    OCR 에 주입되므로, 일반 멤버가 공유 인식 품질을 오염시키지 못하게 한다.
    """
    admin_map = _route_admin_map()
    assert admin_map[("POST", "/meal-invoice-ocr/corrections")] is True
    assert admin_map[("DELETE", "/meal-invoice-ocr/corrections/{correction_id}")] is True
    assert admin_map[("DELETE", "/meal-invoice-ocr/corrections")] is True
    assert admin_map[("POST", "/meal-invoice-ocr/catalog")] is True
    assert admin_map[("DELETE", "/meal-invoice-ocr/catalog")] is True


def test_read_and_extract_routes_are_not_admin_gated() -> None:
    """추출·조회·내보내기는 일반 멤버도 쓸 수 있어야 한다(관리자 전용 아님)."""
    admin_map = _route_admin_map()
    assert admin_map[("POST", "/meal-invoice-ocr/extract")] is False
    assert admin_map[("GET", "/meal-invoice-ocr/corrections")] is False
    assert admin_map[("GET", "/meal-invoice-ocr/catalog")] is False
    assert admin_map[("POST", "/meal-invoice-ocr/export.xlsx")] is False


# --------------------------------------------------------------------------- #
# LLM task 등록(플랫폼 예산/정책/감사 경로)
# --------------------------------------------------------------------------- #
def test_ocr_llm_task_registered_with_token_budget() -> None:
    """추출·재판독 workload가 각각 유일한 토큰 예산을 소유해야 한다."""
    from open_alm_api.domains.ai.registry import get_ai_capability_registry
    from open_alm_api.domains.meal_invoice_ocr.task_kinds import (
        MEAL_INVOICE_OCR_RESCAN_TASK_KIND,
        MEAL_INVOICE_OCR_RESCAN_WORKLOAD_ID,
        MEAL_INVOICE_OCR_TASK_KIND,
        MEAL_INVOICE_OCR_WORKLOAD_ID,
    )

    # 부트스트랩(도메인 import + register_ai_capabilities 실행)을 트리거한다.
    get_ai_capability_registry()
    for workload_id, task_kind in (
        (MEAL_INVOICE_OCR_WORKLOAD_ID, MEAL_INVOICE_OCR_TASK_KIND),
        (MEAL_INVOICE_OCR_RESCAN_WORKLOAD_ID, MEAL_INVOICE_OCR_RESCAN_TASK_KIND),
    ):
        workload = get_ai_capability_registry().resolve_llm_workload(workload_id)
        assert workload.task_kind == task_kind
        assert workload.allowed_routes == ("local",)
        assert workload.local_max_output_tokens == 4096
        assert workload.external_max_output_tokens == 4096


# --------------------------------------------------------------------------- #
# 요청 payload 한도(DoS·저장소 비대화 방지)
# --------------------------------------------------------------------------- #
def test_correction_image_rejects_non_data_url() -> None:
    with pytest.raises(ValidationError):
        MealInvoiceCorrectionItem(크롭이미지="http://evil.example/x.png")


def test_correction_image_rejects_oversized() -> None:
    huge = "data:image/png;base64," + "A" * (_MAX_IMAGE_CHARS + 1)
    with pytest.raises(ValidationError):
        MealInvoiceCorrectionItem(크롭이미지=huge)


def test_correction_image_accepts_empty_and_valid_data_url() -> None:
    assert MealInvoiceCorrectionItem(크롭이미지="").크롭이미지 == ""
    ok = MealInvoiceCorrectionItem(크롭이미지="data:image/png;base64,AAAA")
    assert ok.크롭이미지.startswith("data:image/")


def test_correction_save_request_rejects_too_many_items() -> None:
    item = MealInvoiceCorrectionItem(교정={"품명": "x"})
    with pytest.raises(ValidationError):
        MealInvoiceCorrectionSaveRequest(items=[item] * (_MAX_CORRECTION_ITEMS + 1))


def test_export_request_rejects_too_many_documents() -> None:
    with pytest.raises(ValidationError):
        MealInvoiceExportRequest(documents=[MealInvoiceExportDocument()] * (_MAX_EXPORT_DOCS + 1))


def test_export_dto_rejects_image_and_extra_fields() -> None:
    """내보내기 전용 DTO 는 이미지·미지정 필드를 갖지 않고, 보내오면 '무시'가 아니라 '거절'한다
    (extra='forbid'). 거대 문자열을 extra 필드로 실어 메모리만 소모하는 경로를 차단한다."""
    assert "크롭이미지" not in MealInvoiceExportRow.model_fields
    assert "페이지이미지" not in MealInvoiceExportDocument.model_fields
    with pytest.raises(ValidationError):
        MealInvoiceExportRow.model_validate(
            {"품명": "쌀", "크롭이미지": "data:image/png;base64,AAAA"}
        )
    with pytest.raises(ValidationError):
        MealInvoiceExportDocument.model_validate({"페이지이미지": "data:image/png;base64,AAAA"})
    with pytest.raises(ValidationError):
        MealInvoiceExportRequest.model_validate({"documents": [], "junk": "x" * 1000})


@pytest.fixture
def ws_db(monkeypatch: "pytest.MonkeyPatch") -> str:
    """격리된 SQLite 에 meal_invoice_ocr 상태 테이블을 만들고 스토리지 세션을 그쪽으로 돌린다.

    corrections 이미지(MinIO)는 인메모리로 대체. 공유 인프라 없이 DB 스토리지 경로를 검증한다.
    """
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    from open_alm_api.core.db import Base
    from open_alm_api.core.model_registry import import_all_models
    from open_alm_api.domains.meal_invoice_ocr import corrections as corr
    from open_alm_api.domains.meal_invoice_ocr import state as state_mod
    from open_alm_api.domains.meal_invoice_ocr.models import MealInvoiceOcrState

    import_all_models()  # workspaces 등 FK 대상 테이블 정의를 metadata 에 등록(실제 정의).
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(
        engine, tables=[Base.metadata.tables["workspaces"], MealInvoiceOcrState.__table__]
    )
    factory = sessionmaker(bind=engine, autoflush=False)
    monkeypatch.setattr(state_mod, "get_session_factory", lambda: factory)

    class WorkspaceId(str):
        pass

    objects: dict[str, str] = {}
    legacy_images: dict[str, dict] = {}
    monkeypatch.setattr(corr, "_get_json", lambda key, default: legacy_images.get(key, default))
    monkeypatch.setattr(
        corr, "_put_object_text", lambda key, data, *_a, **_k: objects.__setitem__(key, data)
    )
    monkeypatch.setattr(corr, "_get_object_text", lambda key: objects.get(key))
    monkeypatch.setattr(corr, "_remove_object", lambda key: objects.pop(key, None))
    workspace_id = WorkspaceId("ws-meal-test")
    workspace_id.objects = objects  # type: ignore[attr-defined]
    workspace_id.legacy_images = legacy_images  # type: ignore[attr-defined]
    return workspace_id


def test_db_vocab_accumulates_union(ws_db: str) -> None:
    assert vocab.add_names(ws_db, ["가지", "감자"]) == 2
    assert vocab.add_names(ws_db, ["감자", "고추"]) == 3
    assert vocab.load_vocab(ws_db) == ["가지", "감자", "고추"]
    # 다른 워크스페이스와 격리.
    assert vocab.load_vocab("other-ws") == []


def test_db_catalog_roundtrip_and_clear(ws_db: str) -> None:
    catalog_mod.save_catalog(ws_db, {"김치": {"단위": "k", "단가": "1000", "원산지": "국내산"}})
    catalog_mod.save_units(ws_db, ["k", "박스"])
    assert catalog_mod.load_catalog(ws_db)["김치"]["단가"] == "1000"
    assert catalog_mod.load_units(ws_db) == ["k", "박스"]
    catalog_mod.clear_catalog(ws_db)
    assert catalog_mod.load_catalog(ws_db) == {}
    assert catalog_mod.load_units(ws_db) == []


def test_db_catalog_and_units_replace_in_one_mutation(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    calls: list[tuple[dict, list]] = []

    def fake_mutate_state(workspace_id: str, mutate) -> None:
        assert workspace_id == "ws-catalog"
        state = SimpleNamespace(catalog_json={"old": {}}, units_json=["old"])
        mutate(state)
        calls.append((state.catalog_json, state.units_json))

    monkeypatch.setattr(catalog_mod, "mutate_state", fake_mutate_state)

    catalog_mod.save_catalog_with_units("ws-catalog", {"김치": {"단위": "k"}}, ["k"])

    assert calls == [({"김치": {"단위": "k"}}, ["k"])]


def test_db_state_first_insert_retries_after_unique_race(
    ws_db: str, monkeypatch: "pytest.MonkeyPatch"
) -> None:
    from sqlalchemy.exc import IntegrityError
    from sqlalchemy.orm import Session as OrmSession

    from open_alm_api.domains.meal_invoice_ocr import state as state_mod
    from open_alm_api.domains.meal_invoice_ocr.models import MealInvoiceOcrState

    original_flush = OrmSession.flush
    raised = False

    def flaky_flush(self: OrmSession, *args: object, **kwargs: object) -> None:
        nonlocal raised
        if not raised and any(isinstance(obj, MealInvoiceOcrState) for obj in self.new):
            raised = True
            raise IntegrityError("insert", {}, Exception("duplicate workspace state"))
        original_flush(self, *args, **kwargs)

    monkeypatch.setattr(OrmSession, "flush", flaky_flush)

    def apply_catalog(state: MealInvoiceOcrState) -> str:
        state.catalog_json = {"김치": {"단위": "k"}}
        return "saved"

    assert state_mod.mutate_state(ws_db, apply_catalog) == "saved"
    assert raised is True
    assert catalog_mod.load_catalog(ws_db) == {"김치": {"단위": "k"}}


def test_db_corrections_append_delete_clear(ws_db: str) -> None:
    from open_alm_api.domains.meal_invoice_ocr import corrections as corr

    assert (
        corr.append_corrections(ws_db, [{"원본": {"품명": "오독"}, "교정": {"품명": "정답"}}]) == 1
    )
    records = corr.load_corrections(ws_db)
    assert len(records) == 1 and records[0]["교정"]["품명"] == "정답"
    assert corr.delete_correction(ws_db, records[0]["id"]) == 0
    corr.append_corrections(ws_db, [{"원본": {"품명": "x"}, "교정": {"품명": "y"}}])
    corr.clear_corrections(ws_db)
    assert corr.load_corrections(ws_db) == []


def test_db_corrections_share_one_doc_image_across_rows(ws_db: str) -> None:
    """한 문서의 여러 확인 행이 페이지 이미지를 반복 전송하지 않고, 문서키로 하나의 이미지를 공유 참조한다.

    (리뷰 회귀: 20~30행 명세서에서 행마다 전체 페이지 이미지를 실으면 요청 총량 상한을 넘어 실패했다.)
    """
    from open_alm_api.domains.meal_invoice_ocr import corrections as corr

    image = "data:image/png;base64,AAAABBBBCCCC"
    records = [
        {"문서키": "docA", "원본": {"품명": f"오독{i}"}, "교정": {"품명": f"정답{i}"}}
        for i in range(3)
    ]
    assert corr.append_corrections(ws_db, records, {"docA": image}) == 3
    stored = corr.load_corrections(ws_db)
    # 세 레코드 모두 같은 이미지 참조를 갖고, 저장된 이미지는 (중복 제거되어) 하나뿐이다.
    refs = {r.get("image_ref") for r in stored}
    assert len(refs) == 1 and None not in refs
    assert len(corr.load_images(ws_db)) == 1
    # 각 레코드에서 원본 페이지 이미지를 조회할 수 있다.
    assert corr.get_correction_image(ws_db, stored[0]["id"]) == image
    assert corr.get_correction_image(ws_db, stored[-1]["id"]) == image
    objects = getattr(ws_db, "objects")
    assert len(objects) == 1
    assert corr._legacy_img_key(ws_db) not in objects
    assert next(iter(objects)).startswith(f"meal-invoice-ocr/corrections-images/{ws_db}/")


def test_db_correction_delete_keeps_shared_image_until_last_ref(ws_db: str) -> None:
    """같은 페이지 이미지를 공유하는 교정들: 하나를 지워도 다른 참조가 있으면 이미지가 유지되고,
    마지막 참조를 지우면 이미지도 정리된다(삭제는 커밋 성공 후에만 이미지 정리 → 무결성 보존)."""
    from open_alm_api.domains.meal_invoice_ocr import corrections as corr

    image = "data:image/png;base64,SHAREDPAGE"
    records = [
        {"문서키": "docA", "원본": {"품명": f"o{i}"}, "교정": {"품명": f"c{i}"}} for i in range(2)
    ]
    corr.append_corrections(ws_db, records, {"docA": image})
    stored = corr.load_corrections(ws_db)
    assert len(corr.load_images(ws_db)) == 1
    # 한 교정을 지워도 다른 교정이 같은 이미지를 참조하므로 이미지는 유지.
    corr.delete_correction(ws_db, stored[0]["id"])
    assert len(corr.load_images(ws_db)) == 1
    assert corr.get_correction_image(ws_db, stored[1]["id"]) == image
    # 마지막 참조까지 지우면 이미지도 정리된다.
    corr.delete_correction(ws_db, stored[1]["id"])
    assert corr.load_images(ws_db) == {}
    assert getattr(ws_db, "objects") == {}


def test_db_corrections_reject_workspace_image_quota(
    ws_db: str, monkeypatch: "pytest.MonkeyPatch"
) -> None:
    from open_alm_api.domains.meal_invoice_ocr import corrections as corr

    monkeypatch.setattr(corr, "_MAX_WORKSPACE_IMAGE_BYTES", 10)

    with pytest.raises(corr.CorrectionImageQuotaExceeded):
        corr.append_corrections(
            ws_db,
            [{"문서키": "docA", "원본": {"품명": "o"}, "교정": {"품명": "c"}}],
            {"docA": "data:image/png;base64,TOO-LARGE"},
        )

    assert corr.load_corrections(ws_db) == []
    assert getattr(ws_db, "objects") == {}


def test_correction_save_request_doc_image_total_counted() -> None:
    """문서이미지 맵도 요청 전체 이미지 총량 상한에 합산된다(개별 상한은 지켜도 총량 초과면 거절)."""
    from open_alm_api.domains.meal_invoice_ocr.schemas import (
        _MAX_IMAGE_CHARS,
        _MAX_TOTAL_IMAGE_CHARS,
    )

    per = "data:image/png;base64," + "A" * (_MAX_IMAGE_CHARS - 100)
    count = _MAX_TOTAL_IMAGE_CHARS // len(per) + 2
    문서이미지 = {f"doc{i}": per for i in range(count)}
    with pytest.raises(ValidationError):
        MealInvoiceCorrectionSaveRequest(items=[], 문서이미지=문서이미지)


def test_upload_type_validation_accepts_pdf_and_images_rejects_others() -> None:
    assert service.is_supported_invoice_upload(b"%PDF-1.7\n...") is True
    assert service.is_supported_invoice_upload(b"\x89PNG\r\n\x1a\n....") is True
    assert service.is_supported_invoice_upload(b"\xff\xd8\xff\xe0JFIF") is True
    assert service.is_supported_invoice_upload(b"RIFF\x00\x00\x00\x00WEBPVP8 ") is True
    # 실행 파일/스크립트/텍스트 등은 거절.
    assert service.is_supported_invoice_upload(b"MZ\x90\x00 fake exe") is False
    assert service.is_supported_invoice_upload(b"#!/bin/sh\nrm -rf /") is False
    assert service.is_supported_invoice_upload(b"") is False


def test_content_length_body_limit_predicate() -> None:
    """JSON body 엔드포인트는 파싱 전 content-length 로 과대 요청을 거절한다(프록시 1025MB 허용)."""
    from open_alm_api.domains.meal_invoice_ocr.router import (
        _MAX_EXPORT_BODY_BYTES,
        _content_length_exceeds,
        _content_length_missing_or_invalid,
        export_invoices,
        save_corrections,
        _BODY_LIMIT_BY_ENDPOINT,
    )

    limit = _MAX_EXPORT_BODY_BYTES
    assert _content_length_exceeds(str(limit + 1), limit) is True
    assert _content_length_exceeds(str(limit), limit) is False
    # 숫자 비교 함수는 missing/invalid 를 별도 helper 에 맡긴다.
    assert _content_length_exceeds(None, limit) is False
    assert _content_length_exceeds("not-a-number", limit) is False
    assert _content_length_missing_or_invalid(None) is True
    assert _content_length_missing_or_invalid("not-a-number") is True
    assert _content_length_missing_or_invalid(str(limit)) is False
    # export/corrections 엔드포인트에 파싱-전 상한이 실제로 등록돼 있다.
    assert _BODY_LIMIT_BY_ENDPOINT[export_invoices] == _MAX_EXPORT_BODY_BYTES
    assert save_corrections in _BODY_LIMIT_BY_ENDPOINT


def test_existing_route_app_rejects_unsafe_json_content_length() -> None:
    """APIRoute 생성 후 채워진 endpoint map도 요청 시점에 반영돼야 한다."""
    from open_alm_api.domains.meal_invoice_ocr.router import (
        _MAX_EXPORT_BODY_BYTES,
        export_invoices,
        router,
    )

    route = next(
        route for route in router.routes if getattr(route, "endpoint", None) is export_invoices
    )

    for headers in (
        [(b"content-length", str(_MAX_EXPORT_BODY_BYTES + 1).encode())],
        [],
        [(b"content-length", b"not-a-number")],
    ):
        messages: list[dict] = []
        scope = {
            "type": "http",
            "http_version": "1.1",
            "method": "POST",
            "scheme": "http",
            "path": "/meal-invoice-ocr/export.xlsx",
            "raw_path": b"/meal-invoice-ocr/export.xlsx",
            "query_string": b"",
            "headers": headers,
            "client": ("testclient", 50000),
            "server": ("testserver", 80),
        }

        async def receive() -> dict:
            return {"type": "http.request", "body": b"", "more_body": False}

        async def send(message: dict) -> None:
            messages.append(message)

        try:
            asyncio.run(route.app(scope, receive, send))
        except HTTPException as exc:
            assert exc.status_code == 413
        else:
            assert any(
                message.get("type") == "http.response.start" and message.get("status") == 413
                for message in messages
            )


def test_detect_filetype_uses_magic_bytes_not_suffix() -> None:
    """파서 타입은 파일명 suffix 가 아니라 실제 매직바이트로 정한다(확장자 없거나 틀려도 정상 렌더)."""
    assert service._detect_filetype(b"%PDF-1.7\n...") == "pdf"
    assert service._detect_filetype(b"\x89PNG\r\n\x1a\n....") == "png"
    assert service._detect_filetype(b"\xff\xd8\xff\xe0JFIF") == "jpg"
    assert service._detect_filetype(b"BMxxxxxxxx") == "bmp"
    assert service._detect_filetype(b"II*\x00rest-of-tiff") == "tif"
    assert service._detect_filetype(b"MM\x00*rest-of-tiff") == "tif"
    assert service._detect_filetype(b"RIFF\x00\x00\x00\x00WEBPVP8 ") == "webp"
    assert service._detect_filetype(b"not-an-image-or-pdf") == ""


def test_catalog_upload_type_validation() -> None:
    assert catalog_mod.is_xlsx_upload(b"PK\x03\x04rest-of-zip") is True
    assert catalog_mod.is_xlsx_upload(b"%PDF-1.7") is False
    assert catalog_mod.is_xlsx_upload(b"plain text") is False


def test_correction_dict_value_length_rejected() -> None:
    from open_alm_api.domains.meal_invoice_ocr.schemas import _MAX_TEXT_FIELD_CHARS

    with pytest.raises(ValidationError):
        MealInvoiceCorrectionItem(교정={"품명": "가" * (_MAX_TEXT_FIELD_CHARS + 1)})


def test_correction_dict_too_many_fields_rejected() -> None:
    from open_alm_api.domains.meal_invoice_ocr.schemas import _MAX_CORRECTION_FIELDS

    big = {f"k{i}": "v" for i in range(_MAX_CORRECTION_FIELDS + 1)}
    with pytest.raises(ValidationError):
        MealInvoiceCorrectionItem(원본=big)


def test_correction_save_request_rejects_total_image_overflow() -> None:
    from open_alm_api.domains.meal_invoice_ocr.schemas import (
        _MAX_IMAGE_CHARS,
        _MAX_TOTAL_IMAGE_CHARS,
    )

    # 개별 상한(_MAX_IMAGE_CHARS)은 지키지만, 여러 item 누적이 전체 상한을 넘는 경우.
    per = "data:image/png;base64," + "A" * (_MAX_IMAGE_CHARS - 100)
    count = _MAX_TOTAL_IMAGE_CHARS // len(per) + 2
    items = [MealInvoiceCorrectionItem(크롭이미지=per) for _ in range(count)]
    with pytest.raises(ValidationError):
        MealInvoiceCorrectionSaveRequest(items=items)


def test_build_fewshot_truncates_to_limit() -> None:
    from open_alm_api.domains.meal_invoice_ocr import corrections as corrections_mod

    records = [{"원본": {"품명": f"오독{i}"}, "교정": {"품명": f"정답{i}"}} for i in range(100)]
    text = corrections_mod.build_fewshot(records, limit=10)
    # 각 예시 쌍은 '오독'→'정답' 형태('→' 가 따옴표 사이). 헤더의 '오독 → 정답' 은 제외하고 센다.
    assert text.count("'→'") == 10


def test_build_fewshot_omits_real_item_names_as_misreads() -> None:
    """실재 품목을 '잘못 읽은 글자'로 제시하면 모델이 그 이름을 기피하거나 다른 행에 옮겨 쓴다."""
    from open_alm_api.domains.meal_invoice_ocr import corrections as corrections_mod

    records = [
        {"원본": {"품명": "깔강자"}, "교정": {"품명": "감자"}},
        {"원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}},
    ]

    text = corrections_mod.build_fewshot(records, known_names={"깻잎", "감자"})

    assert "'깔강자'→'감자'" in text
    assert "깻잎" not in text


def test_build_fewshot_omits_names_confirmed_as_correct_elsewhere() -> None:
    """다른 교정에서 정답으로 확정된 이름은 오독 예시로 쓰지 않는다(사전이 비어 있어도)."""
    from open_alm_api.domains.meal_invoice_ocr import corrections as corrections_mod

    records = [
        {"원본": {"품명": "쭉파"}, "교정": {"품명": "쑥갓"}},
        {"원본": {"품명": "쑥갓"}, "교정": {"품명": "파"}},
    ]

    text = corrections_mod.build_fewshot(records)

    assert "'쭉파'→'쑥갓'" in text
    assert "'쑥갓'→'파'" not in text


def test_build_fewshot_omits_server_side_refinement_pairs() -> None:
    """정제만으로 같아지는 쌍은 오독이 아니다.

    ``원본.품명`` 은 정제 전 OCR 원문이므로, 사용자가 품명을 손대지 않은 교정도 '원문 → 정제 결과'
    쌍으로 남는다. 이를 예시로 주면 모델에게 브랜드·규격을 지워서 읽으라고 가르치게 되고(원문 그대로
    읽기 설계와 충돌) 예시 예산도 잠식한다. 실재 품목 필터도 원문 문자열이라 그냥은 걸리지 않는다.
    진짜 오독이 섞인 쌍도 왼쪽을 정제 품명으로 맞춰 규격 제거 지시가 함께 실리지 않게 한다.
    """
    from open_alm_api.domains.meal_invoice_ocr import corrections as corrections_mod

    records = [
        {"원본": {"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA"}, "교정": {"품명": "쇠고기맛냉면육수"}},
        {"원본": {"품명": "근추가루(1kg)"}, "교정": {"품명": "고춧가루"}},
    ]

    text = corrections_mod.build_fewshot(records)

    assert "냉면육수" not in text
    assert "'근추가루'→'고춧가루'" in text
    assert "1kg" not in text


def test_export_request_rejects_too_many_rows_per_doc() -> None:
    from open_alm_api.domains.meal_invoice_ocr.schemas import _MAX_ROWS_PER_DOC

    # 문서당 행 수 상한은 품목 필드(max_length)가 문서 생성 시점에 강제한다.
    with pytest.raises(ValidationError):
        MealInvoiceExportDocument(
            품목=[MealInvoiceExportRow() for _ in range(_MAX_ROWS_PER_DOC + 1)]
        )


def test_export_request_rejects_oversized_text_field() -> None:
    from open_alm_api.domains.meal_invoice_ocr.schemas import _MAX_TEXT_FIELD_CHARS

    # 행 텍스트 필드 길이 상한은 각 필드(max_length)가 행 생성 시점에 강제한다.
    with pytest.raises(ValidationError):
        MealInvoiceExportRow(품명="가" * (_MAX_TEXT_FIELD_CHARS + 1))


def test_clamped_scale_downscales_oversized_page() -> None:
    """거대 MediaBox 는 렌더 픽셀 상한 안으로 다운스케일되고, 정상 페이지는 목표 DPI 를 유지한다."""
    import types

    from open_alm_api.domains.meal_invoice_ocr.service import (
        MAX_RENDER_PIXELS,
        RENDER_DPI,
        _clamped_scale,
    )

    big = types.SimpleNamespace(rect=types.SimpleNamespace(width=50000.0, height=50000.0))
    scale = _clamped_scale(big, RENDER_DPI)
    pixels = (big.rect.width * scale) * (big.rect.height * scale)
    assert pixels <= MAX_RENDER_PIXELS + 1

    normal = types.SimpleNamespace(rect=types.SimpleNamespace(width=595.0, height=842.0))  # A4 pt
    assert _clamped_scale(normal, RENDER_DPI) == RENDER_DPI / 72.0


def test_export_escapes_formula_injection() -> None:
    """사용자 문자열이 '=', '+', '-', '@' 로 시작하면 엑셀 수식으로 실행되지 않게 무해화한다."""
    doc = MealInvoiceDocument(
        원본파일="x.pdf",
        거래처="테스트",
        거래일="2026-06-01",
        품목=[
            MealInvoiceRow(품명='=HYPERLINK("http://evil")', 원산지="국내산", 비고="+cmd"),
        ],
    )
    wb = load_workbook(io.BytesIO(export_documents_xlsx([doc])))
    ws = wb.active
    # 헤더(1행) 다음 첫 데이터 행. 품목=3열, 비고=11열.
    item_cell = ws.cell(row=2, column=3).value
    note_cell = ws.cell(row=2, column=11).value
    assert str(item_cell).startswith("'=")
    assert str(note_cell).startswith("'+")


def test_apply_vendor_aliases_auto_applies_exact_name_and_suggests_unit() -> None:
    corrections = [
        {
            "거래처": "상선유통",
            "원본": {"품명": "깻잎", "수량": "25 k", "원산지": "국내산"},
            "교정": {"품명": "감자", "수량": "25 박스", "원산지": "중국산"},
        },
    ]
    aliases = service.build_vendor_aliases(corrections)
    doc = MealInvoiceDocument(
        거래처="상선유통",
        품목=[MealInvoiceRow(품명="깻잎", 수량="30 k", 원산지="국내산")],
    )

    service._apply_vendor_aliases(doc, aliases)

    # 원본 품명 정확 일치 → 품명·원산지·단위 모두 자동 적용(단위는 수량 문자열에 합쳐짐).
    assert doc.품목[0].품명 == "감자"
    assert doc.품목[0].사전후보 == ""
    assert doc.품목[0].수량 == "30 박스"
    assert doc.품목[0].원산지 == "중국산"
    assert doc.품목[0].단위후보 == ""


def test_vendor_alias_exact_name_match_auto_applies_learned_name() -> None:
    # 사용자가 '정확 일치 시 자동 적용'을 택했으므로, 거래처+원본 품명이 정확히 일치하면
    # 과거 확정 교정 품명을 바로 적용한다(유사 매칭은 아래 fallback 테스트대로 후보 유지).
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상선유통",
                "원본": {"품명": "양파", "수량": "25 k", "원산지": "국내산"},
                "교정": {"품명": "양상추", "수량": "25 박스", "원산지": "중국산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="상선유통",
        품목=[MealInvoiceRow(품명="양파", 수량="30 k", 원산지="국내산")],
    )

    service._apply_vendor_aliases(doc, aliases)

    row = doc.품목[0]
    assert row.품명 == "양상추"  # 정확 일치 → 자동 적용
    assert row.사전후보 == ""
    assert row.수량 == "30 박스"  # 단위도 자동 적용
    assert row.원산지 == "중국산"  # 원산지도 자동 적용
    assert row.단위후보 == ""


def test_vendor_alias_exact_match_does_not_override_known_real_item() -> None:
    # 회귀(실측): '냉동새우살'→'무청시래기' 학습이 쌓인 뒤, OCR 이 '냉동새우살'을 정확히 읽었는데도
    # 정확 일치 자동적용이 품명을 '무청시래기'로 덮고 그 행 원산지까지 바꿨다. 읽은 이름이 실재 품목
    # (사전·기준 카탈로그에 존재)이면 오독으로 단정하지 않고 후보 칩으로만 제시해야 한다.
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "Open ALM",
                "원본": {"품명": "냉동새우살", "수량": "15 ea", "원산지": "국내산"},
                "교정": {"품명": "무청시래기", "수량": "15 pk", "원산지": "중국산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="Open ALM",
        품목=[MealInvoiceRow(품명="냉동새우살", 수량="15 ea", 원산지="국내산")],
    )

    service._apply_vendor_aliases(doc, aliases, vocab=["냉동새우살"], catalog={})

    row = doc.품목[0]
    assert row.품명 == "냉동새우살"  # 정상 판독을 덮지 않는다
    assert row.원산지 == "국내산"  # 다른 품목의 학습 원산지가 이식되지 않는다
    assert row.수량 == "15 ea"  # 단위도 그대로
    # 대신 검수자가 원클릭으로 뒤집을 수 있게 후보로 제시한다.
    assert row.사전후보 == "무청시래기"


def test_vendor_alias_known_item_still_applies_field_only_learning() -> None:
    # 실재 품목 가드는 '이름을 다른 품목으로 바꾸라'는 학습만 막아야 한다. 이름을 바꾸지 않는 학습
    # (같은 품목의 원산지·단위 교정)까지 막으면 '반복 거래처는 첫 1회 교정 후 자동 채움' 계약이
    # 깨진다 — known_names 는 기준 카탈로그 키를 전부 포함해 사실상 모든 행이 해당된다.
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "부추", "수량": "3 ea", "원산지": "국내산"},
                "교정": {"품명": "부추", "수량": "3 k", "원산지": "중국산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="성진유통", 품목=[MealInvoiceRow(품명="부추", 수량="3 ea", 원산지="국내산")]
    )

    service._apply_vendor_aliases(doc, aliases, vocab=["부추"], catalog={})

    row = doc.품목[0]
    assert row.수량 == "3 k"
    assert row.원산지 == "중국산"


def test_vendor_alias_blocked_learning_overrides_weaker_suggestion() -> None:
    # 사람이 확정한 교정은 사전·카탈로그 유사 추정보다 신뢰도가 높다. annotate_vocab/annotate_catalog 가
    # 먼저 후보를 채웠더라도 덮어써야 한다 — 아니면 학습이 적용도 노출도 되지 않아 볼 방법이 없다.
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "Open ALM",
                "원본": {"품명": "냉동새우살", "원산지": "국내산"},
                "교정": {"품명": "무청시래기", "원산지": "중국산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="Open ALM",
        품목=[MealInvoiceRow(품명="냉동새우살", 사전후보="약한추정", 사전점수=0.8)],
    )

    service._apply_vendor_aliases(doc, aliases, vocab=["냉동새우살"], catalog={})

    row = doc.품목[0]
    assert row.품명 == "냉동새우살"
    assert row.원산지 == ""  # 다른 품목의 학습 원산지가 이식되지 않는다
    assert row.사전후보 == "무청시래기"
    assert row.사전점수 == 1.0


def test_vendor_alias_exact_match_still_applies_for_unknown_misread() -> None:
    # 읽은 이름이 실재하지 않는 글자(OCR 오독)면 기존대로 자동 적용한다 — 학습 루프의 본래 목적.
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "Open ALM",
                "원본": {"품명": "실근약산건식품", "수량": "6 ea", "원산지": "국내산"},
                "교정": {"품명": "실곤약", "수량": "6 ea", "원산지": "국내산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="Open ALM",
        품목=[MealInvoiceRow(품명="실근약산건식품", 수량="6 ea", 원산지="국내산")],
    )

    service._apply_vendor_aliases(doc, aliases, vocab=["실곤약"], catalog={})

    assert doc.품목[0].품명 == "실곤약"


def test_vendor_alias_known_item_guard_also_covers_catalog_names() -> None:
    # 실재 품목 판정은 사전(vocab)뿐 아니라 기준 카탈로그(영양사 대장) 키도 함께 본다.
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "부추"},
                "교정": {"품명": "숙주"},
            }
        ]
    )
    doc = MealInvoiceDocument(거래처="성진유통", 품목=[MealInvoiceRow(품명="부추")])

    service._apply_vendor_aliases(doc, aliases, vocab=[], catalog={"부추": {"단위": "k"}})

    assert doc.품목[0].품명 == "부추"
    assert doc.품목[0].사전후보 == "숙주"


def test_rescan_reports_row_count_mismatch(monkeypatch) -> None:
    # pass-1 과 pass-2 가 같은 표를 다르게 분할하면 행 누락·병합 가능성이 있다. 호출측이 경고로
    # 노출할 수 있게 (pass-1, pass-2) 행 수를 보고해야 한다.
    payload = {
        "품목": [
            {"품명": "무", "수량": "1", "box": [100, 130]},
            {"품명": "파", "수량": "2", "box": [140, 170]},
        ]
    }
    rescanned = {"행": [{"품명": "무", "원산지": "국내산", "단위": ""}]}
    monkeypatch.setattr(service, "_render_band_png", lambda png, y0, y1: png)
    monkeypatch.setattr(
        service, "_invoke_vision", lambda png, **kwargs: json.dumps(rescanned, ensure_ascii=False)
    )
    row_counts: list[tuple[int, int]] = []
    service._rescan_band_into_payload(
        b"png",
        payload,
        db=None,
        workspace_id="w",
        actor_user_id=None,
        row_counts=row_counts,
    )
    assert row_counts == [(2, 1)]


def test_build_vendor_aliases_most_recent_wins() -> None:
    """같은 거래처+원본품명이면 뒤(최근) 교정이 앞선 교정을 덮는다."""
    corrections = [
        {"거래처": "상선유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "상추"}},
        {"거래처": "상선유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}},
    ]
    aliases = service.build_vendor_aliases(corrections)
    doc = MealInvoiceDocument(거래처="상선유통", 품목=[MealInvoiceRow(품명="깻잎")])
    service._apply_vendor_aliases(doc, aliases)
    # 정확 일치 → 최근 교정(감자)을 자동 적용.
    assert doc.품목[0].품명 == "감자"
    assert doc.품목[0].사전후보 == ""


def test_build_vendor_aliases_skips_records_without_vendor_or_name() -> None:
    corrections = [
        {"거래처": "", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}},
        {"거래처": "상선유통", "원본": {"품명": ""}, "교정": {"품명": "감자"}},
    ]
    assert service.build_vendor_aliases(corrections) == {}


def test_build_vendor_aliases_skips_amount_only_corrections() -> None:
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상선유통",
                "원본": {"품명": "깻잎", "수량": "25 k", "금액": "25000"},
                "교정": {"품명": "깻잎", "수량": "25 k", "금액": "26000"},
            }
        ]
    )

    assert aliases == {}


def test_latest_unchanged_name_and_unit_retire_an_old_vendor_suggestion() -> None:
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상선유통",
                "원본": {"품명": "깻잎", "수량": "25 k", "금액": "25000"},
                "교정": {"품명": "감자", "수량": "25 박스", "금액": "25000"},
            },
            {
                "거래처": "상선유통",
                "원본": {"품명": "깻잎", "수량": "25 k", "금액": "25000"},
                "교정": {"품명": "깻잎", "수량": "25 k", "금액": "26000"},
            },
        ]
    )

    assert aliases == {}


def test_similar_vendor_name_does_not_inherit_another_vendor_learning() -> None:
    """이름이 비슷한 거래처에도 기존 거래처 학습을 적용하지 않는다.

    '상전유통'은 '성진유통'과 유사도 0.82 로 가깝지만, 손글씨 오독인지 이름이 비슷한 새 거래처인지
    구분할 근거가 없다. 유사도로 치환하면 신규 실제 거래처 문서에 다른 거래처의 확정 교정이 자동
    적용되어 공유 학습이 오염되므로, 상호가 정확히 일치할 때만 적용한다.
    """
    aliases = service.build_vendor_aliases(
        [{"거래처": "성진유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}}]
    )

    for vendor in ("상전유통", "싱진유통", "현방농협상사"):
        doc = MealInvoiceDocument(거래처=vendor, 품목=[MealInvoiceRow(품명="깻잎")])
        service._apply_vendor_aliases(doc, aliases)
        assert doc.품목[0].품명 == "깻잎"

    확정 = MealInvoiceDocument(거래처="성진유통", 품목=[MealInvoiceRow(품명="깻잎")])
    service._apply_vendor_aliases(확정, aliases)
    assert 확정.품목[0].품명 == "감자"


def test_apply_vendor_aliases_scoped_to_vendor() -> None:
    """다른 거래처 문서는 같은 글씨라도 건드리지 않는다."""
    aliases = service.build_vendor_aliases(
        [{"거래처": "상선유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}}]
    )
    doc = MealInvoiceDocument(거래처="한밭농산", 품목=[MealInvoiceRow(품명="깻잎")])
    service._apply_vendor_aliases(doc, aliases)
    assert doc.품목[0].품명 == "깻잎"
    assert doc.품목[0].사전후보 == ""


def test_apply_vendor_aliases_keeps_kg_and_k_vendor_names_distinct() -> None:
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "KG푸드", "원본": {"품명": "양파"}, "교정": {"품명": "감자"}},
            {"거래처": "K푸드", "원본": {"품명": "양파"}, "교정": {"품명": "당근"}},
        ]
    )

    kg_doc = MealInvoiceDocument(거래처="ＫＧ 푸드", 품목=[MealInvoiceRow(품명="양파")])
    k_doc = MealInvoiceDocument(거래처="K푸드", 품목=[MealInvoiceRow(품명="양파")])
    service._apply_vendor_aliases(kg_doc, aliases)
    service._apply_vendor_aliases(k_doc, aliases)

    # 거래처 키가 구분되므로 각자 학습만 정확 일치 자동 적용된다(서로 섞이지 않음).
    assert kg_doc.품목[0].품명 == "감자"
    assert k_doc.품목[0].품명 == "당근"
    assert kg_doc.품목[0].사전후보 == ""
    assert k_doc.품목[0].사전후보 == ""


def test_apply_vendor_aliases_empty_correction_fields_do_not_clobber() -> None:
    """원산지 미수집(구) 교정은 빈 값으로 기존 원산지를 지우지 않는다."""
    aliases = service.build_vendor_aliases(
        [{"거래처": "상선유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}}]
    )
    doc = MealInvoiceDocument(
        거래처="상선유통",
        품목=[MealInvoiceRow(품명="깻잎", 수량="10 k", 원산지="국내산")],
    )
    service._apply_vendor_aliases(doc, aliases)
    row = doc.품목[0]
    assert row.품명 == "감자"  # 정확 일치 → 자동 적용
    assert row.사전후보 == ""
    assert row.원산지 == "국내산"  # 교정에 원산지 없으면 보존
    assert row.수량 == "10 k"  # 교정에 단위 없으면 보존


def test_apply_vendor_aliases_exact_match_overrides_current_unit_and_origin() -> None:
    # 품명이 정확 일치하면 학습한 원산지·단위를 현재 OCR 값과 무관하게 자동 적용한다
    # (반복 거래처+품목은 학습값이 authoritative — OCR 이 마커/단위를 놓쳐도 자동 채움).
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상선유통",
                "원본": {"품명": "깻잎", "수량": "25 k", "원산지": "국내산"},
                "교정": {"품명": "감자", "수량": "25 박스", "원산지": "중국산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="상선유통",
        품목=[MealInvoiceRow(품명="깻잎", 수량="30 통", 원산지="미국산")],
    )

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "감자"
    assert doc.품목[0].사전후보 == ""
    assert doc.품목[0].수량 == "30 박스"  # 학습 단위가 현재 '통'을 덮어씀
    assert doc.품목[0].단위후보 == ""
    assert doc.품목[0].원산지 == "중국산"  # 학습 원산지가 현재 '미국산'을 덮어씀


def test_apply_vendor_alias_suggestions_do_not_chain() -> None:
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "상선유통", "원본": {"품명": "참깨잎"}, "교정": {"품명": "깻잎"}},
            {"거래처": "상선유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}},
        ]
    )
    doc = MealInvoiceDocument(거래처="상선유통", 품목=[MealInvoiceRow(품명="참깨잎")])
    service._apply_vendor_aliases(doc, aliases)
    # 참깨잎→깻잎만 정확 일치로 자동 적용. 깻잎→감자로 연쇄 적용하지 않는다.
    assert doc.품목[0].품명 == "깻잎"
    assert doc.품목[0].사전후보 == ""


def test_extract_documents_adds_vendor_alias_suggestions_last(
    monkeypatch: "pytest.MonkeyPatch",
) -> None:
    """extract_documents는 조립 맨 끝에 거래처별 과거 교정을 적용한다(정확 일치는 자동, 단위는 후보)."""
    import json as _json

    monkeypatch.setattr(service, "_open_document", lambda upload: _FakeDoc(upload.filename, 1))
    monkeypatch.setattr(service, "_render_page_png", lambda page: page.encode())
    monkeypatch.setattr(service, "_render_page_display_data_url", lambda page: "")
    monkeypatch.setattr(service, "_rescan_band_into_payload", lambda *a, **k: False)
    import open_alm_api.core.db as _db

    monkeypatch.setattr(
        _db, "get_session_factory", lambda: (lambda: SimpleNamespace(close=lambda: None))
    )
    monkeypatch.setattr(
        service,
        "_call_vision",
        lambda png, **_k: _json.dumps(
            {"거래처": "상선유통", "품목": [{"품명": "깻잎", "수량": "30k"}]}
        ),
    )
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상선유통",
                "원본": {"품명": "깻잎", "수량": "30 k", "원산지": "국내산"},
                "교정": {"품명": "감자", "수량": "30 박스", "원산지": "중국산"},
            }
        ]
    )
    resp = service.extract_documents(
        [service.MealInvoiceUpload("a.jpg", "image/jpeg", b"x")],
        workspace_id="w",
        aliases=aliases,
    )
    row = resp.documents[0].품목[0]
    assert row.품명 == "감자"  # 원본 품명 정확 일치 → 품명·원산지·단위 자동 적용
    assert row.수량 == "30 박스"
    assert row.원산지 == "중국산"
    assert row.사전후보 == ""
    assert row.단위후보 == ""


def test_apply_vendor_aliases_fuzzy_matches_similar_ocr_name() -> None:
    # '정채'로 학습했는데 다음에 비슷한 '적채'로 읽히면(자모 유사도↑) 그 교정을 자동 적용한다.
    corrections = [
        {
            "거래처": "상전유통",
            "원본": {"품명": "정채", "수량": "7 k"},
            "교정": {"품명": "적채", "수량": "7 통", "원산지": "국내산"},
        }
    ]
    aliases = service.build_vendor_aliases(corrections)
    doc = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="적채", 수량="5 k")])

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "적채"  # 학습 교정 품명 적용
    assert doc.품목[0].수량 == "5 통"  # 유사 매칭 → 단위도 자동 적용
    assert doc.품목[0].단위후보 == ""


def test_apply_vendor_aliases_fuzzy_applies_learned_name_on_ocr_variant() -> None:
    # 사용자 시나리오: '고칫갸루'로 읽어 '고춧가루'로 학습 → 다음에 '고찻가루'로 읽혀도 '고춧가루'로 보정.
    aliases = service.build_vendor_aliases(
        [{"거래처": "상전유통", "원본": {"품명": "고칫갸루"}, "교정": {"품명": "고춧가루"}}]
    )
    doc = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="고찻가루")])
    service._apply_vendor_aliases(doc, aliases)
    assert doc.품목[0].품명 == "고춧가루"


def test_apply_vendor_aliases_fuzzy_matches_name_with_uppercase_ascii() -> None:
    """대문자가 섞인 품명도 초성 게이트를 통과해야 한다.

    별칭 키는 소문자로 정규화되는데 초성 열만 원본 품명에서 뽑으면 'A갑자칩'의 초성('A…')이
    키에서 뽑은 초성('a…')과 달라 모든 대문자 포함 품명이 유사 매칭에서 통째로 탈락했다.
    """
    aliases = service.build_vendor_aliases(
        [{"거래처": "상전유통", "원본": {"품명": "A갑자칩"}, "교정": {"품명": "감자스낵"}}]
    )
    doc = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="A갑쟈칩")])

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "감자스낵"


def test_apply_vendor_aliases_fuzzy_does_not_apply_to_different_item() -> None:
    # '양파' 학습이 '양상추'(자모 유사도 ~0.62 < 0.72)에는 적용되지 않아야 한다(오매칭 방지).
    aliases = service.build_vendor_aliases(
        [{"거래처": "상전유통", "원본": {"품명": "양파"}, "교정": {"품명": "감자"}}]
    )
    doc = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="양상추")])
    service._apply_vendor_aliases(doc, aliases)
    assert doc.품목[0].품명 == "양상추"  # 다른 품목이므로 보정 안 함


def test_apply_vendor_aliases_keeps_name_confirmed_as_correct_elsewhere() -> None:
    """정답으로 확정된 적 있는 이름은 오독 키로 쓰지 않는다(자동적용 대신 후보 칩).

    '얌파'를 '양파'로 교정한 이력이 있으면 '양파'는 이 거래처에 실재하는 품목이다. 그러면 과거의
    '양파'→'감자' 교정은 그때 그 행의 오독이었을 뿐이고, 다음에 읽은 '양파'가 오독이라고 단정할 수
    없다. 자동으로 덮어쓰면 정상 판독을 파괴하므로 검수 후보로만 제시한다.
    """
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "상전유통", "원본": {"품명": "양파"}, "교정": {"품명": "감자"}},
            {"거래처": "상전유통", "원본": {"품명": "얌파"}, "교정": {"품명": "양파"}},
        ]
    )
    doc = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="양파")])

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "양파"
    assert doc.품목[0].사전후보 == "감자"

    # 실제 오독('얌파')은 그대로 자동 교정된다.
    misread = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="얌파")])
    service._apply_vendor_aliases(misread, aliases)
    assert misread.품목[0].품명 == "양파"


def test_apply_vendor_aliases_skips_ambiguous_corrected_name_fallback() -> None:
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상전유통",
                "원본": {"품명": "정채", "수량": "7 k"},
                "교정": {"품명": "적채", "수량": "7 통"},
            },
            {
                "거래처": "상전유통",
                "원본": {"품명": "젓채", "수량": "7 k"},
                "교정": {"품명": "적채", "수량": "7 박스"},
            },
        ]
    )
    doc = MealInvoiceDocument(거래처="상전유통", 품목=[MealInvoiceRow(품명="적채", 수량="5 k")])

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].수량 == "5 k"
    assert doc.품목[0].단위후보 == ""


def test_row_keeps_raw_ocr_values_when_learning_overwrites_them() -> None:
    """자동 보정이 품명·단위·원산지를 덮어써도 OCR 원문은 보존돼야 한다.

    보정된 값을 학습 원본으로 저장하면 '보정값 → 사용자 교정' 쌍이 쌓여, 실재 품목이 오독 키로 굳고
    서로 반대 방향 교정이 짝으로 남는다(깻잎→감자 와 감자→깻잎 공존). 학습 키는 언제나 OCR 원문이다.
    """
    doc = service._document_from_payload(
        {
            "거래처": "성진유통",
            "품목": [{"품명": "근추가루", "수량": "10 k", "원산지": "국내산"}],
        },
        filename="a.pdf",
        page_no=1,
    )
    assert doc.품목[0].원문.품명 == "근추가루"

    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "근추가루", "수량": "10 k", "원산지": "국내산"},
                "교정": {"품명": "고춧가루", "수량": "10 박스", "원산지": "중국산"},
            }
        ]
    )
    service._apply_vendor_aliases(doc, aliases)

    row = doc.품목[0]
    # 화면에는 보정된 값이 보이고,
    assert (row.품명, row.수량, row.원산지) == ("고춧가루", "10 박스", "중국산")
    # 학습 키로 쓸 OCR 원문은 그대로 남는다.
    assert (row.원문.품명, row.원문.수량, row.원문.원산지) == ("근추가루", "10 k", "국내산")


def test_row_source_keeps_unfiltered_ocr_name() -> None:
    """학습 키는 정제 전 OCR 원문이다. 표시 품명만 브랜드·규격·단위를 지운 값이다."""
    doc = service._document_from_payload(
        {
            "거래처": "성진유통",
            "품목": [{"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA", "수량": "3 박스"}],
        },
        filename="a.pdf",
        page_no=1,
    )

    row = doc.품목[0]
    assert row.품명 == "쇠고기맛냉면육수"
    assert row.원문.품명 == "쇠고기맛냉면육수(면사랑/실온/5kg)EA"


def test_apply_vendor_aliases_matches_unfiltered_ocr_name() -> None:
    """OCR 원문으로 학습해 두면 정제 규칙이 못 줄인 품명도 확정 교정으로 출력된다."""
    doc = service._document_from_payload(
        {
            "거래처": "성진유통",
            "품목": [{"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA", "수량": "3 박스"}],
        },
        filename="a.pdf",
        page_no=1,
    )
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA"},
                "교정": {"품명": "냉면육수"},
            }
        ]
    )

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "냉면육수"
    # 학습 키는 그대로 원문을 유지해 다음 저장에서도 같은 키로 쌓인다.
    assert doc.품목[0].원문.품명 == "쇠고기맛냉면육수(면사랑/실온/5kg)EA"


def test_apply_vendor_aliases_keeps_filtered_name_without_learning() -> None:
    """같은 원문에 학습이 없으면 정제된 품명이 그대로 표시 값으로 남는다."""
    doc = service._document_from_payload(
        {
            "거래처": "성진유통",
            "품목": [{"품명": "신라면컵라면(6입) 농심", "수량": "2 박스"}],
        },
        filename="a.pdf",
        page_no=1,
    )
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "오징어짬뽕컵 6입"},
                "교정": {"품명": "오징어짬뽕"},
            }
        ]
    )

    service._apply_vendor_aliases(doc, aliases)

    # 정제는 괄호·단위·앞 영문 브랜드만 지우므로 뒤에 붙은 한글 브랜드('농심')는 남는다. 학습이
    # 없는 원문에서는 이 폴백 값이 그대로 표시되고, 사용자가 한 번 교정하면 원문 키로 학습된다.
    assert doc.품목[0].품명 == "신라면컵라면 농심"


def test_apply_vendor_aliases_raw_source_match_does_not_bypass_known_item_guard() -> None:
    """원문 전체가 학습 키와 일치해도, 읽은 품명이 실재 품목이면 이름을 덮어쓰지 않는다.

    원문 일치는 '이번엔 제대로 읽었다'와 '이번에도 같은 오독이다'를 구분하지 못한다. 학습이 만들어진
    행도 브랜드·규격까지 같은 원문으로 읽혔을 수 있다. 사람이 확정한 교정이므로 버리지는 않고 후보
    칩으로만 제시해 검수자가 원클릭으로 반영하게 한다.
    """
    doc = service._document_from_payload(
        {
            "거래처": "성진유통",
            "품목": [{"품명": "삼다수(무라벨),DC,광동,10L(500ML*20EA)/B", "수량": "4 박스"}],
        },
        filename="a.pdf",
        page_no=1,
    )
    assert doc.품목[0].품명 == "삼다수"
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "삼다수(무라벨),DC,광동,10L(500ML*20EA)/B"},
                "교정": {"품명": "생수"},
            }
        ]
    )

    service._apply_vendor_aliases(doc, aliases, vocab=["삼다수"])

    assert doc.품목[0].품명 == "삼다수"
    assert doc.품목[0].사전후보 == "생수"


def test_apply_vendor_aliases_still_matches_legacy_filtered_corrections() -> None:
    """원문 없이 정제 품명으로만 저장된 기존 교정 데이터도 계속 매칭된다(하위 호환)."""
    doc = service._document_from_payload(
        {"거래처": "성진유통", "품목": [{"품명": "근추가루", "수량": "10 k"}]},
        filename="a.pdf",
        page_no=1,
    )
    aliases = service.build_vendor_aliases(
        [{"거래처": "성진유통", "원본": {"품명": "근추가루"}, "교정": {"품명": "고춧가루"}}]
    )

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "고춧가루"


def test_build_vendor_aliases_ignores_refinement_only_name_difference() -> None:
    """품명을 손대지 않고 원산지만 고치면 이름 교정 엔트리를 만들지 않는다.

    학습 키는 OCR 원문이고 사용자가 화면에서 확인하는 값은 정제 품명이다. 두 값을 직접 비교하면
    품명을 바꾸지 않은 교정이 '이름 변경 학습'으로 분류돼 실재 품목 가드·모순 학습 분류가 오염된다.
    """
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA", "원산지": ""},
                "교정": {"품명": "쇠고기맛냉면육수", "원산지": "국내산"},
            }
        ]
    )

    entry = next(iter(aliases["성진유통"]["raw"].values()))
    assert "name" not in entry
    assert entry["origin"]["fixed"] == "국내산"

    doc = service._document_from_payload(
        {
            "거래처": "성진유통",
            "품목": [{"품명": "쇠고기맛냉면육수(면사랑/실온/5kg)EA", "수량": "3 박스"}],
        },
        filename="a.pdf",
        page_no=1,
    )
    service._apply_vendor_aliases(doc, aliases)

    assert (doc.품목[0].품명, doc.품목[0].원산지) == ("쇠고기맛냉면육수", "국내산")


def test_build_vendor_aliases_reconfirmation_discards_stale_name_learning() -> None:
    """같은 원문을 다시 확인해 품명이 그대로면 오래된 오교정 학습을 폐기한다(자기치유)."""
    raw = "삼다수(무라벨),DC,광동,10L(500ML*20EA)/B"
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "성진유통", "원본": {"품명": raw}, "교정": {"품명": "생수"}},
            # 재확인: 표시 품명('삼다수')을 그대로 두었으므로 바뀐 것이 없다.
            {"거래처": "성진유통", "원본": {"품명": raw}, "교정": {"품명": "삼다수"}},
        ]
    )

    assert aliases == {}


def test_apply_vendor_aliases_fuzzy_matches_misread_variant_of_raw_source() -> None:
    """오독 변형 보정은 원문 키 공간에서 이뤄진다('고칫갸루(1kg)' 학습 → '고찻가루(1kg)' 매칭)."""
    doc = service._document_from_payload(
        {"거래처": "성진유통", "품목": [{"품명": "고찻가루(1kg)", "수량": "2 k"}]},
        filename="a.pdf",
        page_no=1,
    )
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "성진유통",
                "원본": {"품명": "고칫갸루(1kg)"},
                "교정": {"품명": "고춧가루"},
            }
        ]
    )

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "고춧가루"


def test_build_vendor_aliases_blocks_contradicting_raw_source_learning() -> None:
    """양방향 모순 학습은 원문 키로 저장돼도 자동적용에서 빠지고 후보로만 남는다."""
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "성진유통", "원본": {"품명": "감자"}, "교정": {"품명": "깻잎"}},
            {"거래처": "성진유통", "원본": {"품명": "깻잎(국내산)"}, "교정": {"품명": "감자"}},
        ]
    )

    assert aliases["성진유통"]["raw"] == {}
    assert set(aliases["성진유통"]["suggest"]) == {"감자", "깻잎국내산"}

    doc = service._document_from_payload(
        {"거래처": "성진유통", "품목": [{"품명": "깻잎(국내산)", "수량": "1 k"}]},
        filename="a.pdf",
        page_no=1,
    )
    service._apply_vendor_aliases(doc, aliases)

    # 품명은 읽은 그대로 두고, 사람이 확정한 반대 방향 교정은 후보 칩으로만 제시한다.
    assert doc.품목[0].품명 == "깻잎"
    assert doc.품목[0].사전후보 == "감자"


def test_apply_vendor_aliases_fuzzy_skips_different_item_with_other_initials() -> None:
    """'후추'→'고춧가루' 학습이 '부추'(다른 품목)를 덮어쓰지 않는다.

    부추/후추는 자모 유사도 0.75 로 퍼지 임계값(0.72)을 넘지만 서로 다른 품목이다. 짧은 한글 품명은
    유사도만으로 '오독 변형'과 '다른 품목'을 가를 수 없어 초성 열(ㅂㅊ vs ㅎㅊ)을 함께 본다.
    """
    aliases = service.build_vendor_aliases(
        [{"거래처": "상선유통", "원본": {"품명": "후추"}, "교정": {"품명": "고춧가루"}}]
    )
    doc = MealInvoiceDocument(거래처="상선유통", 품목=[MealInvoiceRow(품명="부추", 수량="11 k")])

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "부추"


def test_apply_vendor_aliases_fuzzy_skips_name_present_in_vocab() -> None:
    """읽은 품명이 사전에 있는 실재 품목이면 오독으로 보지 않아 퍼지 자동적용 대상이 아니다."""
    aliases = service.build_vendor_aliases(
        [{"거래처": "상선유통", "원본": {"품명": "정채"}, "교정": {"품명": "적채"}}]
    )
    doc = MealInvoiceDocument(거래처="상선유통", 품목=[MealInvoiceRow(품명="젓채")])

    service._apply_vendor_aliases(doc, aliases, vocab=["젓채", "적채"])

    assert doc.품목[0].품명 == "젓채"


def test_apply_vendor_aliases_drops_reciprocal_corrections() -> None:
    """서로 반대 방향으로 학습된 쌍은 어느 쪽으로 읽어도 뒤집히므로 둘 다 자동적용하지 않는다."""
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "상선유통", "원본": {"품명": "깻잎"}, "교정": {"품명": "감자"}},
            {"거래처": "상선유통", "원본": {"품명": "감자"}, "교정": {"품명": "깻잎"}},
        ]
    )
    for read in ("깻잎", "감자"):
        doc = MealInvoiceDocument(거래처="상선유통", 품목=[MealInvoiceRow(품명=read)])
        service._apply_vendor_aliases(doc, aliases)
        assert doc.품목[0].품명 == read


def test_apply_vendor_aliases_keeps_name_that_is_target_of_another_correction() -> None:
    """'쭉파'→'쑥갓' 학습이 있으면 '쑥갓'은 실재 품목이므로 '쑥갓'→'파' 학습을 자동적용하지 않는다."""
    aliases = service.build_vendor_aliases(
        [
            {"거래처": "상선유통", "원본": {"품명": "쑥갓"}, "교정": {"품명": "파"}},
            {"거래처": "상선유통", "원본": {"품명": "쭉파"}, "교정": {"품명": "쑥갓"}},
        ]
    )
    doc = MealInvoiceDocument(거래처="상선유통", 품목=[MealInvoiceRow(품명="쑥갓")])

    service._apply_vendor_aliases(doc, aliases)

    assert doc.품목[0].품명 == "쑥갓"


def test_apply_vendor_aliases_corrected_name_fallback_never_changes_name() -> None:
    """교정 품명 색인은 이미 정답을 읽은 경우이므로 원산지·단위만 보강하고 품명은 건드리지 않는다."""
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상선유통",
                "원본": {"품명": "정채", "수량": "7 k"},
                "교정": {"품명": "적채", "수량": "7 통", "원산지": "국내산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="상선유통", 품목=[MealInvoiceRow(품명="적채", 수량="5 k", 원산지="")]
    )

    service._apply_vendor_aliases(doc, aliases, vocab=["적채"])

    row = doc.품목[0]
    assert row.품명 == "적채"
    assert row.수량 == "5 통"
    assert row.원산지 == "국내산"


def test_vendor_unit_suggestion_survives_when_ocr_reads_corrected_name() -> None:
    aliases = service.build_vendor_aliases(
        [
            {
                "거래처": "상전유통",
                "원본": {"품명": "정채", "수량": "7 k"},
                "교정": {"품명": "적채", "수량": "7 통", "원산지": "국내산"},
            }
        ]
    )
    doc = MealInvoiceDocument(
        거래처="상전유통",
        품목=[MealInvoiceRow(품명="적채", 수량="5 k", 원산지="국내산")],
    )
    service._apply_vendor_aliases(doc, aliases)
    row = doc.품목[0]
    assert row.품명 == "적채"
    assert row.수량 == "5 통"  # 유사 매칭 → 학습 단위 자동 적용
    assert row.단위후보 == ""
