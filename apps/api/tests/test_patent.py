from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import UploadFile

from open_alm_api.domains.patent import kipris, llm, prompts, router, service


class _FakeResponse:
    def __init__(self, text: str, status_code: int = 200) -> None:
        self.text = text
        self.status_code = status_code
        self.headers = {"content-length": str(len(text.encode("utf-8")))}

    def raise_for_status(self) -> None:
        return None

    def iter_bytes(self):
        yield self.text.encode("utf-8")

    def __enter__(self):
        return self

    def __exit__(self, *_args) -> None:
        return None


def _word_search_xml() -> str:
    return """<?xml version='1.0'?>
    <response><body><items>
      <item>
        <applicationNumber>10-2024-0001234</applicationNumber>
        <registrationNumber>10-2500000</registrationNumber>
        <inventionTitle>전기차 히트펌프 시스템</inventionTitle>
        <applicantName>Open ALM</applicantName>
        <applicationDate>20240115</applicationDate>
        <registerStatus>등록</registerStatus>
        <ipcNumber>F25B30/00</ipcNumber>
        <astrtCont>배터리 폐열을 활용하는 히트펌프</astrtCont>
      </item>
    </items><count><totalCount>42</totalCount></count></body></response>"""


def _bibliography_xml() -> str:
    return """<?xml version='1.0'?>
    <response><body>
      <item><applicationNumber>1020240001234</applicationNumber>
        <inventionTitle>전기차 히트펌프 시스템</inventionTitle>
        <astrtCont>상세 초록: 배터리 폐열을 활용하는 히트펌프 시스템의 통합 제어 구성</astrtCont></item>
      <claimInfo><claim>청구항 1. 히트펌프를 포함하는 시스템.</claim></claimInfo>
      <claimInfo><claim>청구항 2. 제1항에 있어서 배터리 폐열을 회수.</claim></claimInfo>
      <ipcInfo><ipcNumber>F25B30/00</ipcNumber></ipcInfo>
    </body></response>"""


def test_word_search_parses_items_and_total(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _FakeResponse(_word_search_xml()))
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    items, total = client.word_search("히트펌프", page=1, num_rows=15)
    assert total == 42
    assert len(items) == 1
    assert items[0]["app_no"] == "10-2024-0001234"
    assert items[0]["applicant"] == "Open ALM"
    assert items[0]["ipc"] == "F25B30/00"


def test_kipris_http_error_redacts_api_key(monkeypatch: pytest.MonkeyPatch) -> None:
    """Upstream HTTP errors echo the request URL, which carries the API key in
    accessKey/ServiceKey query params — it must never reach the error message."""
    import httpx

    secret = "SUPER-SECRET-KIPRIS-KEY"

    class _ErrResponse:
        def raise_for_status(self) -> None:
            request = httpx.Request(
                "GET",
                f"http://example/kipi/getWordSearch?word=x&accessKey={secret}&ServiceKey={secret}",
            )
            raise httpx.HTTPStatusError(
                f"Server error '500' for url '{request.url}'",
                request=request,
                response=httpx.Response(500, request=request),
            )

    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _ErrResponse())
    client = kipris.KiprisClient(secret, "https://example/kipi")
    with pytest.raises(kipris.KiprisError) as exc_info:
        client.word_search("x")
    message = str(exc_info.value)
    assert secret not in message
    assert "accessKey" not in message and "ServiceKey" not in message
    assert "HTTP 500" in message


def test_kipris_retries_timeouts_three_times_with_bounded_timeouts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    attempts: list[object] = []
    retry_attempts: list[tuple[int, float | None]] = []

    def fake_get(_url, **kwargs):
        attempts.append(kwargs["timeout"])
        if len(attempts) <= 3:
            raise kipris.httpx.ReadTimeout(
                "provider timeout",
                request=kipris.httpx.Request("GET", "https://example/kipi"),
            )
        return _FakeResponse(_word_search_xml())

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    monkeypatch.setattr(kipris.time, "monotonic", lambda: 100.0)
    monkeypatch.setattr(
        kipris,
        "_wait_before_retry",
        lambda attempt, *, deadline_monotonic: retry_attempts.append((attempt, deadline_monotonic)),
    )

    items, total = kipris.KiprisClient("KEY", "https://example/kipi").word_search("히트펌프")

    assert len(items) == 1
    assert total == 42
    assert len(attempts) == 4
    assert retry_attempts == [(0, 115.0), (1, 115.0), (2, 115.0)]
    assert all(timeout.connect == 5.0 for timeout in attempts)
    assert all(timeout.read == 15.0 for timeout in attempts)


@pytest.mark.parametrize(
    ("status_code", "expected_attempts", "error_type"),
    [
        (429, 4, kipris.KiprisRetryableError),
        (502, 4, kipris.KiprisRetryableError),
        (503, 4, kipris.KiprisRetryableError),
        (504, 4, kipris.KiprisRetryableError),
        (500, 1, kipris.KiprisError),
    ],
)
def test_kipris_retries_only_allowlisted_http_statuses(
    monkeypatch: pytest.MonkeyPatch,
    status_code: int,
    expected_attempts: int,
    error_type: type[Exception],
) -> None:
    attempts = 0

    class _StatusResponse:
        def raise_for_status(self) -> None:
            nonlocal attempts
            attempts += 1
            request = kipris.httpx.Request("GET", "https://example/kipi")
            raise kipris.httpx.HTTPStatusError(
                "upstream failure",
                request=request,
                response=kipris.httpx.Response(status_code, request=request),
            )

    monkeypatch.setattr(kipris.httpx, "get", lambda *_args, **_kwargs: _StatusResponse())
    monkeypatch.setattr(kipris, "_wait_before_retry", lambda *_args, **_kwargs: None)

    with pytest.raises(error_type):
        kipris.KiprisClient("KEY", "https://example/kipi").word_search("히트펌프")
    assert attempts == expected_attempts


def test_generic_search_rejects_an_expired_deadline_before_http_access(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    called = False

    def fake_stream(*_args, **_kwargs):
        nonlocal called
        called = True
        return _FakeResponse(_generic_domestic_search_xml())

    monkeypatch.setattr(kipris.time, "monotonic", lambda: 100.0)
    monkeypatch.setattr(kipris.httpx, "stream", fake_stream)

    with pytest.raises(kipris.KiprisTimeoutError, match="제한 시간"):
        kipris.KiprisClient("KEY", "https://example/kipi").search(
            kipris.PatentSearchCriteria(query="storage", countries=("KR",)),
            deadline_monotonic=100.0,
        )
    assert called is False


def test_sync_lookup_calls_can_share_one_operation_deadline(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    now = 100.0
    network_calls = 0

    def fake_get(*_args, **_kwargs):
        nonlocal network_calls
        network_calls += 1
        return _FakeResponse(_word_search_xml())

    monkeypatch.setattr(kipris.time, "monotonic", lambda: now)
    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    deadline = now + 15.0

    client.word_search("히트펌프", deadline_monotonic=deadline)
    now = deadline
    with pytest.raises(kipris.KiprisTimeoutError, match="제한 시간"):
        client.applicant_names("히트펌프", deadline_monotonic=deadline)

    assert network_calls == 1


def test_sync_lookup_rejects_response_that_completes_after_deadline(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    now = 100.0

    def fake_get(*_args, **_kwargs):
        nonlocal now
        now = 116.0
        return _FakeResponse(_word_search_xml())

    monkeypatch.setattr(kipris.time, "monotonic", lambda: now)
    monkeypatch.setattr(kipris.httpx, "get", fake_get)

    with pytest.raises(kipris.KiprisTimeoutError, match="제한 시간"):
        kipris.KiprisClient("KEY", "https://example/kipi").word_search(
            "히트펌프",
            deadline_monotonic=115.0,
        )


def test_generic_search_retries_retryable_stream_status(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    attempts = 0

    class _RetryableResponse:
        headers: dict[str, str] = {}

        def __enter__(self):
            return self

        def __exit__(self, *_args) -> None:
            return None

        def raise_for_status(self) -> None:
            request = kipris.httpx.Request("GET", "https://example/kipi")
            raise kipris.httpx.HTTPStatusError(
                "upstream unavailable",
                request=request,
                response=kipris.httpx.Response(503, request=request),
            )

    def fake_stream(*_args, **_kwargs):
        nonlocal attempts
        attempts += 1
        if attempts <= 3:
            return _RetryableResponse()
        return _FakeResponse(_generic_domestic_search_xml())

    monkeypatch.setattr(kipris.httpx, "stream", fake_stream)
    monkeypatch.setattr(kipris, "_wait_before_retry", lambda *_args, **_kwargs: None)

    page = kipris.KiprisClient("KEY", "https://example/kipi").search(
        kipris.PatentSearchCriteria(query="storage", countries=("KR",))
    )

    assert attempts == 4
    assert page.total_count == 31


@pytest.mark.parametrize(
    "base_url",
    [
        "http://example/kipi",
        "http://127.evil.test/kipi",
    ],
)
def test_kipris_rejects_insecure_non_local_base_url(
    base_url: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    called = False

    def fake_get(*args, **kwargs):
        nonlocal called
        called = True
        return _FakeResponse(_word_search_xml())

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient("KEY", base_url)
    with pytest.raises(kipris.KiprisError, match="HTTPS"):
        client.word_search("히트펌프")
    assert called is False


def test_kipris_allows_local_http_base_url(monkeypatch: pytest.MonkeyPatch) -> None:
    seen_url = ""

    def fake_get(url, params=None, timeout=None, **_kwargs):
        nonlocal seen_url
        seen_url = url
        return _FakeResponse(_word_search_xml())

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient("KEY", "http://127.0.0.1:9001/kipi")
    client.word_search("히트펌프")
    assert seen_url.startswith("http://127.0.0.1:9001/")


def _generic_domestic_search_xml() -> str:
    return """<?xml version='1.0'?>
    <response><body><items>
      <item>
        <applicationNumber>10-2024-0004321</applicationNumber>
        <openNumber>10-2024-0098765</openNumber>
        <inventionTitle>분산 저장 장치</inventionTitle>
        <applicantName>Example Research;Second Applicant</applicantName>
        <applicationDate>20240203</applicationDate>
        <openDate>20240910</openDate>
        <ipcNumber>G06F 3/06|H04L 67/00</ipcNumber>
        <astrtCont>&lt;P&gt;분산 저장을 위한 일반적인 장치.&lt;/P&gt;</astrtCont>
      </item>
    </items><count><totalCount>31</totalCount></count></body></response>"""


def _generic_foreign_search_xml(country: str) -> str:
    return f"""<?xml version='1.0'?>
    <response><body><totalSearchCount>61</totalSearchCount><items>
      <searchResult>
        <ltrtno>000020240012345A1</ltrtno>
        <countryCode>{country}</countryCode>
        <publishrNo>2024012345</publishrNo>
        <applicationNo>2024-0012345</applicationNo>
        <inventionName>Distributed storage device</inventionName>
        <applicant>Example Research|Second Applicant</applicant>
        <applicationDate>20240115</applicationDate>
        <openDate>20240718</openDate>
        <ipc>G06F 3/06|H04L 67/00</ipc>
      </searchResult>
    </items></body></response>"""


def test_generic_search_uses_domestic_word_endpoint_and_typed_result(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    seen: dict[str, object] = {}

    def fake_stream(_method, url, params=None, timeout=None, verify=None):
        seen.update(url=url, params=params, timeout=timeout, verify=verify)
        return _FakeResponse(_generic_domestic_search_xml())

    monkeypatch.setattr(kipris.httpx, "stream", fake_stream)
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    result_page = client.search(
        kipris.PatentSearchCriteria(
            query="distributed storage",
            countries=("kr",),
        ),
        page=1,
        page_size=10,
    )

    assert str(seen["url"]).endswith("/getWordSearch")
    assert seen["verify"] is True
    assert seen["params"] == {
        "word": "distributed storage",
        "accessKey": "KEY",
        "ServiceKey": "KEY",
        "pageNo": 1,
        "numOfRows": 10,
    }
    assert result_page.total_count == 31
    assert result_page.has_next is True
    assert result_page.jurisdictions == (
        kipris.PatentSearchJurisdictionPage(
            jurisdiction="KR",
            total_count=31,
            returned_count=1,
            has_next=True,
        ),
    )
    result = result_page.results[0]
    assert isinstance(result, kipris.PatentSearchResult)
    assert result.canonical_number == "KR1020240098765"
    assert result.publication_number == "10-2024-0098765"
    assert result.application_number == "10-2024-0004321"
    assert result.applicants == ("Example Research", "Second Applicant")
    assert result.jurisdiction == "KR"
    assert result.ipc_codes == ("G06F 3/06", "H04L 67/00")
    assert result.filing_date == date(2024, 2, 3)
    assert result.publication_date == date(2024, 9, 10)
    assert result.provider_document_id == "10-2024-0004321"
    assert not hasattr(result, "raw")


def test_generic_domestic_search_bounds_malicious_provider_fields_and_page_size(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    applicants = "|".join(f"Applicant-{index}-" + ("A" * 400) for index in range(45))
    classifications = "|".join(f"G06F-{index}-" + ("C" * 80) for index in range(70))
    oversized_item = f"""
      <item>
        <applicationNumber>{"1" * 500}</applicationNumber>
        <openNumber>{"2" * 500}</openNumber>
        <inventionTitle>{"T" * 2_000}</inventionTitle>
        <applicantName>{applicants}</applicantName>
        <applicationDate>{"20240101" * 20}</applicationDate>
        <openDate>{"20240202" * 20}</openDate>
        <ipcNumber>{classifications}</ipcNumber>
        <astrtCont>{"A" * 100_000}</astrtCont>
      </item>
    """
    xml = (
        "<response><body><items>"
        + oversized_item
        + oversized_item.replace("T" * 2_000, "SECOND")
        + "</items><count><totalCount>99999999999999999999</totalCount>"
        "</count></body></response>"
    )
    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _FakeResponse(xml),
    )

    page = kipris.KiprisClient("KEY", "https://example/kipi").search(
        kipris.PatentSearchCriteria(query="storage", countries=("KR",)),
        page_size=1,
    )

    assert len(page.results) == 1
    result = page.results[0]
    assert len(result.application_number) == 96
    assert len(result.publication_number) == 96
    assert len(result.canonical_number) <= 128
    assert len(result.provider_document_id) <= 128
    assert len(result.external_url) <= 2_048
    assert len(result.title) == 1_000
    assert len(result.abstract) == 32_000
    assert 0 < len(result.applicants) <= 30
    assert all(len(value) <= 320 for value in result.applicants)
    assert 0 < len(result.ipc_codes) <= 50
    assert all(len(value) <= 64 for value in result.ipc_codes)
    assert result.filing_date is None
    assert result.publication_date is None
    assert page.total_count == 2_147_483_647


def test_generic_search_uses_domestic_advanced_endpoint_with_only_caller_filters(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple[str, dict[str, object]]] = []

    def fake_stream(_method, url, params=None, timeout=None, **_kwargs):
        calls.append((url, dict(params or {})))
        return _FakeResponse(_generic_domestic_search_xml())

    monkeypatch.setattr(kipris.httpx, "stream", fake_stream)
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    client.search(
        kipris.PatentSearchCriteria(
            query="distributed storage",
            countries=("KR",),
            classification_codes=("g06f", "h04l"),
            applicants=("Example Research", "Second Applicant"),
        ),
        page=2,
        page_size=10,
    )

    assert len(calls) == 1
    url, params = calls[0]
    assert url.endswith("/getAdvancedSearch")
    assert params == {
        "astrtCont": "distributed storage",
        "ipcNumber": "G06F|H04L",
        "applicantName": "Example Research|Second Applicant",
        "accessKey": "KEY",
        "ServiceKey": "KEY",
        "pageNo": 2,
        "numOfRows": 10,
    }


def test_generic_search_paginates_each_caller_selected_foreign_country(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple[str, dict[str, object], object]] = []

    def fake_stream(_method, url, params=None, timeout=None, verify=None):
        request_params = dict(params or {})
        calls.append((url, request_params, verify))
        return _FakeResponse(_generic_foreign_search_xml(str(request_params["collectionValues"])))

    monkeypatch.setattr(kipris.httpx, "stream", fake_stream)
    client = kipris.KiprisClient(
        "KEY",
        "https://plus.kipris.or.kr/kipo-api/kipi/patUtiModInfoSearchSevice",
    )
    result_page = client.search(
        kipris.PatentSearchCriteria(
            query="distributed storage",
            countries=("US", "EP"),
            classification_codes=("G06F", "H04L"),
            applicants=("Example Research", "Second Applicant"),
        ),
        page=2,
        page_size=30,
    )

    assert len(calls) == 2
    assert [call[1]["collectionValues"] for call in calls] == ["US", "EP"]
    for url, params, verify in calls:
        assert url == (
            "https://plus.kipris.or.kr/openapi/rest/"
            "ForeignPatentAdvencedSearchService/advancedSearch"
        )
        assert verify is True
        assert params == {
            "free": "distributed storage",
            "collectionValues": params["collectionValues"],
            "currentPage": 31,
            "sortField": "AD",
            "sortState": "true",
            "accessKey": "KEY",
            "ipc": "G06F+H04L",
            "applicant": "Example Research+Second Applicant",
        }

    assert result_page.total_count == 122
    assert result_page.has_next is True
    assert tuple(result.jurisdiction for result in result_page.results) == ("US", "EP")
    assert result_page.results[0].canonical_number == "US2024012345A1"
    assert result_page.results[0].provider_document_id == "000020240012345A1"
    assert result_page.results[0].external_url == (
        "https://patents.google.com/patent/US2024012345A1"
    )


def test_generic_foreign_search_bounds_malicious_provider_projection(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    applicants = ";".join(f"Applicant-{index}-" + ("A" * 400) for index in range(45))
    classifications = "\n".join(f"H04L-{index}-" + ("C" * 80) for index in range(70))
    xml = f"""<response><body><totalSearchCount>7</totalSearchCount><items>
      <searchResult>
        <ltrtno>{"9" * 300}A1</ltrtno>
        <countryCode>{"US" * 20}</countryCode>
        <publishrNo>{"8" * 300}</publishrNo>
        <applicationNo>{"7" * 300}</applicationNo>
        <inventionName>{"T" * 2_000}</inventionName>
        <applicant>{applicants}</applicant>
        <applicationDate>{"20240101" * 20}</applicationDate>
        <openDate>{"20240202" * 20}</openDate>
        <ipc>{classifications}</ipc>
        <abstract>{"A" * 100_000}</abstract>
      </searchResult>
    </items></body></response>"""
    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _FakeResponse(xml),
    )

    page = kipris.KiprisClient("KEY", "https://example/kipi").search(
        kipris.PatentSearchCriteria(query="storage", countries=("US",)),
    )

    result = page.results[0]
    assert result.jurisdiction == "US"
    assert len(result.application_number) == 96
    assert len(result.publication_number) == 96
    assert len(result.canonical_number) <= 128
    assert len(result.provider_document_id) <= 128
    assert len(result.external_url) <= 2_048
    assert len(result.title) == 1_000
    assert len(result.abstract) == 32_000
    assert 0 < len(result.applicants) <= 30
    assert all(len(value) <= 320 for value in result.applicants)
    assert 0 < len(result.ipc_codes) <= 50
    assert all(len(value) <= 64 for value in result.ipc_codes)
    assert result.filing_date is None
    assert result.publication_date is None


def test_generic_search_rejects_unknown_country_before_network(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    called = False

    def fake_stream(*args, **kwargs):
        nonlocal called
        called = True
        return _FakeResponse("<response/>")

    monkeypatch.setattr(kipris.httpx, "stream", fake_stream)
    with pytest.raises(ValueError, match="Unsupported KIPRIS country code.*ZZ"):
        kipris.PatentSearchCriteria(query="storage", countries=("KR", "ZZ"))
    assert called is False


def test_generic_search_rejects_malformed_xml(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _FakeResponse("<response><broken>"),
    )
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    with pytest.raises(kipris.KiprisError, match="응답 파싱 실패"):
        client.search(kipris.PatentSearchCriteria(query="storage", countries=("KR",)))


def test_generic_foreign_search_http_error_redacts_encoded_api_key(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    secret = "FOREIGN/SECRET+KEY"

    class _ForeignErrorResponse(_FakeResponse):
        def __init__(self) -> None:
            super().__init__("")

        def raise_for_status(self) -> None:
            request = kipris.httpx.Request(
                "GET",
                f"https://example/openapi/advancedSearch?free=x&accessKey={secret}",
            )
            raise kipris.httpx.HTTPStatusError(
                f"Server error for url '{request.url}'",
                request=request,
                response=kipris.httpx.Response(500, request=request),
            )

    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _ForeignErrorResponse(),
    )
    client = kipris.KiprisClient(secret, "https://example/kipi")
    with pytest.raises(kipris.KiprisError) as exc_info:
        client.search(kipris.PatentSearchCriteria(query="storage", countries=("US",)))
    message = str(exc_info.value)
    assert secret not in message
    assert "FOREIGN%2FSECRET%2BKEY" not in message
    assert "accessKey" not in message
    assert "HTTP 500" in message


def test_generic_search_http_error_does_not_disclose_private_query_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    private_query = "confidential invention phrase"
    private_applicant = "Private Applicant"

    class _PrivateQueryErrorResponse(_FakeResponse):
        def __init__(self) -> None:
            super().__init__("")

        def raise_for_status(self) -> None:
            request = kipris.httpx.Request(
                "GET",
                "https://example/openapi/advancedSearch",
                params={
                    "free": private_query,
                    "applicant": private_applicant,
                    "accessKey": "KEY",
                },
            )
            response = kipris.httpx.Response(503, request=request)
            raise kipris.httpx.HTTPStatusError(
                f"Server error for url '{request.url}'",
                request=request,
                response=response,
            )

    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _PrivateQueryErrorResponse(),
    )
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    with pytest.raises(kipris.KiprisError) as exc_info:
        client.search(
            kipris.PatentSearchCriteria(
                query=private_query,
                countries=("US",),
                applicants=(private_applicant,),
            )
        )

    message = str(exc_info.value)
    assert message == "KIPRIS 요청 실패 (HTTP 503)."
    assert private_query not in message
    assert private_applicant not in message


@pytest.mark.parametrize(
    ("criteria_kwargs", "message"),
    [
        ({"query": "q" * 8_001, "countries": ("KR",)}, "query must be at most"),
        (
            {
                "query": "storage",
                "countries": ("KR",),
                "applicants": tuple(f"applicant-{index}" for index in range(31)),
            },
            "applicants must contain at most 30",
        ),
        (
            {
                "query": "storage",
                "countries": ("KR",),
                "applicants": ("x" * 257,),
            },
            "applicants values must be at most 256",
        ),
    ],
)
def test_generic_search_criteria_are_bounded(
    criteria_kwargs: dict[str, object],
    message: str,
) -> None:
    with pytest.raises(ValueError, match=message):
        kipris.PatentSearchCriteria(**criteria_kwargs)


def test_generic_search_streams_with_a_bounded_xml_response(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(kipris, "_MAX_SEARCH_XML_BYTES", 16)
    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _FakeResponse("<response>" + ("x" * 100) + "</response>"),
    )
    client = kipris.KiprisClient("KEY", "https://example/kipi")

    with pytest.raises(kipris.KiprisError, match="크기 제한"):
        client.search(kipris.PatentSearchCriteria(query="storage", countries=("KR",)))


def test_generic_foreign_result_uses_unpadded_provider_number_when_publication_missing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    xml = (
        _generic_foreign_search_xml("US")
        .replace(
            "<publishrNo>2024012345</publishrNo>",
            "",
        )
        .replace("<applicationNo>2024-0012345</applicationNo>", "")
    )
    monkeypatch.setattr(
        kipris.httpx,
        "stream",
        lambda *args, **kwargs: _FakeResponse(xml),
    )
    client = kipris.KiprisClient("KEY", "https://example/kipi")

    result = client.search(kipris.PatentSearchCriteria(query="storage", countries=("US",))).results[
        0
    ]

    assert result.canonical_number == "US20240012345A1"
    assert result.external_url.endswith("/US20240012345A1")


def test_platform_kipris_factory_owns_settings_and_fails_closed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from open_alm_api.core import settings as settings_module

    monkeypatch.setattr(
        settings_module,
        "get_settings",
        lambda: SimpleNamespace(
            patent_kipris_api_key="PLATFORM-KEY",
            patent_kipris_base_url="https://example/kipi",
        ),
    )
    client = kipris.get_kipris_client()
    assert client.configured is True
    assert client.base_url == "https://example/kipi"

    monkeypatch.setattr(
        settings_module,
        "get_settings",
        lambda: SimpleNamespace(
            patent_kipris_api_key="",
            patent_kipris_base_url="https://example/kipi",
        ),
    )
    with pytest.raises(kipris.KiprisNotConfiguredError):
        kipris.get_kipris_client()


def test_fetch_full_resolves_claims_from_bibliography(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_get(url, params=None, timeout=None, **_kwargs):
        if url.endswith("getAdvancedSearch"):
            return _FakeResponse(_word_search_xml())
        if url.endswith("getBibliographyDetailInfoSearch"):
            return _FakeResponse(_bibliography_xml())
        return _FakeResponse("<response/>")

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    result = client.fetch_full("10-2024-0001234")
    assert result is not None
    assert result["app_no"] == "10-2024-0001234"
    assert len(result["claims"]) == 2
    assert result["abstract"].startswith("상세 초록")  # bibliography abstract wins (longer)


def test_strip_markup_removes_pseudo_html_tags() -> None:
    assert kipris._strip_markup("<P>청구항 1. 시스템.</P>") == "청구항 1. 시스템."
    cleaned = kipris._strip_markup("앞<br/>뒤 <P>단락</P>")
    assert "<" not in cleaned and ">" not in cleaned


def test_fetch_full_strips_claim_markup(monkeypatch: pytest.MonkeyPatch) -> None:
    biblio = (
        "<response><body>"
        "<item><applicationNumber>1020240001234</applicationNumber></item>"
        "<claimInfo><claim>&lt;P&gt;1. 공조시스템.&lt;/P&gt;</claim></claimInfo>"
        "</body></response>"
    )

    def fake_get(url, params=None, timeout=None, **_kwargs):
        if url.endswith("getAdvancedSearch"):
            return _FakeResponse(_word_search_xml())
        if url.endswith("getBibliographyDetailInfoSearch"):
            return _FakeResponse(biblio)
        return _FakeResponse("<response/>")

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    result = client.fetch_full("10-2024-0001234")
    assert result is not None
    assert result["claims"] == ["1. 공조시스템."]
    assert "<P>" not in result["claims"][0]


def test_fetch_full_returns_none_when_not_found(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _FakeResponse("<response/>"))
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    assert client.fetch_full("99-9999-9999999") is None


@pytest.mark.parametrize(
    "number,expected",
    [
        ("EP 4 344 910 B1", "EP"),
        ("US 10,173,491 B2", "US"),
        ("WO 2023/070606", "WO"),
        ("JP 2020-082912", "JP"),
        ("CN 113022261 A", "CN"),
        ("10-2850274", None),
        ("1020060054014", None),
        ("KR10-2850274", None),
        ("2019970025199", None),
    ],
)
def test_detect_foreign_country(number: str, expected: str | None) -> None:
    assert kipris.detect_foreign_country(number) == expected


def test_word_fallback_rejects_unverified_match(monkeypatch: pytest.MonkeyPatch) -> None:
    """A getWordSearch hit whose number fields don't contain the searched digits
    must be discarded — this is what previously surfaced an unrelated patent."""
    unrelated = (
        "<response><body><items><item>"
        "<applicationNumber>1020247009229</applicationNumber>"
        "<inventionTitle>넥틴-4 항체 및 접합체</inventionTitle>"
        "</item></items></body></response>"
    )

    def fake_get(url, params=None, timeout=None, **_kwargs):
        if url.endswith("getWordSearch"):
            return _FakeResponse(unrelated)
        return _FakeResponse("<response/>")  # advancedSearch finds nothing

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    # digits 4344910 are absent from the returned number → no false match.
    assert client.fetch_full("4344910") is None


@pytest.mark.parametrize(
    "number,expected",
    [
        ("EP 4 344 910 B1", "https://patents.google.com/patent/EP4344910B1"),
        ("US 10,173,491 B2", "https://patents.google.com/patent/US10173491B2"),
        ("WO 2023/070606", "https://patents.google.com/patent/WO2023070606"),
    ],
)
def test_foreign_google_patents_url(number: str, expected: str) -> None:
    country = kipris.detect_foreign_country(number)
    assert kipris.foreign_google_patents_url(country, number) == expected


def test_fetch_foreign_full_resolves_via_bibliographic(monkeypatch: pytest.MonkeyPatch) -> None:
    search_xml = (
        "<response><body><items><searchResult>"
        "<ltrtno>202300070606A1</ltrtno><countryCode>WO</countryCode>"
        "<publishrNo>2023070606</publishrNo>"
        "<inventionName>DEHUMIDIFICATION MODE CONTROL METHOD</inventionName>"
        "<applicant>Geely</applicant>"
        "</searchResult></items></body></response>"
    )
    biblio_xml = (
        "<response><body><items><bibliographicInfo>"
        "<summation><astrtCont>A dehumidification mode control method.</astrtCont></summation>"
        "<ipcInfo><ipcCd>B60H 1/00</ipcCd></ipcInfo>"
        "</bibliographicInfo></items></body></response>"
    )

    def fake_get(url, params=None, timeout=None, headers=None, **_kwargs):
        if url.endswith("/advancedSearch"):
            return _FakeResponse(search_xml)
        if url.endswith("/bibliographicInfo"):
            return _FakeResponse(biblio_xml)
        return _FakeResponse("<response/>")  # Google page has no <claims> block

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient(
        "KEY", "https://plus.kipris.or.kr/kipo-api/kipi/patUtiModInfoSearchSevice"
    )
    result = client.fetch_foreign_full("WO 2023/070606", "WO")
    assert result.get("error") is None
    assert result["title"] == "DEHUMIDIFICATION MODE CONTROL METHOD"
    assert result["ipc"] == "B60H 1/00"
    assert result["foreign"] is True
    # KIPRIS has no claims and Google scrape returns none → abstract fallback.
    assert result["claims"] == ["A dehumidification mode control method."]
    assert result["claims_source"] == "abstract_fallback"
    assert result["google_patents_url"] == "https://patents.google.com/patent/WO2023070606"


def test_fetch_foreign_full_falls_back_to_google_claims(monkeypatch: pytest.MonkeyPatch) -> None:
    """When KIPRIS biblio has no claimText (WO/PCT), claims come from Google Patents."""
    search_xml = (
        "<response><body><items><searchResult>"
        "<ltrtno>202300070606A1</ltrtno><countryCode>WO</countryCode>"
        "<publishrNo>2023070606</publishrNo>"
        "<inventionName>DEHUMIDIFICATION MODE CONTROL METHOD</inventionName>"
        "</searchResult></items></body></response>"
    )
    biblio_no_claims = (
        "<response><body><items><bibliographicInfo>"
        "<summation><astrtCont>A dehumidification method.</astrtCont></summation>"
        "<ipcInfo><ipcCd>B60H 1/00</ipcCd></ipcInfo>"
        "</bibliographicInfo></items></body></response>"
    )
    google_html = (
        '<section itemprop="claims" itemscope>'
        '<div itemprop="content" html><claims lang="EN">'
        '<claim num="1"><claim-text>'
        '<span class="notranslate"><span class="google-src-text">原文</span>'
        "A method for controlling dehumidification.</span></claim-text></claim>"
        '<claim num="2"><claim-text>The method of claim 1, further comprising X.</claim-text></claim>'
        "</claims></div></section>"
    )

    def fake_get(url, params=None, timeout=None, headers=None, **_kwargs):
        if url.endswith("/advancedSearch"):
            return _FakeResponse(search_xml)
        if url.endswith("/bibliographicInfo"):
            return _FakeResponse(biblio_no_claims)
        if "patents.google.com" in url:
            return _FakeResponse(google_html)
        return _FakeResponse("<response/>")

    monkeypatch.setattr(kipris.httpx, "get", fake_get)
    client = kipris.KiprisClient(
        "KEY", "https://plus.kipris.or.kr/kipo-api/kipi/patUtiModInfoSearchSevice"
    )
    result = client.fetch_foreign_full("WO 2023/070606", "WO")
    assert result.get("error") is None
    assert result["claims_source"] == "google_patents"
    assert len(result["claims"]) == 2
    assert result["claims"][0].startswith("1. A method for controlling dehumidification")
    # the source-language span must be stripped out
    assert "原文" not in result["claims"][0]


@pytest.mark.parametrize(
    "texts",
    [
        ["x"] * 101,  # too many items
        ["y" * 10_001],  # single item too long
        ["z" * 9_000] * 30,  # total too large (270k chars)
    ],
)
def test_translate_rejects_oversized_input(texts: list[str]) -> None:
    """The LLM-backed translate endpoint must bound per-request work; oversized
    input is rejected (413) before any LLM call."""

    def _boom(*args: object, **kwargs: object) -> dict:
        raise AssertionError("LLM must not be called for oversized input")

    namespace = SimpleNamespace(id="x")
    with pytest.raises(Exception) as exc_info:  # noqa: PT011 - asserting status below
        service.translate_texts(
            None,  # db unused before the size guard
            workspace=namespace,
            user=namespace,
            texts=texts,
        )
    assert getattr(exc_info.value, "status_code", None) == 413


def test_translate_within_limits_calls_llm(monkeypatch: pytest.MonkeyPatch) -> None:
    """Inputs within limits are translated; order is preserved, empties pass through."""
    monkeypatch.setattr(
        service.llm,
        "generate_json",
        lambda *a, **k: {"0": "초록 번역", "2": "청구항 번역"},
    )
    namespace = SimpleNamespace(id="x")
    result = service.translate_texts(
        None, workspace=namespace, user=namespace, texts=["abstract", "", "claim 1"]
    )
    assert result.translations == ["초록 번역", "", "청구항 번역"]


def test_fetch_foreign_full_not_found_returns_google_link(monkeypatch: pytest.MonkeyPatch) -> None:
    # advancedSearch returns a result whose numbers don't contain the digits.
    mismatch = (
        "<response><body><items><searchResult>"
        "<ltrtno>999999999A1</ltrtno><publishrNo>9999999</publishrNo>"
        "<inventionName>UNRELATED</inventionName>"
        "</searchResult></items></body></response>"
    )
    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _FakeResponse(mismatch))
    client = kipris.KiprisClient(
        "KEY", "https://plus.kipris.or.kr/kipo-api/kipi/patUtiModInfoSearchSevice"
    )
    result = client.fetch_foreign_full("EP 4 344 910 B1", "EP")
    assert result["error"] == "foreign_not_found"
    assert result["google_patents_url"] == "https://patents.google.com/patent/EP4344910B1"


def test_resolve_view_urls_builds_google_patents(monkeypatch: pytest.MonkeyPatch) -> None:
    xml = """<response><body><item>
        <registerNumber>1025000000000</registerNumber>
    </item></body></response>"""
    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _FakeResponse(xml))
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    urls = client.resolve_view_urls("10-2024-0001234")
    assert urls["pdf_url"].startswith("https://patents.google.com/patent/KR")
    assert urls["type"] == "google_patents"


def test_resolve_view_urls_utility_model_uses_u_kind(monkeypatch: pytest.MonkeyPatch) -> None:
    # Utility model (doc-type code "20"): published open number → Google "U" kind,
    # leading "20" dropped (e.g. 2019990010756 → KR19990010756U).
    xml = "<response><body><item><openNumber>2019990010756</openNumber></item></body></response>"
    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _FakeResponse(xml))
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    urls = client.resolve_view_urls("2019970025199")
    assert urls["pdf_url"] == "https://patents.google.com/patent/KR19990010756U"


def test_resolve_view_urls_patent_open_drops_leading_10(monkeypatch: pytest.MonkeyPatch) -> None:
    xml = "<response><body><item><openNumber>1020000061599</openNumber></item></body></response>"
    monkeypatch.setattr(kipris.httpx, "get", lambda *a, **k: _FakeResponse(xml))
    client = kipris.KiprisClient("KEY", "https://example/kipi")
    urls = client.resolve_view_urls("1020000061599")
    assert urls["pdf_url"] == "https://patents.google.com/patent/KR20000061599A"


@pytest.mark.parametrize(
    "raw,expected",
    [
        ('{"a": 1}', {"a": 1}),
        ('```json\n{"a": 1}\n```', {"a": 1}),
        ('{"a": [1, 2', {"a": [1, 2]}),  # repaired: close array + object
        ("not json", {}),
    ],
)
def test_parse_json_lenient(raw: str, expected: dict) -> None:
    assert llm.parse_json_lenient(raw) == expected


def test_search_query_prompt_contains_schema_rules() -> None:
    text = prompts.search_query_prompt("전기차 히트펌프")
    assert "keyword_groups" in text
    assert "search_formula" in text
    assert "전기차 히트펌프" in text


def test_report_prompt_includes_context_and_label() -> None:
    text = prompts.report_prompt("invention-disclosure", "내 발명", "참고특허")
    assert "[발명의 명칭]" in text
    assert "내 발명" in text
    assert "참고특허" in text
    assert prompts.REPORT_LABELS["invention-disclosure"] == "직무발명서"


def test_ai_search_shapes_results_and_tags(monkeypatch: pytest.MonkeyPatch) -> None:
    fake_client = SimpleNamespace(
        word_search=lambda q, page=1, num_rows=15, deadline_monotonic=None: (
            [
                {
                    "app_no": "10-2024-0001234",
                    "reg_no": "",
                    "title": "전기차 히트펌프",
                    "applicant": "Open ALM",
                    "date": "20240115",
                    "status": "등록",
                    "ipc": "F25B30/00",
                    "abstract": "초록",
                }
            ],
            1,
        ),
        applicant_names=lambda q, page=1, num_rows=30, deadline_monotonic=None: ["Open ALM"],
    )
    monkeypatch.setattr(service, "_client", lambda: fake_client)

    def fake_generate_json(db, *, workspace_id, actor_user_id, prompt, max_tokens, temperature=0.2):
        if "검색어를 생성" in prompt:
            return {
                "main_keywords": ["히트펌프"],
                "search_query": "전기차 히트펌프",
                "tech_summary": "요약",
            }
        return {
            "summary": "검색 요약",
            "core_techs": ["히트펌프"],
            "suggested_queries": ["배터리 열관리"],
            "patent_tags": {"1": {"tags": ["냉매"], "relevance": 88}},
        }

    monkeypatch.setattr(service.llm, "generate_json", fake_generate_json)

    result = service.ai_search(
        None,
        workspace=SimpleNamespace(id="w1"),
        user=SimpleNamespace(id="u1"),
        query="전기차 배터리 히트펌프",
        page=1,
        applicant_filter="",
    )
    assert result.total == 1
    assert result.search_query == "전기차 히트펌프"
    assert result.results[0].relevance == 88
    assert result.results[0].ai_tags == ["냉매"]
    assert result.ai_summary == "검색 요약"
    assert result.top_applicants[0].name == "Open ALM"


def test_ai_search_shares_one_short_deadline_across_all_kipris_pages(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple[str, int, int, float | None]] = []

    class _DeadlineClient:
        def word_search(
            self,
            _query: str,
            *,
            page: int,
            num_rows: int,
            deadline_monotonic: float | None,
        ):
            calls.append(("word", page, num_rows, deadline_monotonic))
            return [], 150

        def applicant_names(
            self,
            _query: str,
            *,
            page: int,
            num_rows: int,
            deadline_monotonic: float | None,
        ):
            calls.append(("applicant", page, num_rows, deadline_monotonic))
            return []

    monkeypatch.setattr(service, "_client", _DeadlineClient)
    monkeypatch.setattr(service.time, "monotonic", lambda: 200.0)
    monkeypatch.setattr(
        service.llm,
        "generate_json",
        lambda *_args, **_kwargs: {
            "main_keywords": ["히트펌프"],
            "search_query": "전기차 히트펌프",
            "tech_summary": "요약",
        },
    )

    result = service.ai_search(
        None,
        workspace=SimpleNamespace(id="w1"),
        user=SimpleNamespace(id="u1"),
        query="전기차 배터리 히트펌프",
        page=1,
        applicant_filter="",
    )

    assert result.total == 150
    assert calls == [
        ("word", 1, 15, 215.0),
        *[("applicant", page, 30, 215.0) for page in range(1, 6)],
    ]
    assert service._KIPRIS_INTERACTIVE_SEARCH_TIMEOUT_SECONDS == 15.0


def test_generate_report_rejects_unknown_type() -> None:
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as exc:
        service.generate_report(
            None,
            workspace=SimpleNamespace(id="w1"),
            user=SimpleNamespace(id="u1"),
            content="x",
            report_type="bogus",
            patent_context="",
        )
    assert exc.value.status_code == 400


def _docx_bytes(paragraphs: list[str]) -> bytes:
    import io

    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_extract_text_from_docx() -> None:
    result = service.extract_text(
        filename="invention.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=_docx_bytes(["발명의 명칭: 전기차 히트펌프", "배터리 폐열 회수 구성"]),
    )
    assert "전기차 히트펌프" in result.text
    assert "배터리 폐열 회수" in result.text
    assert result.char_count == len(result.text)


def test_extract_text_from_xlsx() -> None:
    result = service.extract_text(
        filename="specs.xlsx",
        mime_type="",
        content=_xlsx_bytes([["항목", "값"], ["냉매", "R134a"]]),
    )
    assert "냉매" in result.text
    assert "R134a" in result.text


def test_extract_text_rejects_empty_and_unsupported() -> None:
    from fastapi import HTTPException

    with pytest.raises(HTTPException) as empty_exc:
        service.extract_text(filename="x.docx", mime_type="", content=b"")
    assert empty_exc.value.status_code == 400

    with pytest.raises(HTTPException) as unsupported_exc:
        service.extract_text(filename="note.txt", mime_type="text/plain", content=b"hello")
    assert unsupported_exc.value.status_code == 415


@pytest.mark.anyio
async def test_extract_endpoint_offloads_document_parsing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    called: dict[str, object] = {}

    async def fake_run_in_threadpool(func, *args, **kwargs):
        called["func"] = func
        called["kwargs"] = kwargs
        return func(*args, **kwargs)

    def fake_extract_text(*, filename: str, mime_type: str, content: bytes):
        return service.ExtractResponse(
            filename=filename,
            text=content.decode("utf-8"),
            char_count=len(content),
            truncated=False,
        )

    monkeypatch.setattr(router, "run_in_threadpool", fake_run_in_threadpool)
    monkeypatch.setattr(router.service, "extract_text", fake_extract_text)

    upload = UploadFile(file=BytesIO(b"patent text"), filename="idea.docx")
    result = await router.extract(file=upload)

    assert result.text == "patent text"
    assert called["func"] is fake_extract_text
    assert called["kwargs"] == {
        "filename": "idea.docx",
        "mime_type": "",
        "content": b"patent text",
    }
