from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.exceptions import HTTPException

from ai_do_api.core.db import Base
from ai_do_api.domains.auth.models import User, Workspace, utcnow_naive
from ai_do_api.domains.legacy_issues.dataset_records import (
    export_dataset_records_xlsx,
    get_dataset_definition,
)
from ai_do_api.domains.legacy_issues import excel_exports as excel_export_service
from ai_do_api.domains.legacy_issues import router as legacy_issue_router
from ai_do_api.domains.legacy_issues.excel_exports import (
    EXCEL_EXPORT_MAX_ATTEMPTS,
    EXCEL_EXPORT_MAX_ATTACHMENTS_PER_RECORD,
    EXCEL_EXPORT_STALE_RUNNING_AFTER,
    EXCEL_EXPORT_STATUS_COMPLETED,
    EXCEL_EXPORT_STATUS_EXPIRED,
    EXCEL_EXPORT_STATUS_FAILED,
    EXCEL_EXPORT_STATUS_QUEUED,
    EXCEL_EXPORT_STATUS_RUNNING,
    ExcelExportAttachmentRef,
    ExcelExportManifest,
    _build_base_workbook,
    _normalize_excel_export_column_keys,
    cleanup_expired_excel_export_jobs,
    create_dataset_excel_export_job,
    create_vehicle_module_checklist_excel_export_job,
    excel_export_includes_attachments,
    excel_export_job_status,
    get_excel_export_job,
    republish_pending_excel_export_jobs,
    validate_excel_export_attachments,
)
from ai_do_api.domains.legacy_issues.excel_ole_package import embed_ole_attachments
from ai_do_api.domains.legacy_issues.models import LegacyIssueExcelExportJob


class _StorageResponse(BytesIO):
    def __init__(self, payload: bytes) -> None:
        super().__init__(payload)
        self.released = False

    def release_conn(self) -> None:
        self.released = True


class _Storage:
    def __init__(self, payloads: dict[str, bytes]) -> None:
        self.payloads = payloads
        self.responses: list[_StorageResponse] = []

    def get_object(self, _bucket_name: str, object_name: str) -> _StorageResponse:
        response = _StorageResponse(self.payloads[object_name])
        self.responses.append(response)
        return response


class _CleanupStorage:
    def __init__(self, *, fail_remove: bool = False) -> None:
        self.fail_remove = fail_remove
        self.removed: list[str] = []

    def remove_object(self, _bucket_name: str, object_name: str) -> None:
        self.removed.append(object_name)
        if self.fail_remove:
            raise RuntimeError("storage unavailable")

    def list_objects(self, *_args, **_kwargs) -> list[object]:
        return []


def _attachment(record_id: str, key: str, payload: bytes) -> ExcelExportAttachmentRef:
    return ExcelExportAttachmentRef(
        record_id=record_id,
        filename=f"근거-{key}.pdf",
        size_bytes=len(payload),
        storage_key=key,
    )


def test_master_export_lists_filenames_and_embeds_attachments_in_appended_columns(
    tmp_path: Path,
) -> None:
    definition = get_dataset_definition("common-master")
    record = SimpleNamespace(
        id="record-1",
        stable_record_id="stable-1",
        field_values={"legacy_issue_number": "LI-001", "symptom": "현상"},
    )
    payloads = {"one": b"first", "two": b"second"}
    manifest = ExcelExportManifest(
        definition=definition,
        rows=[record],
        attachments_by_record_id={
            record.id: [
                _attachment(record.id, "one", payloads["one"]),
                _attachment(record.id, "two", payloads["two"]),
            ]
        },
        sheet_name="master",
        result_filename="master.xlsx",
        checklist_layout=False,
    )
    storage = _Storage(payloads)
    base_path = tmp_path / "base.xlsx"
    result_path = tmp_path / "result.xlsx"

    placements = _build_base_workbook(
        manifest,
        workbook_path=base_path,
        client=storage,
        bucket_name="bucket",
    )
    base_book = load_workbook(base_path)
    base_sheet = base_book.active
    assert base_sheet.protection.sheet is False
    assert base_sheet.protection.password is None
    attachment_column = base_sheet.max_column
    filename_column = attachment_column - 1
    assert base_sheet.cell(1, filename_column).value == "첨부"
    assert base_sheet.cell(2, filename_column).value == "첨부파일명"
    assert base_sheet.cell(2, attachment_column).value == "첨부파일"
    assert any(
        merged.min_row == 1
        and merged.max_row == 1
        and merged.min_col == filename_column
        and merged.max_col == attachment_column
        for merged in base_sheet.merged_cells.ranges
    )
    assert base_sheet.cell(3, filename_column).value == "근거-one.pdf\n근거-two.pdf"
    assert base_sheet.cell(3, attachment_column).value in (None, "")
    assert base_sheet.cell(3, filename_column).protection.locked is True
    assert base_sheet.cell(3, attachment_column).protection.locked is True
    assert {item.cell_coordinate for item in placements} == {
        base_sheet.cell(3, attachment_column).coordinate
    }

    embed_ole_attachments(base_path, result_path, placements)

    result_book = load_workbook(result_path)
    result_sheet = result_book.active
    assert result_sheet.protection.sheet is False
    assert result_sheet.protection.password is None
    assert result_sheet.cell(3, attachment_column).value in (None, "")
    with ZipFile(result_path) as archive:
        embeddings = [
            name
            for name in archive.namelist()
            if name.startswith("xl/embeddings/") and name.endswith(".bin")
        ]
    assert len(embeddings) == 2
    assert len(storage.responses) == 2
    assert all(response.closed and response.released for response in storage.responses)


def test_checklist_export_places_attachments_below_check_group(tmp_path: Path) -> None:
    definition = get_dataset_definition("common-master")
    record = SimpleNamespace(
        id="check-record-1",
        source_stable_record_id="stable-1",
        field_values={
            "legacy_issue_number": "LI-001",
            "check_plan": "점검",
            "reflection_result": "결과",
        },
    )
    payload = b"evidence"
    manifest = ExcelExportManifest(
        definition=definition,
        rows=[record],
        attachments_by_record_id={record.id: [_attachment(record.id, "evidence", payload)]},
        sheet_name="checklist",
        result_filename="checklist.xlsx",
        checklist_layout=True,
    )
    base_path = tmp_path / "checklist-base.xlsx"

    placements = _build_base_workbook(
        manifest,
        workbook_path=base_path,
        client=_Storage({"evidence": payload}),
        bucket_name="bucket",
    )

    book = load_workbook(base_path)
    sheet = book.active
    assert sheet.protection.sheet is False
    assert sheet.protection.password is None
    filename_columns = [cell.column for cell in sheet[2] if cell.value == "첨부파일명"]
    attachment_columns = [cell.column for cell in sheet[2] if cell.value == "첨부파일"]
    assert len(filename_columns) == 1
    assert len(attachment_columns) == 1
    filename_column = filename_columns[0]
    attachment_column = attachment_columns[0]
    assert filename_column + 1 == attachment_column
    assert sheet.cell(3, filename_column).value == "근거-evidence.pdf"
    assert sheet.cell(3, filename_column).protection.locked is True
    assert sheet.cell(3, attachment_column).protection.locked is True
    assert placements[0].cell_coordinate == sheet.cell(3, attachment_column).coordinate
    assert sheet.cell(3, attachment_column).value in (None, "")
    check_group = next(
        merged
        for merged in sheet.merged_cells.ranges
        if merged.min_row == 1
        and merged.max_row == 1
        and merged.min_col <= attachment_column <= merged.max_col
    )
    assert sheet.cell(1, check_group.min_col).value == "CHECK"


@pytest.mark.parametrize("checklist_layout", [False, True])
def test_export_without_attachments_has_no_attachment_columns_or_storage_reads(
    tmp_path: Path,
    checklist_layout: bool,
) -> None:
    definition = get_dataset_definition("common-master")
    record = SimpleNamespace(
        id="record-1",
        stable_record_id="stable-1",
        source_stable_record_id="stable-1",
        field_values={"legacy_issue_number": "LI-001", "reflection_result": "결과"},
    )
    storage = _Storage({"unused": b"payload"})
    manifest = ExcelExportManifest(
        definition=definition,
        rows=[record],
        attachments_by_record_id={record.id: [_attachment(record.id, "unused", b"payload")]},
        sheet_name="export",
        result_filename="export.xlsx",
        checklist_layout=checklist_layout,
        include_attachments=False,
    )
    workbook_path = tmp_path / f"without-attachments-{checklist_layout}.xlsx"

    placements = _build_base_workbook(
        manifest,
        workbook_path=workbook_path,
        client=storage,
        bucket_name="bucket",
    )

    sheet = load_workbook(workbook_path).active
    header_row = 2
    assert "첨부파일명" not in [cell.value for cell in sheet[header_row]]
    assert "첨부파일" not in [cell.value for cell in sheet[header_row]]
    assert placements == []
    assert storage.responses == []


@pytest.mark.parametrize("checklist_layout", [False, True])
def test_export_uses_only_requested_columns_in_requested_order(
    tmp_path: Path,
    checklist_layout: bool,
) -> None:
    definition = get_dataset_definition("common-master")
    record = SimpleNamespace(
        id="record-1",
        stable_record_id="stable-1",
        source_stable_record_id="stable-1",
        field_values={
            "legacy_issue_number": "LI-001",
            "reflection_result": "검토 결과",
            "symptom": "숨겨진 현상",
        },
    )
    manifest = ExcelExportManifest(
        definition=definition,
        rows=[record],
        attachments_by_record_id={},
        sheet_name="selected",
        result_filename="selected.xlsx",
        checklist_layout=checklist_layout,
        column_keys=("reflection_result", "legacy_issue_number"),
        include_attachments=False,
    )
    workbook_path = tmp_path / f"selected-{checklist_layout}.xlsx"

    _build_base_workbook(
        manifest,
        workbook_path=workbook_path,
        client=_Storage({}),
        bucket_name="bucket",
    )

    sheet = load_workbook(workbook_path).active
    assert [
        sheet.cell(2, column).value or sheet.cell(1, column).value for column in range(1, 4)
    ] == [
        "record_id",
        "반영/검토결과",
        "과거차관리번호",
    ]
    assert sheet.max_column == 3
    assert sheet.cell(3, 2).value == "검토 결과"
    assert sheet.cell(3, 3).value == "LI-001"
    assert "현상" not in [cell.value for cell in sheet[2]]


def test_export_rejects_unknown_or_duplicate_requested_columns() -> None:
    definition = get_dataset_definition("common-master")

    for column_keys in [
        ["legacy_issue_number", "unknown"],
        ["legacy_issue_number", "legacy_issue_number"],
    ]:
        with pytest.raises(HTTPException) as blocked:
            _normalize_excel_export_column_keys(definition, column_keys)

        assert blocked.value.status_code == 422
        assert blocked.value.detail.code == "legacy_issues.excel_export_invalid_columns"


def test_dataset_export_request_preserves_unique_record_id_order() -> None:
    payload = legacy_issue_router.LegacyIssueExcelExportCreateRequest(
        record_ids=[" record-2 ", "record-1"],
    )

    assert payload.record_ids == ["record-2", "record-1"]

    with pytest.raises(ValueError):
        legacy_issue_router.LegacyIssueExcelExportCreateRequest(
            record_ids=["record-1", "record-1"],
        )


def test_dataset_export_row_selection_preserves_order_and_scope() -> None:
    rows = [
        SimpleNamespace(id="record-1"),
        SimpleNamespace(id="record-2"),
    ]

    assert [
        row.id
        for row in excel_export_service._select_dataset_export_rows(
            rows,
            ("record-2", "outside-scope", "record-1"),
        )
    ] == ["record-2", "record-1"]
    assert excel_export_service._select_dataset_export_rows(rows, ()) == []
    assert excel_export_service._select_dataset_export_rows(rows, None) is rows


def test_attachment_filenames_are_newline_separated_and_formula_safe(tmp_path: Path) -> None:
    definition = get_dataset_definition("common-master")
    record = SimpleNamespace(
        id="record-1",
        stable_record_id='=HYPERLINK("record")',
        field_values={"legacy_issue_number": '=HYPERLINK("field")'},
    )
    filenames = ['=HYPERLINK("bad")', "+SUM.xlsx", "정상.pdf"]
    manifest = ExcelExportManifest(
        definition=definition,
        rows=[record],
        attachments_by_record_id={
            record.id: [
                ExcelExportAttachmentRef(record.id, filename, 1, f"key-{index}")
                for index, filename in enumerate(filenames)
            ]
        },
        sheet_name="master",
        result_filename="master.xlsx",
        checklist_layout=False,
    )
    workbook_path = tmp_path / "formula-safe.xlsx"

    _build_base_workbook(
        manifest,
        workbook_path=workbook_path,
        client=_Storage({}),
        bucket_name="bucket",
    )

    sheet = load_workbook(workbook_path, data_only=False).active
    filename_cell = sheet.cell(3, sheet.max_column - 1)
    assert filename_cell.value == "\n".join(filenames)
    assert filename_cell.data_type == "s"
    legacy_issue_column = next(
        column
        for column in range(1, sheet.max_column + 1)
        if (sheet.cell(2, column).value or sheet.cell(1, column).value) == "과거차관리번호"
    )
    assert sheet.cell(3, 1).data_type == "s"
    assert sheet.cell(3, legacy_issue_column).data_type == "s"


def test_master_export_uses_grouped_headers_and_readable_layout(tmp_path: Path) -> None:
    definition = get_dataset_definition("common-master")
    record = SimpleNamespace(
        id="record-1",
        stable_record_id="stable-1",
        field_values={
            "legacy_issue_number": "LI-001",
            "occurrence_date": "2026-07-16",
            "received_date": "2026-07-17",
            "symptom": "긴 현상 설명입니다. " * 18,
            "check_plan": "점검 방안",
            "master_status": "등재",
        },
    )
    manifest = ExcelExportManifest(
        definition=definition,
        rows=[record],
        attachments_by_record_id={},
        sheet_name="master",
        result_filename="master.xlsx",
        checklist_layout=False,
        include_attachments=False,
    )
    workbook_path = tmp_path / "styled-master.xlsx"

    _build_base_workbook(
        manifest,
        workbook_path=workbook_path,
        client=_Storage({}),
        bucket_name="bucket",
    )

    sheet = load_workbook(workbook_path).active
    headers = {
        (sheet.cell(2, column).value or sheet.cell(1, column).value): column
        for column in range(1, sheet.max_column + 1)
    }
    legacy_issue_column = headers["과거차관리번호"]
    occurrence_date_column = headers["발생일"]
    received_date_column = headers["접수일"]
    symptom_column = headers["현상"]
    check_plan_column = headers["점검방안"]
    master_status_column = headers["마스터 상태"]

    assert sheet.freeze_panes == "A3"
    assert sheet.sheet_view.showGridLines is False
    assert sheet.print_title_rows == "$1:$2"
    assert sheet.cell(1, check_plan_column).value == "CHECK"
    assert sheet.cell(1, master_status_column).value == "점검 근거"
    assert any(
        merged.min_row == 1
        and merged.max_row == 1
        and merged.min_col <= check_plan_column <= merged.max_col
        for merged in sheet.merged_cells.ranges
    )
    assert any(
        merged.min_row == 1
        and merged.max_row == 2
        and merged.min_col == legacy_issue_column
        and merged.max_col == legacy_issue_column
        for merged in sheet.merged_cells.ranges
    )
    assert (
        sheet.column_dimensions[sheet.cell(3, symptom_column).column_letter].width
        > sheet.column_dimensions[sheet.cell(3, occurrence_date_column).column_letter].width
    )
    assert sheet.row_dimensions[3].height > 22
    assert sheet.row_dimensions[1].height > 20
    assert sheet.row_dimensions[2].height > 20
    assert sheet.cell(1, check_plan_column).font.bold is True
    assert sheet.cell(1, check_plan_column).font.color.rgb == "00FFFFFF"
    assert sheet.cell(2, check_plan_column).fill.fgColor.rgb != "00000000"
    assert sheet.cell(3, symptom_column).border.bottom.style == "thin"
    assert sheet.cell(3, occurrence_date_column).is_date is True
    assert sheet.cell(3, occurrence_date_column).value == datetime(2026, 7, 16)
    assert sheet.cell(3, occurrence_date_column).number_format == "yyyy-mm-dd"
    assert sheet.cell(3, received_date_column).is_date is True
    assert sheet.cell(3, received_date_column).value == datetime(2026, 7, 17)
    assert sheet.cell(3, received_date_column).number_format == "yyyy-mm-dd"


def test_synchronous_master_export_writes_date_fields_as_excel_dates() -> None:
    definition = get_dataset_definition("common-master")
    workbook = load_workbook(
        BytesIO(
            export_dataset_records_xlsx(
                definition,
                [
                    SimpleNamespace(
                        id="record-1",
                        stable_record_id="stable-1",
                        field_values={
                            "occurrence_date": "2026-07-16",
                            "received_date": "2026-07-17",
                        },
                    )
                ],
            )
        )
    )
    sheet = workbook.active
    assert sheet.protection.sheet is False
    assert sheet.protection.password is None
    headers = {sheet.cell(1, column).value: column for column in range(1, sheet.max_column + 1)}

    occurrence_date_cell = sheet.cell(2, headers["발생일"])
    received_date_cell = sheet.cell(2, headers["접수일"])
    assert occurrence_date_cell.is_date is True
    assert occurrence_date_cell.value == datetime(2026, 7, 16)
    assert occurrence_date_cell.number_format == "yyyy-mm-dd"
    assert received_date_cell.is_date is True
    assert received_date_cell.value == datetime(2026, 7, 17)
    assert received_date_cell.number_format == "yyyy-mm-dd"


def test_attachment_limit_rejects_a_single_oversized_record() -> None:
    attachments = [
        ExcelExportAttachmentRef(
            record_id="record-1",
            filename=f"{index}.bin",
            size_bytes=1,
            storage_key=f"key-{index}",
        )
        for index in range(EXCEL_EXPORT_MAX_ATTACHMENTS_PER_RECORD + 1)
    ]

    with pytest.raises(HTTPException) as blocked:
        validate_excel_export_attachments({"record-1": attachments})

    assert blocked.value.status_code == 413
    assert blocked.value.detail.code == "legacy_issues.excel_export_record_attachment_limit"


def test_dataset_export_reuses_an_identical_active_job(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    monkeypatch.setattr(
        excel_export_service, "list_dataset_records", lambda *_args, **_kwargs: ([], 0)
    )
    monkeypatch.setattr(
        excel_export_service, "_dataset_attachment_refs", lambda *_args, **_kwargs: {}
    )
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        db.add_all([workspace, user])
        db.commit()
        definition = get_dataset_definition("common-master")
        revision = SimpleNamespace(id="revision-1")

        first = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=definition,
            revision=revision,
            view_key="aircon",
            departments=["개발", "개발"],
            column_keys=["symptom", "legacy_issue_number"],
            record_ids=["record-2", "record-1"],
        )
        second = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=definition,
            revision=revision,
            view_key="aircon",
            departments=["개발"],
            column_keys=["symptom", "legacy_issue_number"],
            record_ids=["record-2", "record-1"],
        )
        different_order = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=definition,
            revision=revision,
            view_key="aircon",
            departments=["개발"],
            column_keys=["legacy_issue_number", "symptom"],
            record_ids=["record-2", "record-1"],
        )
        different_row_order = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=definition,
            revision=revision,
            view_key="aircon",
            departments=["개발"],
            column_keys=["symptom", "legacy_issue_number"],
            record_ids=["record-1", "record-2"],
        )

        assert second.id == first.id
        assert different_order.id != first.id
        assert different_row_order.id != first.id
        assert first.request_params["record_ids"] == ["record-2", "record-1"]
        assert db.query(LegacyIssueExcelExportJob).count() == 3


def test_dataset_export_counts_and_loads_attachments_for_selected_rows_only(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    rows = [
        SimpleNamespace(id="record-1"),
        SimpleNamespace(id="record-2"),
        SimpleNamespace(id="record-3"),
    ]
    attachment_row_ids: list[str] = []
    requested_record_id_batches: list[tuple[str, ...]] = []

    def list_records(*_args, **kwargs):
        requested_record_id_batches.append(tuple(kwargs["record_ids"]))
        return rows, len(rows)

    monkeypatch.setattr(
        excel_export_service,
        "DATASET_EXPORT_RECORD_ID_QUERY_BATCH_SIZE",
        2,
    )
    monkeypatch.setattr(
        excel_export_service,
        "list_dataset_records",
        list_records,
    )

    def attachment_refs(*_args, **kwargs):
        attachment_row_ids.extend(row.id for row in kwargs["rows"])
        return {}

    monkeypatch.setattr(excel_export_service, "_dataset_attachment_refs", attachment_refs)
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        db.add_all([workspace, user])
        db.commit()

        job = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=get_dataset_definition("common-master"),
            revision=SimpleNamespace(id="revision-1"),
            view_key="aircon",
            departments=None,
            record_ids=["record-3", "outside-scope", "record-1"],
        )

        assert job.record_count == 2
        assert attachment_row_ids == ["record-3", "record-1"]
        assert requested_record_id_batches == [
            ("record-3", "outside-scope"),
            ("record-1",),
        ]


def test_dataset_export_worker_reapplies_requested_row_order(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    definition = get_dataset_definition("common-master")
    rows = [
        SimpleNamespace(id="record-1"),
        SimpleNamespace(id="record-2"),
        SimpleNamespace(id="record-3"),
    ]
    monkeypatch.setattr(
        excel_export_service,
        "get_dataset_definition_for_view",
        lambda *_args, **_kwargs: definition,
    )
    monkeypatch.setattr(
        excel_export_service,
        "resolve_read_revision",
        lambda *_args, **_kwargs: SimpleNamespace(
            current=SimpleNamespace(revision_no=3),
        ),
    )
    monkeypatch.setattr(
        excel_export_service,
        "list_dataset_records",
        lambda *_args, **_kwargs: (rows, len(rows)),
    )
    job = SimpleNamespace(
        dataset_key=definition.key,
        module_key="aircon",
        request_params={
            "include_attachments": False,
            "record_ids": ["record-3", "record-1"],
            "view_key": "aircon",
        },
        revision_id="revision-1",
        source_kind="dataset",
    )

    manifest = excel_export_service._load_export_manifest(
        SimpleNamespace(),
        workspace=SimpleNamespace(id="workspace-1"),
        job=job,
    )

    assert [row.id for row in manifest.rows] == ["record-3", "record-1"]


def test_dataset_export_attachment_option_creates_distinct_jobs_and_skips_lookup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    attachment_lookups = 0

    def attachment_refs(*_args, **_kwargs):
        nonlocal attachment_lookups
        attachment_lookups += 1
        return {}

    monkeypatch.setattr(
        excel_export_service, "list_dataset_records", lambda *_args, **_kwargs: ([], 0)
    )
    monkeypatch.setattr(excel_export_service, "_dataset_attachment_refs", attachment_refs)
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        db.add_all([workspace, user])
        db.commit()
        definition = get_dataset_definition("common-master")
        revision = SimpleNamespace(id="revision-1")

        without_attachments = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=definition,
            revision=revision,
            view_key="aircon",
            departments=None,
            include_attachments=False,
        )
        with_attachments = create_dataset_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            definition=definition,
            revision=revision,
            view_key="aircon",
            departments=None,
            include_attachments=True,
        )

        assert without_attachments.id != with_attachments.id
        assert without_attachments.request_params["include_attachments"] is False
        assert with_attachments.request_params["include_attachments"] is True
        assert without_attachments.attachment_count == 0
        assert without_attachments.attachment_bytes == 0
        assert attachment_lookups == 1
        assert db.query(LegacyIssueExcelExportJob).count() == 2


def test_old_job_without_attachment_option_defaults_to_included() -> None:
    job = SimpleNamespace(request_params={})

    assert excel_export_includes_attachments(job) is True


def test_checklist_export_route_defaults_omitted_body_to_attachments(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    now = utcnow_naive()
    job = SimpleNamespace(
        id="job-1",
        status=EXCEL_EXPORT_STATUS_QUEUED,
        source_kind="vehicle_module_checklist",
        request_params={},
        record_count=0,
        attachment_count=0,
        attachment_bytes=0,
        processed_attachment_count=0,
        processed_attachment_bytes=0,
        result_filename=None,
        result_size_bytes=None,
        error_code=None,
        created_at=now,
        updated_at=now,
        completed_at=None,
        expires_at=None,
    )
    captured: dict[str, object] = {}

    def create_job(*_args, **kwargs):
        captured.update(kwargs)
        return job

    class DbStub:
        committed = False

        def commit(self) -> None:
            self.committed = True

    db = DbStub()
    monkeypatch.setattr(
        legacy_issue_router,
        "create_vehicle_module_checklist_excel_export_job",
        create_job,
    )
    monkeypatch.setattr(legacy_issue_router, "_enabled_module_keys", lambda: frozenset())

    result = legacy_issue_router.create_legacy_issue_vehicle_module_checklist_excel_export(
        "checklist-1",
        payload=None,
        db=db,
        current_user=SimpleNamespace(id="user-1"),
        current_workspace=SimpleNamespace(id="workspace-1"),
    )

    assert captured["include_attachments"] is True
    assert captured["column_keys"] is None
    assert result.include_attachments is True
    assert db.committed is True


def test_checklist_export_request_accepts_attachment_exclusion() -> None:
    payload = legacy_issue_router.LegacyIssueChecklistExcelExportCreateRequest(
        column_keys=["reflection_result", "check_plan"], include_attachments=False
    )

    assert payload.column_keys == ["reflection_result", "check_plan"]
    assert payload.include_attachments is False


def test_checklist_export_reuses_an_identical_active_job(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    checklist = SimpleNamespace(
        id="checklist-1",
        source_dataset_key="common-master",
        module_key="aircon",
        source_master_revision_id="revision-1",
    )
    definition = get_dataset_definition("common-master")
    monkeypatch.setattr(
        excel_export_service,
        "list_vehicle_module_checklist_records",
        lambda *_args, **_kwargs: (checklist, definition, [], 0),
    )
    monkeypatch.setattr(
        excel_export_service,
        "list_vehicle_module_checklist_attachments_for_records",
        lambda *_args, **_kwargs: {},
    )
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        db.add_all([workspace, user])
        db.commit()

        first = create_vehicle_module_checklist_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            checklist_id=checklist.id,
            module_keys=frozenset({"aircon"}),
        )
        second = create_vehicle_module_checklist_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            checklist_id=checklist.id,
            module_keys=frozenset({"aircon"}),
        )

        assert second.id == first.id
        assert db.query(LegacyIssueExcelExportJob).count() == 1


def test_checklist_export_attachment_option_changes_dedup_and_skips_lookup(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    checklist = SimpleNamespace(
        id="checklist-1",
        source_dataset_key="common-master",
        module_key="aircon",
        source_master_revision_id="revision-1",
    )
    definition = get_dataset_definition("common-master")
    attachment_lookups = 0

    def list_attachments(*_args, **_kwargs):
        nonlocal attachment_lookups
        attachment_lookups += 1
        return {}

    monkeypatch.setattr(
        excel_export_service,
        "list_vehicle_module_checklist_records",
        lambda *_args, **_kwargs: (checklist, definition, [], 0),
    )
    monkeypatch.setattr(
        excel_export_service,
        "list_vehicle_module_checklist_attachments_for_records",
        list_attachments,
    )
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        db.add_all([workspace, user])
        db.commit()

        without_attachments = create_vehicle_module_checklist_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            checklist_id=checklist.id,
            module_keys=frozenset({"aircon"}),
            include_attachments=False,
        )
        with_attachments = create_vehicle_module_checklist_excel_export_job(
            db,
            workspace=workspace,
            user=user,
            checklist_id=checklist.id,
            module_keys=frozenset({"aircon"}),
            include_attachments=True,
        )

        assert without_attachments.id != with_attachments.id
        assert without_attachments.attachment_count == 0
        assert without_attachments.attachment_bytes == 0
        assert attachment_lookups == 1
        assert db.query(LegacyIssueExcelExportJob).count() == 2


def test_completed_job_is_reported_expired_after_ttl() -> None:
    job = SimpleNamespace(
        status=EXCEL_EXPORT_STATUS_COMPLETED,
        expires_at=utcnow_naive() - timedelta(seconds=1),
    )

    assert excel_export_job_status(job) == EXCEL_EXPORT_STATUS_EXPIRED


def test_cleanup_expires_database_pointer_even_when_object_deletion_fails() -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    now = utcnow_naive()
    storage = _CleanupStorage(fail_remove=True)
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        job = LegacyIssueExcelExportJob(
            id="expired-job",
            workspace_id=workspace.id,
            requested_by_id=user.id,
            source_kind="dataset",
            request_params={},
            status=EXCEL_EXPORT_STATUS_COMPLETED,
            result_storage_key="legacy-issues/excel-exports/workspace-1/expired-job.xlsx",
            result_size_bytes=123,
            expires_at=now - timedelta(seconds=1),
            created_at=now,
            updated_at=now,
        )
        db.add_all([workspace, user, job])
        db.commit()

        result = cleanup_expired_excel_export_jobs(
            db,
            client=storage,
            bucket_name="bucket",
        )

        db.refresh(job)
        assert result == {"expired": 1, "deleted_orphans": 0, "failed": 1}
        assert job.status == EXCEL_EXPORT_STATUS_EXPIRED
        assert job.result_storage_key is None
        assert job.result_size_bytes is None
        assert storage.removed == ["legacy-issues/excel-exports/workspace-1/expired-job.xlsx"]


def test_republisher_recovers_stale_running_jobs_and_dead_letters_exhausted_ones(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    published: list[str] = []
    monkeypatch.setattr(
        excel_export_service,
        "publish_excel_export_job",
        lambda job_id: published.append(job_id),
    )
    now = utcnow_naive()
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        user = User(
            id="user-1",
            login_id="user-1",
            email="user-1@example.com",
            full_name="User One",
            password_hash="hash",
            status="active",
        )
        recoverable = _job(
            "recoverable",
            workspace=workspace,
            user=user,
            attempts=1,
            started_at=now - EXCEL_EXPORT_STALE_RUNNING_AFTER - timedelta(minutes=1),
        )
        exhausted = _job(
            "exhausted",
            workspace=workspace,
            user=user,
            attempts=EXCEL_EXPORT_MAX_ATTEMPTS,
            started_at=now - EXCEL_EXPORT_STALE_RUNNING_AFTER - timedelta(minutes=1),
        )
        db.add_all([workspace, user, recoverable, exhausted])
        db.commit()

        assert republish_pending_excel_export_jobs(db) == 1

        assert db.get(LegacyIssueExcelExportJob, recoverable.id).status == (
            EXCEL_EXPORT_STATUS_QUEUED
        )
        assert db.get(LegacyIssueExcelExportJob, exhausted.id).status == (
            EXCEL_EXPORT_STATUS_FAILED
        )
        assert published == [recoverable.id]


def test_export_job_is_visible_only_to_its_requester() -> None:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(
        engine,
        tables=[Workspace.__table__, User.__table__, LegacyIssueExcelExportJob.__table__],
    )
    now = utcnow_naive()
    with Session(engine) as db:
        workspace = Workspace(id="workspace-1", key="workspace-1", name="Workspace")
        owner = User(
            id="owner",
            login_id="owner",
            email="owner@example.com",
            full_name="Owner",
            password_hash="hash",
            status="active",
        )
        outsider = User(
            id="outsider",
            login_id="outsider",
            email="outsider@example.com",
            full_name="Outsider",
            password_hash="hash",
            status="active",
        )
        job = LegacyIssueExcelExportJob(
            id="job-1",
            workspace_id=workspace.id,
            requested_by_id=owner.id,
            source_kind="dataset",
            request_params={},
            status=EXCEL_EXPORT_STATUS_QUEUED,
            created_at=now,
            updated_at=now,
        )
        db.add_all([workspace, owner, outsider, job])
        db.commit()

        assert (
            get_excel_export_job(
                db,
                workspace=workspace,
                user=owner,
                job_id=job.id,
            ).id
            == job.id
        )
        with pytest.raises(HTTPException) as hidden:
            get_excel_export_job(
                db,
                workspace=workspace,
                user=outsider,
                job_id=job.id,
            )

        assert hidden.value.status_code == 404
        assert hidden.value.detail.code == "legacy_issues.excel_export_not_found"


def _job(
    job_id: str,
    *,
    workspace: Workspace,
    user: User,
    attempts: int,
    started_at,
) -> LegacyIssueExcelExportJob:
    return LegacyIssueExcelExportJob(
        id=job_id,
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind="dataset",
        request_params={},
        status=EXCEL_EXPORT_STATUS_RUNNING,
        attempts=attempts,
        started_at=started_at,
        created_at=started_at,
        updated_at=started_at,
    )
