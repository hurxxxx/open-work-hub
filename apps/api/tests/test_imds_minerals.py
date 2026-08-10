from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
import openpyxl
import pytest

from open_alm_api.domains.auth.dependencies import (
    require_current_user,
    require_current_workspace,
)
from open_alm_api.domains.imds_minerals import parser
from open_alm_api.domains.imds_minerals.router import require_imds_minerals_app_enabled, router
from open_alm_api.domains.imds_minerals.target_list import match_oem
from open_alm_api.domains.imds_minerals.workbook import ImdsMeta, write_workbook


_WS_PREFIX = "/api/v1/workspaces/acme/imds-minerals"


def _workbook_bytes(workbook: openpyxl.Workbook) -> bytes:
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _template_bytes(sheet_name: str = "2. 책임광물 사용현황") -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet["B4"] = "old car"
    worksheet["AO4"] = "old note"
    return _workbook_bytes(workbook)


def _target_list_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet["A2"] = "차종"
    worksheet["B2"] = "OEM품번"
    worksheet["C2"] = "DCC품번"
    worksheet["D2"] = "품명"
    worksheet["A3"] = "NX4"
    worksheet["B3"] = "97200X9030"
    worksheet["C3"] = "DCC-001"
    worksheet["D3"] = "HEATER COMPLETE ASSY"
    worksheet["B4"] = "ABC-002"
    worksheet["C4"] = "DCC-002"
    worksheet["D4"] = "SUB PART"
    return _workbook_bytes(workbook)


def _make_client() -> TestClient:
    from starlette.exceptions import HTTPException as StarletteHTTPException

    from open_alm_api.app import localized_http_exception_handler

    app = FastAPI()
    app.include_router(router, prefix="/api/v1/workspaces/{workspace_slug}")
    app.add_exception_handler(StarletteHTTPException, localized_http_exception_handler)
    app.dependency_overrides[require_imds_minerals_app_enabled] = lambda: None
    app.dependency_overrides[require_current_user] = lambda: SimpleNamespace(id="user-1")
    app.dependency_overrides[require_current_workspace] = lambda: SimpleNamespace(id="ws-1")
    return TestClient(app)


def test_parser_wraps_invalid_pdf_as_parse_error() -> None:
    with pytest.raises(parser.ImdsParseError):
        parser.parse_and_classify(b"not a pdf")


def test_match_oem_reads_target_list_and_forward_fills_vehicle() -> None:
    first = match_oem(_target_list_bytes(), "97200x9030")
    second = match_oem(_target_list_bytes(), "abc-002")

    assert first == {
        "car": "NX4",
        "oem": "97200X9030",
        "dcc": "DCC-001",
        "end_name": "HEATER COMPLETE ASSY",
    }
    assert second == {
        "car": "NX4",
        "oem": "ABC-002",
        "dcc": "DCC-002",
        "end_name": "SUB PART",
    }


def test_write_workbook_fills_relevant_rows_and_clears_old_data() -> None:
    rows = [
        {
            "level": 1,
            "name": "HEATER ASSY",
            "code": "PN-1",
            "type": "component",
        },
        {
            "level": 2,
            "name": "Gold",
            "code": "7440-57-5",
            "type": "chemical",
        },
    ]
    content, written = write_workbook(
        _template_bytes(),
        "2. 책임광물 사용현황",
        rows,
        [0, 1],
        ImdsMeta(car="NX4", end_name="HEATER", oem="97200X9030", dcc="DCC-001"),
    )

    worksheet = openpyxl.load_workbook(BytesIO(content))["2. 책임광물 사용현황"]
    assert written == 2
    assert worksheet["B4"].value == "NX4"
    assert worksheet["C4"].value == "HEATER"
    assert worksheet["D4"].value == "97200X9030"
    assert worksheet["E4"].value == "DCC-001"
    assert worksheet["F4"].value == 1
    assert worksheet["G4"].value == 1
    assert worksheet["V4"].value == "HEATER ASSY"
    assert worksheet["W4"].value == "PN-1"
    assert worksheet["AO4"].value is None
    assert worksheet["F5"].value == 2
    assert worksheet["H5"].value == 2
    assert worksheet["Z5"].value == "Gold"
    assert worksheet["AA5"].value == "금"


def test_analyze_rejects_non_pdf_extension() -> None:
    response = _make_client().post(
        f"{_WS_PREFIX}/analyze",
        files={"file": ("report.txt", b"x", "text/plain")},
    )

    assert response.status_code == 400
    assert response.json()["code"] == "imds_minerals.invalid_pdf_type"


def test_imds_router_blocks_disabled_app() -> None:
    from starlette.exceptions import HTTPException as StarletteHTTPException

    from open_alm_api.app import localized_http_exception_handler

    app = FastAPI()
    app.include_router(router, prefix="/api/v1/workspaces/{workspace_slug}")
    app.add_exception_handler(StarletteHTTPException, localized_http_exception_handler)

    def disabled() -> None:
        from open_alm_api.core.i18n import localized_http_exception

        raise localized_http_exception(status_code=403, code="imds_minerals.app_disabled")

    app.dependency_overrides[require_imds_minerals_app_enabled] = disabled
    app.dependency_overrides[require_current_user] = lambda: SimpleNamespace(id="user-1")
    app.dependency_overrides[require_current_workspace] = lambda: SimpleNamespace(id="ws-1")

    response = TestClient(app).post(
        f"{_WS_PREFIX}/analyze",
        files={"file": ("report.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")},
    )

    assert response.status_code == 403
    assert response.json()["code"] == "imds_minerals.app_disabled"


def test_analyze_rejects_corrupt_pdf_as_localized_400() -> None:
    response = _make_client().post(
        f"{_WS_PREFIX}/analyze",
        files={"file": ("report.pdf", b"not a pdf", "application/pdf")},
    )

    assert response.status_code == 400
    assert response.json()["code"] == "imds_minerals.parse_failed"


def test_match_endpoint_returns_metadata() -> None:
    response = _make_client().post(
        f"{_WS_PREFIX}/match",
        data={"oem": "97200x9030"},
        files={
            "list_file": (
                "target-list.xlsx",
                _target_list_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200
    assert response.json() == {
        "found": True,
        "car": "NX4",
        "end_name": "HEATER COMPLETE ASSY",
        "oem": "97200X9030",
        "dcc": "DCC-001",
    }


def test_generate_rejects_corrupt_pdf_as_localized_400() -> None:
    response = _make_client().post(
        f"{_WS_PREFIX}/generate",
        data={
            "sheet": "2. 책임광물 사용현황",
            "car": "NX4",
            "end_name": "HEATER",
            "oem": "97200X9030",
            "dcc": "DCC-001",
        },
        files={
            "pdf": ("report.pdf", b"not a pdf", "application/pdf"),
            "template": (
                "template.xlsx",
                _template_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 400
    assert response.json()["code"] == "imds_minerals.parse_failed"


def test_generate_rejects_corrupt_template_as_localized_400(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "open_alm_api.domains.imds_minerals.service.parse_pdf",
        lambda _content: SimpleNamespace(rows=[], relevant=[]),
    )

    response = _make_client().post(
        f"{_WS_PREFIX}/generate",
        data={
            "sheet": "2. 책임광물 사용현황",
            "car": "NX4",
            "end_name": "HEATER",
            "oem": "97200X9030",
            "dcc": "DCC-001",
        },
        files={
            "pdf": ("report.pdf", b"%PDF-1.4\n%%EOF", "application/pdf"),
            "template": (
                "template.xlsx",
                b"not an xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 400
    assert response.json()["code"] == "imds_minerals.template_parse_failed"
