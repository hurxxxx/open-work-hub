from __future__ import annotations

from io import BytesIO
import zipfile

import pytest

from ai_do_api.domains.document_processing import extract_document
from ai_do_api.domains.document_processing import extractors as extractors_module
from ai_do_api.domains.document_processing import html_extractor as html_extractor_module
from ai_do_api.domains.document_processing.extractors import (
    DocumentExtractBundle,
    EvidenceBlock,
    UnsupportedDocumentType,
)
from ai_do_api.domains.document_processing.html_extractor import extract_html_stream


def _pptx_bytes(entries: dict[str, str]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for path, xml in entries.items():
            archive.writestr(path, xml)
    return buffer.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    buffer = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(buffer)
    return buffer.getvalue()


def _docx_with_malformed_vertical_merge_row() -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("고장과 무관한 본문 절차")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "병합 셀"
    table.cell(0, 1).text = "보존할 표 값"
    table.cell(1, 0).text = "병합 연속"
    table.cell(1, 1).text = "마지막 값"
    table.cell(0, 0).merge(table.cell(1, 0))

    # Reproduce a real-world malformed vertical merge: the continuation cell
    # points to grid offset 0, but the row above declares that offset as absent.
    first_row = table._tbl.tr_lst[0]  # pyright: ignore[reportPrivateUsage]
    first_row.remove(first_row.tc_lst[0])
    grid_before = OxmlElement("w:gridBefore")
    grid_before.set(qn("w:val"), "1")
    first_row.get_or_add_trPr().append(grid_before)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _slide_xml(*paragraphs: str) -> str:
    body = "".join(
        f"<a:p><a:r><a:t>{paragraph}</a:t></a:r></a:p>" for paragraph in paragraphs
    )
    return (
        "<p:sld xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\" "
        "xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\">"
        f"<p:cSld><p:spTree>{body}</p:spTree></p:cSld></p:sld>"
    )


def _notes_xml(*paragraphs: str) -> str:
    return _slide_xml(*paragraphs)


def _slide_with_table_xml() -> str:
    return (
        "<p:sld xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\" "
        "xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\">"
        "<p:cSld><p:spTree>"
        "<a:p><a:r><a:t>제품 사양</a:t></a:r></a:p>"
        "<a:tbl>"
        "<a:tr>"
        "<a:tc><a:txBody><a:p><a:r><a:t>항목</a:t></a:r></a:p></a:txBody></a:tc>"
        "<a:tc><a:txBody><a:p><a:r><a:t>값</a:t></a:r></a:p></a:txBody></a:tc>"
        "</a:tr>"
        "<a:tr>"
        "<a:tc><a:txBody><a:p><a:r><a:t>정격 출력</a:t></a:r></a:p></a:txBody></a:tc>"
        "<a:tc><a:txBody><a:p><a:r><a:t>10 kW</a:t></a:r></a:p></a:txBody></a:tc>"
        "</a:tr>"
        "</a:tbl>"
        "</p:spTree></p:cSld></p:sld>"
    )


class FixedExtractor:
    def supports(self, *, mime_type: str, filename: str) -> bool:
        return filename == "custom.fixture"

    def extract(
        self,
        *,
        document_id: str,
        filename: str,
        mime_type: str,
        content: bytes,
    ) -> DocumentExtractBundle:
        return DocumentExtractBundle(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type,
            evidence_blocks=[
                EvidenceBlock(
                    document_id=document_id,
                    block_id=f"{document_id}:custom",
                    locator_kind="custom",
                    locator_label="Custom",
                    section_path="Custom",
                    block_kind="text",
                    text=content.decode("utf-8"),
                )
            ],
        )


def test_extract_document_uses_supplied_extractor_registry() -> None:
    bundle = extract_document(
        document_id="doc",
        filename="custom.fixture",
        mime_type="application/custom",
        content=b"hello",
        extractors=[FixedExtractor()],
    )

    assert bundle.normalized_text == "[Custom] hello"
    assert bundle.to_dict()["evidence_blocks"][0]["block_id"] == "doc:custom"


def test_pptx_extractor_returns_slide_text_table_and_deduped_notes() -> None:
    bundle = extract_document(
        document_id="doc",
        filename="sample.pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        content=_pptx_bytes(
            {
                "ppt/slides/slide2.xml": _slide_xml("second"),
                "ppt/slides/slide1.xml": _slide_with_table_xml(),
                "ppt/notesSlides/notesSlide1.xml": _notes_xml("note", "note", "other note"),
            }
        ),
    )

    assert [block.block_kind for block in bundle.evidence_blocks] == [
        "text",
        "table",
        "text",
        "text",
        "text",
    ]
    assert [block.locator_label for block in bundle.evidence_blocks] == [
        "Slide 1",
        "Slide 1",
        "Slide 2",
        "Slide 1 notes",
        "Slide 1 notes",
    ]
    assert bundle.evidence_blocks[1].rows == [["항목", "값"], ["정격 출력", "10 kW"]]
    assert [block.text for block in bundle.evidence_blocks[-2:]] == ["note", "other note"]


def test_extract_document_rejects_unsupported_or_malformed_documents() -> None:
    with pytest.raises(UnsupportedDocumentType, match="Unsupported document type"):
        extract_document(
            document_id="doc",
            filename="sample.txt",
            mime_type="text/plain",
            content=b"text",
        )

    with pytest.raises(UnsupportedDocumentType, match="not a valid zip package"):
        extract_document(
            document_id="doc",
            filename="broken.pptx",
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            content=b"not a zip",
        )


def test_docx_extractor_skips_one_malformed_vertical_merge_row() -> None:
    bundle = extract_document(
        document_id="doc",
        filename="malformed-table.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=_docx_with_malformed_vertical_merge_row(),
    )

    assert bundle.evidence_blocks[0].text == "고장과 무관한 본문 절차"
    assert any("보존할 표 값" in block.text for block in bundle.evidence_blocks)
    assert bundle.metadata == {"malformed_table_rows_skipped": 1}


def test_pptx_extractor_stops_at_char_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    from ai_do_api.domains.document_processing import pptx as pptx_module

    monkeypatch.setattr(pptx_module, "_MAX_EXTRACTED_CHARS", 5)

    bundle = extract_document(
        document_id="doc",
        filename="sample.pptx",
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        content=_pptx_bytes(
            {
                "ppt/slides/slide1.xml": _slide_xml("abcdefghij", "klmnop"),
                "ppt/slides/slide2.xml": _slide_xml("second slide text"),
            }
        ),
    )

    assert bundle.evidence_blocks
    assert sum(len(block.text) for block in bundle.evidence_blocks) <= 5
    assert all("second" not in block.text for block in bundle.evidence_blocks)


def test_extract_embedded_office_documents_reads_nested_pptx() -> None:
    from ai_do_api.domains.document_processing.extractors import (
        extract_embedded_office_documents,
    )

    inner = _pptx_bytes({"ppt/slides/slide1.xml": _slide_xml("삽입된 발명 설명 텍스트")})
    outer = BytesIO()
    with zipfile.ZipFile(outer, "w") as archive:
        archive.writestr("ppt/presentation.xml", "<presentation />")
        archive.writestr("ppt/slides/slide1.xml", _slide_xml("바깥 슬라이드"))
        archive.writestr("ppt/embeddings/inner.pptx", inner)

    docs = extract_embedded_office_documents(outer.getvalue())

    assert [name for name, _text in docs] == ["inner.pptx"]
    assert "삽입된 발명 설명 텍스트" in docs[0][1]


def test_extract_embedded_office_documents_skips_decompression_bomb() -> None:
    from ai_do_api.domains.document_processing.extractors import (
        extract_embedded_office_documents,
    )

    # A tiny compressed member that expands ~1000x. It must be rejected by the
    # size/expansion-ratio guard before archive.read() allocates it in full,
    # never decompressed into a result.
    outer = BytesIO()
    with zipfile.ZipFile(outer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/presentation.xml", "<presentation />")
        archive.writestr("ppt/embeddings/bomb.pptx", b"\x00" * (5 * 1024 * 1024))

    docs = extract_embedded_office_documents(outer.getvalue())

    assert docs == []


def test_extract_embedded_office_documents_ignores_non_zip() -> None:
    from ai_do_api.domains.document_processing.extractors import (
        extract_embedded_office_documents,
    )

    assert extract_embedded_office_documents(b"%PDF-1.4 not a zip") == []


def test_xlsx_extractor_stops_at_row_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extractors_module, "_MAX_XLSX_ROWS_PER_SHEET", 2)
    monkeypatch.setattr(extractors_module, "_MAX_XLSX_TOTAL_ROWS", 2)

    bundle = extract_document(
        document_id="doc",
        filename="table.xlsx",
        mime_type="",
        content=_xlsx_bytes([["r1"], ["r2"], ["r3"]]),
    )

    assert len(bundle.evidence_blocks) == 1
    assert bundle.evidence_blocks[0].rows == [["r1"], ["r2"]]


def test_html_extractor_preserves_structured_locators_and_strips_active_content() -> None:
    content = """
    <!doctype html>
    <html>
      <head>
        <title>제동 시스템 시험 결과</title>
        <style>.secret { display: block }</style>
        <script src="https://example.invalid/execute.js">실행 금지 스크립트</script>
      </head>
      <body>
        <h1>요약</h1>
        <p>정상 제동 압력은 12 bar 입니다.</p>
        <ul><li><p>시험 장비 교정 완료</p></li></ul>
        <table>
          <tr><th>항목</th><th>결과</th></tr>
          <tr><td>응답 시간</td><td>120 ms</td></tr>
        </table>
        <iframe src="https://example.invalid/remote">원격 프레임 본문</iframe>
        <div hidden>숨김 본문</div>
        <noscript>대체 실행 본문</noscript>
      </body>
    </html>
    """.encode()

    bundle = extract_document(
        document_id="html-doc",
        filename="report.html",
        mime_type="text/html",
        content=content,
    )

    assert [block.locator_kind for block in bundle.evidence_blocks] == [
        "title",
        "heading",
        "paragraph",
        "list_item",
        "table_row",
        "table_row",
    ]
    assert bundle.evidence_blocks[2].section_path == "요약"
    assert bundle.evidence_blocks[4].rows == [["항목", "결과"]]
    assert bundle.evidence_blocks[5].locator_label == "Table 1, row 2"
    assert "실행 금지" not in bundle.normalized_text
    assert "원격 프레임" not in bundle.normalized_text
    assert "숨김 본문" not in bundle.normalized_text
    assert "대체 실행" not in bundle.normalized_text
    assert "example.invalid" not in bundle.normalized_text
    assert bundle.metadata["format"] == "html"
    assert bundle.metadata["charset"] == "utf-8-sig"
    assert bundle.metadata["table_row_count"] == 2
    assert bundle.metadata["truncated"] is False


def test_html_extractor_handles_cp949_declared_charset() -> None:
    content = (
        '<!doctype html><html><head><meta charset="euc-kr"></head>'
        "<body><h1>시험 결과</h1><p>냉각수 온도 정상 범위 확인 완료</p></body></html>"
    ).encode("cp949")

    bundle = extract_document(
        document_id="html-doc",
        filename="report.htm",
        mime_type="text/html",
        content=content,
    )

    assert "냉각수 온도 정상 범위 확인 완료" in bundle.normalized_text
    assert bundle.metadata["charset"] in {"cp949", "euc_kr"}


def test_html_stream_records_character_limit_truncation() -> None:
    content = (
        "<!doctype html><html><body><p>" + ("표 데이터 " * 100) + "</p></body></html>"
    ).encode()

    result = extract_html_stream(
        document_id="html-doc",
        stream=BytesIO(content),
        max_chars=40,
    )

    assert result.extracted_chars == 40
    assert result.truncated is True
    assert result.truncation_reason == "character_limit"
    assert result.metadata["input_bytes"] == len(content)
    assert result.peak_pending_chars <= 40
    assert result.evidence_blocks


def test_html_stream_bounds_malformed_nested_captures(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(html_extractor_module, "_MAX_HTML_NESTING", 3)
    content = (
        "<!doctype html><html><body>"
        "<p>outer<p>middle<p>inner<p>too deep"
        "</body></html>"
    ).encode()

    result = extract_html_stream(
        document_id="html-doc",
        stream=BytesIO(content),
        max_chars=32,
    )

    assert result.truncated is True
    assert result.truncation_reason == "structure_limit"
    assert result.peak_pending_chars <= 32
    assert result.extracted_chars <= 32


def test_html_stream_caps_table_cell_structure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(html_extractor_module, "_MAX_HTML_TABLE_CELLS", 3)
    content = (
        "<!doctype html><html><body><table><tr>"
        "<td></td><td></td><td></td><td></td>"
        "</tr></table></body></html>"
    ).encode()

    result = extract_html_stream(
        document_id="html-doc",
        stream=BytesIO(content),
    )

    assert result.truncated is True
    assert result.truncation_reason == "structure_limit"
    assert result.peak_pending_chars == 0
