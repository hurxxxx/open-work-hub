"""책임광물 조사표 엑셀 양식 작성.

레거시 ``write_excel`` 을 그대로 옮기되, 디스크 경로 대신 업로드된 양식 엑셀 바이트를
입력으로 받아 채워 넣은 결과를 바이트로 돌려준다.

열 매핑 (레거시 유지):
    B 차종 · C 최종품명 · D OEM품번 · E DCC품번 · F 연번
    G~U 트리 레벨(L1~L15) · V 부품명 · W 부품 품번 · X 제품명 · Y 재료명
    Z 화학물질명 · AA 광물명 · AO 비고(품번 없음)
"""

from __future__ import annotations

import io
from dataclasses import dataclass

import openpyxl

LEVEL_COLUMNS = ["G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U"]
_CLEAR_COLUMNS = [
    "B",
    "C",
    "D",
    "E",
    "F",
    "G",
    "H",
    "I",
    "J",
    "K",
    "L",
    "M",
    "N",
    "O",
    "P",
    "Q",
    "R",
    "S",
    "T",
    "U",
    "V",
    "W",
    "X",
    "Y",
    "Z",
    "AA",
    "AO",
]
_FIRST_DATA_ROW = 4


class ImdsWorkbookError(Exception):
    """양식 엑셀을 열거나 시트를 찾지 못했을 때 발생한다."""


class ImdsTemplateError(ImdsWorkbookError):
    """양식 엑셀 파일을 열 수 없을 때 발생한다."""


class ImdsSheetNotFoundError(ImdsWorkbookError):
    """양식 엑셀에서 요청한 시트를 찾지 못했을 때 발생한다."""


@dataclass
class ImdsMeta:
    """양식 머리 정보. 레거시 CLI 의 ``--car/--end-name/--oem/--dcc`` 에 해당한다."""

    car: str
    end_name: str
    oem: str
    dcc: str


def write_workbook(
    template: bytes,
    sheet: str,
    rows: list[dict],
    indices: list[int],
    meta: ImdsMeta,
) -> tuple[bytes, int]:
    """양식 바이트에 책임광물 가지를 채워 넣고 (결과 바이트, 기록 행수)를 돌려준다."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(template))
    except Exception as error:  # noqa: BLE001 - openpyxl raises broad errors
        raise ImdsTemplateError(f"cannot open template: {error}") from error
    if sheet not in wb.sheetnames:
        raise ImdsSheetNotFoundError(f"sheet not found: {sheet}")
    ws = wb[sheet]

    for ri in range(_FIRST_DATA_ROW, ws.max_row + 1):
        for col in _CLEAR_COLUMNS:
            ws[f"{col}{ri}"].value = None

    no = 1
    from .parser import get_mineral

    for idx in indices:
        r = rows[idx]
        er = _FIRST_DATA_ROW + (no - 1)
        ws[f"B{er}"] = meta.car
        ws[f"C{er}"] = meta.end_name
        ws[f"D{er}"] = meta.oem
        ws[f"E{er}"] = meta.dcc
        ws[f"F{er}"] = no
        lv = r["level"]
        if 1 <= lv <= 15:
            ws[f"{LEVEL_COLUMNS[lv - 1]}{er}"] = lv
        if r["type"] == "component":
            ws[f"V{er}"] = r["name"]
            if r["code"]:
                ws[f"W{er}"] = r["code"]
            else:
                ws[f"AO{er}"] = "품번 없음"
        elif r["type"] == "product":
            ws[f"X{er}"] = r["name"]
        elif r["type"] == "material":
            ws[f"Y{er}"] = r["name"]
        elif r["type"] == "chemical":
            ws[f"Z{er}"] = r["name"]
            mineral = get_mineral(r)
            if mineral:
                ws[f"AA{er}"] = mineral
        no += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), no - 1
