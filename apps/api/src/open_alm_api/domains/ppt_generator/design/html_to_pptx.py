"""HTML 슬라이드 덱 → 편집형 .pptx 하이브리드 변환기.

Claude 가 생성한 자체완결형 HTML(절대좌표 기반 Brandlogy 덱)을 헤드리스 Edge(CDP)로
렌더해 각 요소의 실제 픽셀 좌표·스타일을 추출한 뒤, python-pptx 로

  - 텍스트  → 편집 가능한 텍스트박스(폰트/색/정렬/줄바꿈 보존)
  - 배경/카드 → 네이티브 도형(라운드 사각형, 단색·그라데이션 채우기, 테두리)
  - SVG/차트 → 단독 렌더한 투명 PNG 이미지

로 재구성한다. 결과는 모양이 거의 동일하면서도 텍스트·도형을 편집할 수 있다.
(차트·다이어그램은 이미지로 남는다.)

이미지-only 변환(``_html_deck_to_pptx_bytes``)의 편집형 대안이며, 실패 시 호출측에서
이미지 변환으로 폴백하도록 예외를 던진다.
"""

from __future__ import annotations

import asyncio
import io
import json
import math
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W_CM = 27.0
SLIDE_H_CM = 16.75
EMU_PER_CM = 360000

# 렌더된 슬라이드의 각 요소를 추출하는 페이지 내 스크립트.
# 텍스트는 getClientRects 의 top 변화로 '실제 시각 줄바꿈'을 박제(text)하고,
# 연속 원문(raw)도 함께 담아 본문은 박스 폭에서 자연 재흐름시킨다.
_EXTRACT_JS = r"""
(async () => {
  await document.fonts.ready;
  await new Promise(r => setTimeout(r, 300));
  const slides = [...document.querySelectorAll('.slide')];
  const tops = slides.length ? slides : [...document.querySelectorAll('section')];
  const result = [];
  // 셀의 '논리 줄'(블록 자식/<br> 단위) 텍스트를 추출 — 표 셀 안의 불릿 div 들을 줄로 보존.
  const blockText = (el) => {
    let out = '';
    el.childNodes.forEach(n => {
      if (n.nodeType === 3) { out += n.textContent; return; }
      if (n.nodeName === 'BR') { out += '\n'; return; }
      if (n.nodeType === 1) {
        const d = getComputedStyle(n).display;
        const isBlock = (d==='block'||d==='flex'||d==='list-item'||d==='table'||d==='grid');
        if (isBlock && out && !out.endsWith('\n')) out += '\n';
        out += blockText(n);
        if (isBlock && !out.endsWith('\n')) out += '\n';
      }
    });
    return out;
  };
  const normLines = (s) => s.split('\n').map(v => v.replace(/[\s ]+/g, ' ').trim())
                            .filter(v => v.length).join('\n');
  for (const slide of tops) {
    const sb = slide.getBoundingClientRect();
    const scs = getComputedStyle(slide);
    const items = [{kind:'rect', x:0, y:0, w:sb.width, h:sb.height,
      bg: scs.backgroundColor, bgImage: scs.backgroundImage, radius:0,
      borderColor:'', borderW:0}];
    const skip = new Set();
    for (const el of slide.querySelectorAll('*')) {
      if (skip.has(el)) continue;
      const tag = el.tagName.toLowerCase();
      const cs = getComputedStyle(el);
      if (cs.display==='none' || cs.visibility==='hidden' || parseFloat(cs.opacity)===0) continue;
      const r = el.getBoundingClientRect();
      if (r.width < 1 || r.height < 1) continue;
      const x = r.left - sb.left, y = r.top - sb.top;
      if (tag==='svg' || tag==='canvas' || tag==='img') {
        let svg = '', src = '';
        if (tag==='svg') svg = new XMLSerializer().serializeToString(el);
        else if (tag==='img') src = el.currentSrc || el.src || '';
        items.push({kind:'image', x, y, w:r.width, h:r.height, svg, src});
        el.querySelectorAll('*').forEach(d => skip.add(d));
        continue;
      }
      if (tag==='table') {
        // HTML 표 → 네이티브 PPT 표로 변환. 셀별 기하/스타일/병합(span) 추출.
        const cells = [];
        for (const cell of el.querySelectorAll('td,th')) {
          const ccs = getComputedStyle(cell);
          if (ccs.display==='none' || ccs.visibility==='hidden') continue;
          const cr = cell.getBoundingClientRect();
          if (cr.width < 1 || cr.height < 1) continue;
          // 셀 배경: 자신이 투명하면 가장 가까운 배경색 조상에서 상속 탐색
          let cbg = ccs.backgroundColor;
          if (!cbg || cbg==='rgba(0, 0, 0, 0)' || cbg==='transparent') {
            let p = cell;
            while (p && (cbg==='rgba(0, 0, 0, 0)' || cbg==='transparent' || !cbg)) {
              p = p.parentElement; if (!p) break;
              cbg = getComputedStyle(p).backgroundColor;
            }
          }
          // 텍스트 색/굵기/크기: 첫 텍스트 노드를 가진 자손 기준(없으면 셀)
          let styleEl = cell;
          const tw = document.createTreeWalker(cell, NodeFilter.SHOW_TEXT);
          let tnode; while (tnode = tw.nextNode()) { if (tnode.textContent.trim()) { styleEl = tnode.parentElement || cell; break; } }
          const tcs = getComputedStyle(styleEl);
          cells.push({
            x: cr.left - sb.left, y: cr.top - sb.top, w: cr.width, h: cr.height,
            text: normLines(blockText(cell)),
            colspan: cell.colSpan || 1, rowspan: cell.rowSpan || 1,
            bg: cbg, color: tcs.color, size: parseFloat(tcs.fontSize),
            weight: parseInt(tcs.fontWeight) || (cell.tagName==='TH' ? 700 : 400),
            align: ccs.textAlign, valign: ccs.verticalAlign, fontFamily: tcs.fontFamily,
            lh: tcs.lineHeight,
            borderColor: ccs.borderTopColor || ccs.borderColor || '',
            borderW: parseFloat(ccs.borderTopWidth) || 0,
          });
        }
        if (cells.length) {
          items.push({kind:'table', x, y, w:r.width, h:r.height, cells});
          // 표 셀 안의 <img>/<svg> 는 네이티브 표로 안 들어가므로, 셀 위치에 picture 로 따로 올림.
          for (const im of el.querySelectorAll('img,svg')) {
            const ir = im.getBoundingClientRect();
            if (ir.width < 1 || ir.height < 1) continue;
            let isvg = '', isrc = '';
            if (im.tagName.toLowerCase()==='svg') isvg = new XMLSerializer().serializeToString(im);
            else isrc = im.currentSrc || im.src || '';
            if (!isvg && !isrc) continue;
            items.push({kind:'image', x: ir.left - sb.left, y: ir.top - sb.top,
              w: ir.width, h: ir.height, svg: isvg, src: isrc});
          }
          el.querySelectorAll('*').forEach(d => skip.add(d));
          continue;
        }
      }
      const hasBg = cs.backgroundColor && cs.backgroundColor!=='rgba(0, 0, 0, 0)' && cs.backgroundColor!=='transparent';
      const hasImg = cs.backgroundImage && cs.backgroundImage!=='none';
      const bw = parseFloat(cs.borderTopWidth)||0;
      const bColor = cs.borderTopColor;
      const bVisible = bw>0 && bColor && bColor!=='rgba(0, 0, 0, 0)';
      if (hasBg || hasImg || bVisible) {
        items.push({kind:'rect', x, y, w:r.width, h:r.height, bg:cs.backgroundColor,
          bgImage:cs.backgroundImage, radius: parseFloat(cs.borderTopLeftRadius)||0,
          borderColor: bVisible ? bColor : '', borderW: bVisible ? bw : 0});
      }
      const flines = [];
      const splitNode = (node) => {
        const s = node.textContent; if (!s) return;
        const range = document.createRange();
        let start = 0, lastTop = null;
        for (let i = 0; i < s.length; i++) {
          range.setStart(node, i); range.setEnd(node, i + 1);
          const rects = range.getClientRects();
          if (!rects.length) continue;
          const top = Math.round(rects[0].top);
          if (lastTop === null) lastTop = top;
          else if (Math.abs(top - lastTop) > 2) { flines.push(s.slice(start, i)); start = i; lastTop = top; }
        }
        flines.push(s.slice(start));
      };
      for (const n of el.childNodes) { if (n.nodeType === 3) splitNode(n); }
      const t = flines.map(v => v.replace(/[\s ]+/g, ' ').trim()).filter(v => v.length).join('\n');
      let raw = '';
      for (const n of el.childNodes) {
        if (n.nodeType === 3) raw += n.textContent;
        else if (n.nodeName === 'BR') raw += '\n';
      }
      raw = raw.replace(/[^\S\n]+/g, ' ').replace(/ *\n */g, '\n').trim();
      if (t) {
        items.push({kind:'text', x, y, w:r.width, h:r.height, text:t, raw:raw,
          color:cs.color, size: parseFloat(cs.fontSize), weight: parseInt(cs.fontWeight)||400,
          align: cs.textAlign, lh: cs.lineHeight, transform: cs.textTransform,
          fontFamily: cs.fontFamily, valign: cs.verticalAlign});
      }
    }
    result.push({w: sb.width, h: sb.height, items});
  }
  return {slides: result};
})()
"""

_ALIGN = {
    "left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER,
    "justify": PP_ALIGN.JUSTIFY, "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT,
}
_BROWSER_ENV_VAR = "OPEN_ALM_PPT_BROWSER_PATH"
_BROWSER_PATHS = (
    # Windows
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    # Linux (prod 서버) — PATH 에 없을 때(snap/opt 등) 대비한 절대경로 폴백
    "/usr/bin/chromium",
    "/usr/bin/chromium-browser",
    "/usr/bin/google-chrome",
    "/usr/bin/google-chrome-stable",
    "/usr/bin/microsoft-edge",
    "/usr/bin/microsoft-edge-stable",
    "/snap/bin/chromium",
    "/opt/google/chrome/chrome",
    # macOS
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
)
_BROWSER_COMMANDS = (
    "microsoft-edge",
    "microsoft-edge-stable",
    "msedge",
    "google-chrome",
    "google-chrome-stable",
    "chrome",
    "chromium",
    "chromium-browser",
)
_BROWSER_STDERR_LIMIT = 4_096


def _browser_preference(path: str) -> int:
    """Prefer system Chrome/Edge over Chromium wrappers, which are often Snap-backed."""

    name = os.path.basename(path).lower()
    normalized = path.replace("\\", "/").lower()
    if "chromium" in name or "/snap/" in normalized:
        return 30
    if "google-chrome" in name or name in {"chrome", "chrome.exe"}:
        return 10
    if "edge" in name or name == "msedge.exe":
        return 20
    return 25


def _find_browser() -> str | None:
    """finalize(HTML→PDF/기하 추출)용 헤드리스 브라우저(Edge/Chrome/Chromium) 경로.

    크로스플랫폼: dev finalize 는 Windows PC 워커가, prod finalize 는 Linux 서버
    워커가 처리한다. ① env override → ② 플랫폼별 표준 절대경로(_BROWSER_PATHS) →
    ③ PATH(shutil.which, _BROWSER_COMMANDS) 순으로 탐색한다.
    """
    configured = os.getenv(_BROWSER_ENV_VAR)
    if configured and os.path.exists(configured):
        return configured

    candidates: list[str] = []
    for path in _BROWSER_PATHS:
        if path and os.path.exists(path) and path not in candidates:
            candidates.append(path)
    for name in _BROWSER_COMMANDS:
        found = shutil.which(name)
        if found and found not in candidates:
            candidates.append(found)
    if not candidates:
        return None
    return min(enumerate(candidates), key=lambda item: (_browser_preference(item[1]), item[0]))[1]


def _bounded_browser_stderr(value) -> str:
    """Return a bounded, single-line stderr tail without ever logging HTML input."""

    if hasattr(value, "seek") and hasattr(value, "read"):
        value.flush()
        value.seek(0, os.SEEK_END)
        size = value.tell()
        value.seek(max(0, size - _BROWSER_STDERR_LIMIT))
        raw = value.read(_BROWSER_STDERR_LIMIT)
    elif isinstance(value, str):
        raw = value.encode("utf-8", errors="replace")[-_BROWSER_STDERR_LIMIT:]
    else:
        raw = bytes(value or b"")[-_BROWSER_STDERR_LIMIT:]
    return " ".join(raw.decode("utf-8", errors="replace").split())


def _browser_failure_message(operation: str, returncode: int, stderr) -> str:
    detail = _bounded_browser_stderr(stderr)
    suffix = f": {detail}" if detail else ""
    return f"{operation} 실패(exit={returncode}){suffix}"


def _headless_browser_env() -> dict[str, str]:
    env = os.environ.copy()
    if sys.platform != "win32":
        # The long-lived systemd user manager also runs without a graphical
        # login.  Inheriting its session bus makes headless Chrome activate
        # xdg-desktop-portal, whose GTK backend then waits for a DISPLAY and
        # fails the otherwise healthy user manager. PPT rendering does not
        # use desktop portals, so isolate only the browser child from D-Bus.
        env["DBUS_SESSION_BUS_ADDRESS"] = "unix:path=/dev/null"
    return env


def _run_browser(command: list[str], *, timeout: int, operation: str):
    try:
        result = subprocess.run(
            command,
            timeout=timeout,
            check=False,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
            env=_headless_browser_env(),
        )
    except subprocess.TimeoutExpired as exc:
        detail = _bounded_browser_stderr(exc.stderr)
        suffix = f": {detail}" if detail else ""
        raise RuntimeError(f"{operation} 시간 초과({timeout}초){suffix}") from exc
    if result.returncode != 0:
        raise RuntimeError(_browser_failure_message(operation, result.returncode, result.stderr))
    return result


def _free_port() -> int:
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


_GENERIC_FONTS = {"sans-serif", "serif", "monospace", "system-ui", "-apple-system",
                  "cursive", "fantasy", "ui-sans-serif", "ui-serif", "ui-monospace"}


def _first_font_family(font_family: str | None) -> str | None:
    """computed font-family 스택에서 첫 실제 폰트명을 추출(따옴표/일반패밀리 제외)."""
    if not font_family:
        return None
    for part in font_family.split(","):
        name = part.strip().strip('"').strip("'").strip()
        if name and name.lower() not in _GENERIC_FONTS:
            return name
    return None


def _parse_rgb(s: str | None):
    if not s:
        return None
    m = re.match(r"rgba?\(([^)]+)\)", s)
    if not m:
        return None
    parts = [p.strip() for p in m.group(1).split(",")]
    try:
        r, g, b = int(float(parts[0])), int(float(parts[1])), int(float(parts[2]))
    except (ValueError, IndexError):
        return None
    a = float(parts[3]) if len(parts) > 3 else 1.0
    return (r, g, b, a)


def _parse_gradient(bg: str):
    m = re.search(r"linear-gradient\((.*)\)", bg, re.S)
    if not m:
        return None
    inner = m.group(1)
    angle = 90.0
    am = re.match(r"\s*(-?[\d.]+)deg\s*,", inner)
    if am:
        angle = float(am.group(1))
        inner = inner[am.end():]
    elif inner.lstrip().startswith("to ") and "," in inner:
        head, inner = inner.split(",", 1)
        # CSS 키워드 방향 → CSS 각도(0deg=위쪽, 시계방향). 모서리는 정사각형 근사.
        kw = " ".join(head.strip()[3:].split()).lower()  # 'to ' 제거 후 정규화
        angle = {
            "top": 0.0, "right": 90.0, "bottom": 180.0, "left": 270.0,
            "top right": 45.0, "right top": 45.0,
            "bottom right": 135.0, "right bottom": 135.0,
            "bottom left": 225.0, "left bottom": 225.0,
            "top left": 315.0, "left top": 315.0,
        }.get(kw, 90.0)
    stops = []
    for cm in re.finditer(r"(rgba?\([^)]+\))\s*([\d.]+%)?", inner):
        rgb = _parse_rgb(cm.group(1))
        pos = cm.group(2)
        if rgb:
            stops.append((float(pos[:-1]) / 100.0 if pos else None, rgb))
    if len(stops) < 2:
        return None
    n = len(stops)
    for i, (p, c) in enumerate(stops):
        if p is None:
            stops[i] = (i / (n - 1), c)
    return angle, stops


def _kill_proc_tree(proc: subprocess.Popen) -> None:
    """런처 PID 뿐 아니라 ``--headless=new`` 가 띄운 렌더러/GPU 자식까지 정리한다.

    Windows 에서 launcher 만 terminate 하면 분리된 msedge.exe 자식들이 살아남아
    누적되고 임시 프로필/HTML 디렉터리를 잠근다. 프로세스 트리를 통째로 종료한다.
    """
    if proc.poll() is not None:
        return
    try:
        if sys.platform == "win32":
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", str(proc.pid)],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False,
            )
        else:
            proc.terminate()
    except Exception:
        try:
            proc.kill()
        except Exception:
            pass
    try:
        proc.wait(timeout=5)
    except Exception:
        pass


async def _extract_geometry(html_path: str, browser: str) -> dict:
    import websockets

    port = _free_port()
    with tempfile.TemporaryDirectory(prefix="pptbrowser_") as profile_dir:
        with tempfile.TemporaryFile(mode="w+b") as stderr_file:
            proc = subprocess.Popen(
                [
                    browser,
                    "--headless=new",
                    "--disable-gpu",
                    "--no-sandbox",
                    "--hide-scrollbars",
                    f"--user-data-dir={profile_dir}",
                    f"--remote-debugging-port={port}",
                    "--remote-allow-origins=*",
                    "--window-size=1280,900",
                    "file:///" + html_path.replace("\\", "/"),
                ],
                stdout=subprocess.DEVNULL,
                stderr=stderr_file,
                env=_headless_browser_env(),
            )
            try:
                ws_url = None
                for _ in range(80):
                    returncode = proc.poll()
                    if returncode is not None:
                        raise RuntimeError(
                            _browser_failure_message("CDP 브라우저 시작", returncode, stderr_file)
                        )
                    try:
                        data = json.load(
                            urllib.request.urlopen(f"http://127.0.0.1:{port}/json", timeout=1)
                        )
                        for target in data:
                            if target.get("type") == "page" and target.get(
                                "webSocketDebuggerUrl"
                            ):
                                ws_url = target["webSocketDebuggerUrl"]
                                break
                        if ws_url:
                            break
                    except Exception:
                        pass
                    time.sleep(0.25)
                if not ws_url:
                    _kill_proc_tree(proc)
                    detail = _bounded_browser_stderr(stderr_file)
                    suffix = f": {detail}" if detail else ""
                    raise RuntimeError(f"CDP 페이지 타깃을 찾지 못했습니다{suffix}")
                _id = 0
                async with websockets.connect(ws_url, max_size=None) as ws:

                    async def cmd(method, params=None):
                        nonlocal _id
                        _id += 1
                        mid = _id
                        await ws.send(
                            json.dumps({"id": mid, "method": method, "params": params or {}})
                        )
                        while True:
                            msg = json.loads(await ws.recv())
                            if msg.get("id") == mid:
                                return msg

                    await cmd("Page.enable")
                    await cmd("Runtime.enable")
                    await asyncio.sleep(1.2)
                    res = await cmd(
                        "Runtime.evaluate",
                        {"expression": _EXTRACT_JS, "returnByValue": True, "awaitPromise": True},
                    )
                    try:
                        return res["result"]["result"]["value"]
                    except (KeyError, TypeError) as e:
                        raise RuntimeError(f"기하정보 추출 실패: {json.dumps(res)[:400]}") from e
            finally:
                _kill_proc_tree(proc)


def _render_svg_png(svg: str, w: float, h: float, browser: str) -> io.BytesIO | None:
    if not svg.strip():
        return None
    with tempfile.TemporaryDirectory(prefix="pptsvg_") as directory:
        html = (
            "<!DOCTYPE html><html><head><meta charset='utf-8'>"
            "<style>html,body{margin:0;padding:0;}svg{display:block;}</style></head>"
            f"<body>{svg}</body></html>"
        )
        html_path = os.path.join(directory, "s.html")
        png_path = os.path.join(directory, "s.png")
        profile_dir = os.path.join(directory, "browser-profile")
        with open(html_path, "w", encoding="utf-8") as file:
            file.write(html)
        _run_browser(
            [
                browser,
                "--headless=new",
                "--disable-gpu",
                "--no-sandbox",
                f"--user-data-dir={profile_dir}",
                "--force-device-scale-factor=2",
                "--default-background-color=00000000",
                f"--window-size={int(math.ceil(w))},{int(math.ceil(h))}",
                f"--screenshot={png_path}",
                "file:///" + html_path.replace("\\", "/"),
            ],
            timeout=60,
            operation="SVG PNG 변환",
        )
        if not os.path.exists(png_path):
            return None
        with open(png_path, "rb") as file:
            return io.BytesIO(file.read())


def _decode_img_src(src: str | None):
    """data:image/...;base64,... → BytesIO. 그 외(http/file)는 None(보안상 건너뜀)."""
    if not src or not src.startswith("data:image"):
        return None
    try:
        b64 = src.split(",", 1)[1]
        import base64

        return io.BytesIO(base64.b64decode(b64))
    except Exception:
        return None


def _set_alpha(color, alpha: float) -> None:
    srgb = color._xFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))


class _Builder:
    def __init__(self, sw_px: float, sh_px: float, browser: str,
                 sw_cm: float = SLIDE_W_CM, sh_cm: float = SLIDE_H_CM,
                 ox_cm: float = 0.0, oy_cm: float = 0.0):
        self.sw = sw_px
        self.sh = sh_px
        self.sw_cm = sw_cm   # 추출 캔버스(px) 가 대응하는 물리 폭(cm)
        self.sh_cm = sh_cm
        self.ox_cm = ox_cm   # 배치 원점 오프셋(cm) — 본문 칸에 끼워넣을 때 사용
        self.oy_cm = oy_cm
        self.browser = browser
        self.font_scale = (sw_cm / sw_px) / 2.54 * 72.0

    # EX/EY = 크기 환산(px→EMU, 오프셋 없음)
    def EX(self, px):
        return Emu(int(px * (self.sw_cm / self.sw) * EMU_PER_CM))

    def EY(self, px):
        return Emu(int(px * (self.sh_cm / self.sh) * EMU_PER_CM))

    # PX/PY = 위치 환산(원점 오프셋 포함). ox=oy=0 이면 EX/EY 와 동일.
    def PX(self, px):
        return Emu(int((self.ox_cm + px * (self.sw_cm / self.sw)) * EMU_PER_CM))

    def PY(self, px):
        return Emu(int((self.oy_cm + px * (self.sh_cm / self.sh)) * EMU_PER_CM))

    def add_rect(self, slide, it):
        w = max(1, it["w"])
        h = max(1, it["h"])
        radius = it.get("radius", 0) or 0
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius >= 2 else MSO_SHAPE.RECTANGLE
        sp = slide.shapes.add_shape(kind, self.PX(it["x"]), self.PY(it["y"]), self.EX(w), self.EY(h))
        sp.shadow.inherit = False
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                sp.adjustments[0] = max(0.0, min(0.5, radius / min(w, h)))
            except Exception:
                pass
        bgimg = it.get("bgImage") or "none"
        grad = _parse_gradient(bgimg) if "gradient" in bgimg else None
        rgb = _parse_rgb(it.get("bg"))
        if grad:
            try:
                sp.fill.gradient()
                gs = sp.fill.gradient_stops
                first, last = grad[1][0][1], grad[1][-1][1]
                gs[0].position = 0.0
                gs[0].color.rgb = RGBColor(*first[:3])
                gs[1].position = 1.0
                gs[1].color.rgb = RGBColor(*last[:3])
                sp.fill.gradient_angle = (grad[0] - 90) % 360
            except Exception:
                sp.fill.solid()
                sp.fill.fore_color.rgb = RGBColor(*grad[1][0][1][:3])
        elif rgb and rgb[3] > 0:
            sp.fill.solid()
            sp.fill.fore_color.rgb = RGBColor(rgb[0], rgb[1], rgb[2])
            if rgb[3] < 1:
                _set_alpha(sp.fill.fore_color, rgb[3])
        else:
            sp.fill.background()
        bw = it.get("borderW", 0) or 0
        bc = _parse_rgb(it.get("borderColor"))
        if bw > 0 and bc and bc[3] > 0:
            sp.line.color.rgb = RGBColor(bc[0], bc[1], bc[2])
            sp.line.width = self.EY(bw)
        else:
            sp.line.fill.background()

    def _style_run(self, run, it):
        f = run.font
        f.size = Pt(max(1.0, it["size"] * self.font_scale))
        f.bold = (it.get("weight", 400) or 400) >= 600
        face = _first_font_family(it.get("fontFamily")) or "Pretendard"
        f.name = face
        rPr = run._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            e = rPr.find(qn(tag))
            if e is None:
                e = rPr.makeelement(qn(tag), {})
                rPr.append(e)
            e.set("typeface", face)
        c = _parse_rgb(it.get("color"))
        if c:
            f.color.rgb = RGBColor(c[0], c[1], c[2])

    def add_text(self, slide, it):
        w = max(1, it["w"])
        h = max(1, it["h"])
        txt = it.get("text") or ""
        if it.get("transform") == "uppercase":
            txt = txt.upper()
        size = it.get("size", 12) or 12
        vlines = txt.split("\n")
        # 본문(작은 글자) 여러 줄: 연속 원문(raw)을 박스 폭에서 자연 재흐름. 큰 제목/단일
        # 줄: 박제한 시각 줄바꿈을 그대로 두고 word_wrap 끔(대체폰트 과다 줄바꿈 방지).
        body = (len(vlines) > 1) and (size < 22)
        if body:
            raw = it.get("raw") or txt.replace("\n", " ")
            if it.get("transform") == "uppercase":
                raw = raw.upper()
            lines = raw.split("\n")
        else:
            lines = vlines
        tb = slide.shapes.add_textbox(self.PX(it["x"]), self.PY(it["y"]), self.EX(w), self.EY(h))
        tf = tb.text_frame
        tf.word_wrap = bool(body)
        _va = (it.get("valign") or "").lower()
        tf.vertical_anchor = (
            MSO_ANCHOR.MIDDLE if _va == "middle"
            else MSO_ANCHOR.BOTTOM if _va == "bottom"
            else MSO_ANCHOR.TOP
        )
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, 0)
        try:
            tf.auto_size = None
        except Exception:
            pass
        align = _ALIGN.get((it.get("align") or "left"), PP_ALIGN.LEFT)
        lh = it.get("lh")
        line_pt = None
        try:
            if lh and lh.endswith("px"):
                line_pt = Pt(float(lh[:-2]) * self.font_scale)
        except Exception:
            line_pt = None
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if line_pt is not None:
                p.line_spacing = line_pt
            run = p.add_run()
            run.text = ln
            self._style_run(run, it)

    def add_table(self, slide, it):
        """HTML 표 추출 결과(cells)를 네이티브 PowerPoint 표로 — 셀 병합·테두리 포함."""
        cells = it.get("cells") or []
        if not cells:
            raise ValueError("empty table")
        tol = 3.0

        def cluster(vals):
            out = []
            for v in sorted(vals):
                if not out or abs(v - out[-1]) > tol:
                    out.append(v)
            return out

        xs = cluster([c["x"] for c in cells] + [c["x"] + c["w"] for c in cells])
        ys = cluster([c["y"] for c in cells] + [c["y"] + c["h"] for c in cells])
        ncol, nrow = len(xs) - 1, len(ys) - 1
        if ncol < 1 or nrow < 1:
            raise ValueError("bad grid")

        def nidx(v, bounds):
            best, bd = 0, 1e9
            for i, b in enumerate(bounds):
                d = abs(b - v)
                if d < bd:
                    bd, best = d, i
            return best

        gf = slide.shapes.add_table(
            nrow, ncol, self.PX(xs[0]), self.PY(ys[0]),
            self.EX(xs[-1] - xs[0]), self.EY(ys[-1] - ys[0]))
        tbl = gf.table
        tbl.first_row = tbl.first_col = tbl.horz_banding = tbl.vert_banding = False
        # 표 스타일 제거(No Style, No Grid) → 셀 채움/테두리를 우리가 직접 지정.
        tblPr = tbl._tbl.tblPr
        for sid in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(sid)
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
        tblPr.append(sid)
        for ci in range(ncol):
            tbl.columns[ci].width = self.EX(xs[ci + 1] - xs[ci])
        for ri in range(nrow):
            tbl.rows[ri].height = self.EY(ys[ri + 1] - ys[ri])

        for c in cells:
            c0, c1 = nidx(c["x"], xs), nidx(c["x"] + c["w"], xs)
            r0, r1 = nidx(c["y"], ys), nidx(c["y"] + c["h"], ys)
            c0 = min(c0, ncol - 1)
            r0 = min(r0, nrow - 1)
            cspan = max(1, min(ncol - c0, c1 - c0))
            rspan = max(1, min(nrow - r0, r1 - r0))
            cell = tbl.cell(r0, c0)
            if rspan > 1 or cspan > 1:
                try:
                    cell.merge(tbl.cell(r0 + rspan - 1, c0 + cspan - 1))
                except Exception:
                    pass
            rgb = _parse_rgb(c.get("bg"))
            if rgb and rgb[3] > 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(rgb[0], rgb[1], rgb[2])
            else:
                cell.fill.background()
            cell.margin_left = self.EX(7)
            cell.margin_right = self.EX(7)
            cell.margin_top = self.EY(2)
            cell.margin_bottom = self.EY(2)
            _va = (c.get("valign") or "").lower()
            cell.vertical_anchor = (
                MSO_ANCHOR.BOTTOM if _va == "bottom"
                else MSO_ANCHOR.TOP if _va == "top"
                else MSO_ANCHOR.MIDDLE
            )
            tf = cell.text_frame
            tf.word_wrap = True
            try:
                tf.auto_size = None
            except Exception:
                pass
            align = _ALIGN.get((c.get("align") or "left"), PP_ALIGN.LEFT)
            rstyle = {
                "size": c.get("size", 10) or 10, "weight": c.get("weight", 400),
                "fontFamily": c.get("fontFamily"), "color": c.get("color"),
            }
            _clh = c.get("lh")
            _cline_pt = None
            try:
                if _clh and _clh.endswith("px"):
                    _cline_pt = Pt(float(_clh[:-2]) * self.font_scale)
            except Exception:
                _cline_pt = None
            lines = (c.get("text") or "").split("\n") or [""]
            for i, ln in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = align
                if _cline_pt is not None:
                    p.line_spacing = _cline_pt
                run = p.add_run()
                run.text = ln
                self._style_run(run, rstyle)
            bw = c.get("borderW", 0) or 0
            bc = _parse_rgb(c.get("borderColor"))
            if bw > 0 and bc:
                hexc = "%02X%02X%02X" % (bc[0], bc[1], bc[2])
                w_emu = int(self.EY(max(0.75, bw)))
                tcPr = cell._tc.get_or_add_tcPr()
                for tag in ("lnB", "lnT", "lnR", "lnL"):
                    old = tcPr.find(qn(f"a:{tag}"))
                    if old is not None:
                        tcPr.remove(old)
                    ln_el = tcPr.makeelement(
                        qn(f"a:{tag}"),
                        {"w": str(w_emu), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
                    sf = tcPr.makeelement(qn("a:solidFill"), {})
                    clr = tcPr.makeelement(qn("a:srgbClr"), {"val": hexc})
                    sf.append(clr)
                    ln_el.append(sf)
                    tcPr.insert(0, ln_el)

    def place_items(self, slide, s: dict) -> None:
        """추출된 한 슬라이드의 요소들을 (이미 만들어진) slide 에 배치한다."""
        texts = []
        for it in s.get("items", []):
            k = it.get("kind")
            if k == "rect":
                try:
                    self.add_rect(slide, it)
                except Exception:
                    pass
            elif k == "table":
                try:
                    self.add_table(slide, it)
                except Exception:
                    pass
            elif k == "image":
                try:
                    svg = it.get("svg", "")
                    if svg.strip():
                        png = _render_svg_png(svg, it["w"], it["h"], self.browser)
                        if png:
                            slide.shapes.add_picture(
                                png, self.PX(it["x"]), self.PY(it["y"]),
                                self.EX(it["w"]), self.EY(it["h"]))
                    else:
                        img_io = _decode_img_src(it.get("src", ""))
                        if img_io is not None:
                            slide.shapes.add_picture(
                                img_io, self.PX(it["x"]), self.PY(it["y"]),
                                self.EX(it["w"]), self.EY(it["h"]))
                except Exception:
                    pass
            elif k == "text":
                texts.append(it)
        for it in texts:
            try:
                self.add_text(slide, it)
            except Exception:
                pass

    def build(self, geom: dict) -> bytes:
        prs = Presentation()
        prs.slide_width = Emu(int(self.sw_cm * EMU_PER_CM))
        prs.slide_height = Emu(int(self.sh_cm * EMU_PER_CM))
        blank = prs.slide_layouts[6]
        for s in geom.get("slides", []):
            slide = prs.slides.add_slide(blank)
            self.place_items(slide, s)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()


def html_deck_to_editable_pptx_bytes(html: str, slide_w_px: int = 1280,
                                     slide_h_px: int = 794,
                                     slide_w_cm: float = SLIDE_W_CM,
                                     slide_h_cm: float = SLIDE_H_CM) -> bytes:
    """HTML 덱 → 편집형 .pptx. 실패 시 예외(호출측이 이미지 변환으로 폴백).

    slide_w_cm/slide_h_cm: 출력 pptx 캔버스 물리 크기(기본 27×16.75cm, Open ALM 양식은 27.517×19.05).
    """
    browser = _find_browser()
    if browser is None:
        raise RuntimeError("변환용 브라우저(Edge/Chrome)를 찾을 수 없습니다.")
    with tempfile.TemporaryDirectory(prefix="pptedit_") as directory:
        html_path = os.path.join(directory, "deck.html")
        with open(html_path, "w", encoding="utf-8") as file:
            file.write(html)
        geom = asyncio.run(_extract_geometry(html_path, browser))
    if not geom or not geom.get("slides"):
        raise RuntimeError("추출된 슬라이드가 없습니다.")
    return _Builder(slide_w_px, slide_h_px, browser, slide_w_cm, slide_h_cm).build(geom)


def corporate_deck_to_pptx_bytes(
    cover: dict,
    body_sections: list[dict],
    body_style: str = "",
    *,
    logo_path: str | None = None,
    body_w_px: int = 1244,
    body_h_px: int = 754,
) -> bytes:
    """Open ALM 양식 .pptx — 표지/본문 틀은 builders_a4 로 네이티브 작성하고, 본문 콘텐츠만
    html_to_pptx 로 변환해 본문 칸(cm 오프셋)에 끼워 넣는다.

    배지 가운데 정렬·대각선 DCC 워터마크·로고·날짜 등 틀은 네이티브라 변환 아티팩트가 없다.
    """
    from pptx.util import Cm

    from open_alm_api.domains.ppt_generator.design import builders_a4 as b4
    from open_alm_api.domains.ppt_generator.design import corporate_frame as df

    browser = _find_browser()
    if browser is None:
        raise RuntimeError("변환용 브라우저(Edge/Chrome)를 찾을 수 없습니다.")

    # 본문 전용 덱(틀 없음) 한 번에 렌더 → 슬라이드별 요소 추출
    body_deck = df.compose_body_only_deck(body_sections, body_style)
    with tempfile.TemporaryDirectory(prefix="pptcorporate_") as directory:
        html_path = os.path.join(directory, "body.html")
        with open(html_path, "w", encoding="utf-8") as file:
            file.write(body_deck)
        geom = asyncio.run(_extract_geometry(html_path, browser))
    slides_geom = (geom or {}).get("slides", [])

    logo = logo_path if (logo_path and os.path.exists(logo_path)) else None

    prs = Presentation()
    prs.slide_width = Cm(df.SLIDE_W_CM)
    prs.slide_height = Cm(df.SLIDE_H_CM)

    # DCC 사선 워터마크는 본문마다 그리지 않고 슬라이드 마스터에 1회만 올린다
    # → 모든 본문 슬라이드가 배경으로 상속(선택·이동 불가한 진짜 워터마크). 표지에선 가린다.
    try:
        b4.add_dcc_watermark_to_master(prs)
    except Exception:
        pass

    # 1) 표지 (네이티브)
    cover_data = {"TITLE": cover.get("title") or "제목 미정"}
    if cover.get("author"):
        cover_data["AUTHOR"] = cover["author"]
    if cover.get("version"):
        cover_data["VERSION"] = cover["version"]
    try:
        b4.build_a4_cover_slide(prs, cover_data, logo_path=logo)
        # 표지엔 본문용 DCC 워터마크가 없어야 하므로 마스터 도형 상속을 끈다.
        if len(prs.slides) > 0:
            b4.set_show_master_shapes(prs.slides[-1], False)
    except Exception:
        pass

    # 2) 본문 슬라이드 — 네이티브 틀 + 본문 칸(오프셋) 콘텐츠
    bld = _Builder(
        body_w_px, body_h_px, browser,
        sw_cm=df.BODY_W_CM, sh_cm=df.BODY_H_CM,
        ox_cm=df.BODY_X_CM, oy_cm=df.BODY_Y_CM,
    )
    blank = prs.slide_layouts[6]
    for i, sg in enumerate(slides_geom):
        slide = prs.slides.add_slide(blank)
        title = body_sections[i].get("title") if i < len(body_sections) else ""
        try:
            # 워터마크는 마스터에 위임(watermark=False), 제목은 24pt 고정.
            b4.apply_body_chrome(slide, header_title=title or "", logo_path=logo,
                                 watermark=False, title_fixed_pt=24)
        except Exception:
            pass
        bld.place_items(slide, sg)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
