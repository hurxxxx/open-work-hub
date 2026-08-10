"""특허비용 품의(approval) 본문 HTML 생성기.

그룹웨어 전자결재 본문에 붙여넣을 HTML을 만든다. 데이터 출처:
* 청구서(cost lines) — 1.총비용 / 3.비용내역 상세표 / 헤더 건수
* 업로드된 6번 파일(건수·금액 정리) — 2.국내외 특허비용 정리(누적 월별 표)
* 고정 템플릿 — 4.의뢰업체 / 5.주기 / 사내표준
* 외부 예산값(선택) — 5.예산표

핵심 규칙(사용자 확인):
* 공급가액 = line.supply_amount (부가세별도, 관납료 포함)
* 비용내역 헤더 "N건" = 해당 카테고리에서 부가세(vat>0)가 붙은 항목 수
* 총비용 = 전 카테고리 공급가액 합 (부가세별도)
"""

from __future__ import annotations

from html import escape
from io import BytesIO
from typing import Any

import openpyxl

from open_alm_api.domains.patent_automation.invoice_pipeline import (
    SECTION_OVERSEAS,
    SUMMARY_SECTION_ORDER,
)
from open_alm_api.domains.patent_automation.summary_xlsx import _overseas_desc, korean_title

# 비용내역(3번) 카테고리 라벨 + 상세표 컬럼 유형.
_CATEGORY_LABEL = {
    "국내출원": "국내출원 비용",
    "심사청구": "심사청구 비용",
    "의견제출": "의견제출 비용",
    "재심사": "재심사 비용",
    "특허등록": "특허등록 비용",
    "연차료": "연차료 비용",
    "해외": "해외 비용",
}
# 상세표 컬럼 유형: domestic / registration / annuity / overseas
_CATEGORY_COLS = {
    "국내출원": "domestic",
    "심사청구": "domestic",
    "의견제출": "domestic",
    "재심사": "domestic",
    "특허등록": "registration",
    "연차료": "annuity",
    "해외": "overseas",
}
# 2번 누적표 카테고리 순서(6번 파일 열 순서와 동일).
_CA_ORDER = ["국내출원", "심사청구", "의견제출", "재심사", "특허등록", "연차료", "해외"]
_CA_HEADER = [
    "출원비용",
    "심사청구비용",
    "의견제출비용",
    "재심사청구비용",
    "등록비용",
    "등록유지비용",
    "해외비용",
]

_DEFAULT_VENDOR = "한라 특허법인 (TEL. 02-567-0131)"
_DOC_STANDARD = "DCS-BB-A110"  # 사내표준 및 관련근거 (고정)
# 모든 테두리는 1px 실선(black)으로 통일 — 점선/입체 테두리 방지(border-collapse + 셀별 solid).
_STYLE_TD = "border:1px solid #000; padding:1px 4px; font-size:13px; font-family:'Times New Roman',나눔고딕;"
_STYLE_TABLE = (
    "border-collapse:collapse; border:1px solid #000; font-size:13px; font-family:나눔고딕;"
)
_TABLE_WIDTH = 700  # 사내표준 표 기준 너비(px). 2번 누적표만 컬럼이 많아 예외(가변 너비).
# 너비를 700px로 통일한 표(비용내역 상세·예산·사내표준). 2번 누적표는 _STYLE_TABLE(가변) 사용.
# table-layout:auto 유지 → 명칭 등 긴 열은 내용에 맞춰 넓게, 긴 텍스트는 줄바꿈으로 700px 유지.
_STYLE_TABLE_FIXED = f"{_STYLE_TABLE} width:{_TABLE_WIDTH}px;"


def _won(n: int | float | None) -> str:
    return f"{int(n or 0):,}"


def _date(value: Any) -> str:
    """'2023-02-10' / '2023.02.10' → '2023.02.10'; 없으면 '-'."""
    s = str(value or "").strip()
    if not s:
        return "-"
    return s.replace("-", ".")


def _td(value: str, *, align: str = "center", colspan: int = 1, bold: bool = False) -> str:
    span = f' colspan="{colspan}"' if colspan > 1 else ""
    weight = "font-weight:bold;" if bold else ""
    return f'<td{span} style="{_STYLE_TD}text-align:{align};{weight}">{value}</td>'


def parse_count_amount_xlsx(content: bytes) -> dict[str, Any]:
    """6번 '국내외 특허비용(건수,금액정리)' 파일을 파싱.
    build_count_amount_summary 레이아웃: 행7~18 월별(B=월, C/D~O/P=건수/금액, Q=월별비용), 행19 합계.
    반환: {year, date_label, rows:[(월라벨, [(건수,금액)x7], 월별비용)], total:([(건수,금액)x7], 합계)}.
    데이터가 있는 월만(월별비용>0 또는 건수합>0) 반환한다.
    """
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb.worksheets[0]
    year = str(ws.cell(3, 3).value or "").replace("년", "").strip()
    date_label = str(ws.cell(2, 17).value or "").strip()
    # 카테고리별 (건수 열, 금액 열) — 6번 파일 고정 열.
    cols = [(3, 4), (5, 6), (7, 8), (9, 10), (11, 12), (13, 14), (15, 16)]

    def read_row(r: int) -> tuple[list[tuple[int, int]], int]:
        cats = [(int(ws.cell(r, cc).value or 0), int(ws.cell(r, ac).value or 0)) for cc, ac in cols]
        return cats, int(ws.cell(r, 17).value or 0)

    rows = []
    for r in range(7, 19):  # 1~12월
        label = str(ws.cell(r, 2).value or "").strip()
        cats, monthly = read_row(r)
        if monthly or any(c or a for c, a in cats):
            rows.append((label, cats, monthly))
    total_cats, total_sum = read_row(19)
    return {"year": year, "date_label": date_label, "rows": rows, "total": (total_cats, total_sum)}


def _count_amount_table_html(ca: dict[str, Any]) -> str:
    """2번 누적표 HTML."""
    parts = [f'<table cellspacing="0" cellpadding="0" style="{_STYLE_TABLE}">']
    # 헤더 2행
    parts.append("<tr>")
    parts.append(_td("", colspan=1))
    for h in _CA_HEADER:
        parts.append(_td(h, colspan=2, bold=True))
    parts.append(_td("월별비용", bold=True))
    parts.append("</tr>")
    parts.append("<tr>")
    parts.append(_td("", colspan=1))
    for _ in _CA_HEADER:
        parts.append(_td("건수", bold=True))
        parts.append(_td("금액", bold=True))
    parts.append(_td("", colspan=1))
    parts.append("</tr>")

    def row_html(
        label: str, cats: list[tuple[int, int]], monthly: int, *, bold: bool = False
    ) -> str:
        cells = [_td(label, bold=True)]
        for cnt, amt in cats:
            cells.append(_td(_won(cnt) if (cnt or amt) else "", align="center", bold=bold))
            cells.append(_td(_won(amt) if (cnt or amt) else "", align="right", bold=bold))
        cells.append(_td(_won(monthly) if monthly else "", align="right", bold=bold))
        return "<tr>" + "".join(cells) + "</tr>"

    for label, cats, monthly in ca["rows"]:
        parts.append(row_html(label, cats, monthly))
    tcats, tsum = ca["total"]
    parts.append(row_html("합계", tcats, tsum, bold=True))
    parts.append("</table>")
    return "".join(parts)


def _detail_columns(kind: str) -> list[str]:
    if kind == "registration":
        return [
            "순번",
            "발명의 명칭",
            "출원일",
            "출원번호",
            "등록일",
            "등록번호",
            "발명자",
            "공급가액(원)",
            "업체",
        ]
    if kind == "annuity":
        return [
            "순번",
            "발명의 명칭",
            "출원일",
            "출원번호",
            "등록일",
            "등록번호",
            "연차",
            "공급가액(원)",
            "업체",
        ]
    if kind == "overseas":
        return [
            "순번",
            "발명의 명칭 (업무 내용)",
            "출원일",
            "출원번호",
            "등록일",
            "등록번호",
            "발명자",
            "공급가액(원)",
            "업체",
        ]
    return ["순번", "발명의 명칭", "출원일", "출원번호", "발명자", "공급가액(원)", "업체"]


def _detail_row_cells(kind: str, seq: int, line: Any) -> list[str]:
    title = escape(korean_title(line.title))
    if kind == "overseas":
        # _overseas_desc는 이미 '[...]' 형태로 반환하므로 추가 대괄호를 붙이지 않는다.
        title = f"{title} {escape(_overseas_desc(line))}"
    inventors = escape(line.inventors or "")
    vendor = escape(line.vendor or "")
    appno = escape(line.application_no or "-")
    regno = escape(line.registration_no or "-")
    supply = _won(line.supply_amount)
    base_front = [
        _td(str(seq)),
        _td(title, align="left"),
        _td(_date(line.application_date)),
        _td(appno),
    ]
    if kind == "domestic":
        return base_front + [_td(inventors), _td(supply, align="right"), _td(vendor)]
    reg = [_td(_date(line.registration_date)), _td(regno)]
    if kind == "registration":
        return base_front + reg + [_td(inventors), _td(supply, align="right"), _td(vendor)]
    if kind == "annuity":
        return (
            base_front
            + reg
            + [_td(str(line.annuity_year or "")), _td(supply, align="right"), _td(vendor)]
        )
    # overseas
    return base_front + reg + [_td(inventors), _td(supply, align="right"), _td(vendor)]


def _detail_table_html(kind: str, lines: list[Any]) -> str:
    cols = _detail_columns(kind)
    ncol = len(cols)
    supply_idx = ncol - 2  # 공급가액 열 인덱스
    parts = [f'<table cellspacing="0" cellpadding="0" style="{_STYLE_TABLE_FIXED}">']
    parts.append("<tr>" + "".join(_td(c, bold=True) for c in cols) + "</tr>")
    total = 0
    for seq, line in enumerate(lines, start=1):
        parts.append("<tr>" + "".join(_detail_row_cells(kind, seq, line)) + "</tr>")
        # 항목별 소계(원본 양식 유지): 라벨이 순번+명칭 2칸 병합, 공급가액 칸에 금액.
        subtotal_cells = [_td("소  계", colspan=2, bold=True)]
        for i in range(2, ncol):
            if i == supply_idx:
                subtotal_cells.append(_td(_won(line.supply_amount), align="right", bold=True))
            else:
                subtotal_cells.append(_td(""))
        parts.append("<tr>" + "".join(subtotal_cells) + "</tr>")
        total += line.supply_amount or 0
    # 합계
    sum_cells = [_td("합  계", colspan=2, bold=True)]
    for i in range(2, ncol):
        sum_cells.append(_td(_won(total), align="right", bold=True) if i == supply_idx else _td(""))
    parts.append("<tr>" + "".join(sum_cells) + "</tr>")
    parts.append("</table>")
    return "".join(parts)


def _budget_table_html(total: int) -> str:
    """5번 예산표. 본품의 금액(=이번 총비용)만 자동으로 채우고, 연간/당월예산·기품의·
    당월잔액은 외부 예산값이라 빈칸으로 둔다(그룹웨어에서 수기 입력)."""
    cols = ["예산항목", "연간예산", "당월예산", "기품의 금액", "본품의 금액", "당월잔액"]
    vals = ["지급수수료", "", "", "", _won(total), ""]
    parts = [f'<table cellspacing="0" cellpadding="0" style="{_STYLE_TABLE_FIXED}">']
    parts.append("<tr>" + "".join(_td(c, bold=True) for c in cols) + "</tr>")
    parts.append(
        "<tr>"
        + "".join(_td(v, align="right" if i else "center") for i, v in enumerate(vals))
        + "</tr>"
    )
    parts.append("</table>")
    return "".join(parts)


def _footer_table_html() -> str:
    """사내표준 및 관련근거 표(고정). 좌측 라벨 셀은 회색 음영, 폭 700px 고정."""
    label_td = (
        f'<td style="{_STYLE_TD}text-align:center;font-weight:bold;background:#eeeeee;width:151px;">'
        "사내표준 및 관련근거</td>"
    )
    value_td = f'<td style="{_STYLE_TD}text-align:center;">{_DOC_STANDARD}</td>'
    return (
        f'<table cellspacing="0" cellpadding="0" style="{_STYLE_TABLE_FIXED}">'
        f"<tr>{label_td}{value_td}</tr></table>"
    )


# 원본 품의 양식과 동일하게 문단 여백을 촘촘하게(margin 1px, line-height 130%) — 기본 여백 방지.
_P_STYLE = "margin:1px 0px;line-height:130%;font-family:'Times New Roman',나눔고딕;font-size:13px;"


def _p(text: str) -> str:
    return f'<p style="{_P_STYLE}">{text}</p>'


def _spacer() -> str:
    """섹션 사이 1줄 띄움(원본과 동일한 촘촘한 빈 문단)."""
    return f'<p style="{_P_STYLE}"><br></p>'


def build_approval_html(
    cost_lines: list[Any],
    *,
    period: str,
    count_amount: dict[str, Any] | None = None,
    vendor_label: str = _DEFAULT_VENDOR,
) -> str:
    """품의 본문 HTML 생성. period='YYYY-MM'."""
    parts = (period or "").split("-")
    month = int(parts[1]) if len(parts) == 2 and parts[1].isdigit() else 0

    by_section: dict[str, list[Any]] = {}
    for line in cost_lines:
        sec = SECTION_OVERSEAS if line.region == "해외" else line.section
        by_section.setdefault(sec, []).append(line)

    grand_total = sum(int(line.supply_amount or 0) for line in cost_lines)

    out: list[str] = ['<div style="font-family:나눔고딕;font-size:13px;">']
    # 1. 총비용
    out.append(_p(f"1. 총&nbsp; 비&nbsp; 용 : <b>{_won(grand_total)}원 (부가세별도)</b>"))
    out.append(_spacer())
    # 2. 국내외 특허비용 정리
    out.append(_p("2. 2026년 국내외 특허비용 정리"))
    if count_amount:
        out.append(_count_amount_table_html(count_amount))
    else:
        out.append(_p("&nbsp;&nbsp;(6번 건수·금액 정리 파일 미첨부 — 누적표 생략)"))
    out.append(_spacer())
    # 3. 비용내역
    out.append(_p("3. 비용내역"))
    idx = 0
    for sec in SUMMARY_SECTION_ORDER:
        lines = by_section.get(sec)
        if not lines:
            continue
        idx += 1
        amount = sum(int(line.supply_amount or 0) for line in lines)
        vat_count = sum(1 for line in lines if (line.vat or 0) > 0)
        label = _CATEGORY_LABEL.get(sec, sec)
        out.append(_p(f"&nbsp;({idx}) {label}: {_won(amount)}원 (부가세별도, {vat_count}건)"))
        out.append(_detail_table_html(_CATEGORY_COLS.get(sec, "domestic"), lines))
    out.append(_spacer())
    # 4. 의뢰 업체
    out.append(_p(f"4. 의뢰 업체 : {escape(vendor_label)}"))
    out.append(_spacer())
    # 5. 주기 + 예산표
    out.append(_p(f"5. 주&nbsp;&nbsp;&nbsp; 기 : 예산은 {month}월 지급수수료 사용"))
    out.append(_budget_table_html(grand_total))
    out.append(_spacer())
    # 사내표준 및 관련근거
    out.append(_footer_table_html())
    out.append("</div>")
    return "".join(out)
