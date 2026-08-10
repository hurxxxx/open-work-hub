from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Literal

from openpyxl import load_workbook

from open_alm_api.core.i18n import localized_http_exception


EXPORT_RECORD_ID_HEADER = "record_id"
ImportFileKind = Literal["csv", "xlsx"]


@dataclass(frozen=True)
class LegacyIssueImportPreviewColumn:
    index: int
    header: str
    sample_values: list[str]


@dataclass(frozen=True)
class LegacyIssueImportPreview:
    columns: list[LegacyIssueImportPreviewColumn]
    preview_rows: list[list[str]]
    suggested_mapping: dict[str, int]
    total_preview_rows: int


@dataclass(frozen=True)
class LegacyIssueImportTable:
    headers: list[str]
    data_rows: list[list[str]]


def normalize_header(value: str) -> str:
    return " ".join(value.replace("\ufeff", "").replace("\n", " ").strip().lower().split())


def decode_csv_content(
    content: bytes,
    *,
    invalid_code: str,
    encodings: tuple[str, ...] = ("utf-8-sig", "utf-8", "cp949", "euc-kr"),
) -> str:
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise localized_http_exception(status_code=400, code=invalid_code)


def parse_tabular_upload(
    *,
    filename: str,
    content: bytes,
    invalid_code: str,
    sheet_name: str | None = None,
) -> tuple[list[list[str]], ImportFileKind]:
    lower_filename = filename.lower()
    if lower_filename.endswith(".csv"):
        text = decode_csv_content(content, invalid_code=invalid_code)
        try:
            return [[cell.strip() for cell in row] for row in csv.reader(StringIO(text))], "csv"
        except csv.Error as error:
            raise localized_http_exception(
                status_code=400,
                code=invalid_code,
                detail=str(error),
            ) from error
    if lower_filename.endswith(".xlsx"):
        try:
            workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        except Exception as error:
            raise localized_http_exception(status_code=400, code=invalid_code) from error
        sheet = (
            workbook[sheet_name]
            if sheet_name and sheet_name in workbook.sheetnames
            else workbook.active
        )
        return (
            [
                [tabular_cell_text(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ],
            "xlsx",
        )
    raise localized_http_exception(status_code=400, code=invalid_code)


def tabular_cell_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def combine_header_rows(rows: list[list[str]], *, header_rows: int) -> list[str]:
    if not rows:
        return []
    max_columns = max((len(row) for row in rows[:header_rows]), default=0)
    padded = [row + [""] * (max_columns - len(row)) for row in rows[:header_rows]]
    if header_rows <= 1:
        return [
            clean_header_text(header) or f"column_{index + 1}"
            for index, header in enumerate(padded[0] if padded else [])
        ]

    top_row = padded[0]
    sub_row = padded[1] if len(padded) > 1 else [""] * max_columns
    headers: list[str] = []
    current_group = ""
    for index in range(max_columns):
        top = clean_header_text(top_row[index])
        sub = clean_header_text(sub_row[index])
        if top:
            current_group = top
        group = top or (current_group if sub else "")
        if group and sub and normalize_header(group) != normalize_header(sub):
            headers.append(clean_header_text(f"{group} - {sub}"))
        elif group or sub:
            headers.append(clean_header_text(group or sub))
        else:
            headers.append(f"column_{index + 1}")
    return headers


def clean_header_text(value: str) -> str:
    return " ".join(value.strip().split())


def build_preview_columns(
    *,
    headers: list[str],
    preview_rows: list[list[str]],
) -> list[LegacyIssueImportPreviewColumn]:
    return [
        LegacyIssueImportPreviewColumn(
            index=index,
            header=header,
            sample_values=[
                row[index].strip()
                for row in preview_rows
                if index < len(row) and row[index].strip()
            ][:3],
        )
        for index, header in enumerate(headers)
    ]


def build_import_preview(
    *,
    table: LegacyIssueImportTable,
    suggested_mapping: dict[str, int],
    preview_row_count: int = 5,
) -> LegacyIssueImportPreview:
    preview_rows = table.data_rows[:preview_row_count]
    return LegacyIssueImportPreview(
        columns=build_preview_columns(headers=table.headers, preview_rows=preview_rows),
        preview_rows=preview_rows,
        suggested_mapping=suggested_mapping,
        total_preview_rows=len(preview_rows),
    )


def validate_mapping(
    *,
    allowed_keys: set[str],
    headers: list[str],
    invalid_code: str,
    mapping: dict[str, int],
) -> None:
    if not mapping:
        raise localized_http_exception(status_code=400, code=invalid_code)
    for key, column_index in mapping.items():
        if key not in allowed_keys or column_index < 0 or column_index >= len(headers):
            raise localized_http_exception(status_code=400, code=invalid_code)


def parse_required_mapping_json(*, invalid_code: str, mapping_json: str) -> dict[str, int]:
    return _parse_mapping_json(invalid_code=invalid_code, mapping_json=mapping_json)


def parse_optional_mapping_json(
    *,
    invalid_code: str,
    mapping_json: str | None,
) -> dict[str, int] | None:
    if not mapping_json:
        return None
    return _parse_mapping_json(invalid_code=invalid_code, mapping_json=mapping_json)


def _parse_mapping_json(*, invalid_code: str, mapping_json: str) -> dict[str, int]:
    try:
        parsed = json.loads(mapping_json)
    except json.JSONDecodeError as error:
        raise localized_http_exception(status_code=400, code=invalid_code) from error
    if not isinstance(parsed, dict):
        raise localized_http_exception(status_code=400, code=invalid_code)
    mapping: dict[str, int] = {}
    for key, value in parsed.items():
        if isinstance(key, str) and isinstance(value, int):
            mapping[key] = value
    return mapping
