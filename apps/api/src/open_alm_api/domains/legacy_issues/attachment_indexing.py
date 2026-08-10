from __future__ import annotations

import base64
import logging
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from celery import Celery
from sqlalchemy import delete, event, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from open_alm_api.core.llm import LlmRuntimeError
from open_alm_api.core.settings import get_settings
from open_alm_api.core.storage import get_minio_client
from open_alm_api.core.telemetry import serialize_current_trace_context
from open_alm_api.core.worker_queue_contract import (
    LEGACY_ISSUE_ATTACHMENT_INDEX_QUEUE,
    LEGACY_ISSUE_ATTACHMENT_INDEX_TASK_NAME,
)
from open_alm_api.core.worker_task_publisher import create_fail_fast_celery_publisher
from open_alm_api.domains.auth.models import utcnow_naive
from open_alm_api.domains.auth.security import new_id
from open_alm_api.domains.ai.gateway import (
    AiGatewayContextPack,
    LlmWorkloadContext,
    execute_llm,
)
from open_alm_api.domains.legacy_issues.ai_search import (
    build_legacy_issue_search_terms,
    embed_legacy_issue_ai_chunks,
)
from open_alm_api.domains.legacy_issues.dataset_records import (
    DATASET_DEFINITIONS,
    LegacyIssueDatasetDefinition,
    field_labels,
)
from open_alm_api.domains.legacy_issues.models import (
    LegacyIssueAiChunk,
    LegacyIssueAttachment,
    LegacyIssueAttachmentArtifact,
    LegacyIssueAttachmentIndexJob,
    LegacyIssueRecord,
)
from open_alm_api.domains.legacy_issues.partitioning import (
    ensure_attachment_job_partition,
    ensure_attachment_partition,
)
from open_alm_api.domains.legacy_issues.settings import get_legacy_issue_settings
from open_alm_api.domains.legacy_issues.task_kinds import (
    LEGACY_ISSUE_ATTACHMENT_SUMMARY_TASK_KIND,
)
from open_alm_api.domains.rag.provider_factory import RagProviderFactory
from open_alm_api.domains.rag.providers.local_document_rendering import (
    render_document_pages,
    suffix_for_content_type,
)


logger = logging.getLogger(__name__)

ATTACHMENT_INDEX_VERSION = "legacy_issue_attachment_index.v1"
ATTACHMENT_INDEX_STATUS_NOT_INDEXED = "not_indexed"
ATTACHMENT_INDEX_STATUS_PENDING = "pending"
ATTACHMENT_INDEX_STATUS_PROCESSING = "processing"
ATTACHMENT_INDEX_STATUS_INDEXED = "indexed"
ATTACHMENT_INDEX_STATUS_FAILED = "failed"
ATTACHMENT_SUMMARY_VERSION = "legacy_issue_attachment_summary.v1"
ATTACHMENT_SUMMARY_STATUS_NOT_SUMMARIZED = "not_summarized"
ATTACHMENT_SUMMARY_STATUS_PENDING = "pending"
ATTACHMENT_SUMMARY_STATUS_PROCESSING = "processing"
ATTACHMENT_SUMMARY_STATUS_SUMMARIZED = "summarized"
ATTACHMENT_SUMMARY_STATUS_FAILED = "failed"
JOB_STATUS_PENDING = "pending"
JOB_STATUS_PROCESSING = "processing"
JOB_STATUS_SUCCEEDED = "succeeded"
JOB_STATUS_FAILED = "failed"
JOB_STATUS_CANCELLED = "cancelled"
PENDING_OPERATION_UPSERT = "upsert"
PENDING_OPERATION_DELETE = "delete"
_PENDING_PUBLISHES_KEY = "legacy_issue_attachment_index_publish_after_commit"
_MAX_EXTRACTED_TEXT_CHARS = 250_000
_MAX_ARTIFACT_TEXT_CHARS = 120_000
_CHUNK_MAX_CHARS = 2200
_CHUNK_OVERLAP_CHARS = 240


@dataclass(frozen=True)
class AttachmentExtractionArtifact:
    artifact_kind: str
    text: str
    page_number: int | None = None
    metadata: dict[str, Any] | None = None


def enqueue_legacy_issue_attachment_index_job(
    db: Session,
    *,
    attachment: LegacyIssueAttachment,
    operation: str = PENDING_OPERATION_UPSERT,
    trigger: str,
) -> LegacyIssueAttachmentIndexJob:
    if operation not in {PENDING_OPERATION_UPSERT, PENDING_OPERATION_DELETE}:
        raise ValueError(f"Unsupported legacy issue attachment index operation: {operation}")
    trace_context = serialize_current_trace_context()
    existing = _select_pending_job(
        db, workspace_id=attachment.workspace_id, attachment_id=attachment.id
    )
    if existing is not None:
        ensure_attachment_job_partition(db, existing, attachment=attachment)
        existing.operation = operation
        existing.trigger = trigger[:40] or "manual"
        existing.trace_context = trace_context
        existing.updated_at = utcnow_naive()
        job = existing
    else:
        job = LegacyIssueAttachmentIndexJob(
            id=new_id(),
            workspace_id=attachment.workspace_id,
            dataset_key=attachment.dataset_key,
            revision_id=attachment.revision_id,
            attachment_id=attachment.id,
            operation=operation,
            trigger=trigger[:40] or "manual",
            status=JOB_STATUS_PENDING,
            attempts=0,
            trace_context=trace_context,
            created_at=utcnow_naive(),
            updated_at=utcnow_naive(),
        )
        ensure_attachment_job_partition(db, job, attachment=attachment)
        try:
            with db.begin_nested():
                db.add(job)
                db.flush()
        except IntegrityError:
            db.expire_all()
            current = _select_pending_job(
                db,
                workspace_id=attachment.workspace_id,
                attachment_id=attachment.id,
            )
            if current is None:
                raise
            ensure_attachment_job_partition(db, current, attachment=attachment)
            current.operation = operation
            current.trigger = trigger[:40] or "manual"
            current.trace_context = trace_context
            current.updated_at = utcnow_naive()
            job = current

    if operation == PENDING_OPERATION_UPSERT:
        attachment.index_status = ATTACHMENT_INDEX_STATUS_PENDING
        attachment.index_error = None
        attachment.index_version = ATTACHMENT_INDEX_VERSION
        if get_legacy_issue_settings().ai_attachment_summary_enabled:
            attachment.ai_summary = None
            attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_PENDING
            attachment.ai_summary_error = None
            attachment.ai_summary_model = None
            attachment.ai_summarized_at = None
            attachment.ai_summary_version = ATTACHMENT_SUMMARY_VERSION
        db.add(attachment)
    _schedule_publish_after_commit(db, job_id=job.id)
    return job


def delete_legacy_issue_attachment_index_data(
    db: Session,
    *,
    workspace_id: str,
    attachment_id: str,
) -> None:
    db.execute(
        delete(LegacyIssueAiChunk).where(
            LegacyIssueAiChunk.workspace_id == workspace_id,
            LegacyIssueAiChunk.attachment_id == attachment_id,
        )
    )
    db.execute(
        delete(LegacyIssueAttachmentArtifact).where(
            LegacyIssueAttachmentArtifact.workspace_id == workspace_id,
            LegacyIssueAttachmentArtifact.attachment_id == attachment_id,
        )
    )
    db.flush()


def process_legacy_issue_attachment_index_job(db: Session, job_id: str) -> str:
    job = _claim_job(db, job_id)
    if job is None:
        return "missing_or_closed"
    if job.operation == PENDING_OPERATION_DELETE:
        delete_legacy_issue_attachment_index_data(
            db,
            workspace_id=job.workspace_id,
            attachment_id=job.attachment_id,
        )
        _mark_job_succeeded(db, job_id, chunk_count=0, artifact_count=0)
        db.commit()
        return "deleted"
    try:
        artifact_count, chunk_count = _process_attachment_upsert(db, job_id=job_id)
    except Exception as error:
        db.rollback()
        _mark_job_failed(db, job_id=job_id, error=error, retry=False)
        raise
    _mark_job_succeeded(db, job_id, chunk_count=chunk_count, artifact_count=artifact_count)
    db.commit()
    return "indexed"


def mark_legacy_issue_attachment_index_job_for_retry(
    db: Session,
    *,
    job_id: str,
    error: Exception,
    countdown_seconds: int,
) -> None:
    _mark_job_failed(
        db,
        job_id=job_id,
        error=error,
        retry=True,
        countdown_seconds=countdown_seconds,
    )


def _process_attachment_upsert(db: Session, *, job_id: str) -> tuple[int, int]:
    job = db.get(LegacyIssueAttachmentIndexJob, job_id)
    if job is None:
        return 0, 0
    attachment = db.get(LegacyIssueAttachment, job.attachment_id)
    if attachment is None:
        job.status = JOB_STATUS_CANCELLED
        job.last_error = "attachment_missing"
        job.completed_at = utcnow_naive()
        db.add(job)
        db.commit()
        return 0, 0
    record = db.get(LegacyIssueRecord, attachment.record_id)
    definition = DATASET_DEFINITIONS.get(attachment.dataset_key)
    if record is None or definition is None:
        raise RuntimeError("Attachment record or dataset definition is missing.")
    ensure_attachment_partition(db, attachment, record=record)
    ensure_attachment_job_partition(db, job, attachment=attachment)
    content = _read_attachment_bytes(attachment)
    extracted = _extract_attachment_artifacts(db, attachment=attachment, content=content)
    if not extracted:
        extracted = [
            AttachmentExtractionArtifact(
                artifact_kind="metadata",
                text=_attachment_metadata_text(
                    attachment=attachment, record=record, definition=definition
                ),
                metadata={"fallback": "metadata_only"},
            )
        ]
    delete_legacy_issue_attachment_index_data(
        db,
        workspace_id=attachment.workspace_id,
        attachment_id=attachment.id,
    )
    now = utcnow_naive()
    artifact_rows = [
        LegacyIssueAttachmentArtifact(
            id=new_id(),
            workspace_id=attachment.workspace_id,
            dataset_key=attachment.dataset_key,
            revision_id=attachment.revision_id,
            record_id=attachment.record_id,
            stable_record_id=attachment.stable_record_id,
            attachment_id=attachment.id,
            job_id=job.id,
            artifact_kind=artifact.artifact_kind,
            page_number=artifact.page_number,
            content_text=_compact_text(artifact.text, limit=_MAX_ARTIFACT_TEXT_CHARS),
            artifact_metadata=artifact.metadata or {},
            created_at=now,
        )
        for artifact in extracted
    ]
    chunks = _build_attachment_chunks(
        definition=definition,
        record=record,
        attachment=attachment,
        artifacts=extracted,
        now=now,
    )
    db.add_all(artifact_rows)
    db.add_all(chunks)
    db.flush()
    embed_legacy_issue_ai_chunks(
        db,
        chunks,
        embed=True,
        cap=get_legacy_issue_settings().ai_max_sync_embedding_chunks,
    )
    db.flush()
    summary_chunks = _summarize_attachment_after_index(
        db,
        definition=definition,
        record=record,
        attachment=attachment,
        artifacts=extracted,
        chunks=chunks,
    )
    if summary_chunks:
        db.add_all(summary_chunks)
        db.flush()
        embed_legacy_issue_ai_chunks(
            db,
            summary_chunks,
            embed=True,
            cap=len(summary_chunks),
        )
        chunks.extend(summary_chunks)
    db.flush()
    return len(artifact_rows), len(chunks)


def _extract_attachment_artifacts(
    db: Session,
    *,
    attachment: LegacyIssueAttachment,
    content: bytes,
) -> list[AttachmentExtractionArtifact]:
    artifacts: list[AttachmentExtractionArtifact] = []
    native_text = _extract_native_text(
        content=content,
        filename=attachment.filename,
        content_type=attachment.content_type,
    )
    if native_text:
        artifacts.append(
            AttachmentExtractionArtifact(
                artifact_kind="native_text",
                text=native_text,
                metadata={"extractor": "local_native"},
            )
        )
    ocr_text, ocr_error = _extract_with_configured_ocr(
        content=content,
        content_type=attachment.content_type,
    )
    if ocr_text:
        artifacts.append(
            AttachmentExtractionArtifact(
                artifact_kind="ocr_text",
                text=ocr_text,
                metadata={"extractor": "configured_ocr"},
            )
        )
    elif ocr_error:
        artifacts.append(
            AttachmentExtractionArtifact(
                artifact_kind="ocr_error",
                text=f"OCR extraction failed for {attachment.filename}: {ocr_error}",
                metadata={"extractor": "configured_ocr", "error": ocr_error},
            )
        )
    legacy_settings = get_legacy_issue_settings()
    if legacy_settings.ai_attachment_index_vlm_always:
        vision_artifacts, vision_error = _extract_with_vision_model(
            db,
            attachment=attachment,
            content=content,
            filename=attachment.filename,
            content_type=attachment.content_type,
            max_pages=legacy_settings.ai_attachment_index_vision_max_pages,
        )
        artifacts.extend(vision_artifacts)
        if vision_error and not vision_artifacts:
            artifacts.append(
                AttachmentExtractionArtifact(
                    artifact_kind="vision_error",
                    text=f"Vision extraction failed for {attachment.filename}: {vision_error}",
                    metadata={"extractor": "vision_ocr", "error": vision_error},
                )
            )
    return [artifact for artifact in artifacts if artifact.text.strip()]


def _extract_native_text(
    *,
    content: bytes,
    filename: str,
    content_type: str | None,
) -> str:
    suffix = _attachment_suffix(filename=filename, content_type=content_type)
    try:
        if suffix in {".txt", ".csv", ".md"}:
            return _decode_text(content)
        if suffix == ".pptx":
            return _extract_pptx_text(content)
        if suffix == ".docx":
            return _extract_docx_text(content)
        if suffix == ".pdf":
            return _extract_pdf_text(content)
        if suffix in {".xlsx", ".xlsm"}:
            return _extract_xlsx_text(content)
    except Exception as error:
        logger.info("Legacy issue native extraction failed for %s: %s", filename, error)
    return ""


def _extract_with_configured_ocr(
    *,
    content: bytes,
    content_type: str | None,
) -> tuple[str, str | None]:
    try:
        ocr_client = RagProviderFactory(get_settings()).build_ocr()
    except Exception as error:
        return "", str(error)
    if ocr_client is None:
        return "", "ocr_provider_disabled"
    try:
        text = ocr_client.extract_text(content=content, content_type=content_type)
    except Exception as error:
        return "", str(error)
    return _compact_text(text, limit=_MAX_EXTRACTED_TEXT_CHARS), None


def _extract_with_vision_model(
    db: Session,
    *,
    attachment: LegacyIssueAttachment,
    content: bytes,
    filename: str,
    content_type: str | None,
    max_pages: int,
) -> tuple[list[AttachmentExtractionArtifact], str | None]:
    settings = get_settings()
    if not settings.rag_vision_ocr_enabled:
        return [], "vision_ocr_disabled"
    suffix = _attachment_suffix(filename=filename, content_type=content_type)
    try:
        with tempfile.TemporaryDirectory(prefix="open-alm-legacy-attachment-") as tmp:
            work_dir = Path(tmp)
            input_path = work_dir / f"input{suffix}"
            input_path.write_bytes(content)
            image_paths = render_document_pages(
                input_path,
                work_dir=work_dir,
                max_pages=max_pages,
                dpi=settings.rag_vision_ocr_dpi,
            )
            artifacts: list[AttachmentExtractionArtifact] = []
            for index, image_path in enumerate(image_paths, start=1):
                text = _call_vision_ocr(
                    db=db,
                    attachment=attachment,
                    image_path=image_path,
                    timeout_seconds=settings.rag_vision_ocr_timeout_seconds,
                    max_tokens=settings.rag_vision_ocr_max_new_tokens,
                )
                artifacts.append(
                    AttachmentExtractionArtifact(
                        artifact_kind="vision_text",
                        text=_compact_text(text, limit=_MAX_ARTIFACT_TEXT_CHARS),
                        page_number=index,
                        metadata={
                            "extractor": "vision_ocr",
                            "rendered_image": image_path.name,
                        },
                    )
                )
            return artifacts, None
    except Exception as error:
        return [], str(error)


def _call_vision_ocr(
    *,
    db: Session,
    attachment: LegacyIssueAttachment,
    image_path: Path,
    timeout_seconds: float,
    max_tokens: int,
) -> str:
    prompt = (
        "이미지에 보이는 모든 한국어/영어 텍스트, 표, 도면 키워드를 Markdown으로 추출하세요. "
        "이미지나 표의 의미도 검색 가능한 짧은 설명으로 덧붙이세요. "
        "추측한 값은 '추정:'으로 표시하고, 요약만 하지 말고 키워드를 보존하세요."
    )
    messages = [
        {
            "role": "user",
            "content": [
                {
                    "type": "image_url",
                    "image_url": {
                        "url": _data_url(image_path.read_bytes(), mime_type="image/png"),
                    },
                },
                {"type": "text", "text": prompt},
            ],
        }
    ]
    result = execute_llm(
        "legacy_issues.attachment_vision",
        LlmWorkloadContext(
            source="legacy_issue_attachment.vision",
            workspace_id=attachment.workspace_id,
            actor_user_id=None,
            principal_kind="system",
            principal_id="legacy_issue_attachment_index",
            app_id="legacy-issues",
        ),
        db,
        messages=messages,
        max_tokens=max_tokens,
        temperature=0,
        reasoning_effort="none",
        timeout_seconds=timeout_seconds,
        audit_entity_id=attachment.id,
    )
    return result.completion.text


def _build_attachment_chunks(
    *,
    definition: LegacyIssueDatasetDefinition,
    record: LegacyIssueRecord,
    attachment: LegacyIssueAttachment,
    artifacts: list[AttachmentExtractionArtifact],
    now: datetime,
) -> list[LegacyIssueAiChunk]:
    values = {key: str(value) for key, value in (record.field_values or {}).items()}
    labels = field_labels(definition)
    record_label = _record_label(values, fallback=record.stable_record_id or record.id)
    metadata = {
        "dataset_title_ko": definition.title_ko,
        "dataset_title_en": definition.title_en,
        "hierarchy_ko": list(definition.hierarchy_ko),
        "hierarchy_en": list(definition.hierarchy_en),
        "record_label": record_label,
        "attachment_id": attachment.id,
        "attachment_filename": attachment.filename,
        "attachment_description": attachment.description,
        "attachment_content_type": attachment.content_type,
    }
    chunks: list[LegacyIssueAiChunk] = []
    max_chunks = get_legacy_issue_settings().ai_attachment_index_max_chunks
    for artifact in artifacts:
        if not _is_searchable_attachment_artifact(artifact.artifact_kind):
            continue
        for text_part in _split_text(artifact.text, max_chars=_CHUNK_MAX_CHARS):
            if len(chunks) >= max_chunks:
                break
            chunk_index = len(chunks) + 1
            search_text = _attachment_chunk_search_text(
                definition=definition,
                labels=labels,
                values=values,
                record_label=record_label,
                attachment=attachment,
                artifact=artifact,
                text_part=text_part,
            )
            chunks.append(
                LegacyIssueAiChunk(
                    id=new_id(),
                    workspace_id=attachment.workspace_id,
                    dataset_key=attachment.dataset_key,
                    revision_id=attachment.revision_id,
                    record_id=attachment.record_id,
                    retrieval_partition_id=attachment.retrieval_partition_id,
                    stable_record_id=attachment.stable_record_id,
                    attachment_id=attachment.id,
                    attachment_filename=attachment.filename,
                    attachment_page=artifact.page_number,
                    attachment_artifact_type=artifact.artifact_kind,
                    chunk_key=f"attachment:{attachment.id}:{chunk_index:04d}",
                    chunk_kind="attachment_text",
                    field_key=None,
                    field_label="첨부파일",
                    field_value=_compact_text(text_part, limit=1200),
                    search_text=search_text,
                    search_terms=build_legacy_issue_search_terms(search_text),
                    evidence_metadata=metadata,
                    created_at=now,
                    updated_at=now,
                )
            )
        if len(chunks) >= max_chunks:
            break
    if chunks:
        return chunks
    fallback_text = _attachment_metadata_text(
        attachment=attachment,
        record=record,
        definition=definition,
    )
    return [
        LegacyIssueAiChunk(
            id=new_id(),
            workspace_id=attachment.workspace_id,
            dataset_key=attachment.dataset_key,
            revision_id=attachment.revision_id,
            record_id=attachment.record_id,
            retrieval_partition_id=attachment.retrieval_partition_id,
            stable_record_id=attachment.stable_record_id,
            attachment_id=attachment.id,
            attachment_filename=attachment.filename,
            chunk_key=f"attachment:{attachment.id}:metadata",
            chunk_kind="attachment_text",
            field_key=None,
            field_label="첨부파일",
            field_value=fallback_text,
            search_text=fallback_text,
            search_terms=build_legacy_issue_search_terms(fallback_text),
            evidence_metadata=metadata,
            created_at=now,
            updated_at=now,
        )
    ]


def _summarize_attachment_after_index(
    db: Session,
    *,
    definition: LegacyIssueDatasetDefinition,
    record: LegacyIssueRecord,
    attachment: LegacyIssueAttachment,
    artifacts: list[AttachmentExtractionArtifact],
    chunks: list[LegacyIssueAiChunk],
) -> list[LegacyIssueAiChunk]:
    legacy_settings = get_legacy_issue_settings()
    if not legacy_settings.ai_attachment_summary_enabled:
        attachment.ai_summary = None
        attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_NOT_SUMMARIZED
        attachment.ai_summary_error = None
        attachment.ai_summary_model = None
        attachment.ai_summarized_at = None
        db.add(attachment)
        return []
    try:
        source_text = _attachment_summary_source_text(
            definition=definition,
            record=record,
            attachment=attachment,
            artifacts=artifacts,
            chunks=chunks,
            limit=legacy_settings.ai_attachment_summary_input_max_chars,
        )
        summary, model = _generate_attachment_ai_summary(
            db,
            attachment=attachment,
            source_text=source_text,
        )
    except Exception as error:
        attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_FAILED
        attachment.ai_summary_error = _compact_text(str(error), limit=2000)
        db.add(attachment)
        return []
    attachment.ai_summary = summary
    attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_SUMMARIZED
    attachment.ai_summary_error = None
    attachment.ai_summary_model = model
    attachment.ai_summary_version = ATTACHMENT_SUMMARY_VERSION
    attachment.ai_summarized_at = utcnow_naive()
    db.add(attachment)
    return [
        _build_attachment_summary_chunk(
            definition=definition,
            record=record,
            attachment=attachment,
            summary=summary,
            now=attachment.ai_summarized_at,
        )
    ]


def _generate_attachment_ai_summary(
    db: Session,
    *,
    attachment: LegacyIssueAttachment,
    source_text: str,
) -> tuple[str, str | None]:
    legacy_settings = get_legacy_issue_settings()
    messages = [
        {
            "role": "system",
            "content": (
                "당신은 첨부파일을 자연어 검색에 쓰기 좋게 색인하는 문서 추출 전문가입니다. "
                "보고서 문장을 새로 만들지 말고, 제공된 변환 텍스트와 메타데이터에 실제로 "
                "나오는 원문 단어, 짧은 구, 코드, 숫자, 표 값을 그대로 정리하세요. 외부 지식, "
                "추론, 번역, 약어 풀이, 원인 재해석, 수치 계산은 금지합니다. 원문에 없는 "
                "동의어를 만들지 말고, 원문 표현을 가능한 한 보존하세요. 내부 처리 오류나 "
                "시스템 진단 문구는 문서 내용처럼 포함하지 마세요."
            ),
        },
        {
            "role": "user",
            "content": (
                "아래 첨부파일 변환 텍스트를 Markdown 색인 카드로 정리하세요. "
                "문장형 요약 리포트가 아니라 원문 기반 추출 목록이어야 합니다.\n"
                "규칙:\n"
                "- 원문에 나온 단어/구/코드/숫자를 그대로 사용\n"
                "- 약어, 업체명, 제품명, 차종 코드의 뜻을 임의로 쓰지 않음\n"
                "- 표 값은 원문 행/열/값 그대로 옮기고 직접 계산하지 않음\n"
                "- 이미지/도면/표 설명은 변환 텍스트에 나온 내용만 정리\n"
                "- 원문에 없는 활용 목적, 원인, 의미, 동의어를 생성하지 않음\n\n"
                "반드시 포함할 항목:\n"
                "1. 파일 식별 정보\n"
                "2. 원문 핵심 문구\n"
                "3. 검색 키워드 묶음: 문제/현상, 원인/부품, 대책/재료, 조건/수치, "
                "차종/코드, 조직/인명, 영문/약어\n"
                "4. 표/수치 원문값\n"
                "5. 추출 한계 또는 원본 확인 필요 항목\n\n"
                f"{source_text}"
            ),
        },
    ]
    try:
        completion = execute_llm(
            LEGACY_ISSUE_ATTACHMENT_SUMMARY_TASK_KIND,
            LlmWorkloadContext(
                source="legacy_issue_attachment.index",
                workspace_id=attachment.workspace_id,
                actor_user_id=None,
                principal_kind="system",
                principal_id="legacy_issue_attachment_index",
                app_id="legacy-issues",
            ),
            db,
            messages=messages,
            context_pack=AiGatewayContextPack(
                messages=messages,
                context_strategy="legacy_issue_attachment_index",
                source_kinds=("legacy_issue_attachment",),
                sensitivity_labels=("internal",),
                content_origin="internal_context",
            ),
            temperature=0.1,
            max_tokens=legacy_settings.ai_attachment_summary_max_tokens,
            reasoning_effort="none",
            audit_entity_id=attachment.id,
        ).completion
    except LlmRuntimeError as error:
        raise RuntimeError(str(error)) from error
    summary = _sanitize_attachment_summary(
        _compact_text(completion.text, limit=_MAX_ARTIFACT_TEXT_CHARS)
    )
    if not summary:
        raise RuntimeError("Local attachment summary model returned an empty result.")
    return summary, completion.model


def _build_attachment_summary_chunk(
    *,
    definition: LegacyIssueDatasetDefinition,
    record: LegacyIssueRecord,
    attachment: LegacyIssueAttachment,
    summary: str,
    now: datetime,
) -> LegacyIssueAiChunk:
    values = {key: str(value) for key, value in (record.field_values or {}).items()}
    labels = field_labels(definition)
    record_label = _record_label(values, fallback=record.stable_record_id or record.id)
    artifact = AttachmentExtractionArtifact(artifact_kind="ai_summary", text=summary)
    search_text = _attachment_chunk_search_text(
        definition=definition,
        labels=labels,
        values=values,
        record_label=record_label,
        attachment=attachment,
        artifact=artifact,
        text_part=summary,
    )
    return LegacyIssueAiChunk(
        id=new_id(),
        workspace_id=attachment.workspace_id,
        dataset_key=attachment.dataset_key,
        revision_id=attachment.revision_id,
        record_id=attachment.record_id,
        retrieval_partition_id=attachment.retrieval_partition_id,
        stable_record_id=attachment.stable_record_id,
        attachment_id=attachment.id,
        attachment_filename=attachment.filename,
        attachment_artifact_type="ai_summary",
        chunk_key=f"attachment:{attachment.id}:ai_summary",
        chunk_kind="attachment_summary",
        field_key=None,
        field_label="첨부파일 AI 추출 키워드",
        field_value=_compact_text(summary, limit=1200),
        search_text=search_text,
        search_terms=build_legacy_issue_search_terms(search_text),
        evidence_metadata={
            "dataset_title_ko": definition.title_ko,
            "dataset_title_en": definition.title_en,
            "hierarchy_ko": list(definition.hierarchy_ko),
            "hierarchy_en": list(definition.hierarchy_en),
            "record_label": record_label,
            "attachment_id": attachment.id,
            "attachment_filename": attachment.filename,
            "attachment_description": attachment.description,
            "attachment_content_type": attachment.content_type,
            "summary_version": ATTACHMENT_SUMMARY_VERSION,
        },
        created_at=now,
        updated_at=now,
    )


def _attachment_summary_source_text(
    *,
    definition: LegacyIssueDatasetDefinition,
    record: LegacyIssueRecord,
    attachment: LegacyIssueAttachment,
    artifacts: list[AttachmentExtractionArtifact],
    chunks: list[LegacyIssueAiChunk],
    limit: int,
) -> str:
    lines = [
        "# Attachment metadata",
        f"dataset: {definition.title_ko}",
        f"record_id: {record.stable_record_id or record.id}",
        f"filename: {attachment.filename}",
        f"content_type: {attachment.content_type}",
        f"description: {attachment.description or ''}",
        "",
        "# Converted attachment text",
    ]
    if artifacts:
        for artifact in artifacts:
            if not _is_searchable_attachment_artifact(artifact.artifact_kind):
                continue
            label = artifact.artifact_kind
            if artifact.page_number is not None:
                label = f"{label} page {artifact.page_number}"
            lines.append(f"\n## {label}\n{artifact.text}")
    else:
        for chunk in chunks:
            if chunk.field_value and _is_searchable_attachment_artifact(
                chunk.attachment_artifact_type or ""
            ):
                lines.append(f"\n## chunk {chunk.chunk_key}\n{chunk.field_value}")
    return _compact_text("\n".join(lines), limit=limit)


def _is_searchable_attachment_artifact(artifact_kind: str) -> bool:
    return bool(artifact_kind) and not artifact_kind.endswith("_error")


def _sanitize_attachment_summary(summary: str) -> str:
    lines: list[str] = []
    undefined_term_pattern = re.compile(
        r"^(?P<indent>\s*)[*-]\s+\*\*(?P<term>[^*]+)\*\*:\s+.*"
        r"(?:원문.*(?:정의|뜻).*없음|정의\s*없음).*$"
    )
    ambiguous_average_pattern = re.compile(r"\s*\(평균\s*[^)]*\)")
    inferred_relation_pattern = re.compile(
        r"\s*\([^)]*(?:관련|항목|열|컬럼|column)[^)]*\)",
        flags=re.IGNORECASE,
    )
    for line in summary.splitlines():
        match = undefined_term_pattern.match(line)
        if match:
            lines.append(
                f"{match.group('indent')}* **{match.group('term').strip()}**: (원문 정의 없음)"
            )
            continue
        if "평균" in line and "주석" in line and ("보입니다" in line or "나타내" in line):
            lines.append(
                "* **표 데이터의 평균값 주석**: 표 하단에 평균값 주석이 있으나, "
                "어떤 항목과 연결되는지는 변환 텍스트만으로 확정하지 않습니다. 원본 확인 필요."
            )
            continue
        if "평균" in line:
            line = inferred_relation_pattern.sub("", line)
        line = ambiguous_average_pattern.sub("", line)
        lines.append(line)
    return "\n".join(lines).strip()


def _attachment_chunk_search_text(
    *,
    definition: LegacyIssueDatasetDefinition,
    labels: dict[str, str],
    values: dict[str, str],
    record_label: str,
    attachment: LegacyIssueAttachment,
    artifact: AttachmentExtractionArtifact,
    text_part: str,
) -> str:
    context_fields = []
    for key in ("legacy_issue_number", "issue_no", "vehicle", "item", "problem", "symptom"):
        value = values.get(key)
        if value:
            context_fields.append(f"{labels.get(key, key)}: {value}")
    lines = [
        definition.title_ko,
        f"record: {record_label}",
        f"attachment: {attachment.filename}",
        f"attachment_description: {attachment.description or ''}",
        f"artifact: {artifact.artifact_kind}",
    ]
    if artifact.page_number is not None:
        lines.append(f"page: {artifact.page_number}")
    lines.extend(context_fields)
    lines.extend(("", text_part))
    return _compact_text("\n".join(lines), limit=8000)


def _read_attachment_bytes(attachment: LegacyIssueAttachment) -> bytes:
    settings = get_settings()
    client = get_minio_client()
    response = client.get_object(settings.minio_bucket, attachment.storage_key)
    try:
        chunks = bytearray()
        while True:
            part = response.read(1024 * 1024)
            if not part:
                break
            chunks.extend(part)
        return bytes(chunks)
    finally:
        response.close()
        response.release_conn()


def _claim_job(db: Session, job_id: str) -> LegacyIssueAttachmentIndexJob | None:
    job = db.get(LegacyIssueAttachmentIndexJob, job_id)
    if job is None or job.status != JOB_STATUS_PENDING:
        return None
    now = utcnow_naive()
    job.status = JOB_STATUS_PROCESSING
    job.attempts += 1
    job.started_at = now
    job.updated_at = now
    job.last_error = None
    attachment = db.get(LegacyIssueAttachment, job.attachment_id)
    if attachment is not None and job.operation == PENDING_OPERATION_UPSERT:
        attachment.index_status = ATTACHMENT_INDEX_STATUS_PROCESSING
        attachment.index_error = None
        if get_legacy_issue_settings().ai_attachment_summary_enabled:
            attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_PROCESSING
            attachment.ai_summary_error = None
        db.add(attachment)
    db.add(job)
    db.commit()
    return job


def _mark_job_succeeded(
    db: Session,
    job_id: str,
    *,
    chunk_count: int,
    artifact_count: int,
) -> None:
    job = db.get(LegacyIssueAttachmentIndexJob, job_id)
    if job is None:
        return
    now = utcnow_naive()
    job.status = JOB_STATUS_SUCCEEDED
    job.last_error = None
    job.completed_at = now
    job.updated_at = now
    job.next_retry_at = None
    attachment = db.get(LegacyIssueAttachment, job.attachment_id)
    if attachment is not None and job.operation == PENDING_OPERATION_UPSERT:
        attachment.index_status = ATTACHMENT_INDEX_STATUS_INDEXED
        attachment.index_error = None
        attachment.indexed_at = now
        attachment.index_version = ATTACHMENT_INDEX_VERSION
        attachment.chunk_count = chunk_count
        attachment.artifact_count = artifact_count
        db.add(attachment)
    db.add(job)


def _mark_job_failed(
    db: Session,
    *,
    job_id: str,
    error: Exception,
    retry: bool,
    countdown_seconds: int = 0,
) -> None:
    job = db.get(LegacyIssueAttachmentIndexJob, job_id)
    if job is None:
        return
    message = _compact_text(str(error), limit=2000)
    now = utcnow_naive()
    job.status = JOB_STATUS_PENDING if retry else JOB_STATUS_FAILED
    job.last_error = message
    job.next_retry_at = now + timedelta(seconds=max(countdown_seconds, 1)) if retry else None
    job.updated_at = now
    job.completed_at = None if retry else now
    attachment = db.get(LegacyIssueAttachment, job.attachment_id)
    if attachment is not None:
        attachment.index_status = (
            ATTACHMENT_INDEX_STATUS_PENDING if retry else ATTACHMENT_INDEX_STATUS_FAILED
        )
        attachment.index_error = message
        if retry:
            attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_PENDING
            attachment.ai_summary_error = message
        else:
            attachment.ai_summary_status = ATTACHMENT_SUMMARY_STATUS_FAILED
            attachment.ai_summary_error = message
        db.add(attachment)
    db.add(job)
    db.commit()


def _select_pending_job(
    db: Session,
    *,
    workspace_id: str,
    attachment_id: str,
) -> LegacyIssueAttachmentIndexJob | None:
    return db.scalar(
        select(LegacyIssueAttachmentIndexJob)
        .where(
            LegacyIssueAttachmentIndexJob.workspace_id == workspace_id,
            LegacyIssueAttachmentIndexJob.attachment_id == attachment_id,
            LegacyIssueAttachmentIndexJob.status == JOB_STATUS_PENDING,
        )
        .order_by(LegacyIssueAttachmentIndexJob.created_at.desc())
        .limit(1)
    )


def _schedule_publish_after_commit(db: Session, *, job_id: str) -> None:
    pending = db.info.setdefault(_PENDING_PUBLISHES_KEY, set())
    if not isinstance(pending, set):
        pending = set()
        db.info[_PENDING_PUBLISHES_KEY] = pending
    pending.add(job_id)


@event.listens_for(Session, "after_commit")
def _publish_pending_attachment_index_jobs(session: Session) -> None:
    if session.in_nested_transaction():
        return
    pending = session.info.pop(_PENDING_PUBLISHES_KEY, None)
    if not pending:
        return
    for job_id in sorted(pending):
        try:
            _publish_job(job_id=job_id)
        except Exception:
            logger.warning("Failed to publish legacy issue attachment index job", exc_info=True)


@event.listens_for(Session, "after_rollback")
def _clear_pending_attachment_index_jobs(session: Session) -> None:
    if session.in_nested_transaction():
        return
    session.info.pop(_PENDING_PUBLISHES_KEY, None)


def _publish_job(*, job_id: str) -> None:
    _celery_client().signature(
        LEGACY_ISSUE_ATTACHMENT_INDEX_TASK_NAME,
        args=[job_id],
        immutable=True,
    ).apply_async(
        queue=LEGACY_ISSUE_ATTACHMENT_INDEX_QUEUE,
        retry=False,
    )


def _celery_client() -> Celery:
    settings = get_settings()
    return create_fail_fast_celery_publisher(
        "open_alm_api_legacy_issue_attachment_index",
        broker=settings.worker_broker_url,
        ignore_result=True,
    )


def _extract_pptx_text(content: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(BytesIO(content))
    lines: list[str] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if isinstance(text, str) and text.strip():
                slide_lines.append(text.strip())
            table = getattr(shape, "table", None)
            if table is not None:
                for row in table.rows:
                    values = [
                        cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()
                    ]
                    if values:
                        slide_lines.append(" | ".join(values))
        if slide.has_notes_slide:
            notes = getattr(slide.notes_slide.notes_text_frame, "text", "")
            if notes and notes.strip():
                slide_lines.append(f"notes: {notes.strip()}")
        if slide_lines:
            lines.append(f"## slide {slide_index}\n" + "\n".join(slide_lines))
    return _compact_text("\n\n".join(lines), limit=_MAX_EXTRACTED_TEXT_CHARS)


def _extract_docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                lines.append(" | ".join(values))
    return _compact_text("\n".join(lines), limit=_MAX_EXTRACTED_TEXT_CHARS)


def _extract_pdf_text(content: bytes) -> str:
    import fitz

    document = fitz.open(stream=content, filetype="pdf")
    lines: list[str] = []
    for page_index in range(document.page_count):
        text = document.load_page(page_index).get_text("text").strip()
        if text:
            lines.append(f"## page {page_index + 1}\n{text}")
    return _compact_text("\n\n".join(lines), limit=_MAX_EXTRACTED_TEXT_CHARS)


def _extract_xlsx_text(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets[:20]:
        lines.append(f"## sheet {sheet.title}")
        for row in sheet.iter_rows(max_row=500, values_only=True):
            values = [
                str(value).strip() for value in row if value is not None and str(value).strip()
            ]
            if values:
                lines.append(" | ".join(values))
    return _compact_text("\n".join(lines), limit=_MAX_EXTRACTED_TEXT_CHARS)


def _split_text(text: str, *, max_chars: int) -> list[str]:
    normalized = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not normalized:
        return []
    chunks: list[str] = []
    current = ""
    for paragraph in re.split(r"\n\s*\n", normalized):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= max_chars:
            current = paragraph
            continue
        start = 0
        while start < len(paragraph):
            end = min(start + max_chars, len(paragraph))
            chunks.append(paragraph[start:end].strip())
            if end >= len(paragraph):
                break
            start = max(end - _CHUNK_OVERLAP_CHARS, start + 1)
        current = ""
    if current:
        chunks.append(current)
    return chunks


def _attachment_metadata_text(
    *,
    attachment: LegacyIssueAttachment,
    record: LegacyIssueRecord,
    definition: LegacyIssueDatasetDefinition,
) -> str:
    values = {key: str(value) for key, value in (record.field_values or {}).items()}
    record_label = _record_label(values, fallback=record.stable_record_id or record.id)
    return _compact_text(
        "\n".join(
            (
                definition.title_ko,
                f"record: {record_label}",
                f"attachment: {attachment.filename}",
                f"content_type: {attachment.content_type}",
                f"description: {attachment.description or ''}",
            )
        ),
        limit=8000,
    )


def _record_label(values: dict[str, str], *, fallback: str) -> str:
    for key in (
        "legacy_issue_number",
        "issue_no",
        "row_no",
        "problem",
        "symptom",
        "defect_type",
        "item",
    ):
        value = values.get(key)
        if value:
            return _compact_text(value, limit=120)
    return fallback


def _attachment_suffix(*, filename: str, content_type: str | None) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix:
        return suffix
    return suffix_for_content_type(content_type)


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            return _compact_text(content.decode(encoding), limit=_MAX_EXTRACTED_TEXT_CHARS)
        except UnicodeDecodeError:
            continue
    return _compact_text(content.decode("utf-8", errors="ignore"), limit=_MAX_EXTRACTED_TEXT_CHARS)


def _compact_text(value: str, *, limit: int) -> str:
    compacted = re.sub(r"[ \t\r\f\v]+", " ", value).strip()
    if len(compacted) <= limit:
        return compacted
    return compacted[: max(limit - 3, 0)].rstrip() + "..."


def _data_url(content: bytes, *, mime_type: str) -> str:
    return f"data:{mime_type};base64,{base64.b64encode(content).decode('ascii')}"
