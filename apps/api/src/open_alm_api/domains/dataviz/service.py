from __future__ import annotations

import io
import uuid
from collections import defaultdict
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from open_alm_api.domains.dataviz.compressor_perf import (
    CATEGORY_ELECTRIC,
    CATEGORY_VARIABLE,
    NUMERIC_FIELDS,
    TOLERANCE_FIELDS,
    TOLERANCE_PROFILE_FIELDS,
    PerfRow,
    check_warnings,
    parse_date_from_filename,
    parse_perf_master,
    parse_perf_source,
    parse_tolerance_reference_file,
)
from open_alm_api.domains.dataviz.models import (
    DataVizCarModelRef,
    DataVizPerfData,
    DataVizPerfProcessedFile,
    DataVizPerfTolerance,
    DataVizPerfToleranceProfile,
    DataVizPerfToleranceProfileCar,
)


def _uuid() -> str:
    return str(uuid.uuid4())


def _row_model(
    *,
    workspace_id: str,
    owner_id: str | None,
    category: str,
    refrigerant: str,
    capacity: str,
    row: PerfRow,
) -> DataVizPerfData:
    return DataVizPerfData(
        id=_uuid(),
        workspace_id=workspace_id,
        owner_id=owner_id,
        category=category,
        refrigerant=refrigerant,
        capacity=capacity,
        test_group=row.test_group,
        comp_type=row.comp_type,
        serial_no=row.serial_no,
        test_date=row.test_date,
        car_model=row.car_model,
        engine_spec=row.engine_spec,
        remarks=row.remarks,
        source_file=row.source_file,
        rpm=row.rpm,
        pd=row.pd,
        td=row.td,
        ps=row.ps,
        ts=row.ts,
        pc=row.pc,
        mass_flow=row.mass_flow,
        vol_eff=row.vol_eff,
        ocr=row.ocr,
        cooling_cap_a=row.cooling_cap_a,
        cooling_cap_f=row.cooling_cap_f,
        power_kw=row.power_kw,
        cop_sc=row.cop_sc,
        heat_balance=row.heat_balance,
        torque=row.torque,
    )


def import_master_workbook(
    db: Session,
    *,
    workspace_id: str,
    owner_id: str,
    category: str,
    filename: str,
    content: bytes,
) -> dict:
    imported = parse_perf_master(content, category)
    inserted = 0
    for sheet in imported:
        db.execute(
            delete(DataVizPerfData).where(
                DataVizPerfData.workspace_id == workspace_id,
                DataVizPerfData.category == category,
                DataVizPerfData.refrigerant == sheet.refrigerant,
                DataVizPerfData.capacity == sheet.capacity,
                DataVizPerfData.source_file == "master",
            )
        )
        for row in sheet.rows:
            row.source_file = "master"
            db.add(
                _row_model(
                    workspace_id=workspace_id,
                    owner_id=owner_id,
                    category=category,
                    refrigerant=sheet.refrigerant,
                    capacity=sheet.capacity,
                    row=row,
                )
            )
            inserted += 1
    db.commit()
    return {
        "ok": True,
        "filename": filename,
        "category": category,
        "imported_sheets": [sheet.sheet_name for sheet in imported],
        "row_count": inserted,
    }


def upload_source_workbooks(
    db: Session,
    *,
    workspace_id: str,
    owner_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    files: list[tuple[str, bytes]],
) -> dict:
    ordered = sorted(files, key=lambda item: parse_date_from_filename(item[0]))
    results = []
    success_count = 0
    for filename, content in ordered:
        duplicate = db.scalar(
            select(DataVizPerfProcessedFile.id).where(
                DataVizPerfProcessedFile.workspace_id == workspace_id,
                DataVizPerfProcessedFile.filename == filename,
                DataVizPerfProcessedFile.category == category,
                DataVizPerfProcessedFile.refrigerant == refrigerant,
                DataVizPerfProcessedFile.capacity == capacity,
            )
        )
        if duplicate:
            results.append(
                {
                    "filename": filename,
                    "status": "duplicate",
                    "row_count": 0,
                    "warning_count": 0,
                    "rows": [],
                }
            )
            continue

        try:
            rows = parse_perf_source(content, category, filename)
        except Exception as exc:  # noqa: BLE001
            results.append(
                {
                    "filename": filename,
                    "status": "error",
                    "row_count": 0,
                    "warning_count": 0,
                    "error": str(exc),
                    "rows": [],
                }
            )
            continue

        if not rows:
            results.append(
                {
                    "filename": filename,
                    "status": "error",
                    "row_count": 0,
                    "warning_count": 0,
                    "error": "no data",
                    "rows": [],
                }
            )
            continue

        ref_cache: dict[str, dict[str, dict[str, float | None]]] = {}
        all_profiles = get_all_tolerance_profiles(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
        )["profiles"]
        row_summaries = []
        warning_count = 0
        for row in rows:
            row.source_file = filename
            db.add(
                _row_model(
                    workspace_id=workspace_id,
                    owner_id=owner_id,
                    category=category,
                    refrigerant=refrigerant,
                    capacity=capacity,
                    row=row,
                )
            )
            refs = _ref_values_cached(
                db,
                workspace_id=workspace_id,
                category=category,
                refrigerant=refrigerant,
                capacity=capacity,
                car_model=row.car_model,
                cache=ref_cache,
            )
            row_dict = _serialize_row_values(row)
            warnings = check_warnings(row_dict, refs, all_profiles)
            warning_count += 1 if warnings else 0
            row_summaries.append(
                {
                    "test_group": row.test_group,
                    "serial_no": row.serial_no,
                    "rpm": row.rpm,
                    "pd": row.pd,
                    "ps": row.ps,
                    "cooling_cap_a": row.cooling_cap_a,
                    "power_kw": row.power_kw,
                    "vol_eff": row.vol_eff,
                    "torque": row.torque,
                    "warnings": warnings,
                }
            )
        db.add(
            DataVizPerfProcessedFile(
                id=_uuid(),
                workspace_id=workspace_id,
                filename=filename,
                category=category,
                refrigerant=refrigerant,
                capacity=capacity,
            )
        )
        try:
            db.commit()
        except IntegrityError:
            db.rollback()
            results.append(
                {
                    "filename": filename,
                    "status": "duplicate",
                    "row_count": 0,
                    "warning_count": 0,
                    "rows": [],
                }
            )
            continue
        success_count += 1
        results.append(
            {
                "filename": filename,
                "status": "ok",
                "row_count": len(rows),
                "warning_count": warning_count,
                "rows": row_summaries,
            }
        )
    return {"status": "ok", "success_count": success_count, "results": results}


def list_perf_data(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
) -> dict:
    rows = list(
        db.scalars(
            select(DataVizPerfData)
            .where(
                DataVizPerfData.workspace_id == workspace_id,
                DataVizPerfData.category == category,
                DataVizPerfData.refrigerant == refrigerant,
                DataVizPerfData.capacity == capacity,
            )
            .order_by(DataVizPerfData.created_at.desc(), DataVizPerfData.id.desc())
        )
    )
    ref_cache: dict[str, dict[str, dict[str, float | None]]] = {}
    all_profiles = get_all_tolerance_profiles(
        db,
        workspace_id=workspace_id,
        category=category,
        refrigerant=refrigerant,
        capacity=capacity,
    )["profiles"]
    data = []
    for row in rows:
        item = serialize_perf_data_row(row)
        refs = _ref_values_cached(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            car_model=row.car_model,
            cache=ref_cache,
        )
        item["warnings"] = check_warnings(item, refs, all_profiles)
        item["no_ref"] = not bool(refs)
        data.append(item)
    return {"data": data}


def analyze_perf_data(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
) -> dict:
    rows = list(
        db.scalars(
            select(DataVizPerfData)
            .where(
                DataVizPerfData.workspace_id == workspace_id,
                DataVizPerfData.category == category,
                DataVizPerfData.refrigerant == refrigerant,
                DataVizPerfData.capacity == capacity,
            )
            .order_by(DataVizPerfData.created_at.desc(), DataVizPerfData.id.desc())
        )
    )
    master_rows = [row for row in rows if row.source_file == "master"]
    source_rows = [row for row in rows if row.source_file and row.source_file != "master"]
    recent_files = []
    for row in source_rows:
        if row.source_file not in recent_files:
            recent_files.append(row.source_file)
        if len(recent_files) >= 3:
            break
    latest_rows = [row for row in source_rows if row.source_file in set(recent_files)]

    ref_cache: dict[str, dict[str, dict[str, float | None]]] = {}
    all_profiles = get_all_tolerance_profiles(
        db,
        workspace_id=workspace_id,
        category=category,
        refrigerant=refrigerant,
        capacity=capacity,
    )["profiles"]
    latest = []
    for row in latest_rows:
        item = serialize_perf_data_row(row)
        refs = _ref_values_cached(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            car_model=row.car_model,
            cache=ref_cache,
        )
        item["warnings"] = check_warnings(item, refs, all_profiles)
        item["no_ref"] = not bool(refs)
        latest.append(item)

    return {
        "averages": _averages(master_rows, group_fields=("test_group",)),
        "comp_averages": _averages(master_rows, group_fields=("comp_type", "test_group")),
        "latest": latest,
    }


def update_perf_row(
    db: Session,
    *,
    workspace_id: str,
    row_id: str,
    values: dict,
) -> dict:
    row = _get_workspace_row(db, workspace_id=workspace_id, row_id=row_id)
    allowed_text = {"comp_type", "test_group", "serial_no", "test_date", "car_model", "engine_spec", "remarks"}
    allowed_numbers = set(NUMERIC_FIELDS)
    changed = False
    for key, value in values.items():
        if key in allowed_text:
            setattr(row, key, "" if value is None else str(value))
            changed = True
        elif key in allowed_numbers:
            setattr(row, key, _float_or_none(value))
            changed = True
    if not changed:
        return {"error": "no valid fields"}
    db.commit()
    return {"status": "ok", "data": serialize_perf_data_row(row)}


def delete_perf_row(db: Session, *, workspace_id: str, row_id: str) -> dict:
    row = _get_workspace_row(db, workspace_id=workspace_id, row_id=row_id)
    db.delete(row)
    db.commit()
    return {"status": "ok"}


def reset_perf_data(db: Session, *, workspace_id: str, category: str | None) -> dict:
    conditions = [DataVizPerfData.workspace_id == workspace_id]
    file_conditions = [DataVizPerfProcessedFile.workspace_id == workspace_id]
    if category:
        conditions.append(DataVizPerfData.category == category)
        file_conditions.append(DataVizPerfProcessedFile.category == category)
    db.execute(delete(DataVizPerfData).where(*conditions))
    db.execute(delete(DataVizPerfProcessedFile).where(*file_conditions))
    db.commit()
    return {"status": "ok"}


def list_car_models(db: Session, *, workspace_id: str) -> dict:
    rows = list(
        db.scalars(
            select(DataVizCarModelRef)
            .where(DataVizCarModelRef.workspace_id == workspace_id)
            .order_by(DataVizCarModelRef.code)
        )
    )
    data = [{"id": row.id, "code": row.code, "description": row.description} for row in rows]
    return {"data": data, "models": [row["code"] for row in data]}


def add_car_model(db: Session, *, workspace_id: str, code: str, description: str = "") -> dict:
    db.add(
        DataVizCarModelRef(
            id=_uuid(),
            workspace_id=workspace_id,
            code=code.strip()[:40],
            description=description.strip(),
        )
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return {"error": "duplicate code"}
    return {"status": "ok"}


def delete_car_model(db: Session, *, workspace_id: str, model_id: str) -> dict:
    row = db.get(DataVizCarModelRef, model_id)
    if row and row.workspace_id == workspace_id:
        db.delete(row)
        db.commit()
    return {"status": "ok"}


def get_tolerance(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
) -> dict:
    return {
        "ref_values": _get_ref_values(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            car_model=car_model,
        ),
        "car_models": _get_tolerance_car_models(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
        ),
        "profiles": _get_tol_profiles(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
        ),
    }


def get_tolerance_all(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
) -> dict:
    conditions = [
        DataVizPerfTolerance.workspace_id == workspace_id,
        DataVizPerfTolerance.category == category,
        DataVizPerfTolerance.refrigerant == refrigerant,
    ]
    if capacity:
        conditions.append(DataVizPerfTolerance.capacity == capacity)
    rows = list(db.scalars(select(DataVizPerfTolerance).where(*conditions)))
    result: dict[str, dict] = {}
    cap_models: dict[str, dict[str, str]] = {}
    for row in rows:
        key = f"{row.capacity}|{row.car_model}"
        cap_models.setdefault(key, {"capacity": row.capacity, "model": row.car_model})
        result.setdefault(key, {}).setdefault(row.test_group, {})[row.field_name] = row.ref_value
    return {"models": result, "cap_models": cap_models}


def save_tolerance(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
    ref_values: dict,
) -> dict:
    for test_group, fields in ref_values.items():
        if not isinstance(fields, dict):
            continue
        for field_name, ref_value in fields.items():
            if field_name not in TOLERANCE_FIELDS:
                continue
            _upsert_tolerance(
                db,
                workspace_id=workspace_id,
                category=category,
                refrigerant=refrigerant,
                capacity=capacity,
                car_model=car_model,
                test_group=str(test_group),
                field_name=field_name,
                ref_value=_float_or_none(ref_value),
            )
    db.commit()
    return {"status": "ok"}


def upload_tolerance_file(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    filename: str,
    content: bytes,
) -> dict:
    rows = parse_tolerance_reference_file(content, filename)
    db.execute(
        delete(DataVizPerfTolerance).where(
            DataVizPerfTolerance.workspace_id == workspace_id,
            DataVizPerfTolerance.category == category,
            DataVizPerfTolerance.refrigerant == refrigerant,
        )
    )
    for row in rows:
        db.add(
            DataVizPerfTolerance(
                id=_uuid(),
                workspace_id=workspace_id,
                category=category,
                refrigerant=refrigerant,
                capacity=row.capacity,
                car_model=row.car_model,
                test_group=row.test_group,
                field_name=row.field_name,
                ref_value=row.ref_value,
                tolerance=None,
            )
        )
    db.commit()
    return {"status": "ok", "inserted": len(rows), "filename": filename}


def delete_tolerance(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
) -> dict:
    conditions = [
        DataVizPerfTolerance.workspace_id == workspace_id,
        DataVizPerfTolerance.category == category,
        DataVizPerfTolerance.refrigerant == refrigerant,
        DataVizPerfTolerance.car_model == car_model,
    ]
    if capacity:
        conditions.append(DataVizPerfTolerance.capacity == capacity)
    db.execute(delete(DataVizPerfTolerance).where(*conditions))
    db.commit()
    return {"status": "ok"}


def get_tolerance_profile(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    profile_name: str,
) -> dict:
    return {
        "tol_data": _get_tol_profile_data(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            profile_name=profile_name,
        ),
        "cars": _get_tol_profile_cars(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            profile_name=profile_name,
        )
        if profile_name != "BASE"
        else [],
        "profiles": _get_tol_profiles(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
        ),
    }


def get_all_tolerance_profiles(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
) -> dict:
    # 한 (카테고리·냉매·용량) 범위의 저장된 모든 시험공차 프로필을 한 번에 반환한다.
    # { profiles: { profile_name: { test_group: { field: {ref,sign,tol,lower} } } } }
    rows = list(
        db.scalars(
            select(DataVizPerfToleranceProfile).where(
                DataVizPerfToleranceProfile.workspace_id == workspace_id,
                DataVizPerfToleranceProfile.category == category,
                DataVizPerfToleranceProfile.refrigerant == refrigerant,
                DataVizPerfToleranceProfile.capacity == capacity,
            )
        )
    )
    profiles: dict[str, dict[str, dict[str, dict]]] = {}
    for row in rows:
        profiles.setdefault(row.profile_name, {}).setdefault(row.test_group, {})[
            row.field_name
        ] = {
            "ref": row.ref_value,
            "sign": row.sign or "±",
            "tol": row.tol_value,
            "lower": row.tol_lower,
        }
    return {"profiles": profiles}


def save_tolerance_profile(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    profile_name: str,
    tol_data: dict,
    cars: list[str],
) -> dict:
    for test_group, fields in tol_data.items():
        if not isinstance(fields, dict):
            continue
        for field_name, cell in fields.items():
            if field_name not in TOLERANCE_PROFILE_FIELDS:
                continue
            # 각 셀은 {ref, sign, tol, lower} 형태. tol 은 ±/+/- 의 공차폭이거나
            # '편측'의 상한(+) 편차, lower 는 '편측'의 하한(-) 편차. 과거 형식
            # (스칼라)도 tol 로 받아들인다.
            if isinstance(cell, dict):
                ref_value = _float_or_none(cell.get("ref"))
                tol_value = _float_or_none(cell.get("tol"))
                tol_lower = _float_or_none(cell.get("lower"))
                sign = str(cell.get("sign") or "±")
            else:
                ref_value = None
                tol_value = _float_or_none(cell)
                tol_lower = None
                sign = "±"
            if sign not in ("±", "+", "-", "편측"):
                sign = "±"
            if sign != "편측":
                tol_lower = None
            _upsert_tolerance_profile(
                db,
                workspace_id=workspace_id,
                category=category,
                refrigerant=refrigerant,
                capacity=capacity,
                profile_name=profile_name,
                test_group=str(test_group),
                field_name=field_name,
                ref_value=ref_value,
                sign=sign,
                tol_value=tol_value,
                tol_lower=tol_lower,
            )
    if profile_name != "BASE":
        db.execute(
            delete(DataVizPerfToleranceProfileCar).where(
                DataVizPerfToleranceProfileCar.workspace_id == workspace_id,
                DataVizPerfToleranceProfileCar.category == category,
                DataVizPerfToleranceProfileCar.refrigerant == refrigerant,
                DataVizPerfToleranceProfileCar.capacity == capacity,
                DataVizPerfToleranceProfileCar.profile_name == profile_name,
            )
        )
        for car_model in cars:
            car_model = car_model.strip()
            if car_model:
                db.add(
                    DataVizPerfToleranceProfileCar(
                        id=_uuid(),
                        workspace_id=workspace_id,
                        category=category,
                        refrigerant=refrigerant,
                        capacity=capacity,
                        profile_name=profile_name,
                        car_model=car_model,
                    )
                )
    db.commit()
    return {"status": "ok"}


def set_car_model_for_source(
    db: Session,
    *,
    workspace_id: str,
    source_file: str,
    car_model: str,
) -> dict:
    rows = list(
        db.scalars(
            select(DataVizPerfData).where(
                DataVizPerfData.workspace_id == workspace_id,
                DataVizPerfData.source_file == source_file,
            )
        )
    )
    for row in rows:
        row.car_model = car_model
    db.commit()
    return {"status": "ok", "updated": len(rows)}


def build_result_workbook(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    all_refrigerant: bool,
    no_warning_color: bool,
) -> bytes:
    rows = _download_rows(
        db,
        workspace_id=workspace_id,
        category=category,
        refrigerant=refrigerant,
        capacity=capacity,
        all_refrigerant=all_refrigerant,
    )
    data_rows = [serialize_perf_data_row(row) for row in rows]
    warning_cache: dict[str, dict[str, str]] = {}
    profiles_by_scope: dict[tuple[str, str], dict] = {}
    for item in data_rows:
        item_ref = item.get("refrigerant") or refrigerant
        item_cap = item.get("capacity") or capacity
        refs = _get_ref_values(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=item_ref,
            capacity=item_cap,
            car_model=item.get("car_model") or "",
        )
        scope_key = (item_ref, item_cap)
        if scope_key not in profiles_by_scope:
            profiles_by_scope[scope_key] = get_all_tolerance_profiles(
                db,
                workspace_id=workspace_id,
                category=category,
                refrigerant=item_ref,
                capacity=item_cap,
            )["profiles"]
        warning_cache[str(item["id"])] = check_warnings(
            item, refs, profiles_by_scope[scope_key]
        )

    wb = Workbook()
    for idx, (sheet_title, sheet_rows) in enumerate(_group_download_sheets(data_rows, category, capacity, all_refrigerant)):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = (sheet_title or "전체")[:31]
        _write_result_sheet(
            ws,
            category=category,
            rows=sheet_rows,
            warning_cache=warning_cache,
            no_warning_color=no_warning_color,
        )
    if not data_rows:
        ws = wb.active
        ws.title = "전체"
        _write_result_sheet(ws, category=category, rows=[], warning_cache={}, no_warning_color=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def serialize_perf_data_row(row: DataVizPerfData) -> dict:
    return {
        "id": row.id,
        "category": row.category,
        "refrigerant": row.refrigerant,
        "capacity": row.capacity,
        "comp_type": row.comp_type,
        "test_group": row.test_group,
        "serial_no": row.serial_no,
        "test_date": row.test_date,
        "car_model": row.car_model,
        "engine_spec": row.engine_spec,
        "rpm": row.rpm,
        "pd": row.pd,
        "td": row.td,
        "ps": row.ps,
        "ts": row.ts,
        "pc": row.pc,
        "mass_flow": row.mass_flow,
        "vol_eff": row.vol_eff,
        "ocr": row.ocr,
        "cooling_cap_a": row.cooling_cap_a,
        "cooling_cap_f": row.cooling_cap_f,
        "power_kw": row.power_kw,
        "cop_sc": row.cop_sc,
        "heat_balance": row.heat_balance,
        "torque": row.torque,
        "remarks": row.remarks,
        "source_file": row.source_file,
    }


def _serialize_row_values(row: PerfRow) -> dict:
    return {
        "test_group": row.test_group,
        "rpm": row.rpm,
        "pd": row.pd,
        "td": row.td,
        "ps": row.ps,
        "ts": row.ts,
        "pc": row.pc,
        "mass_flow": row.mass_flow,
        "vol_eff": row.vol_eff,
        "ocr": row.ocr,
        "cooling_cap_a": row.cooling_cap_a,
        "cooling_cap_f": row.cooling_cap_f,
        "power_kw": row.power_kw,
        "cop_sc": row.cop_sc,
        "heat_balance": row.heat_balance,
        "torque": row.torque,
    }


def _get_workspace_row(db: Session, *, workspace_id: str, row_id: str) -> DataVizPerfData:
    row = db.get(DataVizPerfData, row_id)
    if row is None or row.workspace_id != workspace_id:
        raise KeyError(row_id)
    return row


def _float_or_none(value: object) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return None


def _averages(rows: Iterable[DataVizPerfData], *, group_fields: tuple[str, ...]) -> list[dict]:
    buckets: dict[tuple[str, ...], list[DataVizPerfData]] = defaultdict(list)
    for row in rows:
        buckets[tuple(str(getattr(row, field) or "") for field in group_fields)].append(row)
    out: list[dict] = []
    for key in sorted(buckets):
        members = buckets[key]
        avg = {field: key[index] for index, field in enumerate(group_fields)}
        avg["count"] = len(members)
        for field in NUMERIC_FIELDS:
            vals = [getattr(member, field) for member in members if getattr(member, field) is not None]
            avg[field] = round(sum(vals) / len(vals), 4) if vals else None
        out.append(avg)
    return out


def _ref_values_cached(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
    cache: dict[str, dict[str, dict[str, float | None]]],
) -> dict[str, dict[str, float | None]]:
    # 용량별 기준값(승인도)을 car_model 단위로 캐시해 반환한다.
    if car_model not in cache:
        cache[car_model] = _get_ref_values(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            car_model=car_model,
        )
    return cache[car_model]


def _refs_and_tols(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
    ref_cache: dict[str, dict[str, dict[str, float | None]]],
    tol_cache: dict[str, dict[str, dict[str, float | None]]],
    profile_cache: dict[str, str],
) -> tuple[dict[str, dict[str, float | None]], dict[str, dict[str, float | None]]]:
    if car_model not in ref_cache:
        ref_cache[car_model] = _get_ref_values(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            car_model=car_model,
        )
    if car_model not in profile_cache:
        profile_cache[car_model] = _find_profile_for_car(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            car_model=car_model,
        )
    profile_name = profile_cache[car_model]
    if profile_name not in tol_cache:
        tol_cache[profile_name] = _get_tol_profile_data(
            db,
            workspace_id=workspace_id,
            category=category,
            refrigerant=refrigerant,
            capacity=capacity,
            profile_name=profile_name,
        )
    return ref_cache[car_model], tol_cache[profile_name]


def _get_ref_values(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
) -> dict[str, dict[str, float | None]]:
    rows = list(
        db.scalars(
            select(DataVizPerfTolerance).where(
                DataVizPerfTolerance.workspace_id == workspace_id,
                DataVizPerfTolerance.category == category,
                DataVizPerfTolerance.refrigerant == refrigerant,
                DataVizPerfTolerance.capacity == capacity,
                DataVizPerfTolerance.car_model == car_model,
            )
        )
    )
    result: dict[str, dict[str, float | None]] = {}
    for row in rows:
        result.setdefault(row.test_group, {})[row.field_name] = row.ref_value
    return result


def _get_tolerance_car_models(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
) -> list[str]:
    rows = list(
        db.scalars(
            select(DataVizPerfTolerance.car_model)
            .where(
                DataVizPerfTolerance.workspace_id == workspace_id,
                DataVizPerfTolerance.category == category,
                DataVizPerfTolerance.refrigerant == refrigerant,
                DataVizPerfTolerance.capacity == capacity,
                DataVizPerfTolerance.car_model != "",
            )
            .distinct()
            .order_by(DataVizPerfTolerance.car_model)
        )
    )
    return [row for row in rows if row]


def _get_tol_profiles(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
) -> list[str]:
    rows = list(
        db.scalars(
            select(DataVizPerfToleranceProfile.profile_name)
            .where(
                DataVizPerfToleranceProfile.workspace_id == workspace_id,
                DataVizPerfToleranceProfile.category == category,
                DataVizPerfToleranceProfile.refrigerant == refrigerant,
                DataVizPerfToleranceProfile.capacity == capacity,
            )
            .distinct()
            .order_by(DataVizPerfToleranceProfile.profile_name)
        )
    )
    return ["BASE", *[row for row in rows if row != "BASE"]]


def _get_tol_profile_data(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    profile_name: str,
) -> dict[str, dict[str, dict]]:
    conditions = [
        DataVizPerfToleranceProfile.workspace_id == workspace_id,
        DataVizPerfToleranceProfile.category == category,
        DataVizPerfToleranceProfile.refrigerant == refrigerant,
        DataVizPerfToleranceProfile.profile_name == profile_name,
    ]
    conditions.append(DataVizPerfToleranceProfile.capacity == capacity)
    rows = list(db.scalars(select(DataVizPerfToleranceProfile).where(*conditions)))
    result: dict[str, dict[str, dict]] = {}
    for row in rows:
        result.setdefault(row.test_group, {})[row.field_name] = {
            "ref": row.ref_value,
            "sign": row.sign or "±",
            "tol": row.tol_value,
            "lower": row.tol_lower,
        }
    return result


def _get_tol_profile_cars(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    profile_name: str,
) -> list[str]:
    rows = list(
        db.scalars(
            select(DataVizPerfToleranceProfileCar.car_model)
            .where(
                DataVizPerfToleranceProfileCar.workspace_id == workspace_id,
                DataVizPerfToleranceProfileCar.category == category,
                DataVizPerfToleranceProfileCar.refrigerant == refrigerant,
                DataVizPerfToleranceProfileCar.capacity == capacity,
                DataVizPerfToleranceProfileCar.profile_name == profile_name,
            )
            .order_by(DataVizPerfToleranceProfileCar.car_model)
        )
    )
    return rows


def _find_profile_for_car(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
) -> str:
    row = db.scalar(
        select(DataVizPerfToleranceProfileCar.profile_name).where(
            DataVizPerfToleranceProfileCar.workspace_id == workspace_id,
            DataVizPerfToleranceProfileCar.category == category,
            DataVizPerfToleranceProfileCar.refrigerant == refrigerant,
            DataVizPerfToleranceProfileCar.capacity == capacity,
            DataVizPerfToleranceProfileCar.car_model == car_model,
        )
    )
    return row or "BASE"


def _upsert_tolerance(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    car_model: str,
    test_group: str,
    field_name: str,
    ref_value: float | None,
) -> None:
    existing = db.scalar(
        select(DataVizPerfTolerance).where(
            DataVizPerfTolerance.workspace_id == workspace_id,
            DataVizPerfTolerance.category == category,
            DataVizPerfTolerance.refrigerant == refrigerant,
            DataVizPerfTolerance.capacity == capacity,
            DataVizPerfTolerance.car_model == car_model,
            DataVizPerfTolerance.test_group == test_group,
            DataVizPerfTolerance.field_name == field_name,
        )
    )
    if ref_value is None:
        if existing:
            db.delete(existing)
        return
    if existing:
        existing.ref_value = ref_value
    else:
        db.add(
            DataVizPerfTolerance(
                id=_uuid(),
                workspace_id=workspace_id,
                category=category,
                refrigerant=refrigerant,
                capacity=capacity,
                car_model=car_model,
                test_group=test_group,
                field_name=field_name,
                ref_value=ref_value,
            )
        )


def _upsert_tolerance_profile(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    profile_name: str,
    test_group: str,
    field_name: str,
    ref_value: float | None,
    sign: str,
    tol_value: float | None,
    tol_lower: float | None,
) -> None:
    existing = db.scalar(
        select(DataVizPerfToleranceProfile).where(
            DataVizPerfToleranceProfile.workspace_id == workspace_id,
            DataVizPerfToleranceProfile.category == category,
            DataVizPerfToleranceProfile.refrigerant == refrigerant,
            DataVizPerfToleranceProfile.capacity == capacity,
            DataVizPerfToleranceProfile.profile_name == profile_name,
            DataVizPerfToleranceProfile.test_group == test_group,
            DataVizPerfToleranceProfile.field_name == field_name,
        )
    )
    # 기준값·공차(상한·하한)가 모두 비어 있으면 의미 없는 행이므로 삭제한다.
    if ref_value is None and tol_value is None and tol_lower is None:
        if existing:
            db.delete(existing)
        return
    if existing:
        existing.ref_value = ref_value
        existing.sign = sign
        existing.tol_value = tol_value
        existing.tol_lower = tol_lower
    else:
        db.add(
            DataVizPerfToleranceProfile(
                id=_uuid(),
                workspace_id=workspace_id,
                category=category,
                refrigerant=refrigerant,
                capacity=capacity,
                profile_name=profile_name,
                test_group=test_group,
                field_name=field_name,
                ref_value=ref_value,
                sign=sign,
                tol_lower=tol_lower,
                tol_value=tol_value,
            )
        )


def _download_rows(
    db: Session,
    *,
    workspace_id: str,
    category: str,
    refrigerant: str,
    capacity: str,
    all_refrigerant: bool,
) -> list[DataVizPerfData]:
    conditions = [
        DataVizPerfData.workspace_id == workspace_id,
        DataVizPerfData.category == category,
    ]
    if all_refrigerant:
        pass
    elif capacity:
        conditions.extend(
            [
                DataVizPerfData.refrigerant == refrigerant,
                DataVizPerfData.capacity == capacity,
            ]
        )
    else:
        conditions.append(DataVizPerfData.refrigerant == refrigerant)
    return list(
        db.scalars(
            select(DataVizPerfData)
            .where(*conditions)
            .order_by(
                DataVizPerfData.refrigerant,
                DataVizPerfData.capacity,
                DataVizPerfData.created_at.desc(),
                DataVizPerfData.serial_no,
                DataVizPerfData.test_group,
            )
        )
    )


def _group_download_sheets(
    rows: list[dict],
    category: str,
    capacity: str,
    all_refrigerant: bool,
) -> list[tuple[str, list[dict]]]:
    if all_refrigerant:
        grouped: dict[str, list[dict]] = defaultdict(list)
        for row in rows:
            grouped[f"{row.get('capacity', '')}_{row.get('refrigerant', '')}"].append(row)
        return list(grouped.items())
    if not capacity:
        grouped = defaultdict(list)
        for row in rows:
            grouped[str(row.get("capacity", ""))].append(row)
        return list(grouped.items())
    return [(capacity if category == CATEGORY_VARIABLE else capacity, rows)]


def _write_result_sheet(
    ws,
    *,
    category: str,
    rows: list[dict],
    warning_cache: dict[str, dict[str, str]],
    no_warning_color: bool,
) -> None:
    col_def = _excel_columns(category)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    # 시험공차 위반(tol) → 빨간 글자, 승인도 기준값 위반(ref) → 주황 셀.
    tol_fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
    ref_fill = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")
    header_font = Font(bold=True, size=10, name="Malgun Gothic")
    unit_font = Font(size=9, color="666666", name="Malgun Gothic")
    data_font = Font(size=10, name="Malgun Gothic")
    tol_font = Font(size=10, name="Malgun Gothic", color="DC2626", bold=True)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    for col_idx, (name, unit, _field, _decimals, width) in enumerate(col_def, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
        ws.cell(row=2, column=col_idx, value=unit).fill = header_fill
        ws.cell(row=2, column=col_idx).font = unit_font
        ws.cell(row=2, column=col_idx).alignment = center
        ws.cell(row=2, column=col_idx).border = border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_index, row in enumerate(rows, 3):
        warnings = warning_cache.get(str(row.get("id")), {})
        for col_idx, (name, _unit, field, decimals, _width) in enumerate(col_def, 1):
            cell = ws.cell(row=row_index, column=col_idx)
            cell.border = border
            cell.font = data_font
            if name == "No":
                cell.value = row_index - 2
                cell.alignment = center
                continue
            if not field:
                continue
            value = row.get(field)
            if field == "test_date" and value:
                value = str(value).replace(" 00:00:00", "").replace(" 00:00", "")
            cell.value = value
            cell.alignment = left if field == "remarks" else center
            if decimals is not None and isinstance(value, (int, float)):
                cell.number_format = {0: "#,##0", 1: "0.0", 2: "0.00", 3: "0.000"}[decimals]
            kind = warnings.get(field) if isinstance(warnings, dict) else None
            if kind and not no_warning_color:
                if kind == "tol":
                    cell.fill = tol_fill
                    cell.font = tol_font
                else:  # 'ref' — 승인도 기준값 위반(주황 셀)
                    cell.fill = ref_fill


def _excel_columns(category: str) -> list[tuple[str, str, str | None, int | None, int]]:
    if category == CATEGORY_ELECTRIC:
        return [
            ("No", "", None, None, 5),
            ("기종", "", "comp_type", None, 12),
            ("평가그룹", "", "test_group", None, 9),
            ("S/N", "", "serial_no", None, 11),
            ("RPM", "", "rpm", 0, 8),
            ("Pd", "[kPa G]", "pd", 0, 8),
            ("Ps", "[kPa G]", "ps", 0, 8),
            ("Td", "[°C]", "td", 1, 7),
            ("Ts", "[°C]", "ts", 1, 7),
            ("Cooling Capacity", "[kW]", "cooling_cap_a", 2, 14),
            ("Power", "[kW]", "power_kw", 2, 8),
            ("COP(SC)", "", "cop_sc", 2, 8),
            ("Vol.Eff", "[%]", "vol_eff", 1, 8),
            ("Ref.Massflow", "[kg/h]", "mass_flow", 1, 12),
            ("평가일자", "", "test_date", None, 20),
            ("비고", "", "remarks", None, 35),
        ]
    return [
        ("No", "", None, None, 5),
        ("기종", "", "comp_type", None, 12),
        ("평가그룹", "", "test_group", None, 9),
        ("S/N", "", "serial_no", None, 11),
        ("차종", "", "car_model", None, 8),
        ("엔진사양", "", "engine_spec", None, 16),
        ("RPM", "", "rpm", 0, 8),
        ("Pd", "[kPa]", "pd", 0, 8),
        ("Td", "[°C]", "td", 1, 7),
        ("Ps", "[kPa]", "ps", 0, 8),
        ("Ts", "[°C]", "ts", 2, 7),
        ("Pc", "[kPa]", "pc", 0, 8),
        ("MassFlow", "[kg/h]", "mass_flow", 1, 10),
        ("Vol.Eff", "[%]", "vol_eff", 1, 8),
        ("OCR", "[%]", "ocr", 2, 7),
        ("Cap(A)", "[kW]", "cooling_cap_a", 2, 8),
        ("Power", "[kW]", "power_kw", 2, 8),
        ("COP", "", "cop_sc", 2, 7),
        ("Torque", "[Nm]", "torque", 2, 8),
        ("Cap(F)", "[kW]", "cooling_cap_f", 2, 8),
        ("HeatBal", "[%]", "heat_balance", 2, 8),
        ("평가일자", "", "test_date", None, 20),
        ("비고", "", "remarks", None, 35),
    ]
