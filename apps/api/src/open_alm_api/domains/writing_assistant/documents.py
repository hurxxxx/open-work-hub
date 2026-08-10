"""기안/메일 결과 텍스트를 다운로드 가능한 TXT/DOCX/PDF 바이트로 변환.

원본(`C:\\server\\routes\\document.py`)의 ``download_document`` 구조 파서를 포팅하되,
PDF 한글 폰트는 Windows 전용 ``malgun.ttf`` 대신 reportlab 에 내장된 한국어 CID 폰트
(``HYSMyeongJo-Medium`` / ``HYGothic-Medium``)를 써서 Linux 서버에서도 외부 폰트 파일
없이 동작하게 한다.
"""

from __future__ import annotations

import io
import re

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# 평문 라인을 구분하는 정규식(원본 document.py 와 동일한 규칙)
_RE_DIVIDER = re.compile(r"^[─━\-=]{4,}")
_RE_BRACKET = re.compile(r"^\[.+\]")
_RE_NUM_SECTION = re.compile(r"^\d+\.\s")
_RE_META = re.compile(
    r"^(회의명|일시|장소|참석자|회의\s*기본정보|기안명|기안자|작성일|제목|수신|발신)\s*[:：](.*)$"
)
_RE_SUB_INDENT = re.compile(r"^\s{4,}")
_RE_SUB_BULLET = re.compile(r"^\s{2,}[-·•]")
_RE_BULLET = re.compile(r"^[-·•]\s")


def build_txt(content: str) -> bytes:
    return content.encode("utf-8")


def build_docx(content: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401 (parity with original imports)
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    style.paragraph_format.line_spacing = 1.6

    for raw in content.split("\n"):
        trimmed = raw.strip()
        if not trimmed:
            doc.add_paragraph("")
            continue
        if _RE_DIVIDER.match(trimmed):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("─" * 60)
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            continue
        if _RE_BRACKET.match(trimmed):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(trimmed)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x6B)
            continue
        if _RE_NUM_SECTION.match(trimmed):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(trimmed)
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x6B)
            continue
        meta = _RE_META.match(trimmed)
        if meta:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            run_label = p.add_run(meta.group(1) + ": ")
            run_label.bold = True
            run_label.font.size = Pt(11)
            run_val = p.add_run(meta.group(2).strip())
            run_val.font.size = Pt(11)
            continue
        if _RE_SUB_INDENT.match(raw) or _RE_SUB_BULLET.match(raw):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(2)
            run = p.add_run(trimmed)
            run.font.size = Pt(10)
            continue
        if _RE_BULLET.match(trimmed):
            p = doc.add_paragraph(re.sub(r"^[-·•]\s*", "", trimmed), style="List Bullet")
            p.paragraph_format.left_indent = Cm(1)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(trimmed)
        run.font.size = Pt(11)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


_KOREAN_FONTS_REGISTERED = False


def _ensure_korean_fonts() -> tuple[str, str]:
    """reportlab 내장 한국어 CID 폰트를 등록한다. (본문, 제목용) 폰트명 반환."""
    global _KOREAN_FONTS_REGISTERED
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    body, head = "HYSMyeongJo-Medium", "HYGothic-Medium"
    if not _KOREAN_FONTS_REGISTERED:
        registered = pdfmetrics.getRegisteredFontNames()
        if body not in registered:
            pdfmetrics.registerFont(UnicodeCIDFont(body))
        if head not in registered:
            pdfmetrics.registerFont(UnicodeCIDFont(head))
        _KOREAN_FONTS_REGISTERED = True
    return body, head


def build_pdf(content: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    body_font, head_font = _ensure_korean_fonts()
    buffer = io.BytesIO()
    doc_pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )

    s_title = ParagraphStyle(
        "dt", fontName=head_font, fontSize=16, leading=22, spaceAfter=10,
        textColor=colors.HexColor("#1a3a6b"),
    )
    # 대괄호 제목([제목])과 번호 섹션(1. ...)은 동일한 헤딩 스타일을 공유한다.
    s_section = ParagraphStyle(
        "ds", fontName=head_font, fontSize=13, leading=18, spaceBefore=14, spaceAfter=6,
        textColor=colors.HexColor("#1a3a6b"), backColor=colors.HexColor("#eef2ff"),
        borderPadding=(4, 8, 4, 8),
    )
    s_subsec = ParagraphStyle(
        "dsub", fontName=head_font, fontSize=12, leading=17, spaceBefore=10, spaceAfter=4,
        textColor=colors.HexColor("#2c5282"), leftIndent=8,
    )
    s_meta = ParagraphStyle("dm", fontName=body_font, fontSize=11, leading=16, spaceAfter=2, leftIndent=10)
    s_item = ParagraphStyle("di", fontName=body_font, fontSize=11, leading=17, spaceAfter=2, leftIndent=16)
    s_sub = ParagraphStyle("dsu", fontName=body_font, fontSize=11, leading=16, spaceAfter=1, leftIndent=28)
    s_text = ParagraphStyle("dn", fontName=body_font, fontSize=11, leading=17, spaceAfter=2, leftIndent=10)

    def esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def md(s: str) -> str:
        s = esc(s)
        s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
        # 이탤릭은 공백을 사이에 두지 않은 '*텍스트*' 쌍만 변환한다. '5 * 3 * 2'
        # 같은 평문 곱셈/구분 기호가 통째로 이탤릭으로 묶이는 것을 막는다.
        s = re.sub(r"\*(?=\S)(.+?)(?<=\S)\*", r"<i>\1</i>", s)
        s = re.sub(r"^#{1,6}\s*", "", s)
        return s

    story: list = []
    lines = content.split("\n")
    for i, raw in enumerate(lines):
        trimmed = raw.strip()
        if not trimmed:
            story.append(Spacer(1, 3 * mm))
            continue
        if _RE_DIVIDER.match(trimmed):
            story.append(
                HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#3b5bdb"),
                           spaceBefore=6, spaceAfter=6)
            )
            continue
        if _RE_BRACKET.match(trimmed):
            story.append(Paragraph(md(trimmed), s_section))
            continue
        if _RE_NUM_SECTION.match(trimmed) and not re.match(r"^\d+\.\d+", trimmed):
            story.append(Paragraph(md(trimmed), s_section))
            continue
        if re.match(r"^\d+\.\d+[\.\s]", trimmed):
            story.append(Paragraph(md(trimmed), s_subsec))
            continue
        meta = _RE_META.match(trimmed)
        if meta:
            story.append(Paragraph(f"<b>{esc(meta.group(1))}:</b> {md(meta.group(2).strip())}", s_meta))
            continue
        if _RE_SUB_INDENT.match(raw) or _RE_SUB_BULLET.match(raw):
            clean = re.sub(r"^[-·•]\s*", "", trimmed)
            story.append(Paragraph("◦  " + md(clean), s_sub))
            continue
        if _RE_BULLET.match(trimmed):
            clean = re.sub(r"^[-·•]\s*", "", trimmed)
            story.append(Paragraph("•  " + md(clean), s_item))
            continue
        if i <= 2 and ":" not in trimmed and len(trimmed) < 60:
            story.append(Paragraph(md(trimmed), s_title))
            continue
        story.append(Paragraph(md(trimmed), s_text))

    doc_pdf.build(story)
    return buffer.getvalue()


def render_document(content: str, fmt: str) -> tuple[bytes, str, str]:
    """(bytes, media_type, extension) 반환."""
    if fmt == "txt":
        return build_txt(content), "text/plain; charset=utf-8", "txt"
    if fmt == "docx":
        return build_docx(content), _DOCX_MIME, "docx"
    if fmt == "pdf":
        return build_pdf(content), "application/pdf", "pdf"
    raise ValueError(f"지원하지 않는 형식: {fmt}")
