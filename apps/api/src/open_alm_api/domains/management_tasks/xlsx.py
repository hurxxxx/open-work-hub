from __future__ import annotations

import multiprocessing
import re
import time
import zipfile
from dataclasses import dataclass
from io import BytesIO
from multiprocessing.connection import Connection
from pathlib import PurePosixPath
from xml.etree import ElementTree

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.cell import range_boundaries


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
UPLOAD_READ_CHUNK_BYTES = 1024 * 1024
MAX_ZIP_ENTRIES = 500
MAX_DECOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 5_000
MAX_TOTAL_ROWS = 10_000
MAX_TOTAL_CELLS = 200_000
MAX_CELLS_PER_ROW = 100
MAX_CELL_TEXT_LENGTH = 1_000
MAX_PARSE_SECONDS = 10.0
MAX_DETAIL_ROWS = 200
_PROCESS_EXIT_GRACE_SECONDS = 0.5

_DEPT_HEADERS = {"부서", "부서명"}
_NAME_HEADERS = {"성명", "이름"}
_RELATION_HEADERS = {"관계 및 사번", "관계및사번", "본인여부", "관계"}
_EMPLOYEE_CODE_HEADERS = {"사번", "사원번호", "직원번호", "employee code", "employee_code"}
_DEPENDENT_MARKERS = ("배우자", "가족", "자녀", "부모", "부양")
_SELF_MARKERS = {"본인", "self"}
_FORMULA_LEADS = ("=", "+", "-", "@", "\t", "\r")
_MAX_HEADER_SCAN_ROWS = 8
_STANDARD_WORKBOOK_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)


class UnsafeXlsxError(ValueError):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class ParsedPriorExamRow:
    sheet_name: str
    row_ordinal: int
    dept_name: str
    person_name: str
    raw_relation: str | None
    provided_employee_code: str | None
    is_dependent: bool


async def read_upload_limited(upload: UploadFile) -> bytes:
    content = bytearray()
    while True:
        chunk = await upload.read(UPLOAD_READ_CHUNK_BYTES)
        if not chunk:
            break
        if len(content) + len(chunk) > MAX_UPLOAD_BYTES:
            raise UnsafeXlsxError("file_too_large")
        content.extend(chunk)
    if not content:
        raise UnsafeXlsxError("empty_file")
    return bytes(content)


def _check_deadline(started_at: float) -> None:
    if time.monotonic() - started_at > MAX_PARSE_SECONDS:
        raise UnsafeXlsxError("parse_timeout")


def _validate_archive(content: bytes) -> None:
    if not content.startswith(b"PK") or not zipfile.is_zipfile(BytesIO(content)):
        raise UnsafeXlsxError("invalid_xlsx_container")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise UnsafeXlsxError("too_many_archive_entries")
            names: set[str] = set()
            decompressed = 0
            compressed = 0
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts or entry.filename in names:
                    raise UnsafeXlsxError("invalid_archive_entry")
                names.add(entry.filename)
                if entry.flag_bits & 0x1:
                    raise UnsafeXlsxError("encrypted_workbook")
                decompressed += entry.file_size
                compressed += entry.compress_size
                if decompressed > MAX_DECOMPRESSED_BYTES:
                    raise UnsafeXlsxError("decompressed_size_exceeded")
                if (
                    entry.file_size > 0
                    and entry.file_size > max(entry.compress_size, 1) * MAX_COMPRESSION_RATIO
                ):
                    raise UnsafeXlsxError("compression_ratio_exceeded")

            if decompressed > max(compressed, 1) * MAX_COMPRESSION_RATIO:
                raise UnsafeXlsxError("compression_ratio_exceeded")
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset(names):
                raise UnsafeXlsxError("invalid_xlsx_schema")
            lowered = {name.lower() for name in names}
            if any(
                name.startswith("xl/externallinks/")
                or name.endswith("vbaproject.bin")
                or name.startswith("xl/embeddings/")
                for name in lowered
            ):
                raise UnsafeXlsxError("active_or_external_content")

            content_types = archive.read("[Content_Types].xml")
            if (
                _STANDARD_WORKBOOK_CONTENT_TYPE.encode() not in content_types
                or b"macroEnabled" in content_types
            ):
                raise UnsafeXlsxError("unsupported_workbook_type")
            for name in names:
                if not name.endswith(".rels"):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as error:
                    raise UnsafeXlsxError("invalid_relationship_xml") from error
                for relationship in root:
                    if relationship.attrib.get("TargetMode", "").lower() == "external":
                        raise UnsafeXlsxError("active_or_external_content")
    except zipfile.BadZipFile as error:
        raise UnsafeXlsxError("invalid_xlsx_container") from error


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\ufeff", "").strip()


def _normalized_header(value: object) -> str:
    return re.sub(r"\s+", " ", _text(value)).lower()


def _find_headers(rows: list[list[str]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows[:_MAX_HEADER_SCAN_ROWS]):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row):
            header = _normalized_header(value)
            if "dept" not in mapping and header in _DEPT_HEADERS:
                mapping["dept"] = column_index
            elif "name" not in mapping and header in _NAME_HEADERS:
                mapping["name"] = column_index
            elif "relation" not in mapping and header in _RELATION_HEADERS:
                mapping["relation"] = column_index
            elif "employee_code" not in mapping and header in _EMPLOYEE_CODE_HEADERS:
                mapping["employee_code"] = column_index
        if "dept" in mapping and "name" in mapping:
            return row_index, mapping
    return None


def _value_at(row: list[str], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index]


def _relation_is_dependent(value: str, *, self_flag_column: bool) -> bool:
    normalized = value.strip().lower()
    if not normalized:
        return False
    if self_flag_column:
        return normalized not in _SELF_MARKERS
    return any(marker in normalized for marker in _DEPENDENT_MARKERS)


def _relation_employee_code(value: str, *, self_flag_column: bool) -> str | None:
    normalized = value.strip()
    if not normalized or self_flag_column or normalized.casefold() in _SELF_MARKERS:
        return None
    if any(marker in normalized for marker in _DEPENDENT_MARKERS):
        return None
    return normalized


def _validated_merge_ranges(
    content: bytes,
    *,
    worksheet_path: str,
) -> list[tuple[int, int, int, int]]:
    """Read and bound worksheet merge ranges before openpyxl enumerates rows."""

    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            root = ElementTree.fromstring(archive.read(worksheet_path))
    except (KeyError, ElementTree.ParseError, zipfile.BadZipFile) as error:
        raise UnsafeXlsxError("invalid_xlsx_workbook") from error

    ranges: list[tuple[int, int, int, int]] = []
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1] != "mergeCell":
            continue
        reference = element.attrib.get("ref", "")
        try:
            min_column, min_row, max_column, max_row = range_boundaries(reference)
        except (TypeError, ValueError) as error:
            raise UnsafeXlsxError("invalid_xlsx_workbook") from error
        if (
            not all(
                isinstance(boundary, int)
                for boundary in (min_column, min_row, max_column, max_row)
            )
            or min_column < 1
            or min_row < 1
            or max_column < min_column
            or max_row < min_row
        ):
            raise UnsafeXlsxError("invalid_xlsx_workbook")
        # Merge ranges come from worksheet XML and are not constrained by the
        # cells yielded by openpyxl. Reject an out-of-contract range before
        # expanding it, otherwise an absurd max_row can consume unbounded CPU
        # and memory even though the visible sheet itself is tiny.
        if max_row > MAX_ROWS_PER_SHEET:
            raise UnsafeXlsxError("too_many_rows")
        if max_column > MAX_CELLS_PER_ROW:
            raise UnsafeXlsxError("too_many_cells")
        ranges.append((min_column, min_row, max_column, max_row))
    return ranges


def _merged_department_sources(
    merge_ranges: list[tuple[int, int, int, int]],
    *,
    department_column: int,
) -> dict[int, int]:
    """Map covered 1-based row numbers to their merged department source row."""

    sources: dict[int, int] = {}
    expected_column = department_column + 1
    for min_column, min_row, max_column, max_row in merge_ranges:
        if min_column != expected_column or max_column != expected_column or min_row >= max_row:
            continue
        for row_number in range(min_row + 1, max_row + 1):
            if row_number in sources:
                raise UnsafeXlsxError("invalid_xlsx_workbook")
            sources[row_number] = min_row
    return sources


def parse_prior_exam_workbook(content: bytes) -> list[ParsedPriorExamRow]:
    started_at = time.monotonic()
    _validate_archive(content)
    _check_deadline(started_at)
    try:
        # data_only=False keeps each cell's data_type, so we can tell a formula
        # from a plain value. Benign computed columns (row numbers, cost sums) are
        # allowed because we never read them; but a formula in an identity/relation
        # column (부서/이름/사번/본인여부) is rejected below — its cached value can be
        # missing or stale and would silently misclassify a row (e.g. a spouse read
        # as an employee). Identity columns are plain text in real rosters.
        workbook = load_workbook(
            BytesIO(content),
            data_only=False,
            read_only=True,
            keep_links=False,
        )
    except Exception as error:  # noqa: BLE001 - map parser internals to a stable code
        raise UnsafeXlsxError("invalid_xlsx_workbook") from error
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise UnsafeXlsxError("too_many_sheets")
        parsed: list[ParsedPriorExamRow] = []
        total_rows = 0
        total_cells = 0
        row_ordinal = 0
        for worksheet in workbook.worksheets:
            _check_deadline(started_at)
            worksheet_path = getattr(worksheet, "_worksheet_path", "")
            if not isinstance(worksheet_path, str) or not worksheet_path:
                raise UnsafeXlsxError("invalid_xlsx_workbook")
            merge_ranges = _validated_merge_ranges(
                content,
                worksheet_path=worksheet_path,
            )
            worksheet.reset_dimensions()
            rows: list[list[str]] = []
            formula_rows: list[list[bool]] = []
            for sheet_row_number, cells in enumerate(worksheet.iter_rows(), start=1):
                if sheet_row_number > MAX_ROWS_PER_SHEET:
                    raise UnsafeXlsxError("too_many_rows")
                if len(cells) > MAX_CELLS_PER_ROW:
                    raise UnsafeXlsxError("too_many_cells")
                total_rows += 1
                total_cells += len(cells)
                if total_rows > MAX_TOTAL_ROWS or total_cells > MAX_TOTAL_CELLS:
                    raise UnsafeXlsxError("workbook_size_exceeded")
                values: list[str] = []
                is_formula: list[bool] = []
                for cell in cells:
                    formula = cell.data_type == "f"
                    # A formula cell's .value is the formula text here; we only use
                    # it to flag identity columns, never as a real value.
                    value = "" if formula else _text(cell.value)
                    if len(value) > MAX_CELL_TEXT_LENGTH:
                        raise UnsafeXlsxError("cell_text_too_long")
                    values.append(value)
                    is_formula.append(formula)
                rows.append(values)
                formula_rows.append(is_formula)
                if sheet_row_number % 100 == 0:
                    _check_deadline(started_at)

            located = _find_headers(rows)
            if located is None:
                continue
            header_index, mapping = located
            relation_index = mapping.get("relation")
            relation_header = _normalized_header(_value_at(rows[header_index], relation_index))
            self_flag_column = relation_header in {"본인여부", "관계"}
            merged_department_sources = _merged_department_sources(
                merge_ranges,
                department_column=mapping["dept"],
            )
            # Columns whose values decide identity/relation must be plain data; a
            # formula there is rejected fail-closed (see load_workbook note).
            identity_indexes = [
                index
                for index in (
                    mapping["dept"],
                    mapping["name"],
                    relation_index,
                    mapping.get("employee_code"),
                )
                if index is not None
            ]
            for data_offset, row in enumerate(rows[header_index + 1 :]):
                flags = formula_rows[header_index + 1 + data_offset]
                if any(index < len(flags) and flags[index] for index in identity_indexes):
                    raise UnsafeXlsxError("formula_in_identity_column")
                dept_cell = _value_at(row, mapping["dept"])
                person_name = _value_at(row, mapping["name"])
                if not person_name:
                    continue
                relation = _value_at(row, relation_index)
                is_dependent = _relation_is_dependent(relation, self_flag_column=self_flag_column)
                if dept_cell:
                    dept_name = dept_cell
                else:
                    row_number = header_index + 2 + data_offset
                    source_row_number = merged_department_sources.get(row_number)
                    dept_name = (
                        _value_at(rows[source_row_number - 1], mapping["dept"])
                        if (
                            is_dependent
                            and source_row_number is not None
                            and source_row_number > header_index + 1
                            and all(
                                _value_at(rows[covered_row - 1], mapping["name"])
                                for covered_row in range(source_row_number, row_number)
                            )
                        )
                        else ""
                    )
                if not dept_name:
                    raise UnsafeXlsxError("missing_department")
                employee_code = _value_at(row, mapping.get("employee_code")) or None
                if employee_code is None:
                    employee_code = _relation_employee_code(
                        relation, self_flag_column=self_flag_column
                    )
                row_ordinal += 1
                parsed.append(
                    ParsedPriorExamRow(
                        sheet_name=worksheet.title,
                        row_ordinal=row_ordinal,
                        dept_name=dept_name,
                        person_name=person_name,
                        raw_relation=relation or None,
                        provided_employee_code=employee_code,
                        is_dependent=is_dependent,
                    )
                )
        if not parsed:
            raise UnsafeXlsxError("no_supported_rows")
        return parsed
    finally:
        workbook.close()


def _parse_prior_exam_workbook_worker(content: bytes, writer: Connection) -> None:
    """Parse inside a killable process and return only the bounded typed result."""

    try:
        writer.send(("ok", parse_prior_exam_workbook(content)))
    except UnsafeXlsxError as error:
        writer.send(("unsafe", error.code))
    except Exception:  # noqa: BLE001 - do not expose parser internals across the boundary
        writer.send(("error", "parser_failed"))
    finally:
        writer.close()


def _stop_parse_process(
    process: multiprocessing.Process,
    *,
    force: bool,
) -> None:
    if not force:
        process.join(timeout=_PROCESS_EXIT_GRACE_SECONDS)
    if not process.is_alive():
        return
    process.terminate()
    process.join(timeout=_PROCESS_EXIT_GRACE_SECONDS)
    if process.is_alive():
        process.kill()
        process.join(timeout=_PROCESS_EXIT_GRACE_SECONDS)


def parse_prior_exam_workbook_isolated(
    content: bytes,
    *,
    timeout_seconds: float | None = None,
) -> list[ParsedPriorExamRow]:
    """Parse in a spawn process so timeout can terminate workbook load and traversal."""

    timeout = MAX_PARSE_SECONDS if timeout_seconds is None else timeout_seconds
    if timeout <= 0:
        raise ValueError("timeout_seconds must be positive")

    context = multiprocessing.get_context("spawn")
    reader, writer = context.Pipe(duplex=False)
    process = context.Process(
        target=_parse_prior_exam_workbook_worker,
        args=(content, writer),
        daemon=True,
    )
    started = False
    message_received = False
    try:
        process.start()
        started = True
        writer.close()
        if not reader.poll(timeout):
            raise UnsafeXlsxError("parse_timeout")
        try:
            result_kind, payload = reader.recv()
        except EOFError as error:
            raise UnsafeXlsxError("parser_failed") from error
        message_received = True

        if result_kind == "ok":
            return payload
        if result_kind == "unsafe":
            raise UnsafeXlsxError(payload)
        raise UnsafeXlsxError("parser_failed")
    finally:
        reader.close()
        if not started:
            writer.close()
        else:
            _stop_parse_process(process, force=not message_received)


def escape_formula(value: object) -> object:
    # A lone hyphen is used as a "no" placeholder and is inert in a spreadsheet,
    # so keep it clean; every other formula-leading value is still neutralized.
    if isinstance(value, str) and value != "-" and value.startswith(_FORMULA_LEADS):
        return f"'{value}"
    return value


def build_employee_table_xlsx(
    *, sheet_title: str, headers: list[str], rows: list[list[object]]
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title[:31] or "목록"
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="E8EEF7")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        worksheet.append([escape_formula(value) for value in row])
    worksheet.freeze_panes = "A2"
    for column_cells in worksheet.columns:
        length = max((len(str(cell.value or "")) for cell in column_cells), default=0)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(length + 2, 8), 40
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_target_roster_xlsx(
    *,
    target_year: int,
    rows: list[dict],
    settings: dict[str, object],
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "전체명단"
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    worksheet["A1"] = f"{target_year}년 종합검진 대상자 명단"
    worksheet["A1"].font = Font(bold=True, size=16)
    age_label = "한국식 나이" if settings["age_calc_method"] == "korean" else "만 나이"
    worksheet["D1"] = (
        f"{age_label} {settings['adult_age']}세 이상 또는 "
        f"근속 {settings['service_years_threshold']}년 이상: 2년 1회"
    )
    worksheet["D2"] = f"{age_label} {settings['senior_age']}세 이상: 매년"
    for coordinate in ("D1", "D2"):
        worksheet[coordinate].font = Font(bold=True)
        worksheet[coordinate].alignment = Alignment(
            horizontal="left",
            vertical="center",
        )
    header_row = 5
    for column, title in enumerate(["No.", "부서명", "사번", "이름"], start=1):
        cell = worksheet.cell(row=header_row, column=column, value=title)
        cell.fill = PatternFill("solid", fgColor="B8CCE4")
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = box
    no_fill = PatternFill("solid", fgColor="DCE6F1")
    name_fill = PatternFill("solid", fgColor="FFF2CC")
    for index, row in enumerate(rows, start=1):
        values = [index, row["dept_name"], row["employee_code"], row["name"]]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(
                row=header_row + index,
                column=column,
                value=escape_formula(value),
            )
            cell.border = box
            if column == 1:
                cell.fill = no_fill
                cell.alignment = center
            elif column == 3:
                cell.alignment = center
            elif column == 4:
                cell.fill = name_fill
    last_row = max(header_row, header_row + len(rows))
    worksheet.auto_filter.ref = f"A{header_row}:D{last_row}"
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.column_dimensions["A"].width = 6
    worksheet.column_dimensions["B"].width = 30
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 14
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
