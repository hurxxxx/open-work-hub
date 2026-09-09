from __future__ import annotations

import zipfile
from collections.abc import Iterable
from io import BytesIO
from pathlib import Path

from .contracts import (
    DocumentExtractBundle,
    DocumentExtractor,
    EvidenceBlock,
    UnsupportedDocumentType,
    _ExtractionBudget,
)
from .html_extractor import HtmlExtractor
from .pptx import PptxOpenXmlExtractor, archive_exceeds_limits

_MAX_PDF_PAGES = 80
_MAX_DOCX_PARAGRAPHS = 2_000
_MAX_DOCX_TABLE_ROWS = 5_000
_MAX_XLSX_SHEETS = 20
_MAX_XLSX_ROWS_PER_SHEET = 5_000
_MAX_XLSX_TOTAL_ROWS = 20_000
_MAX_CELL_CHARS = 2_000

# Office packages can embed other Office files (OLE objects) under */embeddings/.
_EMBEDDING_PREFIXES = ("ppt/embeddings/", "word/embeddings/", "xl/embeddings/")
_EMBEDDABLE_SUFFIXES = frozenset({".pptx", ".docx", ".xlsx", ".xlsm"})
_MAX_EMBEDDED_DOCUMENTS = 20

__all__ = [
    "DocumentExtractBundle",
    "DocumentExtractor",
    "EvidenceBlock",
    "HtmlExtractor",
    "PptxOpenXmlExtractor",
    "UnsupportedDocumentType",
    "default_extractors",
    "extract_document",
    "extract_embedded_office_documents",
]


def _cell_text(value: object) -> str:
    return ("" if value is None else str(value))[:_MAX_CELL_CHARS]


class PdfPlumberExtractor:
    _CONTENT_TYPES = {"application/pdf"}

    def supports(self, *, mime_type: str, filename: str) -> bool:
        return Path(filename).suffix.lower() == ".pdf" or mime_type.lower() in self._CONTENT_TYPES

    def extract(
        self, *, document_id: str, filename: str, mime_type: str, content: bytes
    ) -> DocumentExtractBundle:
        import pdfplumber

        blocks: list[EvidenceBlock] = []
        budget = _ExtractionBudget()
        try:
            with pdfplumber.open(BytesIO(content)) as pdf:
                for page_number, page in enumerate(pdf.pages[:_MAX_PDF_PAGES], start=1):
                    if not budget.should_continue():
                        break
                    text = (page.extract_text() or "").strip()
                    if not text:
                        continue
                    if not budget.append(
                        blocks,
                        EvidenceBlock(
                            document_id=document_id,
                            block_id=f"{document_id}:p{page_number}",
                            locator_kind="page",
                            locator_label=f"p.{page_number}",
                            section_path=f"Page {page_number}",
                            block_kind="text",
                            text=text,
                        ),
                    ):
                        break
        except UnsupportedDocumentType:
            raise
        except Exception as exc:  # noqa: BLE001 — surface as unsupported/parse failure
            raise UnsupportedDocumentType("PDF could not be parsed") from exc
        return DocumentExtractBundle(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type,
            evidence_blocks=blocks,
        )


class DocxExtractor:
    _CONTENT_TYPES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    }

    def supports(self, *, mime_type: str, filename: str) -> bool:
        return Path(filename).suffix.lower() == ".docx" or mime_type.lower() in self._CONTENT_TYPES

    def extract(
        self, *, document_id: str, filename: str, mime_type: str, content: bytes
    ) -> DocumentExtractBundle:
        import docx

        blocks: list[EvidenceBlock] = []
        try:
            document = docx.Document(BytesIO(content))
        except Exception as exc:  # noqa: BLE001
            raise UnsupportedDocumentType("DOCX could not be parsed") from exc
        budget = _ExtractionBudget()
        for index, paragraph in enumerate(
            document.paragraphs[:_MAX_DOCX_PARAGRAPHS],
            start=1,
        ):
            if not budget.should_continue():
                break
            text = (paragraph.text or "").strip()
            if not text:
                continue
            if not budget.append(
                blocks,
                EvidenceBlock(
                    document_id=document_id,
                    block_id=f"{document_id}:para{index}",
                    locator_kind="paragraph",
                    locator_label=f"¶{index}",
                    section_path="Body",
                    block_kind="text",
                    text=text,
                ),
            ):
                break
        table_rows_seen = 0
        malformed_table_rows_skipped = 0
        for table_index, table in enumerate(document.tables, start=1):
            if not budget.should_continue() or table_rows_seen >= _MAX_DOCX_TABLE_ROWS:
                break
            rows: list[list[str]] = []
            for row in table.rows:
                if table_rows_seen + len(rows) >= _MAX_DOCX_TABLE_ROWS:
                    break
                try:
                    cells = [(cell.text or "").strip()[:_MAX_CELL_CHARS] for cell in row.cells]
                except ValueError:
                    # Some Word producers emit a vertical-merge continuation whose
                    # preceding row omits the referenced grid cell. python-docx
                    # raises while resolving row.cells; retain all other evidence.
                    malformed_table_rows_skipped += 1
                    continue
                if any(cell for cell in cells):
                    rows.append(cells)
            if not rows:
                continue
            table_rows_seen += len(rows)
            if not budget.append(
                blocks,
                EvidenceBlock(
                    document_id=document_id,
                    block_id=f"{document_id}:table{table_index}",
                    locator_kind="table",
                    locator_label=f"Table {table_index}",
                    section_path="Body",
                    block_kind="table",
                    text="\n".join("\t".join(row) for row in rows),
                    rows=rows,
                ),
            ):
                break
        return DocumentExtractBundle(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type,
            evidence_blocks=blocks,
            metadata={
                "malformed_table_rows_skipped": malformed_table_rows_skipped,
            },
        )


class XlsxExtractor:
    _CONTENT_TYPES = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }

    def supports(self, *, mime_type: str, filename: str) -> bool:
        suffix = Path(filename).suffix.lower()
        return suffix in {".xlsx", ".xlsm"} or mime_type.lower() in self._CONTENT_TYPES

    def extract(
        self, *, document_id: str, filename: str, mime_type: str, content: bytes
    ) -> DocumentExtractBundle:
        from openpyxl import load_workbook

        blocks: list[EvidenceBlock] = []
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UnsupportedDocumentType("XLSX could not be parsed") from exc
        budget = _ExtractionBudget()
        total_rows = 0
        try:
            for sheet in workbook.worksheets[:_MAX_XLSX_SHEETS]:
                if not budget.should_continue() or total_rows >= _MAX_XLSX_TOTAL_ROWS:
                    break
                rows: list[list[str]] = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if (
                        row_index > _MAX_XLSX_ROWS_PER_SHEET
                        or total_rows >= _MAX_XLSX_TOTAL_ROWS
                        or not budget.should_continue()
                    ):
                        break
                    cells = [_cell_text(value) for value in row]
                    if any(cell.strip() for cell in cells):
                        rows.append(cells)
                        total_rows += 1
                if not rows:
                    continue
                if not budget.append(
                    blocks,
                    EvidenceBlock(
                        document_id=document_id,
                        block_id=f"{document_id}:sheet:{sheet.title}",
                        locator_kind="sheet",
                        locator_label=sheet.title,
                        section_path=sheet.title,
                        block_kind="table",
                        text="\n".join("\t".join(row) for row in rows),
                        rows=rows,
                    ),
                ):
                    break
        finally:
            workbook.close()
        return DocumentExtractBundle(
            document_id=document_id,
            filename=filename,
            mime_type=mime_type,
            evidence_blocks=blocks,
        )


def default_extractors() -> list[DocumentExtractor]:
    return [
        HtmlExtractor(),
        PptxOpenXmlExtractor(),
        PdfPlumberExtractor(),
        DocxExtractor(),
        XlsxExtractor(),
    ]


def extract_document(
    *,
    document_id: str,
    filename: str,
    mime_type: str,
    content: bytes,
    extractors: Iterable[DocumentExtractor] | None = None,
) -> DocumentExtractBundle:
    for extractor in extractors or default_extractors():
        if extractor.supports(mime_type=mime_type, filename=filename):
            return extractor.extract(
                document_id=document_id,
                filename=filename,
                mime_type=mime_type,
                content=content,
            )
    raise UnsupportedDocumentType(f"Unsupported document type: {filename or mime_type}")


def extract_embedded_office_documents(content: bytes) -> list[tuple[str, str]]:
    """Extract text from Office files embedded (OLE) inside an Office package.

    Scans the top-level package's ``*/embeddings/`` parts for supported Office
    formats (pptx/docx/xlsx) and runs each through the standard extractors. Only
    one level deep — embedded files are not themselves recursed into — and bounded
    by a document-count cap plus each extractor's own char/time budget. A bad or
    unreadable embedded part is skipped, never failing the parent extraction.
    Returns ``[(filename, text), ...]``; empty for non-zip inputs (e.g. PDF).
    """
    results: list[tuple[str, str]] = []
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if archive_exceeds_limits(archive.infolist()):
                return results
            names = sorted(
                name
                for name in archive.namelist()
                if name.startswith(_EMBEDDING_PREFIXES)
                and Path(name).suffix.lower() in _EMBEDDABLE_SUFFIXES
            )
            for name in names[:_MAX_EMBEDDED_DOCUMENTS]:
                filename = Path(name).name
                # Guard each embedded member's uncompressed size / expansion
                # ratio before decompressing it — a small member can expand into
                # a zip bomb, and archive.read() would allocate it in full before
                # any extractor budget applies.
                if archive_exceeds_limits([archive.getinfo(name)]):
                    continue
                try:
                    bundle = extract_document(
                        document_id=f"embedded:{filename}",
                        filename=filename,
                        mime_type="",
                        content=archive.read(name),
                    )
                except UnsupportedDocumentType:
                    continue
                except Exception:  # noqa: BLE001 - a bad embedded part must not fail the parent
                    continue
                text = bundle.normalized_text.strip()
                if text:
                    results.append((filename, text))
    except (zipfile.BadZipFile, zipfile.LargeZipFile, RuntimeError):
        return results
    return results
