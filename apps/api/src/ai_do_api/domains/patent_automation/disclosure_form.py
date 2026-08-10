"""Parse a 직무발명신고서 (employee-invention disclosure) xlsx form into
patent_records field values.

The form is the company's standard 직무발명신고서 template: a labelled grid
where each label cell sits to the LEFT of its (merged) value cell on the same
row — e.g. ``발명의 명칭`` → title, ``소속부서`` → team, ``국문`` → inventor
names (newline-separated for multiple inventors). We anchor on the label text
and read the first non-empty cell to its right, which tolerates minor template
shifts better than hard-coded coordinates.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

# Prefer the sheet whose title looks like the disclosure form itself (not the
# 작성예시 example sheets or the 명세서 body).
_SHEET_HINT = "신고서"
_SCAN_ROWS = 40


def _norm(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("​", "").split())


def _pick_sheet(wb):
    for ws in wb.worksheets:
        if _SHEET_HINT in (ws.title or ""):
            return ws
    return wb.worksheets[0]


def _value_right_of(ws, label_subs: tuple[str, ...]) -> str | None:
    """First non-empty cell to the right of the first label cell that contains
    any of ``label_subs`` (normalized, whitespace-collapsed)."""
    for row in ws.iter_rows(min_row=1, max_row=_SCAN_ROWS):
        for idx, cell in enumerate(row):
            text = _norm(cell.value)
            if text and any(sub in text for sub in label_subs):
                for nxt in row[idx + 1 :]:
                    if nxt.value is not None and str(nxt.value).strip():
                        return str(nxt.value)
    return None


def parse_disclosure_form(content: bytes) -> dict[str, str]:
    """Extract domestic patent_records field values from a disclosure form.

    Returns keys among: invention_title, inventors, submit_team,
    submitter_name. Inventor-only details (주민등록번호/사번/연락처) have no
    column in 특허현황관리 and are intentionally dropped.
    """
    try:
        wb = load_workbook(BytesIO(content), data_only=True)
    except Exception:  # noqa: BLE001
        return {}
    ws = _pick_sheet(wb)

    title = _value_right_of(ws, ("발명의 명칭", "발명의명칭"))
    names_raw = _value_right_of(ws, ("국문",))
    team = _value_right_of(ws, ("소속부서",))

    values: dict[str, str] = {}
    if title:
        values["invention_title"] = _norm(title)
    if names_raw:
        # Multiple inventors are newline-separated; internal names stay plain,
        # external names keep their (괄호) form verbatim.
        parts = [p.strip() for p in str(names_raw).replace("\r", "").split("\n") if p.strip()]
        if parts:
            values["inventors"] = ",".join(parts)
            values["submitter_name"] = parts[0]
    if team:
        values["submit_team"] = _norm(team)
    return values
