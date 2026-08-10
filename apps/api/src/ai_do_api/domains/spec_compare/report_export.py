"""Render spec-compare results as downloadable Word (.docx) and PDF documents.

Both builders work from the structured result payload (summary + comparison
rows) rather than parsing the Markdown report, so the verdict column can be
colored consistently with the web UI's status badges. The AI analysis prose,
when present, is lifted from the Markdown report as plain paragraphs.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import logging
import os
import re
from xml.sax.saxutils import escape as _xml_escape

logger = logging.getLogger(__name__)


_STATUS_LABELS: dict[str, str] = {
    "same": "동일",
    "different": "상이",
    "base_only": "기준만",
    "target_only": "비교만",
    "unknown": "판단 보류",
}

# Fill / text hex pairs kept in step with the web RowStatusBadge palette.
_STATUS_COLORS: dict[str, tuple[str, str]] = {
    "same": ("D1FAE5", "065F46"),
    "different": ("FEE2E2", "991B1B"),
    "base_only": ("DBEAFE", "1E40AF"),
    "target_only": ("EDE9FE", "5B21B6"),
    "unknown": ("FEF3C7", "92400E"),
}
_DEFAULT_COLORS = ("F1F5F9", "334155")

# Detail comparison table columns, shared by the DOCX and PDF builders so both
# outputs stay structurally identical.
_DETAIL_COLUMNS = ["판정", "사양 항목", "기준 문서", "비교 문서", "설명", "근거"]

# Candidate Korean TTF/OTF fonts, in priority order. The env override lets prod
# point at a bundled font; the rest cover Windows dev and common Linux packages.
_KOREAN_FONT_CANDIDATES: tuple[str | None, ...] = (
    os.environ.get("AI_DO_SPEC_COMPARE_PDF_FONT"),
    r"C:\Windows\Fonts\malgun.ttf",
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    "/usr/share/fonts/opentype/noto/NotoSansCJKkr-Regular.otf",
    "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
)
_KOREAN_BOLD_CANDIDATES: tuple[str | None, ...] = (
    os.environ.get("AI_DO_SPEC_COMPARE_PDF_FONT_BOLD"),
    r"C:\Windows\Fonts\malgunbd.ttf",
    "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf",
)


@dataclass(frozen=True)
class RenderedReport:
    content: bytes
    media_type: str
    extension: str


def _status_label(status: str) -> str:
    return _STATUS_LABELS.get(status, status or "판단 보류")


def _summary_int(summary: dict[str, object], key: str) -> int:
    value = summary.get(key)
    return value if isinstance(value, int) else 0


def _rows_from_payload(payload: dict[str, object]) -> list[dict[str, object]]:
    rows = payload.get("comparison_rows")
    return [row for row in rows if isinstance(row, dict)] if isinstance(rows, list) else []


def _evidence_text(row: dict[str, object]) -> str:
    ids: list[str] = []
    for key in ("base_evidence_ids", "target_evidence_ids"):
        value = row.get(key)
        if isinstance(value, list):
            ids.extend(str(item) for item in value)
    return ", ".join(ids)


def _extract_ai_analysis(report_markdown: str) -> list[str]:
    """Return the '## AI 분석 요약' section as cleaned plain-text lines."""
    if not report_markdown or "## AI 분석 요약" not in report_markdown:
        return []
    after = report_markdown.split("## AI 분석 요약", 1)[1]
    section = re.split(r"\n##\s", after, 1)[0]
    lines: list[str] = []
    for raw in section.splitlines():
        text = raw.strip()
        if not text:
            continue
        text = re.sub(r"^[#\-*>\s]+", "", text)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = text.replace("|", " ")
        if text:
            lines.append(text)
    return lines


def _summary_pairs(summary: dict[str, object]) -> list[tuple[str, str]]:
    return [
        ("보고서 표시 항목", str(_summary_int(summary, "total_rows"))),
        ("동일", str(_summary_int(summary, "same"))),
        ("상이", str(_summary_int(summary, "different"))),
        ("기준 문서에만 있음", str(_summary_int(summary, "base_only"))),
        ("비교 문서에만 있음", str(_summary_int(summary, "target_only"))),
        ("판단 보류", str(_summary_int(summary, "unknown"))),
    ]


def render_report_docx(
    *,
    title: str,
    base_filename: str,
    target_filename: str,
    summary: dict[str, object],
    payload: dict[str, object],
) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    def shade_cell(cell, fill_hex: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill_hex)
        tc_pr.append(shd)

    document = Document()
    for section in document.sections:
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
    document.add_heading("규격서 비교 보고서", level=0)
    if title:
        document.add_paragraph(title)

    document.add_heading("비교 대상", level=1)
    document.add_paragraph(f"기준 문서: {base_filename}")
    document.add_paragraph(f"비교 문서: {target_filename}")

    document.add_heading("요약", level=1)
    summary_table = document.add_table(rows=1, cols=2)
    summary_table.style = "Light Grid Accent 1"
    header = summary_table.rows[0].cells
    header[0].text = "항목"
    header[1].text = "수량"
    for label, value in _summary_pairs(summary):
        cells = summary_table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    document.add_heading("상세 비교표", level=1)
    columns = _DETAIL_COLUMNS
    detail = document.add_table(rows=1, cols=len(columns))
    detail.style = "Light Grid Accent 1"
    for index, column in enumerate(detail.rows[0].cells):
        column.text = columns[index]
    for row in _rows_from_payload(payload):
        status = str(row.get("status", "unknown"))
        fill_hex, text_hex = _STATUS_COLORS.get(status, _DEFAULT_COLORS)
        cells = detail.add_row().cells
        values = [
            _status_label(status),
            str(row.get("spec_name", "")),
            str(row.get("base_value", "")),
            str(row.get("target_value", "")),
            str(row.get("summary", "")),
            _evidence_text(row),
        ]
        for index, value in enumerate(values):
            cells[index].text = value
        shade_cell(cells[0], fill_hex)
        verdict_run = cells[0].paragraphs[0].runs[0]
        verdict_run.bold = True
        verdict_run.font.color.rgb = RGBColor.from_string(text_hex)
        verdict_run.font.size = Pt(9)

    analysis = _extract_ai_analysis(str(payload.get("report_markdown", "")))
    if analysis:
        document.add_heading("AI 분석 요약", level=1)
        for line in analysis:
            document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _register_pdf_korean_font() -> tuple[str, str]:
    """Register a Korean font and return (regular_name, bold_name).

    Prefers a system TTF/OTF (nicer glyphs on Windows dev). When none is
    available — e.g. a Linux production image without CJK font packages —
    falls back to reportlab's built-in Korean CID fonts, the same approach
    as writing_assistant.documents, so Korean text always renders without
    external font files. Helvetica is the last resort only if CID
    registration itself fails.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    def _register(name: str, candidates: tuple[str | None, ...]) -> str | None:
        for path in candidates:
            if not path or not os.path.exists(path):
                continue
            try:
                pdfmetrics.registerFont(TTFont(name, path))
                return name
            except Exception:
                logger.warning("spec_compare.pdf: failed to register font %s", path, exc_info=True)
        return None

    regular = _register("SpecCompareKR", _KOREAN_FONT_CANDIDATES)
    if regular is not None:
        bold = _register("SpecCompareKR-Bold", _KOREAN_BOLD_CANDIDATES) or regular
        return regular, bold

    try:
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont

        registered = pdfmetrics.getRegisteredFontNames()
        for cid_name in ("HYSMyeongJo-Medium", "HYGothic-Medium"):
            if cid_name not in registered:
                pdfmetrics.registerFont(UnicodeCIDFont(cid_name))
        return "HYSMyeongJo-Medium", "HYGothic-Medium"
    except Exception:
        logger.warning(
            "spec_compare.pdf: CID font registration failed; Korean text may not render",
            exc_info=True,
        )
        return "Helvetica", "Helvetica-Bold"


def render_report_pdf(
    *,
    title: str,
    base_filename: str,
    target_filename: str,
    summary: dict[str, object],
    payload: dict[str, object],
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    regular_font, bold_font = _register_pdf_korean_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "krBody", parent=styles["Normal"], fontName=regular_font, fontSize=8, leading=11
    )
    heading = ParagraphStyle(
        "krHeading",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=13,
        spaceBefore=10,
        spaceAfter=6,
    )
    doc_title = ParagraphStyle(
        "krTitle", parent=styles["Title"], fontName=bold_font, fontSize=18, alignment=TA_LEFT
    )

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=title or "규격서 비교 보고서",
    )
    flow: list[object] = [Paragraph("규격서 비교 보고서", doc_title)]
    if title:
        flow.append(Paragraph(_xml_escape(title), body))
    flow.append(Spacer(1, 6))
    flow.append(Paragraph(_xml_escape(f"기준 문서: {base_filename}"), body))
    flow.append(Paragraph(_xml_escape(f"비교 문서: {target_filename}"), body))

    flow.append(Paragraph("요약", heading))
    summary_data = [["항목", "수량"], *[[label, value] for label, value in _summary_pairs(summary)]]
    summary_table = Table(summary_data, colWidths=[70 * mm, 30 * mm])
    summary_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), regular_font),
                ("FONTNAME", (0, 0), (-1, 0), bold_font),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    flow.append(summary_table)

    flow.append(Paragraph("상세 비교표", heading))
    header_style = ParagraphStyle("krHeaderCell", parent=body, fontName=bold_font, textColor=colors.white)
    detail_data: list[list[object]] = [[Paragraph(col, header_style) for col in _DETAIL_COLUMNS]]
    style_commands: list[tuple] = [
        ("FONTNAME", (0, 0), (-1, -1), regular_font),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (1, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]
    for row_index, row in enumerate(_rows_from_payload(payload), start=1):
        status = str(row.get("status", "unknown"))
        fill_hex, text_hex = _STATUS_COLORS.get(status, _DEFAULT_COLORS)
        verdict_style = ParagraphStyle(
            f"krVerdict{row_index}",
            parent=body,
            fontName=bold_font,
            textColor=colors.HexColor(f"#{text_hex}"),
        )
        detail_data.append(
            [
                Paragraph(_status_label(status), verdict_style),
                Paragraph(_xml_escape(str(row.get("spec_name", ""))), body),
                Paragraph(_xml_escape(str(row.get("base_value", ""))), body),
                Paragraph(_xml_escape(str(row.get("target_value", ""))), body),
                Paragraph(_xml_escape(str(row.get("summary", ""))), body),
                Paragraph(_xml_escape(_evidence_text(row)), body),
            ]
        )
        style_commands.append(
            ("BACKGROUND", (0, row_index), (0, row_index), colors.HexColor(f"#{fill_hex}"))
        )

    detail_table = Table(
        detail_data,
        colWidths=[18 * mm, 42 * mm, 50 * mm, 50 * mm, 62 * mm, 30 * mm],
        repeatRows=1,
    )
    detail_table.setStyle(TableStyle(style_commands))
    flow.append(detail_table)

    analysis = _extract_ai_analysis(str(payload.get("report_markdown", "")))
    if analysis:
        flow.append(Paragraph("AI 분석 요약", heading))
        for line in analysis:
            flow.append(Paragraph(_xml_escape(line), body))

    document.build(flow)
    return buffer.getvalue()


def render_report(
    fmt: str,
    *,
    title: str,
    base_filename: str,
    target_filename: str,
    summary: dict[str, object],
    payload: dict[str, object],
) -> RenderedReport:
    kwargs = {
        "title": title,
        "base_filename": base_filename,
        "target_filename": target_filename,
        "summary": summary,
        "payload": payload,
    }
    if fmt == "docx":
        return RenderedReport(
            content=render_report_docx(**kwargs),
            media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            extension="docx",
        )
    if fmt == "pdf":
        return RenderedReport(
            content=render_report_pdf(**kwargs),
            media_type="application/pdf",
            extension="pdf",
        )
    raise ValueError(f"Unsupported report format: {fmt}")
