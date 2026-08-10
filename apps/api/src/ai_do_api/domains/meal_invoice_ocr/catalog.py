from __future__ import annotations

import io
import itertools
import zipfile
from collections import Counter
from collections.abc import Iterable, Iterator
from typing import Any

from openpyxl import load_workbook

from ai_do_api.domains.meal_invoice_ocr.state import (
    _as_dict,
    _as_list,
    mutate_state,
    read_state_field,
)
from ai_do_api.domains.meal_invoice_ocr.vocab import _norm, _score

# 기준 카탈로그(영양사 납품 대장 엑셀에서 적재). 품목별 정식명 + 대표 단위/단가/원산지를 담아
# OCR 결과의 품목/단위/단가를 유사매칭해 '칩'으로 제안하는 데 쓴다. 워크스페이스 DB 행에 저장.
_HEADER_ITEM = ("품목", "품명")
_MAX_ITEMS = 5000
_MATCH_MIN = 0.72
# 단위 합산은 변형 품목(…컵/…라면 등)까지 모으도록 조금 느슨한 임계값.
_UNIT_MATCH_MIN = 0.6
# 업로드 XLSX 처리 상한(압축 확장·대량 셀로 인한 메모리 폭발 방어). 정상 주차별 대장(수십 시트,
# 시트당 수백 행)엔 영향 없고 비정상 대용량만 잘라낸다.
_MAX_XLSX_ARCHIVE_ENTRIES = 2_000
_MAX_XLSX_ARCHIVE_COMPRESSED_BYTES = 30 * 1024 * 1024
_MAX_XLSX_ARCHIVE_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
_MAX_XLSX_ARCHIVE_MEMBER_BYTES = 60 * 1024 * 1024
_MAX_XLSX_ARCHIVE_EXPANSION_RATIO = 100
_MAX_SHEETS = 60
_MAX_TOTAL_ROWS = 100_000
# 헤더 행을 찾기 위해 살펴보는 '내용 있는' 행 수. 대장 양식은 제목·결재란 뒤 7행쯤에 헤더가 있다.
# 장식·병합 셀로 생긴 빈 행은 세지 않는다(고정 행 수로 자르면 장식이 긴 시트에서 헤더를 놓친다).
_MAX_HEADER_SCAN_ROWS = 200
# 헤더 탐색 중 버퍼에 담아 두는 행 수 상한(비정상 XLSX 의 메모리 방어). 빈 행만 이어져도 여기서 멈춘다.
_MAX_HEADER_BUFFER_ROWS = 5_000
# 품목 칸이 이만큼 연속으로 비면 그 시트의 데이터가 끝난 것으로 본다.
_MAX_BLANK_RUN = 40
# 셀 문자열은 이 길이에서 자른다(품목명 등에 비정상적으로 긴 값이 와도 메모리·저장소 방어).
_MAX_CELL_CHARS = 200
# 헤더를 못 찾아도 이 고정 열 순서로 폴백한다(대장 양식: 납품일자·원산지·품목·단위·수량·단가·금액·…).
_FALLBACK_COLS = {"원산지": 1, "품목": 2, "단위": 3, "단가": 5}


# xlsx 는 ZIP 컨테이너다. openpyxl 에 넘기기 전에 실제 ZIP 매직바이트인지 확인한다(확장자 불신).
_ZIP_MAGIC = b"PK\x03\x04"


class CatalogArchiveTooLarge(ValueError):
    """카탈로그 XLSX ZIP 컨테이너가 안전 한도를 초과했다."""


def is_xlsx_upload(content: bytes) -> bool:
    """카탈로그 업로드가 실제 xlsx(ZIP) 컨테이너인지 매직바이트로 확인한다."""
    return content[:4] == _ZIP_MAGIC


def xlsx_archive_exceeds_limits(infos: Iterable[zipfile.ZipInfo]) -> bool:
    """XLSX ZIP 중앙 디렉터리만 보고 압축 폭탄/과대 archive 를 차단한다."""
    count = 0
    compressed = 0
    uncompressed = 0
    for info in infos:
        if info.is_dir():
            continue
        count += 1
        compressed += max(0, info.compress_size)
        uncompressed += max(0, info.file_size)
        if (
            count > _MAX_XLSX_ARCHIVE_ENTRIES
            or info.file_size > _MAX_XLSX_ARCHIVE_MEMBER_BYTES
            or compressed > _MAX_XLSX_ARCHIVE_COMPRESSED_BYTES
            or uncompressed > _MAX_XLSX_ARCHIVE_UNCOMPRESSED_BYTES
        ):
            return True
    return bool(compressed and uncompressed / compressed > _MAX_XLSX_ARCHIVE_EXPANSION_RATIO)


def validate_xlsx_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if xlsx_archive_exceeds_limits(archive.infolist()):
                raise CatalogArchiveTooLarge("catalog xlsx archive exceeds safety limits")
    except zipfile.BadZipFile as error:
        raise ValueError("invalid xlsx archive") from error


def load_catalog(workspace_id: str) -> dict[str, dict[str, Any]]:
    return _as_dict(read_state_field(workspace_id, "catalog_json", {}))


def save_catalog(workspace_id: str, catalog: dict[str, dict[str, Any]]) -> None:
    def _apply(state: Any) -> None:
        state.catalog_json = catalog

    mutate_state(workspace_id, _apply)


def load_units(workspace_id: str) -> list[str]:
    """카탈로그에서 뽑은 단위 목록(빈도 내림차순)."""
    return [str(x) for x in _as_list(read_state_field(workspace_id, "units_json", []))]


def save_units(workspace_id: str, units: list[str]) -> None:
    def _apply(state: Any) -> None:
        state.units_json = units

    mutate_state(workspace_id, _apply)


def save_catalog_with_units(
    workspace_id: str, catalog: dict[str, dict[str, Any]], units: list[str]
) -> None:
    def _apply(state: Any) -> None:
        state.catalog_json = catalog
        state.units_json = units

    mutate_state(workspace_id, _apply)


def clear_catalog(workspace_id: str) -> None:
    def _apply(state: Any) -> None:
        state.catalog_json = {}
        state.units_json = []

    mutate_state(workspace_id, _apply)


def _price_str(value: Any) -> str:
    if value is None or value == "":
        return ""
    try:
        return str(int(round(float(str(value).replace(",", "")))))
    except (ValueError, TypeError):
        return ""


def _col_map(header_row: tuple[Any, ...]) -> dict[str, int] | None:
    """헤더 행에서 열 이름→인덱스 매핑. '품목'(또는 '품명')이 없으면 None."""
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        label = str(cell or "").strip()
        if label:
            mapping[label] = idx
    if not any(h in mapping for h in _HEADER_ITEM):
        return None
    return mapping


def _scan_header(
    row_iter: Iterator[tuple[Any, ...]],
) -> tuple[list[tuple[Any, ...]], dict[str, int], int]:
    """헤더 행이 나올 때까지 시트 앞부분을 버퍼에 담으며 훑는다.

    반환: (버퍼한 행들, 열 매핑, 그 버퍼 안에서의 데이터 시작 인덱스).

    제목·결재란·병합 셀 같은 장식 행은 대부분 비어 있으므로 예산으로 세지 않는다. 고정 행 수로
    자르면 장식이 긴 시트에서 헤더를 지나쳐, 잘못된 고정 열 순서로 조용히 오파싱한다. 내용 있는
    행만 세되 버퍼 크기에도 상한을 둬 비정상 XLSX 에서 메모리가 늘어나지 않게 한다. 끝내 못 찾으면
    고정 열 순서로 폴백하고 첫 행부터 데이터로 취급한다(빈 행은 호출부에서 건너뛴다).
    """
    head: list[tuple[Any, ...]] = []
    scanned = 0
    for row in row_iter:
        head.append(row)
        cols = _col_map(row)
        if cols is not None:
            return (head, cols, len(head))
        if any(str(cell or "").strip() for cell in row):
            scanned += 1
        if scanned >= _MAX_HEADER_SCAN_ROWS or len(head) >= _MAX_HEADER_BUFFER_ROWS:
            break
    return (head, dict(_FALLBACK_COLS), 0)


def parse_catalog_xlsx(content: bytes) -> tuple[dict[str, dict[str, Any]], list[str]]:
    """대장 엑셀(주차별 시트)에서 품목별 카탈로그와 단위 목록을 만든다.

    반환: (catalog, units)
    - catalog: { 품목명: {"단위":.., "단가":.., "원산지":.., "count": n} }. 단가는 '최근값'
      (시트가 최신순으로 정렬돼 있다는 가정 — 처음 만난 값이 최근). count 는 전체 등장 횟수.
    - units: 엑셀에서 쓰인 단위 목록을 '많이 나온 순(내림차순)'으로.
    """
    validate_xlsx_archive(content)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    catalog: dict[str, dict[str, Any]] = {}
    unit_counts: Counter[str] = Counter()  # 전체 단위 빈도(드롭다운 후보).
    item_units: dict[str, Counter[str]] = {}  # 품목별 단위 빈도(과거 이력).
    processed_rows = 0  # 전체 물질화 행 수(압축 확장·대량 셀 메모리 방어를 위한 전역 예산).
    try:
        for sheet_index, sheet_name in enumerate(workbook.sheetnames):
            if sheet_index >= _MAX_SHEETS or processed_rows >= _MAX_TOTAL_ROWS:
                break
            sheet = workbook[sheet_name]
            # 시트를 통째로 물질화하지 않고 스트리밍한다. 엑셀 시트는 서식만 있어도 max_row 가
            # 100만 행을 넘기는 일이 흔한데, 예전처럼 남은 예산만큼 잘라 list() 로 받으면 첫 시트의
            # 빈 행이 전역 예산을 전부 먹어 나머지 주차 시트를 아예 읽지 못했다.
            row_iter = sheet.iter_rows(values_only=True)
            head, cols, start = _scan_header(row_iter)
            processed_rows += start
            blank_run = 0
            for row in itertools.chain(head[start:], row_iter):
                processed_rows += 1
                if processed_rows >= _MAX_TOTAL_ROWS:
                    break
                item_idx = cols.get("품목", cols.get("품명", _FALLBACK_COLS["품목"]))
                item = (
                    str(row[item_idx] or "").strip()[:_MAX_CELL_CHARS]
                    if item_idx < len(row)
                    else ""
                )
                if not item:
                    blank_run += 1
                    if blank_run >= _MAX_BLANK_RUN:  # 데이터 끝으로 보고 이 시트 종료.
                        break
                    continue
                blank_run = 0
                unit_idx = cols.get("단위", _FALLBACK_COLS["단위"])
                price_idx = cols.get("단가", _FALLBACK_COLS["단가"])
                origin_idx = cols.get("원산지", _FALLBACK_COLS["원산지"])

                def _cell(i: int) -> Any:
                    return row[i] if i < len(row) else None

                # 영어 단위는 전부 소문자로 통일(한글 단위는 영향 없음). 'kg'는 'k'로 통일.
                unit_value = (
                    str(_cell(unit_idx) or "").strip().lower().replace("kg", "k")[:_MAX_CELL_CHARS]
                )
                if unit_value:
                    unit_counts[unit_value] += 1
                    item_units.setdefault(item, Counter())[unit_value] += 1

                entry = catalog.get(item)
                if entry is None:
                    # 처음(최신 시트/최상단) 만난 값을 대표(최근)로 고정(단가·원산지).
                    catalog[item] = {
                        "단위": "",
                        "단위들": [],
                        "단가": _price_str(_cell(price_idx)),
                        "원산지": str(_cell(origin_idx) or "").strip()[:_MAX_CELL_CHARS],
                        "count": 1,
                    }
                else:
                    entry["count"] = int(entry.get("count", 0)) + 1
    finally:
        workbook.close()
    # 품목별 단위를 빈도 내림차순으로 정리(동률은 이름순). 대표 단위는 최빈값.
    for name, entry in catalog.items():
        counter = item_units.get(name)
        if counter:
            ordered = [u for u, _ in sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))]
            entry["단위들"] = ordered
            entry["단위카운트"] = dict(counter)  # 유사 품목 단위 합산용.
            entry["단위"] = ordered[0]
    if len(catalog) > _MAX_ITEMS:
        # 등장 횟수 많은 순으로 상위만 유지.
        top = sorted(catalog.items(), key=lambda kv: kv[1].get("count", 0), reverse=True)
        catalog = dict(top[:_MAX_ITEMS])
    # 전체 단위 목록(드롭다운 후보): 많이 나온 순. 동률은 이름순.
    units = [u for u, _ in sorted(unit_counts.items(), key=lambda kv: (-kv[1], kv[0]))]
    return (catalog, units)


# OCR 프롬프트에 후보로 넣을 최대 품목 수. 실제 명세표(15행) 실측 결과 300 종이 가장 좋았다.
#   후보 없음 : 정답 10/15, 대장에 없는 지어낸 단어 4개(근추가우·낱생추·참참갓잎·대추약)
#   상위 300종: 정답 12/15, 지어낸 단어 1개
#   전체 771종: 정답  9/15, 지어낸 단어 4개
# 목록이 길어지면 모델이 목록을 사실상 무시한다(주의 분산). 등장 횟수 내림차순 상위만 넣는다.
_MAX_HINT_ITEMS = 300


def build_item_hint(
    catalog: dict[str, dict[str, Any]], *, limit: int = _MAX_HINT_ITEMS
) -> tuple[str, int]:
    """대장 품목을 OCR 프롬프트용 '후보 목록' 문구로 만든다. (문구, 잘려나간 품목 수)를 돌려준다.

    단위는 이미 고정 목록을 주고 "이 중에서 고르라"고 지시해 거의 틀리지 않는다. 품목명은 후보 없이
    백지에서 읽게 해서 '근추가우'·'낱생추'처럼 실재하지 않는 단어가 나온다. 같은 방식으로 실제 납품
    이력이 있는 품목을 후보로 제시한다.

    강제가 아니라 우선순위다. 대장에 아직 없는 품목이 실제로 들어올 수 있으므로, 목록에 없으면 읽은
    대로 적게 둔다(닫힌 집합으로 강제하면 신규 품목을 영영 못 읽는다).
    """
    if not catalog:
        return ("", 0)
    ranked = sorted(
        catalog.items(),
        key=lambda kv: (-int(kv[1].get("count", 0) or 0), kv[0]),
    )
    names = [name for name, _entry in ranked[:limit] if name]
    if not names:
        return ("", 0)
    dropped = max(0, len(ranked) - len(names))
    return (
        "\n\n납품 이력이 있는 품목 목록이다. 손글씨 품명은 먼저 이 목록에서 가장 비슷한 것을 찾아 "
        "그 표기로 적어라. 목록에 없는 새 단어를 지어내지 마라. 다만 정말로 목록 어디에도 없는 품목이면 "
        "읽은 대로 적어라: " + ", ".join(names) + ".",
        dropped,
    )


def best_match(name: str, catalog: dict[str, dict[str, Any]]) -> tuple[str, float]:
    """카탈로그에서 가장 유사한 품목명과 점수(0~1). 비면 ('', 0)."""
    if not _norm(name) or not catalog:
        return ("", 0.0)
    best = ""
    best_score = 0.0
    for key in catalog:
        s = _score(name, key)
        if s > best_score:
            best_score = s
            best = key
    return (best, best_score)


def units_for(
    name: str, catalog: dict[str, dict[str, Any]], min_score: float = _UNIT_MATCH_MIN
) -> list[str]:
    """이름이 유사한(>= min_score) 모든 카탈로그 품목의 단위를 합산해 빈도 내림차순 목록으로.

    같은 개념 품목이 여러 키(예: '오징어짬뽕컵', '오징어짬뽕라면')로 나뉘어 있어도 단위를 모아 준다.
    """
    if not _norm(name) or not catalog:
        return []
    agg: Counter[str] = Counter()
    for key, entry in catalog.items():
        if _score(name, key) < min_score:
            continue
        counts = entry.get("단위카운트")
        if isinstance(counts, dict):
            for unit, count in counts.items():
                if str(unit).strip():
                    agg[str(unit)] += int(count)
    return [u for u, _ in sorted(agg.items(), key=lambda kv: (-kv[1], kv[0]))]
