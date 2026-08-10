from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ai_do_api.domains.meal_invoice_ocr.schemas import MealInvoiceExportDocument

# 영양사 납품 대장 양식과 동일한 컬럼/순서.
_HEADERS = [
    "납품일자",
    "원산지",
    "품목",
    "단위",
    "수량",
    "단가",
    "금액",
    "신선도",
    "불량여부",
    "반품여부",
    "비고",
]
_COLUMN_WIDTHS = [12, 10, 28, 8, 8, 12, 14, 8, 8, 8, 16]
# 일반 콤마 천단위 서식을 적용할 열(1-indexed): 수량.
_NUMERIC_COLS = (5,)
# 회계 형식(₩ 기호 없는 '쉼표 스타일')을 적용할 열(1-indexed): 단가·금액.
_ACCOUNTING_COLS = (6, 7)
# Excel 회계 서식(기호 없음): 양수 / 음수 / 0("-") / 텍스트 4구역. 숫자 오른쪽에 정렬용 여백을 둔다.
_ACCOUNTING_FORMAT = '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-'
# 오른쪽 정렬할 열(1-indexed): 단가·금액. 나머지는 가운데 정렬.
_RIGHT_COLS = (6, 7)
# 맑은 고딕으로 출력할 열(1-indexed): 신선도.
_MALGUN_NAME = "맑은 고딕"
_MALGUN_COLS = (8,)

# 실제 설치 폰트명은 '현대하모니 M'(하모니와 M 사이 공백). 공백 없으면 Excel 이 기본 폰트로 대체된다.
_FONT_NAME = "현대하모니 M"
_FONT_SIZE = 11

# 붙여 쓴 수량/단위("30K", "1/2box", "12~13kg")를 분리한다. 수량 표현은 숫자와
# 부호·분수·범위 기호까지만 받고, 단위는 문자로 시작해야 숫자 일부를 단위로 오인하지 않는다.
_QTY_WITH_UNIT_RE = re.compile(
    r"^\s*([+\-–—~/?.,\d]*\d[+\-–—~/?.,\d]*|[\-–—]+)\s*([^\W\d_].*)$"
)
# 수량 칸에서 실제 숫자로 변환해도 표시 의미가 바뀌지 않는 값. 명시적 '+'와 분수·범위는
# 텍스트 그대로 두며, '-'와 후행 불확실 표시는 기존 집계 계약을 유지한다.
_CANONICAL_QUANTITY_RE = re.compile(
    r"^-?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|\.\d+)\??$"
)
# 다양한 날짜 표기에서 연·월·일을 뽑아 대장 양식(YY.MM.DD)으로 통일.
_DATE_RE = re.compile(r"(\d{2,4})\D+(\d{1,2})\D+(\d{1,2})")


def _format_date(value: str) -> str:
    m = _DATE_RE.search(str(value or ""))
    if not m:
        return str(value or "")
    year, month, day = m.groups()
    return f"{int(year[-2:]):02d}.{int(month):02d}.{int(day):02d}"


def _split_qty(value: str) -> tuple[str, str]:
    # 새 포맷은 "418 10kg"처럼 공백으로 수량/단위를 나눈다(숫자로 시작하는 단위 보존).
    s = str(value or "").strip()
    if any(char.isspace() for char in s):
        num, unit = s.rsplit(maxsplit=1)
        return (num.strip(), unit.strip())
    # 구버전 붙임 표기는 수량 표현과 문자로 시작하는 단위를 분리한다.
    m = _QTY_WITH_UNIT_RE.fullmatch(s)
    if not m:
        return (s, "")
    return (m.group(1), m.group(2).strip())


# 엑셀/스프레드시트 수식 주입 방지: 셀 문자열이 이 문자로 시작하면 수식(=HYPERLINK…)·명령으로
# 해석될 수 있어, 앞에 작은따옴표를 붙여 '텍스트'로 강제한다(CSV/포뮬러 인젝션 방어).
_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")


def _safe_cell(value: object) -> object:
    """사용자 제어 문자열이 스프레드시트 수식으로 실행되지 않게 무해화한다(비문자열은 그대로)."""
    if isinstance(value, str) and value and value[0] in _FORMULA_LEAD:
        # 순수 대시 자리표시값("-", "--")은 수식이 아니라 빈칸 표시일 뿐이다. 여기에 작은따옴표를
        # 붙이면 openpyxl 이 그 따옴표를 실제 문자로 저장해 셀에 '-' 로 보이므로 제외한다.
        if set(value) <= {"-"}:
            return value
        return "'" + value
    return value


def _num(value: str) -> object:
    """숫자를 뽑아 int/float 로 돌려준다(쉼표·단위·불확실 표기 '?' 제거).

    OCR 불확실값 '3,500?' 처럼 기호가 섞여도 숫자 부분을 추출해 엑셀에서 숫자로 집계·정렬되게
    한다(서버 parse_number 와 동일 규칙). 숫자가 전혀 없으면 원문 문자열을 그대로 보존한다.
    """
    text = str(value or "").strip()
    if not text:
        return ""
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if cleaned and cleaned not in {"-", ".", "-."}:
        try:
            f = float(cleaned)
            return int(f) if f.is_integer() else f
        except ValueError:
            pass
    return text


def _quantity_value(value: str) -> object:
    """손실 없이 숫자로 표현할 수 있는 수량만 int/float 로 변환한다."""
    text = str(value or "").strip()
    if not text or not _CANONICAL_QUANTITY_RE.fullmatch(text):
        return text
    cleaned = text.replace(",", "").removesuffix("?")
    try:
        number = float(cleaned)
        return int(number) if number.is_integer() else number
    except ValueError:
        return text


def export_documents_xlsx(documents: list[MealInvoiceExportDocument]) -> bytes:
    """편집된 문서 목록을 대장 양식(현대하모니M 11pt) xlsx 로 만든다."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "거래명세표"
    sheet.append(_HEADERS)

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    # 헤더 텍스트는 열과 무관하게 전부 표준 글씨체(현대하모니 M). 맑은 고딕은 신선도 '본문' 셀에만
    # 적용한다(헤더 '신선도' 글자까지 맑은 고딕이 되지 않도록).
    for cell in sheet[1]:
        cell.font = Font(name=_FONT_NAME, size=_FONT_SIZE, bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for doc in documents:
        for index, row in enumerate(doc.품목):
            qty, unit = _split_qty(row.수량)
            # 대장처럼 납품일자는 문서의 첫 줄에만, YY.MM.DD 형식으로 적는다.
            date = _format_date(doc.거래일) if index == 0 else ""
            # 사용자 제어 문자열은 모두 _safe_cell 로 무해화(수식 주입 방지). 숫자는 그대로 통과.
            sheet.append(
                [
                    _safe_cell(date),
                    _safe_cell(row.원산지),
                    _safe_cell(row.품명),
                    _safe_cell(unit),
                    _safe_cell(_quantity_value(qty)),
                    _safe_cell(_num(row.단가)),
                    _safe_cell(_num(row.금액)),
                    _safe_cell(row.신선도),
                    _safe_cell(row.불량여부),
                    _safe_cell(row.반품여부),
                    _safe_cell((row.비고 or "").strip() or "-"),
                ]
            )

    body_font = Font(name=_FONT_NAME, size=_FONT_SIZE)
    malgun_font = Font(name=_MALGUN_NAME, size=_FONT_SIZE)
    right_align = Alignment(horizontal="right", vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")
    for row_cells in sheet.iter_rows(min_row=2):
        for idx, cell in enumerate(row_cells, start=1):
            cell.font = malgun_font if idx in _MALGUN_COLS else body_font
            cell.alignment = right_align if idx in _RIGHT_COLS else center_align
        for col in _NUMERIC_COLS:
            row_cells[col - 1].number_format = "#,##0"
        for col in _ACCOUNTING_COLS:
            row_cells[col - 1].number_format = _ACCOUNTING_FORMAT

    # 모든 셀에 점선(dotted) 테두리(상하좌우).
    dotted = Side(style="dotted")
    border = Border(left=dotted, right=dotted, top=dotted, bottom=dotted)
    for row_cells in sheet.iter_rows(min_row=1, max_col=len(_HEADERS)):
        for cell in row_cells:
            cell.border = border

    sheet.freeze_panes = "A2"
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
