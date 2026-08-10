# -*- coding: utf-8 -*-
"""Brandlogy 자유 양식 빌더 — 27 × 16.75 cm 자유배치 디자인.

MiniMax 영감 Brandlogy 디자인 시스템(`Brandlogy_PPT_Design_System_27x16_75cm.md`)을
python-pptx 로 구현한다. 핵심:

  - 캔버스 27 × 16.75 cm, 흰 캔버스(+ dark / hero gradient 옵션)
  - 잠금 존: 헤드라인(0.76~2.34cm) · 부제(2.34~3.30cm) · 본문 박스(3.68~16.05cm)
  - 헤드라인/부제/eyebrow 는 자동 배치, 본문 박스 안 ELEMENTS 만 cm 좌표 자유배치
  - 단일 layout "brandlogy". ELEMENTS 컴포넌트: text / kpi / chart / table /
    bullets / pill / steps / divider — 각 컴포넌트가 내부 레이아웃을 스스로 처리
  - 폰트 Pretendard 전용, 가중치로 위계, 브랜드 블루 #1456f0/#3b82f6/#60a5fa

prompts._BRANDLOGY_SCHEMA 가 이 빌더의 입력 계약이다. HTML 렌더러
(PptSlideRenderer.tsx 의 brandlogy 분기)와 좌표·색상이 1:1 대응해야 한다.
"""
from __future__ import annotations

from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree


# ============================================================
# 캔버스 & 잠금 존 (cm)
# ============================================================
SLIDE_W_CM = 27.0
SLIDE_H_CM = 16.75
SLIDE_W = Cm(SLIDE_W_CM)
SLIDE_H = Cm(SLIDE_H_CM)

MARGIN_L = 0.9              # 좌우 여백
HEADLINE_Y = 0.76
HEADLINE_H = 1.58           # 0.76 ~ 2.34
SUBTITLE_Y = 2.34
SUBTITLE_H = 0.96           # 2.34 ~ 3.30
BODY_X0 = 0.9
BODY_Y0 = 3.68
BODY_X1 = 26.1             # 0.9 + 25.2
BODY_Y1 = 16.05
BODY_W = BODY_X1 - BODY_X0  # 25.2
BODY_H = BODY_Y1 - BODY_Y0  # 12.37


# ============================================================
# 색상
# ============================================================
def _rgb(hex_str: str) -> RGBColor:
    h = (hex_str or "").lstrip("#")
    if len(h) != 6:
        return RGBColor(0x22, 0x22, 0x22)
    try:
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    except ValueError:
        return RGBColor(0x22, 0x22, 0x22)


INK = _rgb("#222222")
SUB = _rgb("#45515e")
MUTE = _rgb("#8e8e93")
BLUE = _rgb("#1456f0")
BLUE2 = _rgb("#3b82f6")
BLUE3 = _rgb("#60a5fa")
PINK = _rgb("#ea5ec1")
BORDER = _rgb("#e5e7eb")
BORDER_SOFT = _rgb("#f2f3f5")
WHITE = _rgb("#ffffff")
DARK = _rgb("#181e25")
SUCCESS = _rgb("#16a34a")

FONT = "Pretendard"

_EMU_PER_PX = 9525  # 96dpi


# ============================================================
# 안전 변환
# ============================================================
def _f(v, default=0.0):
    """cm 좌표 값 안전 변환."""
    if isinstance(v, bool):
        return default
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        num = ""
        for ch in s:
            if ch.isdigit() or ch in ".-+":
                num += ch
            elif num:
                break
        try:
            return float(num)
        except ValueError:
            return default
    return default


def _to_float(v):
    """차트용 숫자 추출 — 실패 시 None."""
    r = _f(v, default=None) if not isinstance(v, str) else None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        num = ""
        for ch in s:
            if ch.isdigit() or ch in ".-+":
                num += ch
            elif num:
                break
        try:
            return float(num)
        except ValueError:
            return None
    return r


def _s(v) -> str:
    return "" if v is None else str(v)


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# ============================================================
# 효과 (그림자 / 그래디언트)
# ============================================================
def _kill_shadow(shape):
    try:
        spPr = shape._element.find(".//" + qn("p:spPr"))
        if spPr is None:
            return
        for tag in ("a:effectLst", "a:effectDag"):
            for ex in spPr.findall(qn(tag)):
                spPr.remove(ex)
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    except Exception:
        pass


def _apply_shadow(shape, kind: str):
    """kind: 'standard' (은은한 검정) | 'glow' (브랜드 퍼플 글로우)."""
    try:
        spPr = shape._element.find(".//" + qn("p:spPr"))
        if spPr is None:
            return
        for tag in ("a:effectLst", "a:effectDag"):
            for ex in spPr.findall(qn(tag)):
                spPr.remove(ex)
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        shd = etree.SubElement(eff, qn("a:outerShdw"))
        if kind == "glow":
            shd.set("blurRad", str(15 * _EMU_PER_PX))
            shd.set("dist", "0")
            shd.set("dir", "0")
            clr = etree.SubElement(shd, qn("a:srgbClr"))
            clr.set("val", "2C1E74")
            etree.SubElement(clr, qn("a:alpha")).set("val", "16000")
        else:  # standard
            shd.set("blurRad", str(6 * _EMU_PER_PX))
            shd.set("dist", str(4 * _EMU_PER_PX))
            shd.set("dir", "5400000")  # 90° 아래
            shd.set("rotWithShape", "0")
            clr = etree.SubElement(shd, qn("a:srgbClr"))
            clr.set("val", "000000")
            etree.SubElement(clr, qn("a:alpha")).set("val", "8000")
    except Exception:
        pass


def _apply_gradient(shape):
    """히어로 그래디언트 — 135° #1456f0 → #3b82f6 → #60a5fa."""
    try:
        spPr = shape._element.find(".//" + qn("p:spPr"))
        if spPr is None:
            return
        for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill"):
            for ex in spPr.findall(qn(tag)):
                spPr.remove(ex)
        grad = spPr.makeelement(qn("a:gradFill"), {})
        # solidFill 자리(보통 ln 앞)에 삽입
        ln = spPr.find(qn("a:ln"))
        if ln is not None:
            ln.addprevious(grad)
        else:
            spPr.append(grad)
        gs_lst = etree.SubElement(grad, qn("a:gsLst"))
        for pos, hexv in ((0, "1456F0"), (50000, "3B82F6"), (100000, "60A5FA")):
            gs = etree.SubElement(gs_lst, qn("a:gs"))
            gs.set("pos", str(pos))
            clr = etree.SubElement(gs, qn("a:srgbClr"))
            clr.set("val", hexv)
        lin = etree.SubElement(grad, qn("a:lin"))
        lin.set("ang", str(45 * 60000))  # 좌상 → 우하
        lin.set("scaled", "1")
    except Exception:
        pass


# ============================================================
# 기본 도형 / 텍스트
# ============================================================
def _set_run_font(run, *, size_pt=11, weight=400, color=INK, italic=False):
    run.font.size = Pt(size_pt)
    run.font.bold = weight >= 600
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("ea", "cs"):
        old = rPr.find(qn(f"a:{tag}"))
        if old is not None:
            rPr.remove(old)
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", FONT)
    cs = etree.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", FONT)


def _add_rect(slide, x, y, w, h, *, fill=None, line_color=None, line_pt=0.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_pt)
    _kill_shadow(sp)
    return sp


def _add_round(slide, x, y, w, h, *, fill=None, line_color=None, line_pt=0.0,
               radius_px=13.0, shadow=None, gradient=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    # radius_px → 짧은 변 대비 비율
    radius_cm = radius_px / 96.0 * 2.54
    short = max(0.1, min(w, h))
    try:
        sp.adjustments[0] = max(0.0, min(0.5, radius_cm / short))
    except Exception:
        pass
    if gradient:
        sp.fill.solid()
        sp.fill.fore_color.rgb = BLUE
    elif fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_pt)
    if gradient:
        _apply_gradient(sp)
    if shadow:
        _apply_shadow(sp, shadow)
    else:
        _kill_shadow(sp)
    return sp


def _add_text(slide, x, y, w, h, text, *, size_pt=11, weight=400, color=INK,
              align="left", anchor="top", line_height=None, wrap=True,
              margins=False, italic=False):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    if not margins:
        tf.margin_left = Emu(0)
        tf.margin_right = Emu(0)
        tf.margin_top = Emu(0)
        tf.margin_bottom = Emu(0)
    tf.vertical_anchor = _ANCHOR.get(anchor, MSO_ANCHOR.TOP)
    al = _ALIGN.get(align, PP_ALIGN.LEFT)
    lines = _s(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al
        if line_height is not None:
            p.line_spacing = line_height
        run = p.add_run()
        run.text = line
        _set_run_font(run, size_pt=size_pt, weight=weight, color=color, italic=italic)
    return tb


# ============================================================
# ELEMENTS — 본문 박스 컴포넌트
# ============================================================
def _clamp(x, y, w, h):
    """본문 박스 안으로 보정."""
    x = max(BODY_X0, min(x, BODY_X1 - 0.5))
    y = max(BODY_Y0, min(y, BODY_Y1 - 0.5))
    w = max(0.5, min(w, BODY_X1 - x))
    h = max(0.3, min(h, BODY_Y1 - y))
    return x, y, w, h


def _el_text(slide, el, x, y, w, h, *, dark=False):
    fill_hex = el.get("FILL")
    has_card = bool(fill_hex)
    shadow = el.get("SHADOW")
    if has_card:
        border = _rgb(el["BORDER"]) if el.get("BORDER") else None
        _add_round(
            slide, x, y, w, h,
            fill=_rgb(fill_hex),
            line_color=border, line_pt=0.75 if border else 0.0,
            radius_px=_f(el.get("RADIUS"), 13.0),
            shadow=(shadow if shadow in ("standard", "glow") else None),
        )
    pad = _f(el.get("PAD"), 0.4 if has_card else 0.0)
    color = _rgb(el["COLOR"]) if el.get("COLOR") else (WHITE if dark and not has_card else INK)
    _add_text(
        slide, x + pad, y + pad, w - pad * 2, h - pad * 2,
        el.get("TEXT", ""),
        size_pt=_f(el.get("SIZE"), 11.0),
        weight=int(_f(el.get("WEIGHT"), 400)),
        color=color,
        align=(el.get("ALIGN") or "left"),
        anchor=(el.get("VALIGN") or "top"),
        line_height=_f(el.get("LINE_HEIGHT"), 1.45),
        wrap=True,
    )


def _el_kpi(slide, el, x, y, w, h, *, dark=False):
    featured = bool(el.get("FEATURED"))
    if featured:
        _add_round(slide, x, y, w, h, gradient=True, radius_px=22, shadow="glow")
        val_c = WHITE
        lbl_c = _rgb("#e8efff")
        delta_c = WHITE
    else:
        _add_round(slide, x, y, w, h, fill=WHITE, line_color=BORDER, line_pt=0.75,
                   radius_px=13, shadow="standard")
        val_c = BLUE
        lbl_c = SUB
        up = el.get("UP", True)
        delta_c = BLUE if up else PINK
    pad = 0.45
    # 라벨 (상단)
    _add_text(slide, x + pad, y + pad, w - pad * 2, 0.55,
              el.get("LABEL", ""), size_pt=10.5, weight=500, color=lbl_c,
              align="left", anchor="top", wrap=False)
    # 값 + 단위
    tb = slide.shapes.add_textbox(Cm(x + pad), Cm(y + h * 0.40), Cm(w - pad * 2), Cm(h * 0.40))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    rv = p.add_run()
    rv.text = _s(el.get("VALUE"))
    _set_run_font(rv, size_pt=30, weight=700, color=val_c)
    if el.get("UNIT"):
        ru = p.add_run()
        ru.text = " " + _s(el.get("UNIT"))
        _set_run_font(ru, size_pt=13, weight=500, color=(WHITE if featured else SUB))
    # 델타 (하단)
    if el.get("DELTA"):
        _add_text(slide, x + pad, y + h - pad - 0.5, w - pad * 2, 0.5,
                  el.get("DELTA"), size_pt=11, weight=700, color=delta_c,
                  align="left", anchor="middle", wrap=False)


def _el_pill(slide, el, x, y, w, h, *, dark=False):
    style = (el.get("STYLE") or "light").lower()
    if style == "dark":
        bg, fg = DARK, WHITE
    elif style == "nav":
        bg, fg = BORDER_SOFT, _rgb("#18181b")
    else:
        bg, fg = WHITE, _rgb("#18181b")
    line = None if style == "dark" else BORDER
    _add_round(slide, x, y, w, h, fill=bg, line_color=line, line_pt=0.75,
               radius_px=999)
    _add_text(slide, x, y, w, h, el.get("TEXT", ""), size_pt=10, weight=600,
              color=fg, align="center", anchor="middle", wrap=False)


def _el_divider(slide, el, x, y, w, h):
    color = _rgb(el["COLOR"]) if el.get("COLOR") else BORDER
    _add_rect(slide, x, y, w, 0.04, fill=color)


def _el_bullets(slide, el, x, y, w, h, *, dark=False):
    has_card = bool(el.get("FILL"))
    if has_card:
        _add_round(slide, x, y, w, h, fill=_rgb(el["FILL"]),
                   line_color=(BORDER if el.get("FILL", "").lower() in ("#ffffff", "#fff") else None),
                   line_pt=0.75,
                   radius_px=_f(el.get("RADIUS"), 13.0),
                   shadow=(el.get("SHADOW") if el.get("SHADOW") in ("standard", "glow") else None))
    pad = 0.45 if has_card else 0.0
    cy = y + pad
    if el.get("TITLE"):
        _add_text(slide, x + pad, cy, w - pad * 2, 0.6, el.get("TITLE"),
                  size_pt=14, weight=600, color=(WHITE if dark and not has_card else INK),
                  align="left", anchor="top", wrap=False)
        cy += 0.75
    items = el.get("ITEMS") or []
    size = _f(el.get("SIZE"), 11.0)
    tb = slide.shapes.add_textbox(Cm(x + pad), Cm(cy), Cm(w - pad * 2), Cm(max(0.4, y + h - pad - cy)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    txt_c = WHITE if dark and not has_card else INK
    for i, item in enumerate(items[:8]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        p.line_spacing = 1.35
        rb = p.add_run()
        rb.text = "•  "
        _set_run_font(rb, size_pt=size, weight=700, color=BLUE)
        rt = p.add_run()
        rt.text = _s(item)
        _set_run_font(rt, size_pt=size, weight=400, color=txt_c)


def _el_steps(slide, el, x, y, w, h, *, dark=False):
    items = el.get("ITEMS") or []
    n = max(1, min(len(items), 6))
    gap = 0.4
    cw = (w - gap * (n - 1)) / n
    for i in range(n):
        it = items[i] if i < len(items) else {}
        cx = x + (cw + gap) * i
        _add_round(slide, cx, y, cw, h, fill=WHITE, line_color=BORDER, line_pt=0.75,
                   radius_px=13, shadow="standard")
        # 번호 원
        d = 0.9
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(cx + 0.3), Cm(y + 0.3), Cm(d), Cm(d))
        circle.fill.solid()
        circle.fill.fore_color.rgb = BLUE
        circle.line.fill.background()
        _kill_shadow(circle)
        _add_text(slide, cx + 0.3, y + 0.3, d, d, str(i + 1), size_pt=14, weight=700,
                  color=WHITE, align="center", anchor="middle", wrap=False)
        _add_text(slide, cx + 0.3, y + 1.35, cw - 0.6, 0.6, it.get("LABEL", ""),
                  size_pt=12.5, weight=600, color=INK, align="left", anchor="top", wrap=False)
        _add_text(slide, cx + 0.3, y + 1.95, cw - 0.6, h - 2.2, it.get("BODY", ""),
                  size_pt=10, weight=400, color=SUB, align="left", anchor="top",
                  line_height=1.35)


def _el_table(slide, el, x, y, w, h):
    columns = el.get("COLUMNS") or []
    rows = el.get("ROWS") or []
    if not columns:
        return
    ncol = len(columns)
    nrow = len(rows) + 1
    try:
        gf = slide.shapes.add_table(nrow, ncol, Cm(x), Cm(y), Cm(w), Cm(h))
        table = gf.table
        table.first_row = False
        table.horz_banding = False
        highlight = el.get("HIGHLIGHT_ROW")
        try:
            highlight = int(highlight)
        except (TypeError, ValueError):
            highlight = None
        for c in range(ncol):
            cell = table.cell(0, c)
            _fill_cell(cell, _s(columns[c]), size=11.5, weight=600, color=INK,
                       fill=BORDER_SOFT, align=("left" if c == 0 else "center"))
        for r, row in enumerate(rows):
            row = row if isinstance(row, list) else [row]
            is_hl = (highlight is not None and r == highlight)
            for c in range(ncol):
                val = row[c] if c < len(row) else ""
                cell = table.cell(r + 1, c)
                _fill_cell(cell, _s(val), size=11, weight=(600 if is_hl else 400),
                           color=(BLUE if is_hl else INK),
                           fill=(_rgb("#eff5ff") if is_hl else WHITE),
                           align=("left" if c == 0 else "center"))
    except Exception:
        pass


def _fill_cell(cell, text, *, size, weight, color, fill, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Cm(0.2)
    cell.margin_right = Cm(0.2)
    cell.margin_top = Cm(0.05)
    cell.margin_bottom = Cm(0.05)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = _ALIGN.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    run.text = text
    _set_run_font(run, size_pt=size, weight=weight, color=color)


def _kill_chart_shadows(chart):
    try:
        nsmap = {
            "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        }
        for spPr in chart._chartSpace.findall(".//c:spPr", nsmap):
            for tag in ("a:effectLst", "a:effectDag"):
                for ex in spPr.findall(qn(tag)):
                    spPr.remove(ex)
            spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    except Exception:
        pass


_SERIES_COLORS = [BLUE2, BLUE3, BLUE, _rgb("#bfdbfe"), PINK, MUTE]


def _el_chart(slide, el, x, y, w, h):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    kind = (el.get("CHART") or "bar").lower()
    labels = el.get("LABELS") or []
    series = el.get("SERIES") or []
    if not series:
        return
    title = el.get("TITLE")
    source = el.get("SOURCE")

    # 제목 / 출처 공간 확보
    chart_y = y
    chart_h = h
    if title:
        _add_text(slide, x, y, w, 0.55, title, size_pt=14, weight=600, color=INK,
                  align="left", anchor="top", wrap=False)
        chart_y = y + 0.6
        chart_h = chart_h - 0.6
    if source:
        chart_h = chart_h - 0.5
        _add_text(slide, x, chart_y + chart_h + 0.05, w, 0.45, source,
                  size_pt=9, weight=400, color=MUTE, align="left", anchor="top", wrap=False)

    try:
        cdata = CategoryChartData()
        cdata.categories = labels or [str(i + 1) for i in range(
            len(series[0].get("values") or []))]
        for s in series:
            vals = [(_to_float(v) or 0) for v in (s.get("values") or [])]
            cdata.add_series(_s(s.get("name") or "series"), vals)

        type_map = {
            "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "donut": XL_CHART_TYPE.DOUGHNUT,
        }
        ctype = type_map.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        gframe = slide.shapes.add_chart(
            ctype, Cm(x), Cm(chart_y), Cm(w), Cm(chart_h), cdata
        )
        chart = gframe.chart
        chart.has_title = False
        if len(series) > 1 or kind == "donut":
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(9)
        else:
            chart.has_legend = False

        # series 색
        try:
            for i, plot_series in enumerate(chart.series):
                col = series[i].get("color") if i < len(series) and series[i].get("color") else None
                rgb = _rgb(col) if col else _SERIES_COLORS[i % len(_SERIES_COLORS)]
                fill = plot_series.format.fill
                fill.solid()
                fill.fore_color.rgb = rgb
                if kind == "line":
                    plot_series.format.line.color.rgb = rgb
                    plot_series.format.line.width = Pt(2.25)
        except Exception:
            pass

        try:
            if kind in ("bar", "hbar"):
                chart.plots[0].gap_width = 80
            if hasattr(chart, "category_axis") and chart.category_axis is not None:
                chart.category_axis.tick_labels.font.size = Pt(10)
                chart.category_axis.tick_labels.font.color.rgb = SUB
            if hasattr(chart, "value_axis") and chart.value_axis is not None:
                chart.value_axis.tick_labels.font.size = Pt(10)
                chart.value_axis.tick_labels.font.color.rgb = SUB
        except Exception:
            pass
        _kill_chart_shadows(chart)
    except Exception:
        pass


_RENDERERS = {
    "text": _el_text,
    "kpi": _el_kpi,
    "pill": _el_pill,
    "bullets": _el_bullets,
    "steps": _el_steps,
}


# ============================================================
# 슬라이드 빌더
# ============================================================
def build_brandlogy_slide(prs, data, logo_path=None):
    """단일 Brandlogy 슬라이드 (자유배치)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = data if isinstance(data, dict) else {}
    bg = (data.get("BG") or "white").lower()
    dark = bg in ("dark", "gradient")

    # 배경
    if bg == "dark":
        _add_rect(slide, 0, 0, SLIDE_W_CM, SLIDE_H_CM, fill=DARK)
    elif bg == "gradient":
        r = _add_rect(slide, 0, 0, SLIDE_W_CM, SLIDE_H_CM, fill=BLUE)
        _apply_gradient(r)

    ink = WHITE if dark else INK
    sub_c = _rgb("#cdd6e0") if dark else SUB
    eyebrow_c = _rgb("#cdd6e0") if dark else MUTE

    # eyebrow (잠금 존 위, 표지/구분 슬라이드)
    if data.get("EYEBROW"):
        _add_text(slide, MARGIN_L, 0.30, BODY_W, 0.45, data.get("EYEBROW"),
                  size_pt=10.5, weight=600, color=eyebrow_c, align="left",
                  anchor="middle", wrap=False)

    # 헤드라인 (잠금 존)
    hl = data.get("HEADLINE")
    if hl:
        hl_size = 32 if bg == "gradient" else 27
        _add_text(slide, MARGIN_L, HEADLINE_Y, BODY_W, HEADLINE_H, hl,
                  size_pt=hl_size, weight=700, color=ink, align="left",
                  anchor="middle", line_height=1.15, wrap=True)

    # 부제 (잠금 존)
    st = data.get("SUBTITLE")
    if st:
        _add_text(slide, MARGIN_L, SUBTITLE_Y, BODY_W, SUBTITLE_H, st,
                  size_pt=12.5, weight=500, color=sub_c, align="left",
                  anchor="middle", line_height=1.30, wrap=True)

    # 본문 박스 ELEMENTS
    elements = data.get("ELEMENTS")
    if not isinstance(elements, list):
        elements = []
    for el in elements:
        if not isinstance(el, dict):
            continue
        etype = (el.get("type") or "").lower()
        x, y, w, h = _clamp(
            _f(el.get("x"), BODY_X0), _f(el.get("y"), BODY_Y0),
            _f(el.get("w"), 6.0), _f(el.get("h"), 2.0),
        )
        try:
            if etype == "chart":
                _el_chart(slide, el, x, y, w, h)
            elif etype == "table":
                _el_table(slide, el, x, y, w, h)
            elif etype == "divider":
                _el_divider(slide, el, x, y, w, h)
            elif etype in ("kpi", "pill", "text", "bullets", "steps"):
                _RENDERERS[etype](slide, el, x, y, w, h, dark=dark)
        except Exception:
            # 한 요소 실패해도 슬라이드 전체는 진행
            continue

    return slide


LAYOUT_BUILDERS_BRANDLOGY = {
    "brandlogy": build_brandlogy_slide,
}


def build_slide_brandlogy(prs, layout_type, data, logo_path=None):
    builder = LAYOUT_BUILDERS_BRANDLOGY.get(layout_type) or build_brandlogy_slide
    return builder(prs, data, logo_path=logo_path)


__all__ = [
    "SLIDE_W",
    "SLIDE_H",
    "SLIDE_W_CM",
    "SLIDE_H_CM",
    "LAYOUT_BUILDERS_BRANDLOGY",
    "build_brandlogy_slide",
    "build_slide_brandlogy",
]
