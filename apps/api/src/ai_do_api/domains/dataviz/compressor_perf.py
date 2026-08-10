"""Compressor performance Excel parsing helpers.

The legacy 08_data_viz module stores parsed compressor rows in SQLite. This
module keeps only the workbook parsing and tolerance math so FastAPI services
can persist the rows in Postgres.
"""
from __future__ import annotations

import csv
import io
import re
from collections import OrderedDict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


CATEGORY_VARIABLE = "가변"
CATEGORY_ELECTRIC = "전동"
REFRIGERANT_NEW = "신냉매"
REFRIGERANT_OLD = "구냉매"

VARIABLE_CAPACITIES = ("122cc", "130cc", "140cc", "155cc", "160cc", "180cc", "190cc", "Benchmark")
ELECTRIC_CAPACITIES = ("19차수", "21차수")

# M1 sheet row mappings (1-based rows; columns are per T-sheet).
MAPPING_VARIABLE: dict[str, int] = {
    "rpm": 73,
    "pd": 4,
    "td": 25,
    "ps": 5,
    "ts": 26,
    "pc": 24,
    "mass_flow": 55,
    "vol_eff": 108,
    "ocr": 121,
    "cooling_cap_a": 105,
    "cooling_cap_f": 106,
    "power_kw": 80,
    "cop_sc": 109,
    "heat_balance": 104,
    "torque": 164,
}
MAPPING_ELECTRIC: dict[str, int] = {
    "rpm": 114,
    "pd": 4,
    "ps": 5,
    "td": 25,
    "ts": 26,
    "cooling_cap_a": 94,
    "power_kw": 72,
    "cop_sc": 108,
    "vol_eff": 100,
    "mass_flow": 131,
}

ES_RANGES = {"ES 1": (790, 810), "ES 2": (1980, 2020), "ES 3": (2970, 3030)}

NUMERIC_FIELDS = (
    "rpm",
    "pd",
    "td",
    "ps",
    "ts",
    "pc",
    "mass_flow",
    "vol_eff",
    "ocr",
    "cooling_cap_a",
    "cooling_cap_f",
    "power_kw",
    "cop_sc",
    "heat_balance",
    "torque",
)
TOLERANCE_FIELDS = (
    "rpm",
    "pd",
    "ps",
    "cooling_cap_a",
    "cooling_cap_f",
    "power_kw",
    "vol_eff",
    "cop_sc",
    "torque",
)

# 시험공차(승인도 성능 기준값) 프로필 입력 항목 — 회전수/토출압력/흡입압력/과열도/과냉도.
# 좌측 '용량별 기준값'(TOLERANCE_FIELDS)과 달리 superheat/subcool 을 포함한다.
TOLERANCE_PROFILE_FIELDS = (
    "rpm",
    "pd",
    "ps",
    "superheat",
    "subcool",
)


@dataclass
class PerfRow:
    test_group: str
    comp_type: str = ""
    serial_no: str = ""
    test_date: str = ""
    car_model: str = ""
    engine_spec: str = ""
    remarks: str = ""
    source_file: str = ""
    rpm: float | None = None
    pd: float | None = None
    td: float | None = None
    ps: float | None = None
    ts: float | None = None
    pc: float | None = None
    mass_flow: float | None = None
    vol_eff: float | None = None
    ocr: float | None = None
    cooling_cap_a: float | None = None
    cooling_cap_f: float | None = None
    power_kw: float | None = None
    cop_sc: float | None = None
    heat_balance: float | None = None
    torque: float | None = None


@dataclass
class MasterSheetRows:
    sheet_name: str
    refrigerant: str
    capacity: str
    rows: list[PerfRow]


@dataclass
class ToleranceRow:
    capacity: str
    car_model: str
    test_group: str
    field_name: str
    ref_value: float


def _as_float(value: object) -> float | None:
    try:
        return float(value) if value is not None else None  # type: ignore[arg-type]
    except (ValueError, TypeError):
        return None


def _as_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _format_date(value: object, fallback: str = "") -> str:
    if value is None:
        return fallback
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d %H:%M")  # type: ignore[union-attr]
    text = str(value).strip()
    return text or fallback


def parse_date_from_filename(filename: str) -> datetime:
    match = re.match(r"(\d{8})_(\d{4})", Path(filename).name)
    if match:
        try:
            return datetime.strptime(match.group(1) + match.group(2), "%Y%m%d%H%M")
        except ValueError:
            pass
    return datetime.now()


def date_text_from_filename(filename: str) -> str:
    match = re.match(r"(\d{8})_(\d{4})", Path(filename).name)
    if not match:
        return ""
    day, time_text = match.groups()
    return f"{day[:4]}-{day[4:6]}-{day[6:8]} {time_text[:2]}:{time_text[2:]}"


def match_car_model_from_filename(filename: str) -> str:
    name = Path(filename).stem
    parts = name.split("_")
    return parts[5].strip() if len(parts) >= 6 else ""


def match_remarks_from_filename(filename: str) -> str:
    # 파일명 양식: 날짜_시간_시험기_타입_기종_차종_엔진사양_Lot No_목적_비고
    # 비고는 10번째 토큰(인덱스 9)이며, 비어 있거나 토큰 자체가 없으면 빈 문자열로 둔다.
    name = Path(filename).stem
    parts = name.split("_")
    return parts[9].strip() if len(parts) >= 10 else ""


def parse_sheet_name(sheet_name: str, category: str) -> tuple[str | None, str | None]:
    name = sheet_name.strip()
    if category == CATEGORY_VARIABLE:
        for capacity in VARIABLE_CAPACITIES:
            if capacity in name:
                tokens = name.split("_")[:2]
                refrigerant = REFRIGERANT_OLD if REFRIGERANT_OLD in name or "구" in tokens else REFRIGERANT_NEW
                return refrigerant, capacity
    else:
        for capacity in ELECTRIC_CAPACITIES:
            if capacity in name:
                refrigerant = REFRIGERANT_OLD if REFRIGERANT_OLD in name or "구" in name else REFRIGERANT_NEW
                return refrigerant, capacity
    return None, None


def _merged_value(ws, row: int, col: int) -> object:
    value = ws.cell(row=row, column=col).value
    if value is not None:
        return value
    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= col <= merged_range.max_col
        ):
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
    return None


def _parse_variable_master_sheet(ws) -> list[PerfRow]:
    rows: list[PerfRow] = []
    for row_idx in range(7, ws.max_row + 1):
        if ws.cell(row=row_idx, column=7).value is None:
            continue

        def fv(col: int) -> float | None:
            return _as_float(ws.cell(row=row_idx, column=col).value)

        rows.append(
            PerfRow(
                test_group=_as_text(ws.cell(row=row_idx, column=2).value),
                comp_type=_as_text(_merged_value(ws, row_idx, 1)),
                serial_no=_as_text(_merged_value(ws, row_idx, 3)),
                test_date=_as_text(_merged_value(ws, row_idx, 21)),
                car_model=_as_text(_merged_value(ws, row_idx, 4)),
                engine_spec=_as_text(ws.cell(row=row_idx, column=5).value),
                remarks=_as_text(_merged_value(ws, row_idx, 22)),
                source_file="master",
                rpm=fv(7),
                pd=fv(8),
                td=fv(9),
                ps=fv(10),
                ts=fv(11),
                pc=fv(12),
                mass_flow=fv(13),
                vol_eff=fv(14),
                ocr=fv(15),
                cooling_cap_a=fv(16),
                power_kw=fv(17),
                cop_sc=fv(18),
                torque=fv(19),
            )
        )
    return rows


def _parse_electric_master_sheet(ws) -> list[PerfRow]:
    rows: list[PerfRow] = []
    for row_idx in range(5, ws.max_row + 1):
        if ws.cell(row=row_idx, column=4).value is None:
            continue

        def fv(col: int) -> float | None:
            return _as_float(ws.cell(row=row_idx, column=col).value)

        rows.append(
            PerfRow(
                test_group=_as_text(_merged_value(ws, row_idx, 14)),
                comp_type=_as_text(_merged_value(ws, row_idx, 1)),
                serial_no=_as_text(_merged_value(ws, row_idx, 3)),
                source_file="master",
                rpm=fv(4),
                pd=fv(5),
                ps=fv(6),
                td=fv(7),
                ts=fv(8),
                cooling_cap_a=fv(9),
                power_kw=fv(10),
                cop_sc=fv(11),
                vol_eff=fv(12),
                mass_flow=fv(13),
            )
        )
    return rows


def parse_perf_master(data: bytes, category: str) -> list[MasterSheetRows]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    try:
        imported: list[MasterSheetRows] = []
        for sheet_name in wb.sheetnames:
            refrigerant, capacity = parse_sheet_name(sheet_name, category)
            if not refrigerant or not capacity:
                continue
            ws = wb[sheet_name]
            rows = (
                _parse_variable_master_sheet(ws)
                if category == CATEGORY_VARIABLE
                else _parse_electric_master_sheet(ws)
            )
            if rows:
                imported.append(
                    MasterSheetRows(
                        sheet_name=sheet_name,
                        refrigerant=refrigerant,
                        capacity=capacity,
                        rows=rows,
                    )
                )
        return imported
    finally:
        wb.close()


def parse_perf_source(data: bytes, category: str, filename: str = "") -> list[PerfRow]:
    """Parse one source workbook's M1/T-sheets into per-ES performance rows."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    try:
        t_sheets = sorted(s for s in wb.sheetnames if s.startswith("T-"))
        if "M1" not in wb.sheetnames or not t_sheets:
            return []
        m1 = wb["M1"]
        t1 = wb[t_sheets[0]]
        comp_type = _as_text(t1.cell(row=5, column=6).value)
        serial_no = _as_text(t1.cell(row=6, column=6).value)
        file_date = date_text_from_filename(filename)

        filename_parts = Path(filename).stem.split("_") if filename else []
        if category == CATEGORY_ELECTRIC and len(filename_parts) >= 4:
            comp_type = filename_parts[2].strip() or comp_type
            serial_no = filename_parts[3].strip() or serial_no

        def tsheet_date(idx: int) -> str:
            if idx < len(t_sheets):
                raw = wb[t_sheets[idx]].cell(row=2, column=5).value
                return _format_date(raw, file_date)
            return file_date

        mapping = MAPPING_VARIABLE if category == CATEGORY_VARIABLE else MAPPING_ELECTRIC
        rows: list[PerfRow] = []

        for i in range(1 if category == CATEGORY_VARIABLE else 0, len(t_sheets)):
            col = 8 + i

            def fv(m1_row: int) -> float | None:
                return _as_float(m1.cell(row=m1_row, column=col).value)

            if category == CATEGORY_VARIABLE:
                rpm_val = fv(mapping["rpm"])
                test_group = None
                if rpm_val is not None:
                    for name, (lo, hi) in ES_RANGES.items():
                        if lo <= rpm_val <= hi:
                            test_group = name
                            break
                if not test_group:
                    continue
            else:
                test_group = f"ES {i + 1}"

            row = PerfRow(
                test_group=test_group,
                comp_type=comp_type,
                serial_no=serial_no,
                test_date=tsheet_date(i),
                car_model=match_car_model_from_filename(filename),
                remarks=match_remarks_from_filename(filename),
                source_file=filename,
            )
            for field, m1_row in mapping.items():
                if field == "mass_flow" and category == CATEGORY_VARIABLE:
                    r55, r56 = fv(55), fv(56)
                    setattr(row, field, r55 if (r55 is not None and r55 > 10) else r56)
                else:
                    setattr(row, field, fv(m1_row))
            rows.append(row)
        return rows
    finally:
        wb.close()


def rows_to_dicts(rows: Iterable[PerfRow]) -> list[dict]:
    return [asdict(row) for row in rows]


def average_by_group(rows: Iterable[PerfRow]) -> list[dict]:
    buckets: dict[str, list[PerfRow]] = {}
    for row in rows:
        buckets.setdefault(row.test_group, []).append(row)
    out: list[dict] = []
    for group in sorted(buckets):
        members = buckets[group]
        avg: dict[str, object] = {"test_group": group, "count": len(members)}
        for field in NUMERIC_FIELDS:
            vals = [getattr(member, field) for member in members if getattr(member, field) is not None]
            avg[field] = round(sum(vals) / len(vals), 4) if vals else None
        out.append(avg)
    return out


# 승인도 성능 기준값(주황) — 냉방성능·체적효율은 기준값 이상이어야 하고(min),
# 소요동력은 기준값 이하여야 한다(max). 위반 시 'ref' 로 표시.
REF_MIN_FIELDS = ("cooling_cap_a", "vol_eff")
REF_MAX_FIELDS = ("power_kw",)
# 시험공차(빨강) — 회전수·토출압력·흡입압력을 시험공차 범위와 비교. 위반 시 'tol'.
TOL_RANGE_FIELDS = ("rpm", "pd", "ps")


def profile_key_from_test_group(test_group: str) -> tuple[str | None, str | None]:
    # 시험군 'ES 1' → (카테고리 'ES', 구분 'ES1'). 끝 숫자를 떼어 접두어=카테고리,
    # 접두어+숫자=구분 키로 만든다(공백 제거).
    match = re.search(r"(\d+)\s*$", test_group)
    if not match:
        return None, None
    prefix = test_group[: match.start()].strip()
    if not prefix:
        return None, None
    return prefix, f"{prefix}{match.group(1)}"


def check_warnings(
    row: dict,
    ref_values: dict[str, dict[str, float | None]],
    all_profiles: dict[str, dict[str, dict[str, dict]]],
) -> dict[str, str]:
    # 반환: { field: 'ref' | 'tol' }. 'ref'=승인도 기준값 위반(주황 셀),
    # 'tol'=시험공차 위반(빨간 글자). 두 경고는 항목 집합이 겹치지 않는다.
    result: dict[str, str] = {}
    test_group = str(row.get("test_group") or "")

    # 승인도 기준값 — 시험군(test_group) 그대로 매칭(용량별 기준값).
    refs = ref_values.get(test_group, {}) if ref_values else {}
    for field in REF_MIN_FIELDS:
        value, ref = row.get(field), refs.get(field)
        if value is not None and ref is not None and float(value) < float(ref):
            result[field] = "ref"
    for field in REF_MAX_FIELDS:
        value, ref = row.get(field), refs.get(field)
        if value is not None and ref is not None and float(value) > float(ref):
            result[field] = "ref"

    # 시험공차 — 시험군 접두어로 프로필(카테고리)·구분을 찾아 기준값±공차와 비교.
    prefix, group_key = profile_key_from_test_group(test_group)
    cells: dict[str, dict] = {}
    if prefix and group_key:
        cells = (all_profiles.get(prefix) or {}).get(group_key) or {}
    for field in TOL_RANGE_FIELDS:
        value = row.get(field)
        cell = cells.get(field)
        if value is None or not isinstance(cell, dict):
            continue
        ref = cell.get("ref")
        if ref is None:
            continue
        measured = float(value)
        ref_value = float(ref)
        sign = cell.get("sign") or "±"
        up = float(cell["tol"]) if cell.get("tol") is not None else 0.0
        down = float(cell["lower"]) if cell.get("lower") is not None else 0.0
        if sign == "+":
            low, high = ref_value, ref_value + up
        elif sign == "-":
            low, high = ref_value - up, ref_value
        elif sign == "편측":
            low, high = ref_value - down, ref_value + up
        else:  # '±'
            low, high = ref_value - up, ref_value + up
        if measured < low or measured > high:
            result[field] = "tol"
    return result


def parse_tolerance_reference_file(data: bytes, filename: str) -> list[ToleranceRow]:
    rows = _read_tolerance_rows(data, filename)
    grouped: "OrderedDict[tuple[str, str], list[list[object]]]" = OrderedDict()
    for row in rows:
        if len(row) < 9:
            continue
        capacity = _as_text(row[0])
        car_model = _as_text(row[1])
        if not capacity or not car_model:
            continue
        if capacity and not capacity.endswith("cc") and "차수" not in capacity:
            capacity = f"{capacity}cc"
        grouped.setdefault((capacity, car_model), []).append(row)

    field_map = {
        2: "rpm",
        3: "pd",
        4: "ps",
        5: "cooling_cap_a",
        6: "power_kw",
        7: "torque",
        8: "vol_eff",
    }
    es_labels = ["ES 1", "ES 2", "ES 3", "ES 4", "ES 5", "ES 6"]
    parsed: list[ToleranceRow] = []
    for (capacity, car_model), group_rows in grouped.items():
        for index, row in enumerate(group_rows[: len(es_labels)]):
            test_group = es_labels[index]
            for col_index, field_name in field_map.items():
                value = _as_float(row[col_index] if col_index < len(row) else None)
                if value is None:
                    continue
                parsed.append(
                    ToleranceRow(
                        capacity=capacity,
                        car_model=car_model,
                        test_group=test_group,
                        field_name=field_name,
                        ref_value=value,
                    )
                )
    return parsed


def _read_tolerance_rows(data: bytes, filename: str) -> list[list[object]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = _decode_csv(data)
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        return [row for idx, row in enumerate(all_rows) if idx >= 2]

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        ws = wb.active
        return [list(row) for idx, row in enumerate(ws.iter_rows(values_only=True)) if idx >= 2]
    finally:
        wb.close()


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")
