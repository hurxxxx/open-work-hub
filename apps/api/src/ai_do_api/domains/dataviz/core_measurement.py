"""Parse the legacy "코어 측정값 정리" core-measurement Excel.

Layout (per data sheet, e.g. 스테이터 / 로터):
  - col A: 기종 (model, only on the first row of each group; may be multi-line)
  - col B: 검사항목 (inspection item)
  - col D: 규격 (spec string, e.g. "96 +0, -0.05", "0.3 이하", "5Kgf 이상일 것")
  - col E..: measured values; a date-header row sits above each model group
    with dates spaced across column blocks.

This module is pure I/O + parsing (no FastAPI imports) so it stays testable.
"""
from __future__ import annotations

import datetime as _dt
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

SAMPLE_DIR = Path(__file__).resolve().parent / "sample_data"


def data_dir() -> Path:
    """Directory holding the core-measurement workbook(s).

    Override with AI_DO_API_DATAVIZ_DATA_DIR; defaults to the committed sample.
    """
    override = os.environ.get("AI_DO_API_DATAVIZ_DATA_DIR", "").strip()
    return Path(override) if override else SAMPLE_DIR


@dataclass
class Spec:
    raw: str
    nominal: float | None = None
    lsl: float | None = None
    usl: float | None = None


@dataclass
class Point:
    index: int
    value: float
    date: str | None = None


@dataclass
class MeasurementItem:
    item: str
    spec: Spec
    points: list[Point] = field(default_factory=list)


@dataclass
class ModelGroup:
    file: str
    sheet: str
    model: str
    items: list[MeasurementItem] = field(default_factory=list)

    @property
    def id(self) -> str:
        return f"{self.file}::{self.sheet}::{self.model}"


_NUM = re.compile(r"[+-]?\d+(?:\.\d+)?")


def parse_spec(raw: object) -> Spec:
    """Turn a spec string into nominal/LSL/USL. Best-effort, tolerant of formats."""
    s = "" if raw is None else str(raw).strip()
    nums = [float(x) for x in _NUM.findall(s)]
    if not nums:
        return Spec(raw=s)
    first = nums[0]
    if "이상" in s:  # minimum bound, e.g. "5Kgf 이상일 것"
        return Spec(raw=s, lsl=first)
    if "이하" in s or "이내" in s:  # maximum bound, e.g. "0.3 이하"
        return Spec(raw=s, usl=first, lsl=0.0)
    if "±" in s or "+-" in s:  # nominal ± tol
        tol = abs(nums[1]) if len(nums) > 1 else 0.0
        return Spec(raw=s, nominal=first, usl=first + tol, lsl=first - tol)
    # nominal + signed offsets, e.g. "96 +0, -0.05", "59.5 +0.02/-0.03"
    offsets = nums[1:]
    upper = max(offsets) if offsets else 0.0
    lower = min(offsets) if offsets else 0.0
    return Spec(raw=s, nominal=first, usl=first + upper, lsl=first + lower)


def _norm_model(value: object) -> str:
    return " ".join(str(value).split())


def _as_date(value: object) -> str | None:
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    return None


def _is_number(value: object) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    return False


def _parse_sheet(file: str, sheet: str, rows: list[tuple]) -> list[ModelGroup]:
    groups: list[ModelGroup] = []
    current: ModelGroup | None = None
    col_dates: dict[int, str] = {}
    for row in rows:
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        d = row[3] if len(row) > 3 else None
        tail = row[4:]

        # Date-header row: no item/spec but dates across the measurement columns.
        if (b is None or str(b).strip() == "") and (d is None or str(d).strip() == ""):
            found = {i + 4: _as_date(v) for i, v in enumerate(tail) if _as_date(v)}
            if found:
                col_dates = {}
                last: str | None = None
                for c in range(4, 4 + len(tail)):
                    if c in found:
                        last = found[c]
                    col_dates[c] = last
                continue

        if a is not None and str(a).strip():
            model = _norm_model(a)
            if current is None or current.model != model:
                current = ModelGroup(file=file, sheet=sheet, model=model)
                groups.append(current)

        # Item row: needs an item name, a spec, and a model context.
        if current is None or b is None or str(b).strip() == "":
            continue
        if d is None or str(d).strip() == "":
            continue

        item = MeasurementItem(item=str(b).strip(), spec=parse_spec(d))
        seq = 0
        for i, v in enumerate(tail):
            if _is_number(v):
                seq += 1
                item.points.append(
                    Point(index=seq, value=float(v), date=col_dates.get(i + 4))
                )
        if item.points:
            current.items.append(item)
    return groups


def parse_workbook(path: Path) -> list[ModelGroup]:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        groups: list[ModelGroup] = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            sheet_groups = _parse_sheet(path.name, sheet, rows)
            groups.extend(g for g in sheet_groups if g.items)
        return groups
    finally:
        wb.close()


def load_groups() -> list[ModelGroup]:
    """Parse every .xlsx in the data dir (sample by default)."""
    groups: list[ModelGroup] = []
    base = data_dir()
    if not base.exists():
        return groups
    for path in sorted(base.glob("*.xlsx")):
        if path.name.startswith("~$"):  # Excel lock file
            continue
        try:
            groups.extend(parse_workbook(path))
        except Exception:  # noqa: BLE001 — skip unreadable workbooks, keep the rest
            continue
    return groups
