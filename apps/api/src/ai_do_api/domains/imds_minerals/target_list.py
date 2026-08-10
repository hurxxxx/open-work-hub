"""(현대차/기아 지정) 책임광물 조사대상품목 리스트 파싱 + OEM품번 매칭.

조사대상품목 리스트 엑셀에서 OEM품번 → (차종, DCC품번, 품명) 매핑을 추출한다.
헤더 행 위치와 열 순서는 양식마다 다를 수 있어 라벨로 탐지하고, 차종은 병합셀이라
아래로 forward-fill 한다. 특정 차종/품번 전용 분기 없이 라벨/열 기준으로만 동작한다.
"""

from __future__ import annotations

import io

import openpyxl

_CAR_LABEL = "차종"
_OEM_LABEL = "OEM품번"
_DCC_LABEL = "DCC품번"
_NAME_LABEL = "품명"
_HEADER_SCAN_ROWS = 20


class ImdsListError(Exception):
    """조사대상품목 리스트를 인식하지 못했을 때 발생한다."""


def _norm(value: object) -> str:
    return str(value).strip().upper() if value not in (None, "") else ""


def parse_target_list(content: bytes) -> dict[str, dict]:
    """리스트 엑셀 → {정규화된 OEM품번: {car, oem, dcc, end_name}} 매핑."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as error:  # noqa: BLE001 - openpyxl raises broad errors
        raise ImdsListError(f"cannot open target list: {error}") from error

    for ws in wb.worksheets:
        header_row = None
        cols: dict[str, int | None] = {}
        for r in range(1, min(ws.max_row, _HEADER_SCAN_ROWS) + 1):
            labels = {
                str(ws.cell(r, c).value).strip(): c
                for c in range(1, ws.max_column + 1)
                if ws.cell(r, c).value not in (None, "")
            }
            if not any(_OEM_LABEL in key for key in labels):
                continue

            def find(target: str) -> int | None:
                for key, col in labels.items():
                    if target in key:
                        return col
                return None

            oem_col = find(_OEM_LABEL)
            if oem_col is None:
                continue
            cols = {
                "car": find(_CAR_LABEL),
                "oem": oem_col,
                "dcc": find(_DCC_LABEL),
                "name": find(_NAME_LABEL),
            }
            header_row = r
            break

        if header_row is None:
            continue

        table: dict[str, dict] = {}
        last_car = ""
        for r in range(header_row + 1, ws.max_row + 1):
            oem = ws.cell(r, cols["oem"]).value if cols["oem"] else None
            car = ws.cell(r, cols["car"]).value if cols["car"] else None
            if car not in (None, ""):
                last_car = str(car).replace("\n", " ").strip()
            if oem in (None, ""):
                continue
            table[_norm(oem)] = {
                "car": last_car,
                "oem": str(oem).strip(),
                "dcc": str(ws.cell(r, cols["dcc"]).value or "").strip() if cols["dcc"] else "",
                "end_name": str(ws.cell(r, cols["name"]).value or "").strip()
                if cols["name"]
                else "",
            }
        if table:
            return table

    raise ImdsListError("no target-list rows recognized")


def match_oem(content: bytes, oem: str) -> dict | None:
    """OEM품번으로 리스트에서 일치 행을 찾는다(없으면 None)."""
    return parse_target_list(content).get(_norm(oem))
