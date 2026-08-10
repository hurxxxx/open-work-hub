from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from urllib.parse import quote

from fastapi import status
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from sqlalchemy import Text, cast, func, or_, select
from sqlalchemy.orm import Session

from ai_do_api.core.i18n import localized_http_exception
from ai_do_api.core.settings import get_settings
from ai_do_api.core.storage import ensure_bucket, get_minio_client
from ai_do_api.domains.auth.models import OrgUnit, User, Workspace, utcnow_naive
from ai_do_api.domains.auth.security import new_id
from ai_do_api.domains.legacy_issues.models import (
    LegacyIssueAiChunk,
    LegacyIssueAttachment,
    LegacyIssueAttachmentArtifact,
    LegacyIssueModuleField,
    LegacyIssueRecord,
    LegacyIssueSystemFieldSetting,
)
from ai_do_api.domains.legacy_issues.module_access import (
    LEGACY_ISSUE_MODULE_KEYS,
)
from ai_do_api.domains.legacy_issues.partitioning import (
    ensure_attachment_partition,
    ensure_chunk_partition,
    ensure_record_partition,
    ensure_revision_partition,
)
from ai_do_api.domains.legacy_issues.history import (
    LEGACY_ISSUE_RECORD_KIND,
    add_attachment_history_entry,
    add_record_history_entries,
    collect_create_changes,
    collect_value_changes,
    normalize_history_value,
)
from ai_do_api.domains.legacy_issues.revisioning import (
    LegacyIssueDataRevision,
    REVISION_STATUS_PUBLISHED,
    add_revision_event,
    create_draft_revision,
    ensure_initial_published_revision,
    get_latest_published_revision,
    get_revision,
    legacy_issue_dataset_revision_key,
)
from ai_do_api.domains.legacy_issues.tabular_import import (
    EXPORT_RECORD_ID_HEADER,
    LegacyIssueImportPreview,
    LegacyIssueImportTable,
    build_import_preview as build_tabular_import_preview,
    combine_header_rows,
    normalize_header,
    parse_required_mapping_json,
    parse_tabular_upload,
    validate_mapping as validate_tabular_mapping,
)


DATASET_ATTACHMENT_MAX_BYTES = 1024 * 1024 * 1024
DATASET_ATTACHMENT_CHUNK_SIZE = 1024 * 1024
DATASET_ATTACHMENT_DESCRIPTION_MAX_CHARS = 500
DATASET_RECORD_EXPORT_LIMIT = 100_000
DRAFT_COPY_ATTACHMENT_INDEX_BATCH_SIZE = 25
DEFAULT_CONTENT_TYPE = "application/octet-stream"
EXCEL_DATE_NUMBER_FORMAT = "yyyy-mm-dd"
DATASET_IMPORT_INVALID_CODE = "legacy_issues.dataset_import_file_invalid"
DATASET_RECORD_EMPTY_CODE = "legacy_issues.dataset_record_empty"
DATASET_REQUIRED_FIELD_MISSING_CODE = "legacy_issues.dataset_required_field_missing"
INTRODUCED_REVISION_FIELD_KEY = "introduced_revision_no"
REFLECTED_REVISION_FIELD_KEY = "reflected_revision"
REGION_FIELD_KEY = "region_zone"
LEGACY_CLAIM_REGION_FIELD_KEY = "claim_region"
COUNTERMEASURE_TYPE_OPTIONS = ("실제", "대외용")
OEM_DISCLOSURE_STATUS_OPTIONS = ("공개", "비공개", "검토중")
CLAIM_REGION_OPTIONS = ("북미", "아중동", "아태", "유럽", "인도", "중국", "중남미", "한국")
OEM_DISCLOSURE_STATUS_FIELD_KEY = "oem_disclosure_status"
LEGACY_OEM_OPEN_FIELD_KEY = "oem_open"
MASTER_STATUS_FIELD_KEY = "evidence_legacy_issue"
MASTER_STATUS_OPTIONS = ("심사대기", "심사제외", "등재", "미등재", "보류")
LEGACY_ISSUE_EVIDENCE_FIELD_KEY = "evidence_past_vehicle_issue"
OX_SELECT_OPTIONS = ("O", "X")
OX_SELECT_FIELD_KEYS = frozenset(
    {
        "applied",
        LEGACY_ISSUE_EVIDENCE_FIELD_KEY,
        "evidence_quality_spec",
    }
)
ATTACHMENT_COMPARE_FIELD_KEY = "primary_attachment"
ATTACHMENT_COMPARE_FIELD_LABEL = "첨부파일"


@dataclass(frozen=True)
class DatasetFieldDefinition:
    key: str
    label_ko: str
    label_en: str
    group_key: str | None = None
    aliases: tuple[str, ...] = ()
    field_type: str = "text"
    options: tuple[str, ...] = ()
    allow_multiple: bool = False
    required: bool = False
    source: str = "system"
    module_key: str | None = None
    field_id: str | None = None
    active: bool = True
    readonly: bool = False
    settings_readonly: bool = False


@dataclass(frozen=True)
class DatasetSystemFieldSettingView:
    id: str | None
    dataset_key: str
    field: DatasetFieldDefinition
    updated_at: Any | None = None


@dataclass(frozen=True)
class LegacyIssueDatasetDefinition:
    key: str
    table_model: type
    title_ko: str
    title_en: str
    hierarchy_ko: tuple[str, ...]
    hierarchy_en: tuple[str, ...]
    fields: tuple[DatasetFieldDefinition, ...]
    header_rows: int = 1
    group_labels_ko: dict[str, str] = field(default_factory=dict)
    group_labels_en: dict[str, str] = field(default_factory=dict)
    import_sheet_name: str | None = None


@dataclass(frozen=True)
class DatasetAttachmentUpload:
    filename: str | None
    content_type: str | None
    content: bytes
    is_primary: bool = False
    description: str | None = None


@dataclass(frozen=True)
class DatasetAttachmentContent:
    body: Iterable[bytes]
    media_type: str
    headers: dict[str, str]


def dataset_attachment_content_headers(filename: str | None) -> dict[str, str]:
    encoded_filename = quote(filename or "attachment", safe="")
    return {
        "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}",
    }


@dataclass(frozen=True)
class DatasetImportResult:
    created: int
    updated: int
    skipped: int
    total_rows: int


@dataclass(frozen=True)
class DatasetRevisionCompareCell:
    field_key: str
    field_label: str
    left_value: str | None
    right_value: str | None
    changed: bool


@dataclass(frozen=True)
class DatasetRevisionCompareRow:
    stable_record_id: str
    status: str
    label: str
    cells: list[DatasetRevisionCompareCell]


def _field(
    key: str,
    label_ko: str,
    label_en: str,
    *,
    group_key: str | None = None,
    aliases: tuple[str, ...] = (),
    field_type: str = "text",
    options: tuple[str, ...] = (),
    allow_multiple: bool = False,
    required: bool = False,
    source: str = "system",
    module_key: str | None = None,
    field_id: str | None = None,
    active: bool = True,
    readonly: bool = False,
    settings_readonly: bool | None = None,
) -> DatasetFieldDefinition:
    return DatasetFieldDefinition(
        key=key,
        label_ko=label_ko,
        label_en=label_en,
        group_key=group_key,
        aliases=aliases,
        field_type=field_type,
        options=options,
        allow_multiple=allow_multiple,
        required=required,
        source=source,
        module_key=module_key,
        field_id=field_id,
        active=active,
        readonly=readonly,
        settings_readonly=readonly if settings_readonly is None else settings_readonly,
    )


COMMON_MASTER_DATASET_KEY = "common-master"
MODULE_CUSTOM_GROUP_KEY = "module_custom"
LEGACY_ISSUE_MODULE_DEPARTMENT_LABELS = {
    "샤시(ACON)": "aircon",
    "컴프레서(전동)": "compressor-electric",
    "컴프레서(기계)": "compressor-mechanical",
    "열교환기": "heat-exchanger",
    "의장(HVAC)": "interior",
    "쿨링모듈": "cooling-module",
    "전장(기구)": "electrical-mechanical",
    "전장(제어-HW)": "electrical-control-hw",
    "전장(제어-SW)": "electrical-control-sw",
}
LEGACY_ISSUE_MODULE_FIELD_TYPES = frozenset(
    {"text", "longText", "number", "date", "select", "boolean", "user", "orgUnit"}
)
LEGACY_ISSUE_MULTIPLE_FIELD_TYPES = frozenset({"select", "user", "orgUnit"})

COMMON_MASTER_FIELDS = (
    _field(
        INTRODUCED_REVISION_FIELD_KEY,
        "반영Rev",
        "Reflected Rev.",
        aliases=("반영Rev.", "리비전", "Revision"),
        field_type="number",
        settings_readonly=True,
    ),
    _field("department", "부서", "Department"),
    _field(
        "registrant",
        "대책 작성자",
        "Countermeasure Author",
        aliases=("등록자", "Registrant"),
    ),
    _field(
        "legacy_issue_number",
        "과거차관리번호",
        "Past Vehicle Issue No.",
        aliases=("과거차 관리번호",),
    ),
    _field("major_category", "대분류", "Major Category"),
    _field("middle_category", "중분류", "Middle Category"),
    _field(
        REGION_FIELD_KEY,
        "권역",
        "Region",
        aliases=("클레임 권역", "Claim Region", LEGACY_CLAIM_REGION_FIELD_KEY),
        field_type="select",
        options=CLAIM_REGION_OPTIONS,
        required=False,
    ),
    _field("occurrence_stage", "발생단계", "Occurrence Stage"),
    _field("occurrence_type", "발생구분", "Occurrence Type"),
    _field(
        OEM_DISCLOSURE_STATUS_FIELD_KEY,
        "OEM 공개여부",
        "OEM Disclosure",
        aliases=("OEM오픈", "OEM 오픈", "OEM Open"),
        field_type="select",
        options=OEM_DISCLOSURE_STATUS_OPTIONS,
    ),
    _field("vehicle_model", "차종", "Vehicle Model"),
    _field("occurrence_date", "발생일", "Occurrence Date", field_type="date"),
    _field("received_date", "접수일", "Received Date", field_type="date"),
    _field("issue_type", "유형", "Issue Type"),
    _field("cause_type", "원인구분", "Cause Type"),
    _field("supplier", "협력사", "Supplier"),
    _field("part_number", "부품번호", "Part No."),
    _field("process_name", "공정명", "Process Name"),
    _field("symptom", "현상", "Symptom"),
    _field("cause", "원인", "Cause"),
    _field("countermeasure", "개선대책", "Countermeasure"),
    _field(
        "countermeasure_type",
        "대책 구분",
        "Countermeasure Type",
        field_type="select",
        options=COUNTERMEASURE_TYPE_OPTIONS,
    ),
    _field("action", "조치", "Action"),
    _field("severity_grade", "중요도/등급", "Severity/Grade"),
    _field("confirmation_content", "확인내용", "Confirmation"),
    _field(
        "check_plan", "점검방안", "Check Plan", group_key="check", aliases=("CHECK - 점검방안",)
    ),
    _field(
        "applied",
        "적용유무",
        "Applied",
        group_key="check",
        aliases=("CHECK - 적용유무",),
        field_type="select",
        options=OX_SELECT_OPTIONS,
    ),
    _field(
        "reflection_result",
        "반영/검토결과",
        "Reflection/Review Result",
        group_key="check",
        aliases=("CHECK - 반영/검토결과",),
    ),
    _field(
        MASTER_STATUS_FIELD_KEY,
        "마스터 상태",
        "Master Status",
        group_key="evidence",
        aliases=("점검 근거 - 마스터 상태",),
        field_type="select",
        options=MASTER_STATUS_OPTIONS,
    ),
    _field(
        LEGACY_ISSUE_EVIDENCE_FIELD_KEY,
        "과거차 문제점",
        "Past Vehicle Issue",
        group_key="evidence",
        aliases=(
            "과거차문제점",
            "점검 근거 - 과거차문제점",
            "점검 근거 - 과거차 문제점",
        ),
        field_type="select",
        options=OX_SELECT_OPTIONS,
    ),
    _field(
        "evidence_design_check_sheet",
        "설계체크시트",
        "Design Check Sheet",
        group_key="evidence",
        aliases=("점검 근거 - 설계체크시트",),
    ),
    _field(
        "evidence_design_fmea",
        "설계FMEA",
        "Design FMEA",
        group_key="evidence",
        aliases=("점검 근거 - 설계FMEA",),
    ),
    _field(
        "evidence_design_standard_guide",
        "설계표준·가이드",
        "Design Standard/Guide",
        group_key="evidence",
        aliases=("점검 근거 - 설계표준·가이드", "설계표준/가이드"),
    ),
    _field(
        "evidence_quality_spec",
        "품질규격",
        "Quality Spec",
        group_key="evidence",
        aliases=("점검 근거 - 품질규격",),
        field_type="select",
        options=OX_SELECT_OPTIONS,
    ),
    _field("notes", "비고", "Notes"),
)

DATASET_DEFINITIONS = {
    COMMON_MASTER_DATASET_KEY: LegacyIssueDatasetDefinition(
        key=COMMON_MASTER_DATASET_KEY,
        table_model=LegacyIssueRecord,
        title_ko="과거차 전체",
        title_en="Past Vehicle Issue All",
        hierarchy_ko=("전체",),
        hierarchy_en=("All",),
        fields=COMMON_MASTER_FIELDS,
        header_rows=2,
        group_labels_ko={"check": "CHECK", "evidence": "점검 근거"},
        group_labels_en={"check": "CHECK", "evidence": "Evidence"},
        import_sheet_name="통합DB",
    ),
}

PROJECTED_RECORD_FIELD_KEYS = tuple(
    field_definition.key for field_definition in COMMON_MASTER_FIELDS
)
DEPRECATED_PROJECTED_RECORD_FIELD_KEYS = (
    LEGACY_OEM_OPEN_FIELD_KEY,
    REFLECTED_REVISION_FIELD_KEY,
)
PROJECTED_RECORD_LABELS = {
    field_definition.key: field_definition.label_ko for field_definition in COMMON_MASTER_FIELDS
}
SEARCH_TEXT_PRIORITY_FIELDS = (
    INTRODUCED_REVISION_FIELD_KEY,
    "department",
    "legacy_issue_number",
    "major_category",
    "middle_category",
    REGION_FIELD_KEY,
    "occurrence_stage",
    "occurrence_type",
    OEM_DISCLOSURE_STATUS_FIELD_KEY,
    "vehicle_model",
    "occurrence_date",
    "received_date",
    "issue_type",
    "cause_type",
    "supplier",
    "part_number",
    "process_name",
    "severity_grade",
    "applied",
    "reflected_revision",
    "symptom",
    "cause",
    "countermeasure",
    "countermeasure_type",
    "action",
    "confirmation_content",
    "check_plan",
    "reflection_result",
    MASTER_STATUS_FIELD_KEY,
    "evidence_design_check_sheet",
    "evidence_design_fmea",
    "evidence_design_standard_guide",
    "evidence_quality_spec",
    "notes",
)


def build_record_projection(
    values: dict[str, Any],
    *,
    raw_fields: dict[str, Any] | None = None,
    value_labels: dict[str, str] | None = None,
) -> dict[str, Any]:
    values = canonicalize_dataset_values(values)
    projection: dict[str, Any] = {}
    for key in PROJECTED_RECORD_FIELD_KEYS:
        if key == INTRODUCED_REVISION_FIELD_KEY:
            projection[key] = _clean_projection_integer(values.get(key))
        else:
            projection[key] = _clean_projection_value(values.get(key))
    projection["search_text"] = build_record_search_text(
        values,
        raw_fields=raw_fields,
        value_labels=value_labels,
    )
    return projection


def apply_record_projection(
    record: LegacyIssueRecord,
    values: dict[str, Any],
    *,
    raw_fields: dict[str, Any] | None = None,
    value_labels: dict[str, str] | None = None,
) -> None:
    for key, value in build_record_projection(
        values,
        raw_fields=raw_fields,
        value_labels=value_labels,
    ).items():
        setattr(record, key, value)
    for key in DEPRECATED_PROJECTED_RECORD_FIELD_KEYS:
        if hasattr(record, key):
            setattr(record, key, None)


def build_record_search_text(
    values: dict[str, Any],
    *,
    raw_fields: dict[str, Any] | None = None,
    value_labels: dict[str, str] | None = None,
) -> str | None:
    values = canonicalize_dataset_values(values)
    lines: list[str] = []
    seen: set[str] = set()
    labels = value_labels or {}
    for field_key in SEARCH_TEXT_PRIORITY_FIELDS:
        value = searchable_dataset_value_text(values.get(field_key))
        if not value:
            continue
        label = labels.get(field_key) or PROJECTED_RECORD_LABELS.get(field_key, field_key)
        _append_search_text_line(lines, seen, f"{label}: {value}")
    for field_key in PROJECTED_RECORD_FIELD_KEYS:
        if field_key in SEARCH_TEXT_PRIORITY_FIELDS:
            continue
        value = searchable_dataset_value_text(values.get(field_key))
        if not value:
            continue
        label = labels.get(field_key) or PROJECTED_RECORD_LABELS.get(field_key, field_key)
        _append_search_text_line(lines, seen, f"{label}: {value}")
    for field_key, raw_value in values.items():
        if field_key in PROJECTED_RECORD_FIELD_KEYS:
            continue
        value = searchable_dataset_value_text(raw_value)
        if not value:
            continue
        label = labels.get(field_key) or field_key
        _append_search_text_line(lines, seen, f"{label}: {value}")
    for key, value in (raw_fields or {}).items():
        clean_key = _clean_projection_value(key)
        clean_value = _clean_projection_value(value)
        if not clean_key or not clean_value:
            continue
        _append_search_text_line(lines, seen, f"{clean_key}: {clean_value}")
    return "\n".join(lines) or None


def _append_search_text_line(lines: list[str], seen: set[str], line: str) -> None:
    normalized = " ".join(line.split())
    if not normalized or normalized in seen:
        return
    seen.add(normalized)
    lines.append(normalized)


def display_dataset_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, list):
        parts = [display_dataset_value(item) for item in value]
        text = "; ".join(part for part in parts if part)
        return text or None
    if isinstance(value, dict):
        for key in ("label", "name", "path", "email", "id"):
            text = _clean_scalar_text(value.get(key))
            if text:
                return text
        return None
    return _clean_scalar_text(value)


def searchable_dataset_value_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, list):
        parts = [searchable_dataset_value_text(item) for item in value]
        text = "; ".join(part for part in parts if part)
        return text or None
    if isinstance(value, dict):
        parts: list[str] = []
        seen: set[str] = set()
        for key in ("label", "name", "email", "path", "id"):
            text = _clean_scalar_text(value.get(key))
            if text and text not in seen:
                seen.add(text)
                parts.append(text)
        return " ".join(parts) or None
    return _clean_scalar_text(value)


def _clean_projection_value(value: Any) -> str | None:
    return display_dataset_value(value)


def _clean_scalar_text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = " ".join(str(value).strip().split())
    return normalized or None


def _clean_projection_integer(value: Any) -> int | None:
    clean_value = _clean_projection_value(value)
    if clean_value is None:
        return None
    try:
        return int(clean_value)
    except ValueError:
        return None


def canonicalize_dataset_values(values: dict[str, Any] | None) -> dict[str, Any]:
    canonical = dict(values or {})
    legacy_reflected_revision = canonical.pop(REFLECTED_REVISION_FIELD_KEY, None)
    if _clean_projection_value(legacy_reflected_revision):
        canonical[INTRODUCED_REVISION_FIELD_KEY] = legacy_reflected_revision
    legacy_region_value = canonical.pop(LEGACY_CLAIM_REGION_FIELD_KEY, None)
    if not _clean_projection_value(canonical.get(REGION_FIELD_KEY)) and _clean_projection_value(
        legacy_region_value
    ):
        canonical[REGION_FIELD_KEY] = legacy_region_value
    legacy_oem_open_value = canonical.pop(LEGACY_OEM_OPEN_FIELD_KEY, None)
    if not _clean_projection_value(
        canonical.get(OEM_DISCLOSURE_STATUS_FIELD_KEY)
    ) and _clean_projection_value(legacy_oem_open_value):
        canonical[OEM_DISCLOSURE_STATUS_FIELD_KEY] = legacy_oem_open_value
    oem_disclosure_status = _normalize_oem_disclosure_status(
        canonical.get(OEM_DISCLOSURE_STATUS_FIELD_KEY)
    )
    if oem_disclosure_status is not None:
        canonical[OEM_DISCLOSURE_STATUS_FIELD_KEY] = oem_disclosure_status
    master_status_value = _normalize_legacy_master_status(canonical.get(MASTER_STATUS_FIELD_KEY))
    if master_status_value is not None:
        canonical[MASTER_STATUS_FIELD_KEY] = master_status_value
    for field_key in OX_SELECT_FIELD_KEYS:
        ox_value = _normalize_ox_select_value(canonical.get(field_key))
        if ox_value is not None:
            canonical[field_key] = ox_value
    return canonical


def _normalize_oem_disclosure_status(value: Any) -> str | None:
    clean_value = _clean_projection_value(value)
    if clean_value is None:
        return None
    if clean_value in OEM_DISCLOSURE_STATUS_OPTIONS:
        return clean_value
    normalized = clean_value.casefold()
    if normalized in {"○", "●", "o", "open", "y", "yes", "true", "1"}:
        return "공개"
    if normalized in {"x", "closed", "n", "no", "false", "0"}:
        return "비공개"
    return clean_value


def _normalize_legacy_master_status(value: Any) -> str | None:
    clean_value = _clean_projection_value(value)
    if clean_value is None:
        return None
    if clean_value in MASTER_STATUS_OPTIONS:
        return clean_value
    normalized = clean_value.casefold()
    if normalized in {"○", "●", "o", "y", "yes", "true", "1", "등재"}:
        return "등재"
    if normalized in {"x", "n", "no", "false", "0", "미등재"}:
        return "미등재"
    return clean_value


def _normalize_ox_select_value(value: Any) -> str | None:
    clean_value = _clean_projection_value(value)
    if clean_value is None:
        return None
    if clean_value in OX_SELECT_OPTIONS:
        return clean_value
    normalized = clean_value.casefold()
    if normalized in {
        "○",
        "●",
        "o",
        "y",
        "yes",
        "true",
        "1",
        "유",
        "적용",
        "있음",
        "등재",
    }:
        return "O"
    if normalized in {
        "x",
        "n",
        "no",
        "false",
        "0",
        "무",
        "미적용",
        "없음",
        "미등재",
    }:
        return "X"
    return clean_value


def get_dataset_definition(dataset_key: str) -> LegacyIssueDatasetDefinition:
    definition = DATASET_DEFINITIONS.get(dataset_key)
    if definition is None:
        raise localized_http_exception(status_code=404, code="legacy_issues.dataset_not_found")
    return definition


def list_system_field_setting_views(
    db: Session,
    *,
    workspace: Workspace,
    dataset_key: str,
) -> list[DatasetSystemFieldSettingView]:
    definition = get_dataset_definition(dataset_key)
    settings_by_key = _system_field_settings_by_key(
        db,
        workspace=workspace,
        dataset_key=definition.key,
    )
    effective_definition = _definition_with_system_field_settings(
        definition,
        settings_by_key,
    )
    views: list[DatasetSystemFieldSettingView] = []
    for field_definition in effective_definition.fields:
        if field_definition.source != "system":
            continue
        setting = settings_by_key.get(field_definition.key)
        views.append(
            DatasetSystemFieldSettingView(
                id=setting.id if setting is not None else None,
                dataset_key=definition.key,
                field=field_definition,
                updated_at=setting.updated_at if setting is not None else None,
            )
        )
    return views


def upsert_system_field_setting(
    db: Session,
    *,
    workspace: Workspace,
    user: User,
    dataset_key: str,
    field_key: str,
    label_ko: str | None = None,
    label_en: str | None = None,
    field_type: str | None = None,
    options: list[str] | None = None,
    allow_multiple: bool | None = None,
    required: bool | None = None,
) -> DatasetSystemFieldSettingView:
    definition = get_dataset_definition(dataset_key)
    base_field = next(
        (
            field_definition
            for field_definition in definition.fields
            if field_definition.key == field_key
        ),
        None,
    )
    if base_field is None or base_field.source != "system":
        raise localized_http_exception(status_code=404, code="legacy_issues.module_field_not_found")
    if base_field.settings_readonly:
        raise localized_http_exception(status_code=400, code="legacy_issues.module_field_invalid")

    row = db.scalar(
        select(LegacyIssueSystemFieldSetting).where(
            LegacyIssueSystemFieldSetting.workspace_id == workspace.id,
            LegacyIssueSystemFieldSetting.dataset_key == definition.key,
            LegacyIssueSystemFieldSetting.field_key == base_field.key,
        )
    )
    current_field = _system_field_with_setting(base_field, row) if row is not None else base_field
    next_field_type = normalize_system_field_type(field_type or current_field.field_type)
    next_options = normalize_system_field_options(
        next_field_type,
        options if options is not None else list(current_field.options),
    )
    next_allow_multiple = normalize_system_field_allow_multiple(
        next_field_type,
        current_field.allow_multiple if allow_multiple is None else allow_multiple,
    )
    now = utcnow_naive()
    if row is None:
        row = LegacyIssueSystemFieldSetting(
            id=new_id(),
            workspace_id=workspace.id,
            dataset_key=definition.key,
            field_key=base_field.key,
            created_at=now,
        )
    row.label_ko = normalize_system_field_label(label_ko or current_field.label_ko)
    row.label_en = normalize_system_field_label(
        label_en if label_en is not None and label_en.strip() else current_field.label_en
    )
    row.field_type = next_field_type
    row.options = next_options
    row.allow_multiple = next_allow_multiple
    row.required = current_field.required if required is None else required
    row.updated_by_id = user.id
    row.updated_at = now
    db.add(row)
    db.flush()
    return DatasetSystemFieldSettingView(
        id=row.id,
        dataset_key=definition.key,
        field=_system_field_with_setting(base_field, row),
        updated_at=row.updated_at,
    )


def get_dataset_definition_for_view(
    db: Session,
    *,
    dataset_key: str,
    workspace: Workspace,
    view_key: str | None = None,
    include_inactive_module_fields: bool = False,
) -> LegacyIssueDatasetDefinition:
    definition = get_dataset_definition_with_system_field_settings(
        db,
        dataset_key=dataset_key,
        workspace=workspace,
    )
    normalized_view_key = normalize_legacy_issue_module_key(view_key)
    if normalized_view_key is None:
        return definition
    module_rows = list_module_field_rows(
        db,
        workspace=workspace,
        module_key=normalized_view_key,
        include_inactive=include_inactive_module_fields,
    )
    return _definition_with_module_rows(definition, module_rows)


def get_dataset_definition_with_all_module_fields(
    db: Session,
    *,
    dataset_key: str,
    workspace: Workspace,
    include_inactive_module_fields: bool = False,
) -> LegacyIssueDatasetDefinition:
    definition = get_dataset_definition_with_system_field_settings(
        db,
        dataset_key=dataset_key,
        workspace=workspace,
    )
    statement = select(LegacyIssueModuleField).where(
        LegacyIssueModuleField.workspace_id == workspace.id,
    )
    if not include_inactive_module_fields:
        statement = statement.where(LegacyIssueModuleField.active.is_(True))
    rows = list(
        db.scalars(
            statement.order_by(
                LegacyIssueModuleField.module_key.asc(),
                LegacyIssueModuleField.sort_order.asc(),
                LegacyIssueModuleField.created_at.asc(),
            )
        )
    )
    if not rows:
        return definition
    return _definition_with_module_rows(definition, rows)


def get_dataset_definition_with_system_field_settings(
    db: Session,
    *,
    dataset_key: str,
    workspace: Workspace,
) -> LegacyIssueDatasetDefinition:
    definition = get_dataset_definition(dataset_key)
    settings_by_key = _system_field_settings_by_key(
        db,
        workspace=workspace,
        dataset_key=definition.key,
    )
    if not settings_by_key:
        return definition
    return _definition_with_system_field_settings(definition, settings_by_key)


def _system_field_settings_by_key(
    db: Session,
    *,
    workspace: Workspace,
    dataset_key: str,
) -> dict[str, LegacyIssueSystemFieldSetting]:
    return {
        row.field_key: row
        for row in db.scalars(
            select(LegacyIssueSystemFieldSetting).where(
                LegacyIssueSystemFieldSetting.workspace_id == workspace.id,
                LegacyIssueSystemFieldSetting.dataset_key == dataset_key,
            )
        )
    }


def _definition_with_system_field_settings(
    definition: LegacyIssueDatasetDefinition,
    settings_by_key: dict[str, LegacyIssueSystemFieldSetting],
) -> LegacyIssueDatasetDefinition:
    return replace(
        definition,
        fields=tuple(
            _system_field_with_setting(field_definition, settings_by_key[field_definition.key])
            if field_definition.key in settings_by_key
            else field_definition
            for field_definition in definition.fields
        ),
    )


def _system_field_with_setting(
    field_definition: DatasetFieldDefinition,
    setting: LegacyIssueSystemFieldSetting | None,
) -> DatasetFieldDefinition:
    if setting is None:
        return field_definition
    return replace(
        field_definition,
        label_ko=setting.label_ko,
        label_en=setting.label_en,
        field_type=setting.field_type,
        options=tuple(setting.options or ()),
        allow_multiple=setting.allow_multiple,
        required=setting.required,
    )


def _definition_with_module_rows(
    definition: LegacyIssueDatasetDefinition,
    module_rows: list[LegacyIssueModuleField],
) -> LegacyIssueDatasetDefinition:
    fields = list(definition.fields)
    fields.extend(module_field_definition(row) for row in module_rows)
    return LegacyIssueDatasetDefinition(
        key=definition.key,
        table_model=definition.table_model,
        title_ko=definition.title_ko,
        title_en=definition.title_en,
        hierarchy_ko=definition.hierarchy_ko,
        hierarchy_en=definition.hierarchy_en,
        fields=tuple(fields),
        header_rows=definition.header_rows,
        group_labels_ko={
            **definition.group_labels_ko,
            MODULE_CUSTOM_GROUP_KEY: "모듈 전용",
        },
        group_labels_en={
            **definition.group_labels_en,
            MODULE_CUSTOM_GROUP_KEY: "Module Fields",
        },
        import_sheet_name=definition.import_sheet_name,
    )


def refresh_dataset_record_projections(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision | None = None,
) -> int:
    model = definition.table_model
    statement = select(model).where(
        model.workspace_id == workspace.id,
        model.dataset_key == definition.key,
    )
    if revision is not None:
        statement = statement.where(model.revision_id == revision.id)
    labels = field_labels(definition)
    count = 0
    for record in db.scalars(statement):
        values = canonicalize_dataset_values(record.field_values)
        record.field_values = values
        apply_record_projection(
            record,
            values,
            raw_fields=dict(record.raw_fields or {}),
            value_labels=labels,
        )
        db.add(record)
        count += 1
    if count:
        db.flush()
    return count


def normalize_legacy_issue_module_key(view_key: str | None) -> str | None:
    normalized = (view_key or "").strip()
    if not normalized or normalized == COMMON_MASTER_DATASET_KEY:
        return None
    if normalized not in LEGACY_ISSUE_MODULE_KEYS:
        raise localized_http_exception(status_code=404, code="legacy_issues.module_not_found")
    return normalized


def normalize_system_field_type(field_type: str) -> str:
    normalized = field_type.strip()
    if normalized not in LEGACY_ISSUE_MODULE_FIELD_TYPES:
        raise localized_http_exception(status_code=400, code="legacy_issues.module_field_invalid")
    return normalized


def normalize_system_field_label(label: str) -> str:
    normalized = " ".join(label.strip().split())
    if not normalized or len(normalized) > 120:
        raise localized_http_exception(status_code=400, code="legacy_issues.module_field_invalid")
    return normalized


def normalize_system_field_options(field_type: str, options: list[str] | None) -> list[str] | None:
    if field_type != "select":
        return None
    normalized: list[str] = []
    seen: set[str] = set()
    for option in options or []:
        value = " ".join(str(option).strip().split())
        if not value or value in seen:
            continue
        seen.add(value)
        normalized.append(value)
    if not normalized:
        raise localized_http_exception(status_code=400, code="legacy_issues.module_field_invalid")
    return normalized[:100]


def normalize_system_field_allow_multiple(field_type: str, allow_multiple: bool) -> bool:
    if field_type not in LEGACY_ISSUE_MULTIPLE_FIELD_TYPES:
        return False
    return bool(allow_multiple)


def module_key_from_legacy_department_label(value: str | None) -> str | None:
    normalized = _clean_projection_value(value)
    if normalized is None:
        return None
    return LEGACY_ISSUE_MODULE_DEPARTMENT_LABELS.get(normalized)


def list_module_field_rows(
    db: Session,
    *,
    workspace: Workspace,
    module_key: str,
    include_inactive: bool = False,
) -> list[LegacyIssueModuleField]:
    statement = select(LegacyIssueModuleField).where(
        LegacyIssueModuleField.workspace_id == workspace.id,
        LegacyIssueModuleField.module_key == module_key,
    )
    if not include_inactive:
        statement = statement.where(LegacyIssueModuleField.active.is_(True))
    return list(
        db.scalars(
            statement.order_by(
                LegacyIssueModuleField.sort_order.asc(),
                LegacyIssueModuleField.created_at.asc(),
            )
        )
    )


def module_field_definition(row: LegacyIssueModuleField) -> DatasetFieldDefinition:
    return _field(
        row.field_key,
        row.label_ko,
        row.label_en,
        group_key=MODULE_CUSTOM_GROUP_KEY,
        field_type=row.field_type,
        options=tuple(row.options or ()),
        allow_multiple=row.allow_multiple,
        required=row.required,
        source="module",
        module_key=row.module_key,
        field_id=row.id,
        active=row.active,
    )


def normalize_upload_filename(filename: str | None) -> str:
    normalized = (filename or "").strip()
    return normalized or "unnamed"


def normalize_attachment_description(description: str | None) -> str | None:
    if description is None:
        return None
    normalized = " ".join(description.strip().split())
    if not normalized:
        return None
    return normalized[:DATASET_ATTACHMENT_DESCRIPTION_MAX_CHARS]


def build_dataset_import_rows(
    definition: LegacyIssueDatasetDefinition,
    *,
    filename: str,
    content: bytes,
) -> LegacyIssueImportTable:
    rows, _file_kind = parse_tabular_upload(
        filename=filename,
        content=content,
        invalid_code=DATASET_IMPORT_INVALID_CODE,
        sheet_name=definition.import_sheet_name,
    )
    lookup = field_lookup(definition)
    for header_index, row in enumerate(rows[:5]):
        normalized_cells = [normalize_header(cell) for cell in row]
        if EXPORT_RECORD_ID_HEADER in [lookup.get(cell, cell) for cell in normalized_cells]:
            return collapse_continuation_columns(
                definition,
                LegacyIssueImportTable(
                    headers=combine_header_rows(rows[header_index:], header_rows=1),
                    data_rows=rows[header_index + 1 :],
                ),
            )
    if len(rows) < definition.header_rows:
        raise localized_http_exception(status_code=400, code=DATASET_IMPORT_INVALID_CODE)
    return collapse_continuation_columns(
        definition,
        LegacyIssueImportTable(
            headers=combine_header_rows(rows, header_rows=definition.header_rows),
            data_rows=rows[definition.header_rows :],
        ),
    )


def collapse_continuation_columns(
    definition: LegacyIssueDatasetDefinition,
    table: LegacyIssueImportTable,
) -> LegacyIssueImportTable:
    lookup = field_lookup(definition)
    labels = field_labels(definition)
    collapsed_headers: list[str] = []
    collapsed_rows: list[list[str]] = [[] for _row in table.data_rows]
    merge_target_index: int | None = None

    for column_index, header in enumerate(table.headers):
        normalized_header = normalize_header(header)
        mapped_field = lookup.get(normalized_header)
        if normalized_header == f"column_{column_index + 1}" and merge_target_index is not None:
            for output_row, row in zip(collapsed_rows, table.data_rows, strict=False):
                value = row[column_index].strip() if column_index < len(row) else ""
                if value:
                    output_row[merge_target_index] = f"{output_row[merge_target_index]}{value}"
            continue

        display_header = (
            labels.get(mapped_field, header) if mapped_field and " - " not in header else header
        )
        collapsed_headers.append(display_header)
        for output_row, row in zip(collapsed_rows, table.data_rows, strict=False):
            output_row.append(row[column_index].strip() if column_index < len(row) else "")
        merge_target_index = (
            len(collapsed_headers) - 1
            if mapped_field and mapped_field != EXPORT_RECORD_ID_HEADER
            else None
        )

    return LegacyIssueImportTable(headers=collapsed_headers, data_rows=collapsed_rows)


def field_lookup(definition: LegacyIssueDatasetDefinition) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for field_definition in definition.fields:
        labels = (
            field_definition.key,
            field_definition.label_ko,
            field_definition.label_en,
            *field_definition.aliases,
        )
        for label in labels:
            lookup[normalize_header(label)] = field_definition.key
    lookup[normalize_header(EXPORT_RECORD_ID_HEADER)] = EXPORT_RECORD_ID_HEADER
    lookup[normalize_header("Record ID")] = EXPORT_RECORD_ID_HEADER
    lookup[normalize_header("Record ID (locked)")] = EXPORT_RECORD_ID_HEADER
    return lookup


def suggested_mapping(
    definition: LegacyIssueDatasetDefinition,
    headers: list[str],
) -> dict[str, int]:
    lookup = field_lookup(definition)
    mapping: dict[str, int] = {}
    used_columns: set[int] = set()
    for index, header in enumerate(headers):
        key = lookup.get(normalize_header(header))
        if key and key not in mapping:
            mapping[key] = index
            used_columns.add(index)

    for field_definition in definition.fields:
        if field_definition.key in mapping:
            continue
        candidate_labels = (
            field_definition.label_ko,
            field_definition.label_en,
            field_definition.key,
            *field_definition.aliases,
        )
        normalized_candidates = [normalize_header(label) for label in candidate_labels]
        for index, header in enumerate(headers):
            if index in used_columns:
                continue
            normalized_header = normalize_header(header)
            if any(
                candidate and candidate in normalized_header for candidate in normalized_candidates
            ):
                mapping[field_definition.key] = index
                used_columns.add(index)
                break
    return mapping


def build_import_preview(
    definition: LegacyIssueDatasetDefinition,
    *,
    filename: str,
    content: bytes,
) -> LegacyIssueImportPreview:
    import_rows = build_dataset_import_rows(definition, filename=filename, content=content)
    return build_tabular_import_preview(
        table=import_rows,
        suggested_mapping=suggested_mapping(definition, import_rows.headers),
    )


def clean_dataset_values(
    definition: LegacyIssueDatasetDefinition,
    values: dict[str, Any],
    *,
    db: Session | None = None,
    workspace: Workspace | None = None,
    include_readonly: bool = False,
) -> dict[str, Any]:
    values = canonicalize_dataset_values(values)
    fields_by_key = {
        field_definition.key: field_definition for field_definition in definition.fields
    }
    result: dict[str, Any] = {}
    for key, value in values.items():
        field_definition = fields_by_key.get(key)
        if field_definition is None:
            continue
        if field_definition.readonly and not include_readonly:
            continue
        normalized = normalize_dataset_field_value(
            field_definition,
            value,
            db=db,
            workspace=workspace,
        )
        if normalized is not None:
            result[key] = normalized
    return result


def normalize_dataset_field_value(
    field_definition: DatasetFieldDefinition,
    value: Any,
    *,
    db: Session | None = None,
    workspace: Workspace | None = None,
) -> Any | None:
    if value is None:
        return None
    field_type = field_definition.field_type
    if field_definition.allow_multiple and field_type in LEGACY_ISSUE_MULTIPLE_FIELD_TYPES:
        return normalize_multiple_dataset_field_value(
            field_definition,
            value,
            db=db,
            workspace=workspace,
        )
    if isinstance(value, list):
        compact_items = [item for item in value if display_dataset_value(item)]
        if not compact_items:
            return None
        if len(compact_items) == 1:
            value = compact_items[0]
        else:
            raise_invalid_field_value(field_definition)
    if field_type == "user":
        return normalize_user_field_value(field_definition, value, db=db)
    if field_type == "orgUnit":
        return normalize_org_unit_field_value(field_definition, value, db=db)
    if field_type == "date":
        return normalize_dataset_date_value(field_definition, value)
    if isinstance(value, bool):
        raw_value = "true" if value else "false"
    else:
        raw_value = str(value).strip()
    if not raw_value:
        return None
    if field_type in {"text", "longText"}:
        return raw_value
    if field_type == "number":
        try:
            Decimal(raw_value)
        except InvalidOperation as error:
            raise_invalid_field_value(field_definition, cause=error)
        return raw_value
    if field_type == "select":
        return normalize_select_field_value(field_definition, raw_value)
    if field_type == "boolean":
        normalized = raw_value.casefold()
        if normalized in {"true", "1", "yes", "y", "on", "예", "사용"}:
            return "true"
        if normalized in {"false", "0", "no", "n", "off", "아니오", "미사용"}:
            return "false"
        raise_invalid_field_value(field_definition)
    return raw_value


def normalize_dataset_date_value(
    field_definition: DatasetFieldDefinition,
    value: Any,
) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    raw_value = str(value).strip()
    if not raw_value:
        return None

    try:
        if len(raw_value) == 8 and raw_value.isdigit():
            parsed = date(
                int(raw_value[0:4]),
                int(raw_value[4:6]),
                int(raw_value[6:8]),
            )
            return parsed.isoformat()

        for separator in ("-", ".", "/"):
            parts = raw_value.split(separator)
            if len(parts) == 3 and len(parts[0]) == 4 and all(part.isdigit() for part in parts):
                parsed = date(int(parts[0]), int(parts[1]), int(parts[2]))
                return parsed.isoformat()

        if "T" in raw_value or " " in raw_value:
            parsed_datetime = datetime.fromisoformat(raw_value.replace("Z", "+00:00"))
            return parsed_datetime.date().isoformat()
    except ValueError as error:
        raise_invalid_field_value(field_definition, cause=error)

    raise_invalid_field_value(
        field_definition,
        cause=ValueError("Unsupported date value."),
    )


def normalize_multiple_dataset_field_value(
    field_definition: DatasetFieldDefinition,
    value: Any,
    *,
    db: Session | None,
    workspace: Workspace | None,
) -> list[Any] | None:
    items = split_multiple_field_input(value)
    normalized: list[Any] = []
    seen: set[str] = set()
    for item in items:
        normalized_item = normalize_dataset_field_value(
            replace(field_definition, allow_multiple=False),
            item,
            db=db,
            workspace=workspace,
        )
        display = display_dataset_value(normalized_item)
        if normalized_item is None or not display:
            continue
        identity = reference_identity(normalized_item) or display
        if identity in seen:
            continue
        seen.add(identity)
        normalized.append(normalized_item)
    return normalized or None


def split_multiple_field_input(value: Any) -> list[Any]:
    if isinstance(value, list):
        return value
    if isinstance(value, tuple):
        return list(value)
    if isinstance(value, str):
        return [item.strip() for item in value.replace("\n", ";").split(";") if item.strip()]
    return [value]


def reference_identity(value: Any) -> str | None:
    if isinstance(value, dict):
        kind = _clean_scalar_text(value.get("kind")) or "ref"
        ref_id = _clean_scalar_text(value.get("id"))
        if ref_id:
            return f"{kind}:{ref_id}"
    return None


def normalize_select_field_value(
    field_definition: DatasetFieldDefinition,
    raw_value: str,
) -> str | None:
    options = tuple(option.strip() for option in field_definition.options if option.strip())
    if options and raw_value not in options:
        raise_invalid_field_value(field_definition)
    return raw_value


def normalize_user_field_value(
    field_definition: DatasetFieldDefinition,
    value: Any,
    *,
    db: Session | None,
) -> dict[str, str] | None:
    if db is None:
        raise_invalid_field_value(field_definition)
    user = resolve_user_field_reference(db, value)
    if user is None:
        return None
    return {
        "kind": "user",
        "id": user.id,
        "label": user.display_name or user.full_name or user.email,
        "email": user.email,
    }


def resolve_user_field_reference(db: Session, value: Any) -> User | None:
    if isinstance(value, dict):
        user_id = _clean_scalar_text(value.get("id"))
        if user_id:
            return db.scalar(
                select(User).where(
                    User.id == user_id,
                    User.status == "active",
                )
            )
        value = value.get("email") or value.get("label") or value.get("name")
    text = _clean_scalar_text(value)
    if not text:
        return None
    normalized = text.casefold()
    if "@" in text:
        users = list(
            db.scalars(
                select(User).where(
                    User.status == "active",
                    func.lower(User.email) == normalized,
                )
            )
        )
    else:
        users = list(
            db.scalars(
                select(User).where(
                    User.status == "active",
                    or_(
                        func.lower(User.full_name) == normalized,
                        func.lower(User.display_name) == normalized,
                    ),
                )
            )
        )
    if len(users) != 1:
        return None
    return users[0]


def normalize_org_unit_field_value(
    field_definition: DatasetFieldDefinition,
    value: Any,
    *,
    db: Session | None,
) -> dict[str, str] | None:
    if db is None:
        raise_invalid_field_value(field_definition)
    org_unit, path = resolve_org_unit_field_reference(db, value)
    if org_unit is None:
        return None
    return {
        "kind": "orgUnit",
        "id": org_unit.id,
        "label": org_unit.name,
        "path": path or org_unit.name,
    }


def resolve_org_unit_field_reference(
    db: Session,
    value: Any,
) -> tuple[OrgUnit | None, str | None]:
    if isinstance(value, dict):
        org_unit_id = _clean_scalar_text(value.get("id"))
        if org_unit_id:
            org_unit = db.scalar(
                select(OrgUnit).where(
                    OrgUnit.id == org_unit_id,
                    OrgUnit.active.is_(True),
                )
            )
            if org_unit is None:
                return None, None
            units_by_id = org_units_by_id(db)
            return org_unit, org_unit_path(org_unit, units_by_id)
        value = value.get("path") or value.get("label") or value.get("name")
    text = _clean_scalar_text(value)
    if not text:
        return None, None
    normalized = text.casefold()
    units_by_id = org_units_by_id(db)
    matches: list[tuple[OrgUnit, str]] = []
    for org_unit in units_by_id.values():
        path = org_unit_path(org_unit, units_by_id)
        candidates = (org_unit.name, org_unit.slug, path)
        if any(candidate.casefold() == normalized for candidate in candidates if candidate):
            matches.append((org_unit, path))
    if len(matches) != 1:
        return None, None
    return matches[0]


def org_units_by_id(db: Session) -> dict[str, OrgUnit]:
    return {
        org_unit.id: org_unit
        for org_unit in db.scalars(
            select(OrgUnit).where(OrgUnit.active.is_(True)).order_by(OrgUnit.name.asc())
        )
    }


def org_unit_path(org_unit: OrgUnit, units_by_id: dict[str, OrgUnit]) -> str:
    parts: list[str] = []
    current: OrgUnit | None = org_unit
    visited: set[str] = set()
    while current is not None and current.id not in visited:
        visited.add(current.id)
        parts.append(current.name)
        current = units_by_id.get(current.parent_id) if current.parent_id else None
    return " / ".join(reversed(parts))


def raise_invalid_field_value(
    field_definition: DatasetFieldDefinition,
    *,
    cause: Exception | None = None,
) -> None:
    exception = localized_http_exception(
        status_code=400,
        code="legacy_issues.module_field_value_invalid",
        detail=field_definition.key,
    )
    if cause is not None:
        raise exception from cause
    raise exception


def require_dataset_values(values: dict[str, Any]) -> None:
    if values:
        return
    raise localized_http_exception(status_code=400, code=DATASET_RECORD_EMPTY_CODE)


def list_dataset_records(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision | None = None,
    revisions: list[LegacyIssueDataRevision] | None = None,
    module_key: str | None = None,
    departments: list[str] | None = None,
    record_ids: Iterable[str] | None = None,
    query: str | None = None,
    limit: int | None = None,
    offset: int = 0,
) -> tuple[list[Any], int]:
    model = definition.table_model
    statement = select(model).where(
        model.workspace_id == workspace.id,
        model.dataset_key == definition.key,
    )
    revision_ids = [item.id for item in revisions or []]
    if revision is not None:
        revision_ids.append(revision.id)
    if not revision_ids:
        return [], 0
    statement = statement.where(model.revision_id.in_(revision_ids))
    normalized_module_key = normalize_legacy_issue_module_key(module_key)
    if normalized_module_key is not None:
        statement = statement.where(model.module_key == normalized_module_key)
    normalized_departments = [
        department.strip() for department in (departments or []) if department.strip()
    ]
    if normalized_departments:
        statement = statement.where(model.department.in_(normalized_departments))
    normalized_record_ids = tuple(
        dict.fromkeys(item.strip() for item in record_ids or [] if item.strip())
    )
    if record_ids is not None:
        if not normalized_record_ids:
            return [], 0
        statement = statement.where(model.id.in_(normalized_record_ids))
    normalized_query = (query or "").strip()
    if normalized_query:
        pattern = f"%{normalized_query}%"
        statement = statement.where(
            or_(
                func.coalesce(model.search_text, "").ilike(pattern),
                cast(model.field_values, Text).ilike(pattern),
                cast(model.raw_fields, Text).ilike(pattern),
            )
        )
    total = int(db.scalar(select(func.count()).select_from(statement.subquery())) or 0)
    row_statement = statement.order_by(model.updated_at.desc(), model.created_at.desc()).offset(
        max(offset, 0)
    )
    if limit is not None:
        row_statement = row_statement.limit(max(limit, 1))
    rows = list(db.scalars(row_statement))
    return rows, total


def field_labels(definition: LegacyIssueDatasetDefinition) -> dict[str, str]:
    return {
        field_definition.key: field_definition.label_ko for field_definition in definition.fields
    }


def field_keys(definition: LegacyIssueDatasetDefinition) -> set[str]:
    return {field_definition.key for field_definition in definition.fields}


def merge_dataset_values(
    definition: LegacyIssueDatasetDefinition,
    *,
    db: Session | None = None,
    workspace: Workspace | None = None,
    current_values: dict[str, Any],
    incoming_values: dict[str, Any],
    include_readonly: bool = False,
) -> tuple[dict[str, Any], dict[str, Any | None]]:
    current_values = canonicalize_dataset_values(current_values)
    incoming_values = canonicalize_dataset_values(incoming_values)
    fields_by_key = {
        field_definition.key: field_definition for field_definition in definition.fields
    }
    next_values = dict(current_values)
    changed_values: dict[str, Any | None] = {}
    for key, value in incoming_values.items():
        field_definition = fields_by_key.get(key)
        if field_definition is None:
            continue
        if field_definition.readonly and not include_readonly:
            continue
        normalized_value = normalize_dataset_field_value(
            field_definition,
            value,
            db=db,
            workspace=workspace,
        )
        if normalized_value is not None:
            next_values[key] = normalized_value
            changed_values[key] = normalized_value
        else:
            next_values.pop(key, None)
            changed_values[key] = None
    return next_values, changed_values


def get_dataset_record(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    record_id: str,
    revision: LegacyIssueDataRevision | None = None,
    module_key: str | None = None,
) -> Any:
    model = definition.table_model
    statement = select(model).where(
        model.id == record_id,
        model.workspace_id == workspace.id,
        model.dataset_key == definition.key,
    )
    if revision is not None:
        statement = statement.where(model.revision_id == revision.id)
    normalized_module_key = normalize_legacy_issue_module_key(module_key)
    if normalized_module_key is not None:
        statement = statement.where(model.module_key == normalized_module_key)
    record = db.scalar(statement)
    if record is None:
        raise localized_http_exception(
            status_code=404, code="legacy_issues.dataset_record_not_found"
        )
    return record


def create_dataset_record(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    revision: LegacyIssueDataRevision,
    values: dict[str, Any],
    module_key: str | None = None,
) -> Any:
    ensure_revision_partition(db, revision)
    clean_values = clean_dataset_values(
        definition,
        values,
        db=db,
        workspace=workspace,
    )
    require_dataset_values(clean_values)
    validate_required_values(definition, clean_values)
    clean_values = apply_introduced_revision_value(
        db,
        definition,
        workspace=workspace,
        revision=revision,
        values=clean_values,
    )
    record = definition.table_model(
        id=new_id(),
        workspace_id=workspace.id,
        dataset_key=definition.key,
        module_key=normalize_legacy_issue_module_key(module_key),
        revision_id=revision.id,
        stable_record_id=new_id(),
        created_by_id=user.id,
        field_values=clean_values,
        raw_fields={},
    )
    ensure_record_partition(db, record, revision=revision)
    apply_record_projection(
        record,
        clean_values,
        raw_fields={},
        value_labels=field_labels(definition),
    )
    db.add(record)
    db.flush()
    add_record_history_entries(
        db,
        workspace=workspace,
        user=user,
        record_kind=LEGACY_ISSUE_RECORD_KIND,
        dataset_key=definition.key,
        record_id=record.stable_record_id or record.id,
        action="create",
        changes=collect_create_changes(
            values=clean_values,
            field_labels=field_labels(definition),
        ),
        revision_id=revision.id,
    )
    db.refresh(record)
    return record


def apply_introduced_revision_value(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision,
    values: dict[str, Any],
) -> dict[str, Any]:
    if INTRODUCED_REVISION_FIELD_KEY not in {field.key for field in definition.fields}:
        return values
    if _clean_projection_value(values.get(INTRODUCED_REVISION_FIELD_KEY)):
        return values
    introduced_revision_no = resolve_introduced_revision_no(
        db,
        workspace=workspace,
        revision=revision,
    )
    if introduced_revision_no is None:
        return values
    return {
        **values,
        INTRODUCED_REVISION_FIELD_KEY: str(introduced_revision_no),
    }


def resolve_introduced_revision_no(
    db: Session,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision,
) -> int | None:
    if revision.revision_no is not None:
        return int(revision.revision_no)
    latest_published_revision_no = db.scalar(
        select(func.max(LegacyIssueDataRevision.revision_no)).where(
            LegacyIssueDataRevision.workspace_id == workspace.id,
            LegacyIssueDataRevision.dataset_key == revision.dataset_key,
            LegacyIssueDataRevision.status == REVISION_STATUS_PUBLISHED,
        )
    )
    return int(latest_published_revision_no or 0) + 1


def update_dataset_record(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    revision: LegacyIssueDataRevision,
    record_id: str,
    values: dict[str, Any],
    module_key: str | None = None,
) -> Any:
    record = get_dataset_record(
        db,
        definition,
        workspace=workspace,
        record_id=record_id,
        revision=revision,
        module_key=module_key,
    )
    normalized_module_key = normalize_legacy_issue_module_key(module_key)
    if normalized_module_key is not None:
        record.module_key = normalized_module_key
    old_values = canonicalize_dataset_values(record.field_values)
    next_values, changed_values = merge_dataset_values(
        definition,
        db=db,
        workspace=workspace,
        current_values=old_values,
        incoming_values=values,
    )
    require_dataset_values(next_values)
    record.field_values = next_values
    apply_record_projection(
        record,
        next_values,
        raw_fields=dict(record.raw_fields or {}),
        value_labels=field_labels(definition),
    )
    record.updated_at = utcnow_naive()
    db.add(record)
    db.flush()
    add_record_history_entries(
        db,
        workspace=workspace,
        user=user,
        record_kind=LEGACY_ISSUE_RECORD_KIND,
        dataset_key=definition.key,
        record_id=record.stable_record_id or record.id,
        action="update",
        changes=collect_value_changes(
            old_values=old_values,
            new_values=changed_values,
            field_labels=field_labels(definition),
        ),
        revision_id=revision.id,
    )
    db.refresh(record)
    return record


def delete_dataset_record(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision,
    record_id: str,
) -> None:
    record = get_dataset_record(
        db, definition, workspace=workspace, record_id=record_id, revision=revision
    )
    for attachment in list_dataset_attachments(
        db, definition, workspace=workspace, record_id=record_id
    ):
        _remove_attachment_object(attachment)
        db.delete(attachment)
    db.delete(record)
    db.flush()


def validate_required_values(
    definition: LegacyIssueDatasetDefinition,
    values: dict[str, Any],
) -> None:
    missing_labels = [
        field_definition.label_ko
        for field_definition in definition.fields
        if field_definition.required
        and field_definition.active
        and not field_definition.readonly
        and not _clean_projection_value(values.get(field_definition.key))
    ]
    if missing_labels:
        raise localized_http_exception(
            status_code=status.HTTP_400_BAD_REQUEST,
            code=DATASET_REQUIRED_FIELD_MISSING_CODE,
            fields=", ".join(missing_labels),
        )


def stamp_introduced_revision_numbers(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision,
) -> int:
    if revision.revision_no is None:
        return 0
    model = definition.table_model
    base_stable_ids: set[str] = set()
    if revision.base_revision_id:
        base_stable_ids = {
            row.stable_record_id or row.id
            for row in db.scalars(
                select(model).where(
                    model.workspace_id == workspace.id,
                    model.dataset_key == definition.key,
                    model.revision_id == revision.base_revision_id,
                )
            )
        }
    stamped = 0
    rows = list(
        db.scalars(
            select(model).where(
                model.workspace_id == workspace.id,
                model.dataset_key == definition.key,
                model.revision_id == revision.id,
            )
        )
    )
    for row in rows:
        stable_record_id = row.stable_record_id or row.id
        if stable_record_id in base_stable_ids:
            continue
        values = canonicalize_dataset_values(row.field_values)
        if _clean_projection_value(values.get(INTRODUCED_REVISION_FIELD_KEY)):
            continue
        values[INTRODUCED_REVISION_FIELD_KEY] = str(revision.revision_no)
        row.field_values = values
        apply_record_projection(
            row,
            values,
            raw_fields=dict(row.raw_fields or {}),
            value_labels=field_labels(definition),
        )
        row.updated_at = utcnow_naive()
        db.add(row)
        stamped += 1
    if stamped:
        db.flush()
    return stamped


def create_dataset_draft_revision(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    base_revision_id: str | None = None,
    note: str | None = None,
    module_key: str | None = None,
) -> LegacyIssueDataRevision:
    normalized_module_key = normalize_legacy_issue_module_key(module_key)
    if normalized_module_key is None:
        raise localized_http_exception(
            status_code=409,
            code="legacy_issues.aggregate_readonly",
        )
    dataset_key = legacy_issue_dataset_revision_key(
        definition.key,
        normalized_module_key,
    )
    revision, created = create_draft_revision(
        db,
        workspace=workspace,
        dataset_key=dataset_key,
        user=user,
        base_revision_id=base_revision_id,
        note=note,
    )
    if not created:
        return revision
    source_revision = get_revision(
        db,
        workspace=workspace,
        dataset_key=dataset_key,
        revision_id=revision.base_revision_id
        or ensure_initial_published_revision(
            db,
            workspace=workspace,
            dataset_key=dataset_key,
            user=user,
        ).id,
    )
    copy_dataset_revision_records(
        db,
        definition,
        workspace=workspace,
        source_revision=source_revision,
        target_revision=revision,
        module_key=normalized_module_key,
    )
    from ai_do_api.domains.legacy_issues.ai_search import (
        reindex_legacy_issue_revision_ai_chunks,
    )

    reindex_definition = get_dataset_definition_with_all_module_fields(
        db,
        dataset_key=definition.key,
        workspace=workspace,
    )
    reindex_legacy_issue_revision_ai_chunks(
        db,
        reindex_definition,
        workspace=workspace,
        revision=revision,
        embed=False,
        reuse_embeddings_from_revision=source_revision,
    )
    latest_published_revision = get_latest_published_revision(
        db,
        workspace=workspace,
        dataset_key=dataset_key,
    )
    if latest_published_revision is not None and latest_published_revision.id != source_revision.id:
        add_dataset_restore_history_entries(
            db,
            definition,
            workspace=workspace,
            user=user,
            baseline_revision=latest_published_revision,
            source_revision=source_revision,
            target_revision=revision,
            module_key=normalized_module_key,
        )
    add_revision_event(
        db,
        workspace=workspace,
        revision=revision,
        user=user,
        action="draft_copy",
        details={"source_revision_id": source_revision.id},
    )
    return revision


def copy_dataset_revision_records(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    source_revision: LegacyIssueDataRevision,
    target_revision: LegacyIssueDataRevision,
    module_key: str | None = None,
) -> None:
    ensure_revision_partition(db, source_revision)
    ensure_revision_partition(db, target_revision)
    in_flight_attachment_targets: list[LegacyIssueAttachment] = []
    model = definition.table_model
    statement = select(model).where(
        model.workspace_id == workspace.id,
        model.dataset_key == definition.key,
        model.revision_id == source_revision.id,
    )
    normalized_module_key = normalize_legacy_issue_module_key(module_key)
    if normalized_module_key is not None:
        statement = statement.where(model.module_key == normalized_module_key)
    source_records = list(db.scalars(statement.order_by(model.created_at.asc())))
    for source in source_records:
        stable_record_id = source.stable_record_id or source.id
        source_values = canonicalize_dataset_values(source.field_values)
        target = model(
            id=new_id(),
            workspace_id=workspace.id,
            dataset_key=definition.key,
            module_key=source.module_key,
            revision_id=target_revision.id,
            stable_record_id=stable_record_id,
            field_values=source_values,
            raw_fields=dict(source.raw_fields or {}),
            imported_source_filename=source.imported_source_filename,
            imported_at=source.imported_at,
            created_by_id=source.created_by_id,
            created_at=source.created_at,
            updated_at=source.updated_at,
        )
        ensure_record_partition(db, target, revision=target_revision)
        apply_record_projection(
            target,
            source_values,
            raw_fields=dict(source.raw_fields or {}),
            value_labels=field_labels(definition),
        )
        db.add(target)
        db.flush()
        for attachment in list_dataset_attachments(
            db,
            definition,
            workspace=workspace,
            record_id=source.id,
        ):
            target_attachment = _copy_dataset_attachment_for_revision(
                attachment,
                workspace_id=workspace.id,
                dataset_key=definition.key,
                target_revision_id=target_revision.id,
                target_record_id=target.id,
                stable_record_id=stable_record_id,
            )
            ensure_attachment_partition(db, target_attachment, record=target)
            db.add(target_attachment)
            db.flush([target_attachment])
            if attachment.index_status in {"pending", "processing"}:
                in_flight_attachment_targets.append(target_attachment)
            _copy_dataset_attachment_index_projection(
                db,
                source_attachment=attachment,
                target_attachment=target_attachment,
            )
    db.flush()
    if in_flight_attachment_targets:
        from ai_do_api.domains.legacy_issues.attachment_indexing import (
            enqueue_legacy_issue_attachment_index_job,
        )
        from ai_do_api.domains.legacy_issues.settings import get_legacy_issue_settings

        if get_legacy_issue_settings().ai_attachment_index_enabled:
            for attachment in in_flight_attachment_targets[:DRAFT_COPY_ATTACHMENT_INDEX_BATCH_SIZE]:
                enqueue_legacy_issue_attachment_index_job(
                    db,
                    attachment=attachment,
                    trigger="draft_copy_inflight",
                )


def _copy_dataset_attachment_for_revision(
    attachment: LegacyIssueAttachment,
    *,
    workspace_id: str,
    dataset_key: str,
    target_revision_id: str,
    target_record_id: str,
    stable_record_id: str,
) -> LegacyIssueAttachment:
    index_status = _copied_attachment_index_status(attachment.index_status)
    summary_status = _copied_attachment_summary_status(attachment.ai_summary_status)
    return LegacyIssueAttachment(
        id=new_id(),
        workspace_id=workspace_id,
        dataset_key=dataset_key,
        revision_id=target_revision_id,
        stable_record_id=stable_record_id,
        record_id=target_record_id,
        filename=attachment.filename,
        content_type=attachment.content_type,
        size_bytes=attachment.size_bytes,
        description=attachment.description,
        storage_key=attachment.storage_key,
        is_primary=attachment.is_primary,
        index_status=index_status,
        index_error=attachment.index_error if index_status == "failed" else None,
        indexed_at=attachment.indexed_at if index_status == "indexed" else None,
        index_version=attachment.index_version,
        chunk_count=attachment.chunk_count if index_status == "indexed" else 0,
        artifact_count=attachment.artifact_count if index_status == "indexed" else 0,
        ai_summary=attachment.ai_summary if summary_status == "summarized" else None,
        ai_summary_status=summary_status,
        ai_summary_error=attachment.ai_summary_error if summary_status == "failed" else None,
        ai_summary_model=attachment.ai_summary_model if summary_status == "summarized" else None,
        ai_summary_version=attachment.ai_summary_version,
        ai_summarized_at=attachment.ai_summarized_at if summary_status == "summarized" else None,
        uploaded_by_id=attachment.uploaded_by_id,
        created_at=attachment.created_at,
    )


def _copy_dataset_attachment_index_projection(
    db: Session,
    *,
    source_attachment: LegacyIssueAttachment,
    target_attachment: LegacyIssueAttachment,
) -> None:
    if target_attachment.index_status != "indexed":
        return
    artifacts = list(
        db.scalars(
            select(LegacyIssueAttachmentArtifact).where(
                LegacyIssueAttachmentArtifact.attachment_id == source_attachment.id
            )
        )
    )
    for artifact in artifacts:
        db.add(
            LegacyIssueAttachmentArtifact(
                id=new_id(),
                workspace_id=target_attachment.workspace_id,
                dataset_key=target_attachment.dataset_key,
                revision_id=target_attachment.revision_id,
                record_id=target_attachment.record_id,
                stable_record_id=target_attachment.stable_record_id,
                attachment_id=target_attachment.id,
                job_id=None,
                artifact_kind=artifact.artifact_kind,
                page_number=artifact.page_number,
                storage_key=artifact.storage_key,
                content_text=artifact.content_text,
                artifact_metadata=_copy_json_dict(artifact.artifact_metadata),
                created_at=artifact.created_at,
            )
        )
    chunks = list(
        db.scalars(
            select(LegacyIssueAiChunk).where(
                LegacyIssueAiChunk.attachment_id == source_attachment.id
            )
        )
    )
    for chunk in chunks:
        copied_chunk = _copy_dataset_attachment_ai_chunk(
            chunk,
            source_attachment=source_attachment,
            target_attachment=target_attachment,
        )
        ensure_chunk_partition(
            db,
            copied_chunk,
            attachment=target_attachment,
        )
        db.add(copied_chunk)


def _copy_dataset_attachment_ai_chunk(
    chunk: LegacyIssueAiChunk,
    *,
    source_attachment: LegacyIssueAttachment,
    target_attachment: LegacyIssueAttachment,
) -> LegacyIssueAiChunk:
    evidence_metadata = _copy_json_dict(chunk.evidence_metadata)
    if evidence_metadata is not None:
        evidence_metadata.update(
            {
                "attachment_id": target_attachment.id,
                "attachment_filename": target_attachment.filename,
                "attachment_description": target_attachment.description,
                "attachment_content_type": target_attachment.content_type,
            }
        )
    return LegacyIssueAiChunk(
        id=new_id(),
        workspace_id=target_attachment.workspace_id,
        dataset_key=target_attachment.dataset_key,
        revision_id=target_attachment.revision_id,
        record_id=target_attachment.record_id,
        stable_record_id=target_attachment.stable_record_id,
        attachment_id=target_attachment.id,
        attachment_filename=target_attachment.filename,
        attachment_page=chunk.attachment_page,
        attachment_artifact_type=chunk.attachment_artifact_type,
        chunk_key=chunk.chunk_key.replace(source_attachment.id, target_attachment.id),
        chunk_kind=chunk.chunk_kind,
        field_key=chunk.field_key,
        field_label=chunk.field_label,
        field_value=chunk.field_value,
        search_text=chunk.search_text,
        search_terms=list(chunk.search_terms) if chunk.search_terms is not None else None,
        embedding_vector=chunk.embedding_vector,
        embedding_model=chunk.embedding_model,
        embedding_dimensions=chunk.embedding_dimensions,
        embedding_status=chunk.embedding_status,
        embedding_error=chunk.embedding_error,
        evidence_metadata=evidence_metadata,
        created_at=chunk.created_at,
        updated_at=chunk.updated_at,
    )


def _copied_attachment_index_status(status_value: str) -> str:
    if status_value in {"indexed", "failed", "not_indexed"}:
        return status_value
    return "not_indexed"


def _copied_attachment_summary_status(status_value: str) -> str:
    if status_value in {"summarized", "failed", "not_summarized"}:
        return status_value
    return "not_summarized"


def _copy_json_dict(value: dict | None) -> dict | None:
    return dict(value) if value is not None else None


def add_dataset_restore_history_entries(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    baseline_revision: LegacyIssueDataRevision,
    source_revision: LegacyIssueDataRevision,
    target_revision: LegacyIssueDataRevision,
    module_key: str | None = None,
) -> None:
    baseline_records = _dataset_records_by_stable_id(
        db,
        definition,
        workspace=workspace,
        revision=baseline_revision,
        module_key=module_key,
    )
    target_records = _dataset_records_by_stable_id(
        db,
        definition,
        workspace=workspace,
        revision=target_revision,
        module_key=module_key,
    )
    labels = field_labels(definition)
    field_keys = tuple(field.key for field in definition.fields)
    for stable_record_id, target in target_records.items():
        baseline = baseline_records.get(stable_record_id)
        baseline_values = canonicalize_dataset_values(baseline.field_values) if baseline else {}
        target_values = canonicalize_dataset_values(target.field_values)
        old_values = {key: baseline_values.get(key) for key in field_keys}
        new_values = {key: target_values.get(key) for key in field_keys}
        add_record_history_entries(
            db,
            workspace=workspace,
            user=user,
            record_kind=LEGACY_ISSUE_RECORD_KIND,
            dataset_key=definition.key,
            record_id=stable_record_id,
            action="restore",
            changes=collect_value_changes(
                old_values=old_values,
                new_values=new_values,
                field_labels=labels,
            ),
            details={
                "baseline_revision_id": baseline_revision.id,
                "source_revision_id": source_revision.id,
            },
            revision_id=target_revision.id,
        )
    db.flush()


def import_dataset_records(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    revision: LegacyIssueDataRevision,
    filename: str,
    content: bytes,
    mapping: dict[str, int],
    module_key: str | None = None,
) -> DatasetImportResult:
    ensure_revision_partition(db, revision)
    import_rows = build_dataset_import_rows(definition, filename=filename, content=content)
    headers = import_rows.headers
    validate_mapping(definition, headers=headers, mapping=mapping)
    data_rows = import_rows.data_rows
    record_id_column = mapping.get(EXPORT_RECORD_ID_HEADER)
    created = 0
    updated = 0
    skipped = 0
    imported_at = utcnow_naive()
    source_filename = normalize_upload_filename(filename)
    field_keys = {field_definition.key for field_definition in definition.fields}
    parsed_rows: list[tuple[list[str], dict[str, Any], str, dict[str, str]]] = []
    normalized_module_key = normalize_legacy_issue_module_key(module_key)

    for row in data_rows:
        if not any(cell.strip() for cell in row):
            parsed_rows.append((row, {}, "", {}))
            skipped += 1
            continue
        values: dict[str, str] = {}
        for key, column_index in mapping.items():
            if key == EXPORT_RECORD_ID_HEADER or key not in field_keys:
                continue
            values[key] = row[column_index].strip() if column_index < len(row) else ""
        clean_values = clean_dataset_values(
            definition,
            values,
            db=db,
            workspace=workspace,
        )
        record_id = (
            row[record_id_column].strip()
            if record_id_column is not None and record_id_column < len(row)
            else ""
        )
        if not clean_values:
            parsed_rows.append((row, {}, record_id, values))
            skipped += 1
            continue
        parsed_rows.append((row, clean_values, record_id, values))

    record_ids = [
        record_id
        for _row, clean_values, record_id, _values in parsed_rows
        if clean_values and record_id
    ]
    duplicate_ids = sorted(
        {record_id for record_id in record_ids if record_ids.count(record_id) > 1}
    )
    if duplicate_ids:
        raise localized_http_exception(
            status_code=400,
            code="legacy_issues.dataset_import_pk_invalid",
            detail=", ".join(duplicate_ids[:5]),
        )

    model = definition.table_model
    existing_by_id: dict[str, Any] = {}
    if record_ids:
        existing_statement = select(model).where(
            model.workspace_id == workspace.id,
            model.dataset_key == definition.key,
            model.revision_id == revision.id,
            model.stable_record_id.in_(record_ids),
        )
        if normalized_module_key is not None:
            existing_statement = existing_statement.where(model.module_key == normalized_module_key)
        existing = list(db.scalars(existing_statement))
        existing_by_id = {record.stable_record_id or record.id: record for record in existing}
        missing_ids = sorted(set(record_ids).difference(existing_by_id))
        if missing_ids:
            raise localized_http_exception(
                status_code=400,
                code="legacy_issues.dataset_import_pk_invalid",
                detail=", ".join(missing_ids[:5]),
            )

    for row, clean_values, record_id, values in parsed_rows:
        if not clean_values:
            continue
        raw_fields = {
            header: row[index].strip()
            for index, header in enumerate(headers)
            if index < len(row) and row[index].strip()
        }
        if record_id:
            record = existing_by_id[record_id]
            ensure_record_partition(db, record, revision=revision)
            if normalized_module_key is not None:
                record.module_key = normalized_module_key
            old_values = canonicalize_dataset_values(record.field_values)
            next_values, changed_values = merge_dataset_values(
                definition,
                db=db,
                workspace=workspace,
                current_values=old_values,
                incoming_values=values,
            )
            require_dataset_values(next_values)
            record.field_values = next_values
            record.raw_fields = raw_fields
            apply_record_projection(
                record,
                next_values,
                raw_fields=raw_fields,
                value_labels=field_labels(definition),
            )
            record.imported_source_filename = source_filename
            record.imported_at = imported_at
            record.updated_at = imported_at
            db.add(record)
            add_record_history_entries(
                db,
                workspace=workspace,
                user=user,
                record_kind=LEGACY_ISSUE_RECORD_KIND,
                dataset_key=definition.key,
                record_id=record.stable_record_id or record.id,
                action="import",
                changes=collect_value_changes(
                    old_values=old_values,
                    new_values=changed_values,
                    field_labels=field_labels(definition),
                ),
                details={"source_filename": source_filename},
                revision_id=revision.id,
            )
            updated += 1
        else:
            validate_required_values(definition, clean_values)
            clean_values = apply_introduced_revision_value(
                db,
                definition,
                workspace=workspace,
                revision=revision,
                values=clean_values,
            )
            record = model(
                id=new_id(),
                workspace_id=workspace.id,
                dataset_key=definition.key,
                module_key=normalized_module_key,
                revision_id=revision.id,
                stable_record_id=new_id(),
                created_by_id=user.id,
                field_values=clean_values,
                raw_fields=raw_fields,
                imported_source_filename=source_filename,
                imported_at=imported_at,
            )
            ensure_record_partition(db, record, revision=revision)
            apply_record_projection(
                record,
                clean_values,
                raw_fields=raw_fields,
                value_labels=field_labels(definition),
            )
            db.add(record)
            db.flush()
            add_record_history_entries(
                db,
                workspace=workspace,
                user=user,
                record_kind=LEGACY_ISSUE_RECORD_KIND,
                dataset_key=definition.key,
                record_id=record.stable_record_id or record.id,
                action="import",
                changes=collect_create_changes(
                    values=clean_values,
                    field_labels=field_labels(definition),
                ),
                details={"source_filename": source_filename},
                revision_id=revision.id,
            )
            created += 1
    db.flush()
    return DatasetImportResult(
        created=created,
        updated=updated,
        skipped=skipped,
        total_rows=len(data_rows),
    )


def validate_mapping(
    definition: LegacyIssueDatasetDefinition,
    *,
    headers: list[str],
    mapping: dict[str, int],
) -> None:
    allowed_keys = {field_definition.key for field_definition in definition.fields}
    allowed_keys.add(EXPORT_RECORD_ID_HEADER)
    validate_tabular_mapping(
        allowed_keys=allowed_keys,
        headers=headers,
        invalid_code=DATASET_IMPORT_INVALID_CODE,
        mapping=mapping,
    )


def export_dataset_records_xlsx(
    definition: LegacyIssueDatasetDefinition,
    records: list[Any],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = definition.key[:31]
    headers = [
        EXPORT_RECORD_ID_HEADER,
        *(field_definition.label_ko for field_definition in definition.fields),
    ]
    sheet.append(headers)
    for record in records:
        values = canonicalize_dataset_values(record.field_values)
        sheet.append(
            [
                record.stable_record_id or record.id,
                *(
                    dataset_excel_value(
                        field_definition,
                        values.get(field_definition.key),
                    )
                    for field_definition in definition.fields
                ),
            ]
        )
    sheet.freeze_panes = "A2"
    locked_fill = PatternFill("solid", fgColor="E5E7EB")
    editable_fill = PatternFill("solid", fgColor="ECFDF5")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.protection = Protection(locked=True)
        cell.fill = locked_fill
    sheet.column_dimensions["A"].width = 40
    for row in sheet.iter_rows(min_row=2):
        row[0].protection = Protection(locked=True)
        row[0].fill = locked_fill
        for cell, field_definition in zip(row[1:], definition.fields, strict=False):
            readonly = field_definition.readonly
            cell.protection = Protection(locked=readonly)
            cell.fill = locked_fill if readonly else editable_fill
            if field_definition.field_type == "date" and cell.is_date:
                cell.number_format = EXCEL_DATE_NUMBER_FORMAT
    for column_index in range(2, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 18
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def dataset_excel_value(
    field_definition: DatasetFieldDefinition,
    value: Any,
) -> str | date:
    display_value = display_dataset_value(value) or ""
    if field_definition.field_type != "date" or not display_value:
        return display_value
    try:
        return date.fromisoformat(display_value)
    except ValueError:
        # Preserve unexpected historical values in exports instead of hiding them.
        return display_value


def compare_dataset_revisions(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    left_revision: LegacyIssueDataRevision,
    right_revision: LegacyIssueDataRevision,
    module_key: str | None = None,
) -> list[DatasetRevisionCompareRow]:
    left_records = _dataset_records_by_stable_id(
        db,
        definition,
        workspace=workspace,
        revision=left_revision,
        module_key=module_key,
    )
    right_records = _dataset_records_by_stable_id(
        db,
        definition,
        workspace=workspace,
        revision=right_revision,
        module_key=module_key,
    )
    left_attachments = list_dataset_attachments_for_records(
        db,
        definition,
        workspace=workspace,
        record_ids=[record.id for record in left_records.values()],
    )
    right_attachments = list_dataset_attachments_for_records(
        db,
        definition,
        workspace=workspace,
        record_ids=[record.id for record in right_records.values()],
    )
    labels = field_labels(definition)
    rows: list[DatasetRevisionCompareRow] = []
    for stable_record_id in sorted(set(left_records) | set(right_records)):
        left = left_records.get(stable_record_id)
        right = right_records.get(stable_record_id)
        left_values = canonicalize_dataset_values(left.field_values) if left else {}
        right_values = canonicalize_dataset_values(right.field_values) if right else {}
        cells: list[DatasetRevisionCompareCell] = []
        changed = False
        for dataset_field in definition.fields:
            left_value = normalize_history_value(left_values.get(dataset_field.key))
            right_value = normalize_history_value(right_values.get(dataset_field.key))
            cell_changed = left_value != right_value
            changed = changed or cell_changed
            cells.append(
                DatasetRevisionCompareCell(
                    field_key=dataset_field.key,
                    field_label=labels.get(dataset_field.key, dataset_field.key),
                    left_value=left_value,
                    right_value=right_value,
                    changed=cell_changed,
                )
            )
        left_attachment_value = _compare_attachment_summary(
            left_attachments.get(left.id, []) if left else []
        )
        right_attachment_value = _compare_attachment_summary(
            right_attachments.get(right.id, []) if right else []
        )
        attachment_changed = left_attachment_value != right_attachment_value
        changed = changed or attachment_changed
        if attachment_changed or left_attachment_value or right_attachment_value:
            cells.append(
                DatasetRevisionCompareCell(
                    field_key=ATTACHMENT_COMPARE_FIELD_KEY,
                    field_label=ATTACHMENT_COMPARE_FIELD_LABEL,
                    left_value=left_attachment_value,
                    right_value=right_attachment_value,
                    changed=attachment_changed,
                )
            )
        if left is None:
            status_value = "added"
        elif right is None:
            status_value = "removed"
        elif changed:
            status_value = "modified"
        else:
            status_value = "unchanged"
        display_values = right_values or left_values
        label = (
            normalize_history_value(display_values.get("legacy_issue_number"))
            or normalize_history_value(display_values.get("symptom"))
            or normalize_history_value(display_values.get("problem"))
            or stable_record_id
        )
        rows.append(
            DatasetRevisionCompareRow(
                stable_record_id=stable_record_id,
                status=status_value,
                label=label,
                cells=cells,
            )
        )
    return rows


def _dataset_records_by_stable_id(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    revision: LegacyIssueDataRevision,
    module_key: str | None = None,
) -> dict[str, Any]:
    model = definition.table_model
    statement = select(model).where(
        model.workspace_id == workspace.id,
        model.dataset_key == definition.key,
        model.revision_id == revision.id,
    )
    normalized_module_key = normalize_legacy_issue_module_key(module_key)
    if normalized_module_key is not None:
        statement = statement.where(model.module_key == normalized_module_key)
    rows = list(db.scalars(statement))
    return {row.stable_record_id or row.id: row for row in rows}


def _compare_attachment_summary(attachments: list[LegacyIssueAttachment]) -> str | None:
    if not attachments:
        return None
    return "\n".join(
        _compare_attachment_line(attachment)
        for attachment in sorted(
            attachments,
            key=lambda item: (
                not item.is_primary,
                item.filename.lower(),
                item.content_type,
                item.size_bytes,
                item.description or "",
            ),
        )
    )


def _compare_attachment_line(attachment: LegacyIssueAttachment) -> str:
    primary_prefix = "대표: " if attachment.is_primary else ""
    size_label = f"{attachment.size_bytes} bytes"
    description = (attachment.description or "").strip()
    suffix = f" - {description}" if description else ""
    return f"{primary_prefix}{attachment.filename} ({size_label}){suffix}"


def list_dataset_attachments(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    record_id: str,
) -> list[LegacyIssueAttachment]:
    return list(
        db.scalars(
            select(LegacyIssueAttachment)
            .where(
                LegacyIssueAttachment.workspace_id == workspace.id,
                LegacyIssueAttachment.dataset_key == definition.key,
                LegacyIssueAttachment.record_id == record_id,
            )
            .order_by(
                LegacyIssueAttachment.is_primary.desc(),
                LegacyIssueAttachment.created_at.desc(),
            )
        )
    )


def list_dataset_attachments_for_records(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    record_ids: list[str],
) -> dict[str, list[LegacyIssueAttachment]]:
    if not record_ids:
        return {}
    attachments = list(
        db.scalars(
            select(LegacyIssueAttachment)
            .where(
                LegacyIssueAttachment.workspace_id == workspace.id,
                LegacyIssueAttachment.dataset_key == definition.key,
                LegacyIssueAttachment.record_id.in_(record_ids),
            )
            .order_by(
                LegacyIssueAttachment.record_id,
                LegacyIssueAttachment.is_primary.desc(),
                LegacyIssueAttachment.created_at.desc(),
            )
        )
    )
    grouped: dict[str, list[LegacyIssueAttachment]] = {}
    for attachment in attachments:
        grouped.setdefault(attachment.record_id, []).append(attachment)
    return grouped


def upload_dataset_attachment(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    record_id: str,
    upload: DatasetAttachmentUpload,
) -> LegacyIssueAttachment:
    record = get_dataset_record(db, definition, workspace=workspace, record_id=record_id)
    ensure_record_partition(db, record)
    if len(upload.content) > DATASET_ATTACHMENT_MAX_BYTES:
        raise localized_http_exception(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            code="legacy_issues.dataset_attachment_too_large",
        )
    attachment_id = new_id()
    filename = normalize_upload_filename(upload.filename)
    storage_key = (
        f"legacy-issues/{definition.key}/{workspace.id}/{record_id}/{attachment_id}/{filename}"
    )
    settings = get_settings()
    client = get_minio_client()
    ensure_bucket()
    client.put_object(
        settings.minio_bucket,
        storage_key,
        BytesIO(upload.content),
        length=len(upload.content),
        content_type=upload.content_type or DEFAULT_CONTENT_TYPE,
    )
    if upload.is_primary:
        _clear_primary_attachments(
            db, definition=definition, workspace=workspace, record_id=record_id
        )
    attachment = LegacyIssueAttachment(
        id=attachment_id,
        workspace_id=workspace.id,
        dataset_key=definition.key,
        revision_id=record.revision_id,
        stable_record_id=record.stable_record_id or record.id,
        record_id=record_id,
        filename=filename,
        content_type=upload.content_type or DEFAULT_CONTENT_TYPE,
        size_bytes=len(upload.content),
        description=normalize_attachment_description(upload.description),
        storage_key=storage_key,
        is_primary=upload.is_primary,
        uploaded_by_id=user.id,
    )
    ensure_attachment_partition(db, attachment, record=record)
    db.add(attachment)
    db.flush()
    add_attachment_history_entry(
        db,
        workspace=workspace,
        user=user,
        record_kind=LEGACY_ISSUE_RECORD_KIND,
        dataset_key=definition.key,
        record_id=record.stable_record_id or record.id,
        action="attachment_upload",
        field_label="첨부",
        old_value=None,
        new_value=filename,
        details={"attachment_id": attachment.id},
        revision_id=record.revision_id,
    )
    db.refresh(attachment)
    return attachment


def set_dataset_attachment_primary(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    attachment_id: str,
    is_primary: bool,
) -> LegacyIssueAttachment:
    attachment = get_dataset_attachment(
        db, definition, workspace=workspace, attachment_id=attachment_id
    )
    old_primary = next(
        (
            item.filename
            for item in list_dataset_attachments(
                db,
                definition,
                workspace=workspace,
                record_id=attachment.record_id,
            )
            if item.is_primary
        ),
        None,
    )
    was_primary = attachment.is_primary
    if is_primary:
        _clear_primary_attachments(
            db, definition=definition, workspace=workspace, record_id=attachment.record_id
        )
        new_primary = attachment.filename
    else:
        new_primary = None if was_primary else old_primary
    attachment.is_primary = is_primary
    db.add(attachment)
    db.flush()
    if old_primary != new_primary:
        add_attachment_history_entry(
            db,
            workspace=workspace,
            user=user,
            record_kind=LEGACY_ISSUE_RECORD_KIND,
            dataset_key=definition.key,
            record_id=attachment.stable_record_id or attachment.record_id,
            action="attachment_primary",
            field_label="대표 첨부",
            old_value=old_primary,
            new_value=new_primary,
            details={"attachment_id": attachment.id},
            revision_id=attachment.revision_id,
        )
    db.refresh(attachment)
    return attachment


def update_dataset_attachment_description(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    attachment_id: str,
    description: str | None,
) -> LegacyIssueAttachment:
    attachment = get_dataset_attachment(
        db, definition, workspace=workspace, attachment_id=attachment_id
    )
    old_description = attachment.description
    new_description = normalize_attachment_description(description)
    attachment.description = new_description
    db.add(attachment)
    db.flush()
    if old_description != new_description:
        add_attachment_history_entry(
            db,
            workspace=workspace,
            user=user,
            record_kind=LEGACY_ISSUE_RECORD_KIND,
            dataset_key=definition.key,
            record_id=attachment.stable_record_id or attachment.record_id,
            action="attachment_description",
            field_label="첨부 설명",
            old_value=old_description,
            new_value=new_description,
            details={"attachment_id": attachment.id},
            revision_id=attachment.revision_id,
        )
    db.refresh(attachment)
    return attachment


def delete_dataset_attachment(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    user: User,
    attachment_id: str,
) -> None:
    attachment = get_dataset_attachment(
        db, definition, workspace=workspace, attachment_id=attachment_id
    )
    record_id = attachment.stable_record_id or attachment.record_id
    filename = attachment.filename
    attachment_id_for_history = attachment.id
    _remove_attachment_object(attachment)
    db.delete(attachment)
    add_attachment_history_entry(
        db,
        workspace=workspace,
        user=user,
        record_kind=LEGACY_ISSUE_RECORD_KIND,
        dataset_key=definition.key,
        record_id=record_id,
        action="attachment_delete",
        field_label="첨부",
        old_value=filename,
        new_value=None,
        details={"attachment_id": attachment_id_for_history},
        revision_id=attachment.revision_id,
    )
    db.flush()


def get_dataset_attachment(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    attachment_id: str,
) -> LegacyIssueAttachment:
    attachment = db.scalar(
        select(LegacyIssueAttachment).where(
            LegacyIssueAttachment.id == attachment_id,
            LegacyIssueAttachment.workspace_id == workspace.id,
            LegacyIssueAttachment.dataset_key == definition.key,
        )
    )
    if attachment is None:
        raise localized_http_exception(
            status_code=404, code="legacy_issues.dataset_attachment_not_found"
        )
    return attachment


def open_dataset_attachment_content(
    db: Session,
    definition: LegacyIssueDatasetDefinition,
    *,
    workspace: Workspace,
    attachment_id: str,
) -> DatasetAttachmentContent:
    attachment = get_dataset_attachment(
        db, definition, workspace=workspace, attachment_id=attachment_id
    )
    settings = get_settings()
    client = get_minio_client()
    try:
        response = client.get_object(settings.minio_bucket, attachment.storage_key)
    except Exception as error:  # pragma: no cover - defensive storage wrapper
        raise localized_http_exception(
            status_code=404,
            code="legacy_issues.dataset_attachment_download_failed",
        ) from error

    def body() -> Iterable[bytes]:
        try:
            while True:
                chunk = response.read(DATASET_ATTACHMENT_CHUNK_SIZE)
                if not chunk:
                    break
                yield chunk
        finally:
            response.close()
            response.release_conn()

    return DatasetAttachmentContent(
        body=body(),
        media_type=attachment.content_type or DEFAULT_CONTENT_TYPE,
        headers=dataset_attachment_content_headers(attachment.filename),
    )


def _clear_primary_attachments(
    db: Session,
    *,
    definition: LegacyIssueDatasetDefinition,
    workspace: Workspace,
    record_id: str,
) -> None:
    for attachment in list_dataset_attachments(
        db, definition, workspace=workspace, record_id=record_id
    ):
        if attachment.is_primary:
            attachment.is_primary = False
            db.add(attachment)


def _remove_attachment_object(attachment: LegacyIssueAttachment) -> None:
    # Revision copies can share the same storage object; physical cleanup needs
    # reference counting and is intentionally left to storage lifecycle jobs.
    return


def parse_mapping_json(mapping_json: str) -> dict[str, int]:
    return parse_required_mapping_json(
        invalid_code=DATASET_IMPORT_INVALID_CODE,
        mapping_json=mapping_json,
    )
