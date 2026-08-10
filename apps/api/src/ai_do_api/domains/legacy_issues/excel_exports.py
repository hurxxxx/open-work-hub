from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import timedelta
from pathlib import Path
import hashlib
import json
import logging
import math
import re
import shutil
import tempfile
from typing import Any, BinaryIO, Protocol
import unicodedata
from urllib.parse import quote

from celery import Celery
from fastapi import status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from ai_do_api.core.i18n import localized_http_exception
from ai_do_api.core.settings import get_settings
from ai_do_api.core.worker_queue_contract import (
    LEGACY_ISSUE_EXCEL_EXPORT_QUEUE,
    LEGACY_ISSUE_EXCEL_EXPORT_TASK_NAME,
)
from ai_do_api.core.worker_task_publisher import create_fail_fast_celery_publisher
from ai_do_api.domains.auth.models import User, Workspace, utcnow_naive
from ai_do_api.domains.auth.security import new_id
from ai_do_api.domains.legacy_issues.dataset_records import (
    DATASET_RECORD_EXPORT_LIMIT,
    EXCEL_DATE_NUMBER_FORMAT,
    EXPORT_RECORD_ID_HEADER,
    LegacyIssueDatasetDefinition,
    canonicalize_dataset_values,
    dataset_excel_value,
    get_dataset_definition_for_view,
    list_dataset_attachments_for_records,
    list_dataset_records,
    normalize_legacy_issue_module_key,
)
from ai_do_api.domains.legacy_issues.excel_ole_package import (
    OleAttachmentPlacement,
    embed_ole_attachments,
)
from ai_do_api.domains.legacy_issues.models import (
    LegacyIssueAttachment,
    LegacyIssueExcelExportJob,
    LegacyIssueVehicleModuleChecklistAttachment,
)
from ai_do_api.domains.legacy_issues.revisioning import (
    LegacyIssueDataRevision,
    legacy_issue_dataset_revision_key,
    resolve_read_revision,
)
from ai_do_api.domains.legacy_issues.vehicle_module_checklist_attachments import (
    list_vehicle_module_checklist_attachments_for_records,
)
from ai_do_api.domains.legacy_issues.vehicle_module_checklists import (
    list_vehicle_module_checklist_records,
)


logger = logging.getLogger(__name__)

EXCEL_EXPORT_SOURCE_DATASET = "dataset"
EXCEL_EXPORT_SOURCE_CHECKLIST = "vehicle_module_checklist"
EXCEL_EXPORT_STATUS_QUEUED = "queued"
EXCEL_EXPORT_STATUS_RUNNING = "running"
EXCEL_EXPORT_STATUS_COMPLETED = "completed"
EXCEL_EXPORT_STATUS_FAILED = "failed"
EXCEL_EXPORT_STATUS_EXPIRED = "expired"

DATASET_EXPORT_RECORD_ID_QUERY_BATCH_SIZE = 5_000
EXCEL_EXPORT_MAX_ATTACHMENT_BYTES = 5 * 1024 * 1024 * 1024
EXCEL_EXPORT_MAX_ATTACHMENTS = 1000
EXCEL_EXPORT_MAX_ATTACHMENTS_PER_RECORD = 64
EXCEL_EXPORT_RESULT_TTL = timedelta(hours=24)
EXCEL_EXPORT_RETRY_DELAY = timedelta(minutes=2)
EXCEL_EXPORT_STALE_RUNNING_AFTER = timedelta(hours=2)
EXCEL_EXPORT_MAX_ATTEMPTS = 3
EXCEL_EXPORT_PROGRESS_BYTES = 64 * 1024 * 1024
EXCEL_EXPORT_TEMP_OVERHEAD_BYTES = 1024 * 1024 * 1024
EXCEL_EXPORT_RESULT_PREFIX = "legacy-issues/excel-exports/"
EXCEL_EXPORT_ORPHAN_GRACE = timedelta(hours=2)
EXCEL_EXPORT_ATTACHMENT_COLUMN_WIDTH = 15
EXCEL_EXPORT_ATTACHMENT_FILENAME_COLUMN_WIDTH = 36
EXCEL_EXPORT_ICONS_PER_LINE = 4
EXCEL_EXPORT_ICON_LINE_HEIGHT_POINTS = 21
EXCEL_EXPORT_HEADER_ROW_HEIGHT = 24
EXCEL_EXPORT_SUBHEADER_ROW_HEIGHT = 27
EXCEL_EXPORT_DATA_ROW_HEIGHT = 22
EXCEL_EXPORT_MAX_TEXT_ROW_HEIGHT = 112
EXCEL_EXPORT_WIDTH_SCAN_ROWS = 1000

_EXCEL_EXPORT_LONG_TEXT_FIELD_KEYS = frozenset(
    {
        "symptom",
        "cause",
        "countermeasure",
        "action",
        "confirmation_content",
        "check_plan",
        "reflection_result",
        "notes",
    }
)
_EXCEL_EXPORT_NARROW_FIELD_TYPES = frozenset({"boolean", "date", "number"})
_EXCEL_EXPORT_HEADER_BORDER = Border(
    left=Side(style="thin", color="AFC2D4"),
    right=Side(style="thin", color="AFC2D4"),
    top=Side(style="thin", color="AFC2D4"),
    bottom=Side(style="thin", color="AFC2D4"),
)
_EXCEL_EXPORT_DATA_BORDER = Border(
    left=Side(style="thin", color="D7DEE8"),
    right=Side(style="thin", color="D7DEE8"),
    top=Side(style="thin", color="D7DEE8"),
    bottom=Side(style="thin", color="D7DEE8"),
)

_SAFE_RESULT_FILENAME = re.compile(r"[^0-9A-Za-z가-힣._-]+")
_PENDING_PUBLISHES_KEY = "legacy_issue_excel_export_pending_publish_ids"


class ExcelExportStorageObject(Protocol):
    def read(self, size: int = -1) -> bytes: ...

    def close(self) -> None: ...

    def release_conn(self) -> None: ...


class ExcelExportStorageClient(Protocol):
    def get_object(self, bucket_name: str, object_name: str) -> ExcelExportStorageObject: ...

    def put_object(
        self,
        bucket_name: str,
        object_name: str,
        data: BinaryIO,
        length: int,
        content_type: str,
    ) -> Any: ...

    def remove_object(self, bucket_name: str, object_name: str) -> None: ...

    def list_objects(
        self,
        bucket_name: str,
        *,
        prefix: str,
        recursive: bool,
    ) -> Iterable[Any]: ...


@dataclass(frozen=True)
class ExcelExportAttachmentRef:
    record_id: str
    filename: str
    size_bytes: int
    storage_key: str


@dataclass(frozen=True)
class ExcelExportManifest:
    definition: LegacyIssueDatasetDefinition
    rows: list[Any]
    attachments_by_record_id: dict[str, list[ExcelExportAttachmentRef]]
    sheet_name: str
    result_filename: str
    checklist_layout: bool
    column_keys: tuple[str, ...] | None = None
    include_attachments: bool = True


@dataclass(frozen=True)
class ExcelExportDownload:
    body: Iterable[bytes]
    media_type: str
    headers: dict[str, str]


def _normalize_excel_export_column_keys(
    definition: LegacyIssueDatasetDefinition,
    column_keys: list[str] | None,
) -> tuple[str, ...] | None:
    if column_keys is None:
        return None
    normalized = tuple(item.strip() for item in column_keys)
    allowed_keys = {field.key for field in definition.fields}
    if (
        any(not item or len(item) > 120 for item in normalized)
        or len(normalized) != len(set(normalized))
        or any(item not in allowed_keys for item in normalized)
    ):
        raise localized_http_exception(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            code="legacy_issues.excel_export_invalid_columns",
        )
    return normalized


def _normalize_excel_export_record_ids(
    record_ids: list[str] | None,
) -> tuple[str, ...] | None:
    if record_ids is None:
        return None
    return tuple(dict.fromkeys(item.strip() for item in record_ids if item.strip()))


def _select_dataset_export_rows(
    rows: list[Any],
    record_ids: tuple[str, ...] | None,
) -> list[Any]:
    if record_ids is None:
        return rows
    rows_by_id = {str(row.id): row for row in rows}
    return [rows_by_id[record_id] for record_id in record_ids if record_id in rows_by_id]


def _load_dataset_export_rows(
    db: Session,
    *,
    workspace: Workspace,
    definition: LegacyIssueDatasetDefinition,
    revision: LegacyIssueDataRevision,
    module_key: str,
    departments: list[str] | None,
    record_ids: tuple[str, ...] | None,
) -> list[Any]:
    if record_ids is None:
        rows, _ = list_dataset_records(
            db,
            definition,
            workspace=workspace,
            revision=revision,
            module_key=module_key,
            departments=departments,
            limit=DATASET_RECORD_EXPORT_LIMIT,
            offset=0,
        )
        return rows

    scoped_rows: list[Any] = []
    for start in range(0, len(record_ids), DATASET_EXPORT_RECORD_ID_QUERY_BATCH_SIZE):
        batch = record_ids[start : start + DATASET_EXPORT_RECORD_ID_QUERY_BATCH_SIZE]
        rows, _ = list_dataset_records(
            db,
            definition,
            workspace=workspace,
            revision=revision,
            module_key=module_key,
            departments=departments,
            record_ids=batch,
            limit=len(batch),
            offset=0,
        )
        scoped_rows.extend(rows)
    return _select_dataset_export_rows(scoped_rows, record_ids)


def create_dataset_excel_export_job(
    db: Session,
    *,
    workspace: Workspace,
    user: User,
    definition: LegacyIssueDatasetDefinition,
    revision: LegacyIssueDataRevision,
    view_key: str,
    departments: list[str] | None,
    column_keys: list[str] | None = None,
    record_ids: list[str] | None = None,
    include_attachments: bool = True,
) -> LegacyIssueExcelExportJob:
    module_key = normalize_legacy_issue_module_key(view_key)
    if module_key is None:
        raise localized_http_exception(
            status_code=status.HTTP_409_CONFLICT,
            code="legacy_issues.aggregate_readonly",
        )
    normalized_column_keys = _normalize_excel_export_column_keys(
        definition,
        column_keys,
    )
    normalized_record_ids = _normalize_excel_export_record_ids(record_ids)
    request_params: dict[str, Any] = {
        "view_key": view_key,
        "departments": sorted({item for item in departments or [] if item}),
        "include_attachments": include_attachments,
    }
    if normalized_column_keys is not None:
        request_params["column_keys"] = list(normalized_column_keys)
    if normalized_record_ids is not None:
        request_params["record_ids"] = list(normalized_record_ids)
    _acquire_excel_export_creation_lock(
        db,
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind=EXCEL_EXPORT_SOURCE_DATASET,
        source_identity={
            "dataset_key": definition.key,
            "module_key": module_key,
            "revision_id": revision.id,
            "request_params": request_params,
        },
    )
    existing = _find_active_excel_export_job(
        db,
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind=EXCEL_EXPORT_SOURCE_DATASET,
        dataset_key=definition.key,
        module_key=module_key,
        revision_id=revision.id,
        checklist_id=None,
        request_params=request_params,
    )
    if existing is not None:
        return existing
    rows = _load_dataset_export_rows(
        db,
        workspace=workspace,
        definition=definition,
        revision=revision,
        module_key=module_key,
        departments=departments,
        record_ids=normalized_record_ids,
    )
    grouped = (
        _dataset_attachment_refs(
            db,
            workspace=workspace,
            definition=definition,
            rows=rows,
        )
        if include_attachments
        else {}
    )
    attachment_count, attachment_bytes = validate_excel_export_attachments(grouped)
    now = utcnow_naive()
    job = LegacyIssueExcelExportJob(
        id=new_id(),
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind=EXCEL_EXPORT_SOURCE_DATASET,
        dataset_key=definition.key,
        module_key=module_key,
        revision_id=revision.id,
        request_params=request_params,
        status=EXCEL_EXPORT_STATUS_QUEUED,
        record_count=len(rows),
        attachment_count=attachment_count,
        attachment_bytes=attachment_bytes,
        next_retry_at=now,
        created_at=now,
        updated_at=now,
    )
    db.add(job)
    db.flush()
    schedule_excel_export_publish_after_commit(db, job_id=job.id)
    return job


def create_vehicle_module_checklist_excel_export_job(
    db: Session,
    *,
    workspace: Workspace,
    user: User,
    checklist_id: str,
    module_keys: frozenset[str] | None,
    column_keys: list[str] | None = None,
    include_attachments: bool = True,
) -> LegacyIssueExcelExportJob:
    request_params: dict[str, Any] = {"include_attachments": include_attachments}
    if column_keys is not None:
        request_params["column_keys"] = column_keys
    _acquire_excel_export_creation_lock(
        db,
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind=EXCEL_EXPORT_SOURCE_CHECKLIST,
        source_identity={
            "checklist_id": checklist_id,
            "request_params": request_params,
        },
    )
    existing = _find_active_excel_export_job(
        db,
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind=EXCEL_EXPORT_SOURCE_CHECKLIST,
        dataset_key=None,
        module_key=None,
        revision_id=None,
        checklist_id=checklist_id,
        request_params=request_params,
    )
    if existing is not None:
        return existing
    checklist, _definition, rows, total = list_vehicle_module_checklist_records(
        db,
        workspace=workspace,
        checklist_id=checklist_id,
        limit=None,
        offset=0,
        module_keys=module_keys,
    )
    normalized_column_keys = _normalize_excel_export_column_keys(
        _definition,
        column_keys,
    )
    if normalized_column_keys is not None:
        request_params["column_keys"] = list(normalized_column_keys)
    grouped: dict[str, list[ExcelExportAttachmentRef]] = {}
    if include_attachments:
        attachments = list_vehicle_module_checklist_attachments_for_records(
            db,
            workspace=workspace,
            checklist=checklist,
            record_ids=[row.id for row in rows],
        )
        grouped = {
            record_id: [_checklist_attachment_ref(item) for item in items]
            for record_id, items in attachments.items()
        }
    attachment_count, attachment_bytes = validate_excel_export_attachments(grouped)
    now = utcnow_naive()
    job = LegacyIssueExcelExportJob(
        id=new_id(),
        workspace_id=workspace.id,
        requested_by_id=user.id,
        source_kind=EXCEL_EXPORT_SOURCE_CHECKLIST,
        dataset_key=checklist.source_dataset_key,
        module_key=checklist.module_key,
        revision_id=checklist.source_master_revision_id,
        checklist_id=checklist.id,
        request_params=request_params,
        status=EXCEL_EXPORT_STATUS_QUEUED,
        record_count=total,
        attachment_count=attachment_count,
        attachment_bytes=attachment_bytes,
        next_retry_at=now,
        created_at=now,
        updated_at=now,
    )
    db.add(job)
    db.flush()
    schedule_excel_export_publish_after_commit(db, job_id=job.id)
    return job


def get_excel_export_job(
    db: Session,
    *,
    workspace: Workspace,
    user: User,
    job_id: str,
) -> LegacyIssueExcelExportJob:
    row = db.scalar(
        select(LegacyIssueExcelExportJob).where(
            LegacyIssueExcelExportJob.id == job_id,
            LegacyIssueExcelExportJob.workspace_id == workspace.id,
            LegacyIssueExcelExportJob.requested_by_id == user.id,
        )
    )
    if row is None:
        raise localized_http_exception(
            status_code=status.HTTP_404_NOT_FOUND,
            code="legacy_issues.excel_export_not_found",
        )
    return row


def open_excel_export_result(
    db: Session,
    *,
    workspace: Workspace,
    user: User,
    job_id: str,
    client: ExcelExportStorageClient | None = None,
    bucket_name: str | None = None,
) -> ExcelExportDownload:
    job = get_excel_export_job(
        db,
        workspace=workspace,
        user=user,
        job_id=job_id,
    )
    effective_status = excel_export_job_status(job)
    if effective_status != EXCEL_EXPORT_STATUS_COMPLETED or not job.result_storage_key:
        raise localized_http_exception(
            status_code=status.HTTP_409_CONFLICT,
            code=(
                "legacy_issues.excel_export_expired"
                if effective_status == EXCEL_EXPORT_STATUS_EXPIRED
                else "legacy_issues.excel_export_not_ready"
            ),
        )
    storage, resolved_bucket = _resolve_storage(client=client, bucket_name=bucket_name)
    try:
        response = storage.get_object(resolved_bucket, job.result_storage_key)
    except Exception as error:
        raise localized_http_exception(
            status_code=status.HTTP_404_NOT_FOUND,
            code="legacy_issues.excel_export_download_failed",
        ) from error

    def body() -> Iterable[bytes]:
        try:
            while True:
                chunk = response.read(1024 * 1024)
                if not chunk:
                    break
                yield chunk
        finally:
            response.close()
            response.release_conn()

    filename = job.result_filename or "legacy-issues.xlsx"
    return ExcelExportDownload(
        body=body(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename, safe='')}",
            "Cache-Control": "private, no-store",
        },
    )


def validate_excel_export_attachments(
    attachments_by_record_id: dict[str, list[ExcelExportAttachmentRef]],
) -> tuple[int, int]:
    attachment_count = 0
    attachment_bytes = 0
    for attachments in attachments_by_record_id.values():
        if len(attachments) > EXCEL_EXPORT_MAX_ATTACHMENTS_PER_RECORD:
            raise localized_http_exception(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                code="legacy_issues.excel_export_record_attachment_limit",
            )
        attachment_count += len(attachments)
        attachment_bytes += sum(max(item.size_bytes, 0) for item in attachments)
    if attachment_count > EXCEL_EXPORT_MAX_ATTACHMENTS:
        raise localized_http_exception(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            code="legacy_issues.excel_export_attachment_count_limit",
        )
    if attachment_bytes > EXCEL_EXPORT_MAX_ATTACHMENT_BYTES:
        raise localized_http_exception(
            status_code=status.HTTP_413_CONTENT_TOO_LARGE,
            code="legacy_issues.excel_export_attachment_size_limit",
        )
    return attachment_count, attachment_bytes


def _acquire_excel_export_creation_lock(
    db: Session,
    *,
    workspace_id: str,
    requested_by_id: str,
    source_kind: str,
    source_identity: dict[str, Any],
) -> None:
    bind = db.get_bind()
    if bind.dialect.name != "postgresql":
        return
    fingerprint = json.dumps(
        [workspace_id, requested_by_id, source_kind, source_identity],
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    lock_key = int.from_bytes(
        hashlib.blake2b(fingerprint, digest_size=8).digest(),
        byteorder="big",
        signed=True,
    )
    db.execute(select(func.pg_advisory_xact_lock(lock_key)))


def _find_active_excel_export_job(
    db: Session,
    *,
    workspace_id: str,
    requested_by_id: str,
    source_kind: str,
    dataset_key: str | None,
    module_key: str | None,
    revision_id: str | None,
    checklist_id: str | None,
    request_params: dict[str, Any],
) -> LegacyIssueExcelExportJob | None:
    statement = select(LegacyIssueExcelExportJob).where(
        LegacyIssueExcelExportJob.workspace_id == workspace_id,
        LegacyIssueExcelExportJob.requested_by_id == requested_by_id,
        LegacyIssueExcelExportJob.source_kind == source_kind,
        LegacyIssueExcelExportJob.status.in_(
            (EXCEL_EXPORT_STATUS_QUEUED, EXCEL_EXPORT_STATUS_RUNNING)
        ),
    )
    if source_kind == EXCEL_EXPORT_SOURCE_DATASET:
        statement = statement.where(
            LegacyIssueExcelExportJob.dataset_key == dataset_key,
            LegacyIssueExcelExportJob.module_key == module_key,
            LegacyIssueExcelExportJob.revision_id == revision_id,
        )
    else:
        statement = statement.where(LegacyIssueExcelExportJob.checklist_id == checklist_id)
    rows = db.scalars(statement.order_by(LegacyIssueExcelExportJob.created_at.desc()))
    normalized_request_params = _normalized_export_request_params(request_params)
    return next(
        (
            row
            for row in rows
            if _normalized_export_request_params(row.request_params) == normalized_request_params
        ),
        None,
    )


def excel_export_includes_attachments(job: LegacyIssueExcelExportJob) -> bool:
    return _include_attachments_from_request_params(job.request_params)


def _include_attachments_from_request_params(request_params: dict[str, Any]) -> bool:
    # Jobs created before this option existed have no flag and included attachments.
    return request_params.get("include_attachments", True) is not False


def _normalized_export_request_params(request_params: dict[str, Any]) -> dict[str, Any]:
    return {
        **request_params,
        "include_attachments": _include_attachments_from_request_params(request_params),
    }


def process_excel_export_job(
    db: Session,
    *,
    job_id: str,
    client: ExcelExportStorageClient | None = None,
    bucket_name: str | None = None,
    temp_root: str | None = None,
) -> str:
    job = _claim_excel_export_job(db, job_id=job_id)
    if job is None:
        return "ignored"
    storage, resolved_bucket = _resolve_storage(client=client, bucket_name=bucket_name)
    workspace = db.get(Workspace, job.workspace_id)
    if workspace is None:
        raise RuntimeError("Excel export workspace is missing.")
    manifest = _load_export_manifest(db, workspace=workspace, job=job)
    attachment_count, attachment_bytes = validate_excel_export_attachments(
        manifest.attachments_by_record_id
    )
    job.record_count = len(manifest.rows)
    job.attachment_count = attachment_count
    job.attachment_bytes = attachment_bytes
    job.processed_attachment_count = 0
    job.processed_attachment_bytes = 0
    job.updated_at = utcnow_naive()
    db.add(job)
    db.commit()

    _require_temp_space(attachment_bytes, temp_root=temp_root)
    output_storage_key = f"{EXCEL_EXPORT_RESULT_PREFIX}{workspace.id}/{job.id}.xlsx"
    last_progress_count = 0
    last_progress_bytes = 0

    with tempfile.TemporaryDirectory(prefix="legacy-issue-export-", dir=temp_root) as temp_dir:
        base_path = Path(temp_dir) / "base.xlsx"
        output_path = Path(temp_dir) / "result.xlsx"
        placements = _build_base_workbook(
            manifest,
            workbook_path=base_path,
            client=storage,
            bucket_name=resolved_bucket,
        )

        def on_progress(completed_count: int, completed_bytes: int) -> None:
            nonlocal last_progress_count, last_progress_bytes
            if (
                completed_count < attachment_count
                and completed_count - last_progress_count < 10
                and completed_bytes - last_progress_bytes < EXCEL_EXPORT_PROGRESS_BYTES
            ):
                return
            _update_excel_export_progress(
                db,
                job_id=job.id,
                completed_count=completed_count,
                completed_bytes=completed_bytes,
            )
            last_progress_count = completed_count
            last_progress_bytes = completed_bytes

        if manifest.include_attachments:
            embed_ole_attachments(
                base_path,
                output_path,
                placements,
                progress_callback=on_progress,
            )
        else:
            shutil.copyfile(base_path, output_path)
        result_size = output_path.stat().st_size
        with output_path.open("rb") as output:
            storage.put_object(
                resolved_bucket,
                output_storage_key,
                output,
                length=result_size,
                content_type=("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            )

    now = utcnow_naive()
    completed_job = db.get(LegacyIssueExcelExportJob, job.id)
    if completed_job is None:
        _remove_result_object(storage, resolved_bucket, output_storage_key)
        return "cancelled"
    completed_job.status = EXCEL_EXPORT_STATUS_COMPLETED
    completed_job.processed_attachment_count = attachment_count
    completed_job.processed_attachment_bytes = attachment_bytes
    completed_job.result_storage_key = output_storage_key
    completed_job.result_filename = manifest.result_filename
    completed_job.result_size_bytes = result_size
    completed_job.error_code = None
    completed_job.last_error = None
    completed_job.next_retry_at = None
    completed_job.completed_at = now
    completed_job.expires_at = now + EXCEL_EXPORT_RESULT_TTL
    completed_job.updated_at = now
    db.add(completed_job)
    try:
        db.commit()
    except Exception:
        db.rollback()
        _remove_result_object(storage, resolved_bucket, output_storage_key)
        raise
    return "completed"


def mark_excel_export_job_for_retry(
    db: Session,
    *,
    job_id: str,
    error: Exception,
) -> bool:
    db.rollback()
    job = db.get(LegacyIssueExcelExportJob, job_id)
    if job is None or job.status in {
        EXCEL_EXPORT_STATUS_COMPLETED,
        EXCEL_EXPORT_STATUS_EXPIRED,
    }:
        return False
    now = utcnow_naive()
    retry = job.attempts < EXCEL_EXPORT_MAX_ATTEMPTS
    job.status = EXCEL_EXPORT_STATUS_QUEUED if retry else EXCEL_EXPORT_STATUS_FAILED
    job.error_code = _excel_export_error_code(error)
    job.last_error = _safe_error_text(error)
    job.next_retry_at = now + EXCEL_EXPORT_RETRY_DELAY if retry else None
    job.completed_at = None if retry else now
    job.updated_at = now
    db.add(job)
    db.commit()
    return retry


def republish_pending_excel_export_jobs(db: Session, *, limit: int = 100) -> int:
    now = utcnow_naive()
    candidates = list(
        db.scalars(
            select(LegacyIssueExcelExportJob)
            .where(
                or_(
                    and_(
                        LegacyIssueExcelExportJob.status == EXCEL_EXPORT_STATUS_QUEUED,
                        or_(
                            LegacyIssueExcelExportJob.next_retry_at.is_(None),
                            LegacyIssueExcelExportJob.next_retry_at <= now,
                        ),
                    ),
                    and_(
                        LegacyIssueExcelExportJob.status == EXCEL_EXPORT_STATUS_RUNNING,
                        LegacyIssueExcelExportJob.started_at.is_not(None),
                        LegacyIssueExcelExportJob.started_at
                        <= now - EXCEL_EXPORT_STALE_RUNNING_AFTER,
                    ),
                ),
            )
            .order_by(LegacyIssueExcelExportJob.created_at.asc())
            .limit(max(1, min(limit, 1000)))
            .with_for_update(skip_locked=True)
        )
    )
    jobs_to_publish: list[LegacyIssueExcelExportJob] = []
    for job in candidates:
        if job.status == EXCEL_EXPORT_STATUS_RUNNING and job.attempts >= EXCEL_EXPORT_MAX_ATTEMPTS:
            job.status = EXCEL_EXPORT_STATUS_FAILED
            job.error_code = "legacy_issues.excel_export_failed"
            job.last_error = "stale_running_job_exhausted"
            job.next_retry_at = None
            job.completed_at = now
            job.updated_at = now
            db.add(job)
            continue
        job.status = EXCEL_EXPORT_STATUS_QUEUED
        job.next_retry_at = now + EXCEL_EXPORT_RETRY_DELAY
        job.updated_at = now
        db.add(job)
        jobs_to_publish.append(job)
    db.commit()
    published = 0
    for job in jobs_to_publish:
        try:
            publish_excel_export_job(job.id)
        except Exception:
            logger.warning("Failed to republish Excel export job", exc_info=True)
        else:
            published += 1
    return published


def cleanup_expired_excel_export_jobs(
    db: Session,
    *,
    client: ExcelExportStorageClient | None = None,
    bucket_name: str | None = None,
    limit: int = 100,
) -> dict[str, int]:
    storage, resolved_bucket = _resolve_storage(client=client, bucket_name=bucket_name)
    now = utcnow_naive()
    rows = list(
        db.scalars(
            select(LegacyIssueExcelExportJob)
            .where(
                LegacyIssueExcelExportJob.status == EXCEL_EXPORT_STATUS_COMPLETED,
                LegacyIssueExcelExportJob.expires_at.is_not(None),
                LegacyIssueExcelExportJob.expires_at <= now,
            )
            .order_by(LegacyIssueExcelExportJob.expires_at.asc())
            .limit(max(1, min(limit, 1000)))
            .with_for_update(skip_locked=True)
        )
    )
    storage_keys: list[str] = []
    for row in rows:
        if row.result_storage_key:
            storage_keys.append(row.result_storage_key)
        row.status = EXCEL_EXPORT_STATUS_EXPIRED
        row.result_storage_key = None
        row.result_size_bytes = None
        row.updated_at = now
        db.add(row)
    db.commit()

    expired = len(rows)
    failed = 0
    for storage_key in storage_keys:
        try:
            storage.remove_object(resolved_bucket, storage_key)
        except Exception:
            failed += 1
            logger.warning("Failed to remove expired Excel export", exc_info=True)
    orphan_deleted, orphan_failed = _cleanup_orphan_export_objects(
        db,
        client=storage,
        bucket_name=resolved_bucket,
        now=now,
        limit=max(1, min(limit, 1000)),
    )
    return {
        "expired": expired,
        "deleted_orphans": orphan_deleted,
        "failed": failed + orphan_failed,
    }


def schedule_excel_export_publish_after_commit(db: Session, *, job_id: str) -> None:
    pending = db.info.setdefault(_PENDING_PUBLISHES_KEY, set())
    if not isinstance(pending, set):
        pending = set()
        db.info[_PENDING_PUBLISHES_KEY] = pending
    pending.add(job_id)


def publish_excel_export_job(job_id: str) -> None:
    _celery_client().signature(
        LEGACY_ISSUE_EXCEL_EXPORT_TASK_NAME,
        args=[job_id],
        immutable=True,
    ).apply_async(queue=LEGACY_ISSUE_EXCEL_EXPORT_QUEUE, retry=False)


def _claim_excel_export_job(db: Session, *, job_id: str) -> LegacyIssueExcelExportJob | None:
    job = db.scalar(
        select(LegacyIssueExcelExportJob)
        .where(LegacyIssueExcelExportJob.id == job_id)
        .with_for_update()
    )
    if job is None or job.status != EXCEL_EXPORT_STATUS_QUEUED:
        db.rollback()
        return None
    now = utcnow_naive()
    job.status = EXCEL_EXPORT_STATUS_RUNNING
    job.attempts += 1
    job.started_at = now
    job.error_code = None
    job.last_error = None
    job.next_retry_at = None
    job.updated_at = now
    db.add(job)
    db.commit()
    return job


def _load_export_manifest(
    db: Session,
    *,
    workspace: Workspace,
    job: LegacyIssueExcelExportJob,
) -> ExcelExportManifest:
    include_attachments = excel_export_includes_attachments(job)
    requested_column_keys = (
        [str(item) for item in job.request_params.get("column_keys", [])]
        if "column_keys" in job.request_params
        else None
    )
    requested_record_ids = (
        _normalize_excel_export_record_ids(
            [str(item) for item in job.request_params.get("record_ids", [])]
        )
        if "record_ids" in job.request_params
        else None
    )
    if job.source_kind == EXCEL_EXPORT_SOURCE_DATASET:
        if not job.dataset_key or not job.module_key or not job.revision_id:
            raise RuntimeError("Dataset export source was deleted.")
        view_key = str(job.request_params.get("view_key") or job.module_key)
        definition = get_dataset_definition_for_view(
            db,
            dataset_key=job.dataset_key,
            workspace=workspace,
            view_key=view_key,
        )
        column_keys = _normalize_excel_export_column_keys(
            definition,
            requested_column_keys,
        )
        revision = resolve_read_revision(
            db,
            workspace=workspace,
            dataset_key=legacy_issue_dataset_revision_key(job.dataset_key, job.module_key),
            revision_id=job.revision_id,
        ).current
        departments = [
            str(item) for item in job.request_params.get("departments", []) if str(item).strip()
        ]
        rows = _load_dataset_export_rows(
            db,
            workspace=workspace,
            definition=definition,
            revision=revision,
            module_key=job.module_key,
            departments=departments or None,
            record_ids=requested_record_ids,
        )
        grouped = (
            _dataset_attachment_refs(
                db,
                workspace=workspace,
                definition=definition,
                rows=rows,
            )
            if include_attachments
            else {}
        )
        revision_label = revision.revision_no if revision.revision_no is not None else "draft"
        result_filename = _result_filename(
            f"legacy-issue-{job.module_key}-rev-{revision_label}.xlsx"
        )
        return ExcelExportManifest(
            definition=definition,
            rows=rows,
            attachments_by_record_id=grouped,
            sheet_name=job.module_key[:31],
            result_filename=result_filename,
            checklist_layout=False,
            column_keys=column_keys,
            include_attachments=include_attachments,
        )
    if job.source_kind == EXCEL_EXPORT_SOURCE_CHECKLIST:
        if not job.checklist_id:
            raise RuntimeError("Checklist export source was deleted.")
        checklist, definition, rows, _ = list_vehicle_module_checklist_records(
            db,
            workspace=workspace,
            checklist_id=job.checklist_id,
            limit=None,
            offset=0,
            module_keys=None,
        )
        column_keys = _normalize_excel_export_column_keys(
            definition,
            requested_column_keys,
        )
        grouped: dict[str, list[ExcelExportAttachmentRef]] = {}
        if include_attachments:
            attachments = list_vehicle_module_checklist_attachments_for_records(
                db,
                workspace=workspace,
                checklist=checklist,
                record_ids=[row.id for row in rows],
            )
            grouped = {
                record_id: [_checklist_attachment_ref(item) for item in items]
                for record_id, items in attachments.items()
            }
        result_filename = _result_filename(
            "legacy-issue-checklist-"
            f"{checklist.module_key}-rev-{checklist.source_master_revision_no or 0}.xlsx"
        )
        return ExcelExportManifest(
            definition=definition,
            rows=rows,
            attachments_by_record_id=grouped,
            sheet_name="checklist",
            result_filename=result_filename,
            checklist_layout=True,
            column_keys=column_keys,
            include_attachments=include_attachments,
        )
    raise RuntimeError("Unknown Excel export source kind.")


def _build_base_workbook(
    manifest: ExcelExportManifest,
    *,
    workbook_path: Path,
    client: ExcelExportStorageClient,
    bucket_name: str,
) -> list[OleAttachmentPlacement]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = manifest.sheet_name[:31]
    placements: list[OleAttachmentPlacement] = []
    if manifest.checklist_layout:
        fields, attachment_filename_index, attachment_index = _checklist_fields(
            manifest.definition,
            column_keys=manifest.column_keys,
            include_attachments=manifest.include_attachments,
        )
        _append_checklist_headers(
            sheet,
            manifest.definition,
            fields,
            attachment_filename_index=attachment_filename_index,
            attachment_index=attachment_index,
        )
        data_start_row = 3
    else:
        fields = _selected_export_fields(
            manifest.definition,
            column_keys=manifest.column_keys,
        )
        attachment_filename_index = len(fields) + 1 if manifest.include_attachments else None
        attachment_index = len(fields) + 2 if manifest.include_attachments else None
        _append_dataset_headers(
            sheet,
            manifest.definition,
            fields,
            include_attachments=manifest.include_attachments,
        )
        data_start_row = 3
    locked_fill = PatternFill("solid", fgColor="EEF2F7")
    editable_fill = PatternFill("solid", fgColor="ECFDF5")
    top_header_fill = PatternFill("solid", fgColor="244A66")
    subheader_fill = PatternFill("solid", fgColor="DCE8F2")
    header_rows = data_start_row - 1
    for header_row, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=header_rows),
        start=1,
    ):
        for cell in row:
            cell.font = Font(
                name="맑은 고딕",
                size=10,
                bold=True,
                color="FFFFFF" if header_row == 1 else "183247",
            )
            cell.protection = Protection(locked=True)
            cell.fill = top_header_fill if header_row == 1 else subheader_fill
            cell.border = _EXCEL_EXPORT_HEADER_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = EXCEL_EXPORT_HEADER_ROW_HEIGHT
    sheet.row_dimensions[2].height = EXCEL_EXPORT_SUBHEADER_ROW_HEIGHT

    for row_offset, record in enumerate(manifest.rows):
        values = canonicalize_dataset_values(record.field_values)
        row_values: list[Any] = [
            getattr(record, "source_stable_record_id", None)
            or getattr(record, "stable_record_id", None)
            or record.id
        ]
        if manifest.checklist_layout:
            for field in fields:
                row_values.append(
                    "" if field is None else dataset_excel_value(field, values.get(field.key))
                )
        else:
            row_values.extend(dataset_excel_value(field, values.get(field.key)) for field in fields)
            if manifest.include_attachments:
                row_values.extend(["", ""])
        sheet.append(row_values)
        excel_row = data_start_row + row_offset
        _force_literal_string_cells(sheet[excel_row])
        for column_index, field in enumerate(fields, start=2):
            if field is None or field.field_type != "date":
                continue
            cell = sheet.cell(excel_row, column_index)
            if cell.is_date:
                cell.number_format = EXCEL_DATE_NUMBER_FORMAT
        record_attachments = manifest.attachments_by_record_id.get(record.id, [])
        if attachment_filename_index is not None:
            filename_cell = sheet.cell(excel_row, attachment_filename_index + 1)
            _set_attachment_filenames(
                filename_cell,
                [attachment.filename for attachment in record_attachments],
            )
        if attachment_index is not None:
            attachment_column = attachment_index + 1
            cell_coordinate = f"{get_column_letter(attachment_column)}{excel_row}"
            for attachment in record_attachments:
                placements.append(
                    OleAttachmentPlacement(
                        sheet_name=sheet.title,
                        cell_coordinate=cell_coordinate,
                        filename=attachment.filename,
                        size_bytes=attachment.size_bytes,
                        stream_factory=_storage_stream_factory(
                            client,
                            bucket_name=bucket_name,
                            storage_key=attachment.storage_key,
                        ),
                    )
                )

    sheet.freeze_panes = f"A{data_start_row}"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.sheet_properties.tabColor = "244A66"
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = f"1:{header_rows}"
    sheet.column_dimensions["A"].width = 28
    for field_index, field in enumerate(fields, start=2):
        if field is None:
            continue
        sheet.column_dimensions[get_column_letter(field_index)].width = _fit_column_width(
            sheet,
            column_index=field_index,
            field=field,
            data_start_row=data_start_row,
        )
    if attachment_filename_index is not None:
        sheet.column_dimensions[
            get_column_letter(attachment_filename_index + 1)
        ].width = EXCEL_EXPORT_ATTACHMENT_FILENAME_COLUMN_WIDTH
    if attachment_index is not None:
        sheet.column_dimensions[
            get_column_letter(attachment_index + 1)
        ].width = EXCEL_EXPORT_ATTACHMENT_COLUMN_WIDTH
    for row in sheet.iter_rows(min_row=data_start_row):
        row[0].protection = Protection(locked=True)
        row[0].fill = locked_fill
        row[0].font = Font(name="맑은 고딕", size=10, color="334155")
        row[0].border = _EXCEL_EXPORT_DATA_BORDER
        row[0].alignment = Alignment(vertical="top", wrap_text=True)
        for cell_index, cell in enumerate(row[1:], start=1):
            if cell_index in {attachment_filename_index, attachment_index}:
                readonly = True
            elif manifest.checklist_layout:
                field = fields[cell_index - 1] if cell_index - 1 < len(fields) else None
                readonly = field is None or field.key not in {
                    "check_plan",
                    "applied",
                    "reflection_result",
                }
            else:
                readonly = fields[cell_index - 1].readonly
            cell.protection = Protection(locked=readonly)
            cell.fill = locked_fill if readonly else editable_fill
            cell.font = Font(name="맑은 고딕", size=10, color="1F2937")
            cell.border = _EXCEL_EXPORT_DATA_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_offset, record in enumerate(manifest.rows):
        excel_row = data_start_row + row_offset
        sheet.row_dimensions[excel_row].height = _fit_row_height(
            sheet,
            excel_row=excel_row,
            attachment_count=len(manifest.attachments_by_record_id.get(record.id, [])),
        )
    workbook.save(workbook_path)
    return placements


def _checklist_fields(
    definition: LegacyIssueDatasetDefinition,
    *,
    column_keys: tuple[str, ...] | None = None,
    include_attachments: bool,
) -> tuple[list[Any | None], int | None, int | None]:
    fields: list[Any | None] = []
    attachment_filename_index: int | None = None
    attachment_index: int | None = None
    inserted = False
    for field in _selected_export_fields(definition, column_keys=column_keys):
        fields.append(field)
        if include_attachments and field.key == "reflection_result":
            fields.extend([None, None])
            attachment_filename_index = len(fields) - 1
            attachment_index = len(fields)
            inserted = True
    if include_attachments and not inserted:
        fields.extend([None, None])
        attachment_filename_index = len(fields) - 1
        attachment_index = len(fields)
    return fields, attachment_filename_index, attachment_index


def _selected_export_fields(
    definition: LegacyIssueDatasetDefinition,
    *,
    column_keys: tuple[str, ...] | None,
) -> list[Any]:
    if column_keys is None:
        return list(definition.fields)
    fields_by_key = {field.key: field for field in definition.fields}
    return [fields_by_key[column_key] for column_key in column_keys]


def _append_dataset_headers(
    sheet: Any,
    definition: LegacyIssueDatasetDefinition,
    fields: list[Any],
    *,
    include_attachments: bool,
) -> None:
    columns = [
        (EXPORT_RECORD_ID_HEADER, None),
        *((field.label_ko, field.group_key) for field in fields),
    ]
    group_labels = dict(definition.group_labels_ko)
    if include_attachments:
        group_labels["attachment"] = "첨부"
        columns.extend(
            [
                ("첨부파일명", "attachment"),
                ("첨부파일", "attachment"),
            ]
        )
    _append_grouped_headers(
        sheet,
        columns=columns,
        group_labels=group_labels,
    )


def _append_checklist_headers(
    sheet: Any,
    definition: LegacyIssueDatasetDefinition,
    fields: list[Any | None],
    *,
    attachment_filename_index: int | None,
    attachment_index: int | None,
) -> None:
    columns: list[tuple[str, str | None]] = [(EXPORT_RECORD_ID_HEADER, None)]
    for field_index, field in enumerate(fields, start=1):
        if field is None:
            if field_index == attachment_filename_index:
                label = "첨부파일명"
            elif field_index == attachment_index:
                label = "첨부파일"
            else:
                raise RuntimeError("Unknown checklist export column.")
            columns.append(
                (
                    label,
                    "check",
                )
            )
        else:
            columns.append((field.label_ko, field.group_key))
    _append_grouped_headers(
        sheet,
        columns=columns,
        group_labels=dict(definition.group_labels_ko),
    )


def _append_grouped_headers(
    sheet: Any,
    *,
    columns: list[tuple[str, str | None]],
    group_labels: dict[str, str],
) -> None:
    top_headers = [
        group_labels.get(group_key, label) if group_key else label for label, group_key in columns
    ]
    bottom_headers = [label for label, _group_key in columns]
    group_keys = [group_key for _label, group_key in columns]
    sheet.append(top_headers)
    sheet.append(bottom_headers)
    _force_literal_string_cells(sheet[1])
    _force_literal_string_cells(sheet[2])
    start = 1
    while start <= len(group_keys):
        group_key = group_keys[start - 1]
        if group_key is None:
            sheet.merge_cells(start_row=1, start_column=start, end_row=2, end_column=start)
            start += 1
            continue
        end = start
        while end < len(group_keys) and group_keys[end] == group_key:
            end += 1
        if end > start:
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        start = end + 1


def _force_literal_string_cells(cells: Iterable[Any]) -> None:
    for cell in cells:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def _fit_column_width(
    sheet: Any,
    *,
    column_index: int,
    field: Any,
    data_start_row: int,
) -> float:
    if field.key in _EXCEL_EXPORT_LONG_TEXT_FIELD_KEYS or field.field_type == "longText":
        minimum, maximum = 28.0, 44.0
    elif field.field_type in _EXCEL_EXPORT_NARROW_FIELD_TYPES:
        minimum, maximum = 12.0, 16.0
    elif field.field_type in {"select", "orgUnit", "user"}:
        minimum, maximum = 14.0, 26.0
    else:
        minimum, maximum = 14.0, 30.0
    observed_width = _display_width(field.label_ko) + 2
    max_row = min(sheet.max_row, data_start_row + EXCEL_EXPORT_WIDTH_SCAN_ROWS - 1)
    for row_number in range(data_start_row, max_row + 1):
        observed_width = max(
            observed_width,
            _display_width(sheet.cell(row_number, column_index).value) + 2,
        )
    return min(maximum, max(minimum, float(observed_width)))


def _fit_row_height(
    sheet: Any,
    *,
    excel_row: int,
    attachment_count: int,
) -> float:
    wrapped_lines = 1
    for column_index in range(1, sheet.max_column + 1):
        column_width = sheet.column_dimensions[get_column_letter(column_index)].width or 8.43
        wrapped_lines = max(
            wrapped_lines,
            _wrapped_line_count(sheet.cell(excel_row, column_index).value, column_width),
        )
    text_height = min(
        EXCEL_EXPORT_MAX_TEXT_ROW_HEIGHT,
        EXCEL_EXPORT_DATA_ROW_HEIGHT + max(0, wrapped_lines - 1) * 15,
    )
    icon_lines = math.ceil(attachment_count / EXCEL_EXPORT_ICONS_PER_LINE)
    icon_height = icon_lines * EXCEL_EXPORT_ICON_LINE_HEIGHT_POINTS + 6 if icon_lines else 0
    return min(409.0, max(EXCEL_EXPORT_DATA_ROW_HEIGHT, text_height, icon_height))


def _wrapped_line_count(value: Any, column_width: float) -> int:
    text = "" if value is None else str(value)
    if not text:
        return 1
    usable_width = max(6, int(column_width) - 2)
    return sum(
        max(1, math.ceil(_display_width(line) / usable_width)) for line in text.splitlines() or [""]
    )


def _display_width(value: Any) -> int:
    text = "" if value is None else str(value)
    return max(
        (
            sum(
                2 if unicodedata.east_asian_width(character) in {"F", "W"} else 1
                for character in line
            )
            for line in text.splitlines() or [""]
        ),
        default=0,
    )


def _set_attachment_filenames(cell: Any, filenames: list[str]) -> None:
    if not filenames:
        cell.value = ""
        return
    cell.value = "\n".join(filenames)
    # Keep formula-like filenames as literal text in the generated workbook.
    cell.data_type = "s"


def _dataset_attachment_refs(
    db: Session,
    *,
    workspace: Workspace,
    definition: LegacyIssueDatasetDefinition,
    rows: list[Any],
) -> dict[str, list[ExcelExportAttachmentRef]]:
    grouped: dict[str, list[ExcelExportAttachmentRef]] = {}
    for batch in _batches([row.id for row in rows], 500):
        attachments = list_dataset_attachments_for_records(
            db,
            definition,
            workspace=workspace,
            record_ids=batch,
        )
        for record_id, items in attachments.items():
            grouped.setdefault(record_id, []).extend(
                _dataset_attachment_ref(item) for item in items
            )
    return grouped


def _dataset_attachment_ref(item: LegacyIssueAttachment) -> ExcelExportAttachmentRef:
    return ExcelExportAttachmentRef(
        record_id=item.record_id,
        filename=item.filename,
        size_bytes=item.size_bytes,
        storage_key=item.storage_key,
    )


def _checklist_attachment_ref(
    item: LegacyIssueVehicleModuleChecklistAttachment,
) -> ExcelExportAttachmentRef:
    return ExcelExportAttachmentRef(
        record_id=item.record_id,
        filename=item.filename,
        size_bytes=item.size_bytes,
        storage_key=item.storage_key,
    )


def _storage_stream_factory(
    client: ExcelExportStorageClient,
    *,
    bucket_name: str,
    storage_key: str,
):
    def open_stream() -> BinaryIO:
        return client.get_object(bucket_name, storage_key)  # type: ignore[return-value]

    return open_stream


def _update_excel_export_progress(
    db: Session,
    *,
    job_id: str,
    completed_count: int,
    completed_bytes: int,
) -> None:
    job = db.get(LegacyIssueExcelExportJob, job_id)
    if job is None or job.status != EXCEL_EXPORT_STATUS_RUNNING:
        return
    job.processed_attachment_count = min(completed_count, job.attachment_count)
    job.processed_attachment_bytes = min(completed_bytes, job.attachment_bytes)
    job.updated_at = utcnow_naive()
    db.add(job)
    db.commit()


def _require_temp_space(attachment_bytes: int, *, temp_root: str | None) -> None:
    temp_dir = temp_root or tempfile.gettempdir()
    required = attachment_bytes + EXCEL_EXPORT_TEMP_OVERHEAD_BYTES
    if shutil.disk_usage(temp_dir).free < required:
        raise localized_http_exception(
            status_code=status.HTTP_507_INSUFFICIENT_STORAGE,
            code="legacy_issues.excel_export_insufficient_space",
        )


def excel_export_job_status(job: LegacyIssueExcelExportJob) -> str:
    if (
        job.status == EXCEL_EXPORT_STATUS_COMPLETED
        and job.expires_at is not None
        and job.expires_at <= utcnow_naive()
    ):
        return EXCEL_EXPORT_STATUS_EXPIRED
    return job.status


def _cleanup_orphan_export_objects(
    db: Session,
    *,
    client: ExcelExportStorageClient,
    bucket_name: str,
    now: Any,
    limit: int,
) -> tuple[int, int]:
    referenced = set(
        db.scalars(
            select(LegacyIssueExcelExportJob.result_storage_key).where(
                LegacyIssueExcelExportJob.result_storage_key.is_not(None)
            )
        )
    )
    deleted = 0
    failed = 0
    try:
        objects = client.list_objects(
            bucket_name,
            prefix=EXCEL_EXPORT_RESULT_PREFIX,
            recursive=True,
        )
        for item in objects:
            if deleted + failed >= limit:
                break
            storage_key = str(getattr(item, "object_name", "") or "")
            modified = getattr(item, "last_modified", None)
            if not storage_key or storage_key in referenced or modified is None:
                continue
            modified_naive = modified.replace(tzinfo=None) if modified.tzinfo else modified
            if modified_naive > now - EXCEL_EXPORT_ORPHAN_GRACE:
                continue
            try:
                client.remove_object(bucket_name, storage_key)
            except Exception:
                failed += 1
            else:
                deleted += 1
    except Exception:
        logger.warning("Failed to reconcile Excel export objects", exc_info=True)
        failed += 1
    return deleted, failed


def _resolve_storage(
    *,
    client: ExcelExportStorageClient | None,
    bucket_name: str | None,
) -> tuple[ExcelExportStorageClient, str]:
    if client is not None and bucket_name:
        return client, bucket_name
    from ai_do_api.core.storage import ensure_bucket, get_minio_client

    ensure_bucket()
    settings = get_settings()
    return get_minio_client(), settings.minio_bucket


def _remove_result_object(
    client: ExcelExportStorageClient,
    bucket_name: str,
    storage_key: str,
) -> None:
    try:
        client.remove_object(bucket_name, storage_key)
    except Exception:
        logger.warning("Failed to remove incomplete Excel export", exc_info=True)


def _excel_export_error_code(error: Exception) -> str:
    detail = getattr(error, "detail", None)
    code = getattr(detail, "code", None)
    if isinstance(code, str) and code:
        return code
    return "legacy_issues.excel_export_failed"


def _safe_error_text(error: Exception) -> str:
    return str(error).replace("\n", " ")[:1000]


def _result_filename(value: str) -> str:
    stem, dot, suffix = value.rpartition(".")
    safe_stem = _SAFE_RESULT_FILENAME.sub("-", stem or value).strip("-._")
    safe_suffix = _SAFE_RESULT_FILENAME.sub("", suffix if dot else "xlsx") or "xlsx"
    return f"{safe_stem or 'legacy-issues'}.{safe_suffix}"


def _batches(items: list[str], size: int) -> Iterable[list[str]]:
    for index in range(0, len(items), size):
        yield items[index : index + size]


def _celery_client() -> Celery:
    return create_fail_fast_celery_publisher(
        "ai_do_api_legacy_issue_excel_export",
        broker=get_settings().worker_broker_url,
        ignore_result=True,
    )


from sqlalchemy import event  # noqa: E402


@event.listens_for(Session, "after_commit")
def _publish_pending_excel_export_jobs(session: Session) -> None:
    if session.in_nested_transaction():
        return
    pending = session.info.pop(_PENDING_PUBLISHES_KEY, None)
    if not pending:
        return
    for job_id in sorted(pending):
        try:
            publish_excel_export_job(job_id)
        except Exception:
            logger.warning("Failed to publish Excel export job", exc_info=True)


@event.listens_for(Session, "after_rollback")
def _clear_pending_excel_export_jobs(session: Session) -> None:
    if not session.in_nested_transaction():
        session.info.pop(_PENDING_PUBLISHES_KEY, None)
