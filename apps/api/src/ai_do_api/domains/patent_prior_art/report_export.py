"""Self-contained HTML, PDF, and DOCX renderers for prior-art reports."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from html import escape as html_escape
from io import BytesIO
import re
from xml.sax.saxutils import escape as xml_escape

from ai_do_api.domains.patent_prior_art.reporting import (
    PatentPriorArtDetailedReport,
    ReportCandidate,
    ReportCount,
    build_detailed_report,
)
from ai_do_api.domains.patent_prior_art.schemas import PatentPriorArtReportFormat

_SUPPORTED_FORMATS = frozenset({"html", "pdf", "docx", "summary_pdf", "summary_docx"})
_DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True, slots=True)
class RenderedPatentPriorArtReport:
    content: bytes
    media_type: str
    extension: str
    filename: str


def render_report_export(
    payload: Mapping[str, object],
    *,
    report_format: PatentPriorArtReportFormat,
    job_id: str,
) -> RenderedPatentPriorArtReport:
    """Render one bounded report without network, external assets, or private metrics."""

    if report_format not in _SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported prior-art report format: {report_format}")
    summary = report_format.startswith("summary_")
    extension = report_format.removeprefix("summary_")

    # All exports use the self-contained neutral projection (build_detailed_report).
    # The full-engine rich report.html is intentionally NOT used here: it loaded
    # CDN JavaScript and provider-controlled (KIPRIS) image URLs through a
    # headless browser, which breaks the no-network self-contained report
    # contract (SSRF / RCE boundary) and re-exposed neutralised legal conclusions.
    report = build_detailed_report(payload)
    if extension == "html":
        content = _render_html(report)
        media_type = "text/html; charset=utf-8"
    elif extension == "pdf":
        content = _render_pdf(report, summary=summary)
        media_type = "application/pdf"
    else:
        content = _render_docx(report, summary=summary)
        media_type = _DOCX_MEDIA_TYPE
    safe_job_id = re.sub(r"[^A-Za-z0-9_-]", "-", job_id).strip("-")[:64] or "report"
    report_kind = "summary" if summary else "detailed"
    return RenderedPatentPriorArtReport(
        content=content,
        media_type=media_type,
        extension=extension,
        filename=f"patent-prior-art-{report_kind}-{safe_job_id}.{extension}",
    )


def _render_html(report: PatentPriorArtDetailedReport) -> bytes:
    sections = [
        _html_section(
            1,
            "과제 개요",
            _html_metadata(report)
            + _html_subheading("기술 요약")
            + _html_paragraph(report.technology_summary)
            + _html_subheading("입력 발명 내용")
            + _html_paragraph(report.invention_text)
            + (
                '<p class="notice">입력 내용은 보고서 크기 제한에 따라 일부만 수록되었습니다.</p>'
                if report.invention_text_truncated
                else ""
            ),
        ),
        _html_section(
            2,
            "검색식 도출 근거",
            _html_table(
                ("항목", "도출 내용"),
                report.search_rationale,
            )
            + _html_subheading("표시 검색식")
            + f"<pre>{_h(report.display_query)}</pre>"
            + _html_subheading("실행 질의")
            + _html_table(
                ("순번", "출처", "국가/권역", "질의", "결과 수"),
                tuple(
                    (
                        str(index),
                        query.source_label,
                        query.jurisdiction,
                        query.query_text,
                        query.result_display,
                    )
                    for index, query in enumerate(report.executed_queries, start=1)
                ),
            ),
        ),
        _html_section(3, "선별 프로세스", _html_list(report.selection_process)),
        _html_section(4, "핵심 특허 표", _html_candidate_table(report.candidates)),
        _html_section(5, "후보별 요지시트", _html_candidate_sheets(report)),
        _html_section(
            6,
            "정량 분석",
            _html_subheading("관련도 분포")
            + _html_count_table(report.relevance_distribution)
            + _html_subheading("출원인 분포")
            + _html_count_table(report.applicant_distribution)
            + _html_subheading("공개연도 분포")
            + _html_count_table(report.year_distribution),
        ),
        _html_section(
            7,
            "분류코드 기반 기술 클러스터",
            _html_table(
                ("분류 클러스터", "후보 수", "포함 공개번호"),
                tuple(
                    (
                        cluster.code,
                        str(cluster.candidate_count),
                        ", ".join(cluster.publications),
                    )
                    for cluster in report.technology_clusters
                ),
            ),
        ),
        _html_section(8, "종합 기술 검토", _html_list(report.technical_review)),
        _html_section(
            9,
            "출처·한계 및 고지",
            _html_list(report.sources_and_limits)
            + _html_subheading("고지")
            + f'<p class="disclaimer">{_h(report.disclaimer)}</p>',
        ),
    ]
    document = f"""<!doctype html>
<html lang="ko">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; img-src 'none'; media-src 'none'; font-src 'none'; connect-src 'none'; form-action 'none'; base-uri 'none'">
  <title>{_h(report.title)}</title>
  <style>
    :root {{ color-scheme: light; font-family: system-ui, sans-serif; color: #172033; background: #fff; }}
    body {{ max-width: 1120px; margin: 0 auto; padding: 36px 28px 72px; line-height: 1.55; }}
    header {{ border-bottom: 3px solid #31557a; padding-bottom: 18px; margin-bottom: 30px; }}
    h1 {{ margin: 0 0 8px; color: #17324d; font-size: 30px; }}
    h2 {{ margin-top: 34px; padding-bottom: 7px; border-bottom: 1px solid #b8c8d8; color: #244967; }}
    h3 {{ margin-top: 22px; color: #31557a; }}
    table {{ width: 100%; border-collapse: collapse; margin: 12px 0 22px; font-size: 14px; }}
    th, td {{ border: 1px solid #cbd5df; padding: 8px 10px; text-align: left; vertical-align: top; overflow-wrap: anywhere; }}
    th {{ background: #e8eef4; color: #17324d; }}
    tbody tr:nth-child(even) {{ background: #f7f9fb; }}
    pre {{ white-space: pre-wrap; overflow-wrap: anywhere; background: #f2f5f8; border: 1px solid #d5dfe8; padding: 12px; }}
    article {{ border: 1px solid #d5dfe8; border-radius: 6px; padding: 14px 16px; margin: 14px 0; break-inside: avoid; }}
    dl {{ display: grid; grid-template-columns: 130px 1fr; gap: 5px 12px; }}
    dt {{ font-weight: 700; color: #31557a; }} dd {{ margin: 0; overflow-wrap: anywhere; }}
    .notice, .disclaimer {{ border-left: 4px solid #c38a2d; background: #fff8e8; padding: 10px 12px; }}
    .muted {{ color: #526477; }}
    @media print {{ body {{ max-width: none; padding: 0; }} h2 {{ break-after: avoid; }} table {{ break-inside: auto; }} }}
  </style>
</head>
<body>
  <header><h1>{_h(report.title)}</h1><p class="muted">기술 선행문헌 상세 조사보고서</p></header>
  <main>{"".join(sections)}</main>
</body>
</html>"""
    return document.encode("utf-8")


def _html_metadata(report: PatentPriorArtDetailedReport) -> str:
    return _html_table(
        ("항목", "내용"),
        (
            ("보고서 유형", report.report_type),
            ("스키마 버전", report.schema_version),
            ("조사 국가/권역", ", ".join(report.jurisdictions) or "-"),
            ("선택 카테고리", ", ".join(report.category_ids) or "-"),
        ),
    )


def _html_candidate_table(candidates: Sequence[ReportCandidate]) -> str:
    return _html_table(
        ("순위", "공개번호", "명칭", "출원인", "공개연도", "분류코드", "관련도"),
        tuple(
            (
                str(candidate.rank),
                candidate.publication_number,
                candidate.title,
                ", ".join(candidate.assignees) or "-",
                candidate.publication_year,
                ", ".join(candidate.classification_codes) or "-",
                candidate.relevance_label,
            )
            for candidate in candidates
        ),
    )


def _html_candidate_sheets(report: PatentPriorArtDetailedReport) -> str:
    articles: list[str] = []
    for candidate in report.candidates:
        details = (
            ("명칭", candidate.title),
            ("국가/권역", candidate.jurisdiction),
            ("출원인", ", ".join(candidate.assignees) or "-"),
            ("출원일", candidate.filing_date),
            ("공개일", candidate.publication_date),
            ("분류코드", ", ".join(candidate.classification_codes) or "-"),
            ("기술 관련도", candidate.relevance_label),
            ("검토 근거", "; ".join(candidate.match_reasons) or "-"),
            ("원문 위치", candidate.external_url),
        )
        detail_html = "".join(
            f"<dt>{_h(label)}</dt><dd>{_h(value)}</dd>" for label, value in details
        )
        narrative = candidate.summary or candidate.abstract or "-"
        articles.append(
            f"<article><h3>{candidate.rank}. {_h(candidate.publication_number)}</h3>"
            f"<dl>{detail_html}</dl><p>{_h(narrative)}</p></article>"
        )
    if report.candidates_truncated:
        articles.append(
            '<p class="notice">보고서 후보 수 제한에 따라 일부 후보만 수록되었습니다.</p>'
        )
    return "".join(articles) or '<p class="muted">선별된 후보가 없습니다.</p>'


def _html_section(number: int, title: str, body: str) -> str:
    return f"<section><h2>{number}. {_h(title)}</h2>{body}</section>"


def _html_subheading(title: str) -> str:
    return f"<h3>{_h(title)}</h3>"


def _html_paragraph(value: str) -> str:
    return f"<p>{_h(value)}</p>"


def _html_list(items: Sequence[str]) -> str:
    if not items:
        return '<p class="muted">-</p>'
    return "<ul>" + "".join(f"<li>{_h(item)}</li>" for item in items) + "</ul>"


def _html_count_table(items: Sequence[ReportCount]) -> str:
    return _html_table(
        ("구분", "후보 수"),
        tuple((item.label, str(item.count)) for item in items),
    )


def _html_table(headers: Sequence[str], rows: Sequence[Sequence[str]]) -> str:
    head = "".join(f'<th scope="col">{_h(header)}</th>' for header in headers)
    body_rows = rows or (tuple("-" for _ in headers),)
    body = "".join(
        "<tr>" + "".join(f"<td>{_h(value)}</td>" for value in row) + "</tr>" for row in body_rows
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def _h(value: object) -> str:
    return html_escape(str(value), quote=True).replace("\n", "<br>")


def _render_pdf(report: PatentPriorArtDetailedReport, *, summary: bool) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import (
        LongTable,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )

    for font_name in ("HYSMyeongJo-Medium", "HYGothic-Medium"):
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    regular_font = "HYSMyeongJo-Medium"
    bold_font = "HYGothic-Medium"
    sample = getSampleStyleSheet()
    body = ParagraphStyle(
        "PriorArtBody",
        parent=sample["BodyText"],
        fontName=regular_font,
        fontSize=8.5,
        leading=12,
        spaceAfter=5,
    )
    bullet = ParagraphStyle(
        "PriorArtBullet",
        parent=body,
        leftIndent=12,
        firstLineIndent=-7,
    )
    heading = ParagraphStyle(
        "PriorArtHeading",
        parent=sample["Heading2"],
        fontName=bold_font,
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=7,
        keepWithNext=True,
    )
    subheading = ParagraphStyle(
        "PriorArtSubheading",
        parent=sample["Heading3"],
        fontName=bold_font,
        fontSize=10.5,
        leading=14,
        spaceBefore=8,
        spaceAfter=5,
        keepWithNext=True,
    )
    title_style = ParagraphStyle(
        "PriorArtTitle",
        parent=sample["Title"],
        fontName=bold_font,
        fontSize=20,
        leading=25,
        alignment=TA_LEFT,
    )

    def paragraph(value: object, style=body):
        escaped = xml_escape(str(value)).replace("\n", "<br/>")
        return Paragraph(escaped or "-", style)

    def table(
        headers: Sequence[str],
        rows: Sequence[Sequence[object]],
        *,
        widths: Sequence[float] | None = None,
    ):
        data = [[paragraph(value, subheading) for value in headers]]
        data.extend([paragraph(value) for value in row] for row in rows)
        if len(data) == 1:
            data.append([paragraph("-") for _ in headers])
        report_table = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
        report_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), regular_font),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6EF")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#17324D")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#AABAC8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F7F9FB")],
                    ),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return report_table

    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=(f"{report.title} 요약보고서" if summary else report.title),
        author="AI-DO",
    )
    flow: list[object] = [
        paragraph(
            f"{report.title} 요약보고서" if summary else report.title,
            title_style,
        ),
        Spacer(1, 5 * mm),
    ]
    if summary:
        _append_pdf_summary(
            flow,
            report,
            paragraph=paragraph,
            table=table,
            heading=heading,
            bullet=bullet,
        )
    else:
        _append_pdf_detailed(
            flow,
            report,
            paragraph=paragraph,
            table=table,
            heading=heading,
            subheading=subheading,
            bullet=bullet,
            page_break=PageBreak,
            spacer=Spacer,
            mm=mm,
        )
    document.build(flow)
    return buffer.getvalue()


def _append_pdf_summary(
    flow: list[object],
    report: PatentPriorArtDetailedReport,
    *,
    paragraph,
    table,
    heading,
    bullet,
) -> None:
    flow.extend(
        [
            paragraph("1. 조사 범위", heading),
            table(
                ("항목", "내용"),
                (
                    ("국가/권역", ", ".join(report.jurisdictions) or "-"),
                    ("카테고리", ", ".join(report.category_ids) or "-"),
                    ("기술 요약", report.technology_summary),
                ),
            ),
            paragraph("2. 후보 요약", heading),
            table(
                ("순위", "공개번호", "명칭", "출원인", "공개연도", "관련도"),
                _candidate_rows(report.candidates, include_codes=False),
            ),
            paragraph("3. 종합 기술 검토", heading),
        ]
    )
    flow.extend(paragraph(item, bullet) for item in report.technical_review)
    flow.extend([paragraph("4. 고지", heading), paragraph(report.disclaimer)])


def _append_pdf_detailed(
    flow: list[object],
    report: PatentPriorArtDetailedReport,
    *,
    paragraph,
    table,
    heading,
    subheading,
    bullet,
    page_break,
    spacer,
    mm,
) -> None:
    flow.extend(
        [
            paragraph("1. 과제 개요", heading),
            table(
                ("항목", "내용"),
                (
                    ("보고서 유형", report.report_type),
                    ("스키마 버전", report.schema_version),
                    ("국가/권역", ", ".join(report.jurisdictions) or "-"),
                    ("카테고리", ", ".join(report.category_ids) or "-"),
                ),
            ),
            paragraph("기술 요약", subheading),
            paragraph(report.technology_summary),
            paragraph("입력 발명 내용", subheading),
            paragraph(report.invention_text),
            paragraph("2. 검색식 도출 근거", heading),
            table(("항목", "도출 내용"), report.search_rationale),
            paragraph("표시 검색식", subheading),
            paragraph(report.display_query),
            paragraph("실행 질의", subheading),
            table(
                ("순번", "출처", "국가/권역", "질의", "결과 수"),
                tuple(
                    (
                        index,
                        query.source_label,
                        query.jurisdiction,
                        query.query_text,
                        query.result_display,
                    )
                    for index, query in enumerate(report.executed_queries, start=1)
                ),
            ),
            paragraph("3. 선별 프로세스", heading),
        ]
    )
    flow.extend(paragraph(item, bullet) for item in report.selection_process)
    flow.extend(
        [
            paragraph("4. 핵심 특허 표", heading),
            table(
                ("순위", "공개번호", "명칭", "출원인", "공개연도", "분류코드", "관련도"),
                _candidate_rows(report.candidates, include_codes=True),
            ),
            page_break(),
            paragraph("5. 후보별 요지시트", heading),
        ]
    )
    for candidate in report.candidates:
        flow.extend(
            [
                paragraph(f"{candidate.rank}. {candidate.publication_number}", subheading),
                table(
                    ("항목", "내용"),
                    _candidate_detail_rows(candidate),
                    widths=(38 * mm, 205 * mm),
                ),
                paragraph(candidate.summary or candidate.abstract or "-"),
                spacer(1, 3 * mm),
            ]
        )
    flow.extend(
        [
            paragraph("6. 정량 분석", heading),
            paragraph("관련도 분포", subheading),
            table(("구분", "후보 수"), _count_rows(report.relevance_distribution)),
            paragraph("출원인 분포", subheading),
            table(("구분", "후보 수"), _count_rows(report.applicant_distribution)),
            paragraph("공개연도 분포", subheading),
            table(("구분", "후보 수"), _count_rows(report.year_distribution)),
            paragraph("7. 분류코드 기반 기술 클러스터", heading),
            table(
                ("분류 클러스터", "후보 수", "포함 공개번호"),
                tuple(
                    (cluster.code, cluster.candidate_count, ", ".join(cluster.publications))
                    for cluster in report.technology_clusters
                ),
            ),
            paragraph("8. 종합 기술 검토", heading),
        ]
    )
    flow.extend(paragraph(item, bullet) for item in report.technical_review)
    flow.append(paragraph("9. 출처·한계 및 고지", heading))
    flow.extend(paragraph(item, bullet) for item in report.sources_and_limits)
    flow.extend([paragraph("고지", subheading), paragraph(report.disclaimer)])


def _render_docx(report: PatentPriorArtDetailedReport, *, summary: bool) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    document.core_properties.title = f"{report.title} 요약보고서" if summary else report.title
    document.core_properties.author = "AI-DO"
    for section in document.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10 if style_name == "Normal" else 14)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    def heading(value: str, level: int = 1) -> None:
        document.add_heading(value, level=level)

    def paragraph(value: object, *, bullet: bool = False) -> None:
        document.add_paragraph(str(value) or "-", style="List Bullet" if bullet else None)

    def table(headers: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
        body_rows = tuple(rows)
        report_table = document.add_table(rows=1, cols=len(headers))
        report_table.style = "Table Grid"
        for index, header in enumerate(headers):
            cell = report_table.rows[0].cells[index]
            cell.text = header
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in body_rows or (tuple("-" for _ in headers),):
            cells = report_table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)

    document.add_heading(
        f"{report.title} 요약보고서" if summary else report.title,
        level=0,
    )
    if summary:
        _append_docx_summary(
            report,
            heading=heading,
            paragraph=paragraph,
            table=table,
        )
    else:
        _append_docx_detailed(
            document,
            report,
            heading=heading,
            paragraph=paragraph,
            table=table,
        )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _append_docx_summary(
    report: PatentPriorArtDetailedReport,
    *,
    heading,
    paragraph,
    table,
) -> None:
    heading("1. 조사 범위")
    table(
        ("항목", "내용"),
        (
            ("국가/권역", ", ".join(report.jurisdictions) or "-"),
            ("카테고리", ", ".join(report.category_ids) or "-"),
            ("기술 요약", report.technology_summary),
        ),
    )
    heading("2. 후보 요약")
    table(
        ("순위", "공개번호", "명칭", "출원인", "공개연도", "관련도"),
        _candidate_rows(report.candidates, include_codes=False),
    )
    heading("3. 종합 기술 검토")
    for item in report.technical_review:
        paragraph(item, bullet=True)
    heading("4. 고지")
    paragraph(report.disclaimer)


def _append_docx_detailed(document, report, *, heading, paragraph, table) -> None:
    heading("1. 과제 개요")
    table(
        ("항목", "내용"),
        (
            ("보고서 유형", report.report_type),
            ("스키마 버전", report.schema_version),
            ("국가/권역", ", ".join(report.jurisdictions) or "-"),
            ("카테고리", ", ".join(report.category_ids) or "-"),
        ),
    )
    heading("기술 요약", level=2)
    paragraph(report.technology_summary)
    heading("입력 발명 내용", level=2)
    paragraph(report.invention_text)
    heading("2. 검색식 도출 근거")
    table(("항목", "도출 내용"), report.search_rationale)
    heading("표시 검색식", level=2)
    paragraph(report.display_query)
    heading("실행 질의", level=2)
    table(
        ("순번", "출처", "국가/권역", "질의", "결과 수"),
        tuple(
            (
                index,
                query.source_label,
                query.jurisdiction,
                query.query_text,
                query.result_display,
            )
            for index, query in enumerate(report.executed_queries, start=1)
        ),
    )
    heading("3. 선별 프로세스")
    for item in report.selection_process:
        paragraph(item, bullet=True)
    heading("4. 핵심 특허 표")
    table(
        ("순위", "공개번호", "명칭", "출원인", "공개연도", "분류코드", "관련도"),
        _candidate_rows(report.candidates, include_codes=True),
    )
    document.add_page_break()
    heading("5. 후보별 요지시트")
    for candidate in report.candidates:
        heading(f"{candidate.rank}. {candidate.publication_number}", level=2)
        table(("항목", "내용"), _candidate_detail_rows(candidate))
        paragraph(candidate.summary or candidate.abstract or "-")
    heading("6. 정량 분석")
    for title, values in (
        ("관련도 분포", report.relevance_distribution),
        ("출원인 분포", report.applicant_distribution),
        ("공개연도 분포", report.year_distribution),
    ):
        heading(title, level=2)
        table(("구분", "후보 수"), _count_rows(values))
    heading("7. 분류코드 기반 기술 클러스터")
    table(
        ("분류 클러스터", "후보 수", "포함 공개번호"),
        tuple(
            (cluster.code, cluster.candidate_count, ", ".join(cluster.publications))
            for cluster in report.technology_clusters
        ),
    )
    heading("8. 종합 기술 검토")
    for item in report.technical_review:
        paragraph(item, bullet=True)
    heading("9. 출처·한계 및 고지")
    for item in report.sources_and_limits:
        paragraph(item, bullet=True)
    heading("고지", level=2)
    paragraph(report.disclaimer)


def _candidate_rows(
    candidates: Sequence[ReportCandidate],
    *,
    include_codes: bool,
) -> tuple[tuple[object, ...], ...]:
    rows: list[tuple[object, ...]] = []
    for candidate in candidates:
        values: list[object] = [
            candidate.rank,
            candidate.publication_number,
            candidate.title,
            ", ".join(candidate.assignees) or "-",
            candidate.publication_year,
        ]
        if include_codes:
            values.append(", ".join(candidate.classification_codes) or "-")
        values.append(candidate.relevance_label)
        rows.append(tuple(values))
    return tuple(rows)


def _candidate_detail_rows(candidate: ReportCandidate) -> tuple[tuple[str, str], ...]:
    return (
        ("명칭", candidate.title),
        ("국가/권역", candidate.jurisdiction),
        ("출원인", ", ".join(candidate.assignees) or "-"),
        ("출원일", candidate.filing_date),
        ("공개일", candidate.publication_date),
        ("분류코드", ", ".join(candidate.classification_codes) or "-"),
        ("기술 관련도", candidate.relevance_label),
        ("검토 근거", "; ".join(candidate.match_reasons) or "-"),
        ("원문 위치", candidate.external_url),
    )


def _count_rows(items: Sequence[ReportCount]) -> tuple[tuple[str, int], ...]:
    return tuple((item.label, item.count) for item in items)


__all__ = [
    "RenderedPatentPriorArtReport",
    "render_report_export",
]
