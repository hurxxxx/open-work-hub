from __future__ import annotations

import io

from ai_do_worker.runtime import ensure_api_src_on_path


_MARKER = "AI-DO PPT browser smoke"


def run_smoke() -> None:
    ensure_api_src_on_path()

    from pptx import Presentation

    from ai_do_api.domains.ppt_generator.design.html_to_pptx import (
        html_deck_to_editable_pptx_bytes,
    )

    html = (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        "<style>html,body{margin:0}.slide{width:640px;height:360px}</style>"
        f"</head><body><section class='slide'><div>{_MARKER}</div></section></body></html>"
    )
    payload = html_deck_to_editable_pptx_bytes(html, 640, 360)
    presentation = Presentation(io.BytesIO(payload))
    if len(presentation.slides) != 1:
        raise RuntimeError("PPT browser smoke produced an unexpected slide count")
    texts = [
        shape.text
        for shape in presentation.slides[0].shapes
        if hasattr(shape, "text") and shape.text
    ]
    if not any(_MARKER in text for text in texts):
        raise RuntimeError("PPT browser smoke marker was not preserved")


def main() -> None:
    run_smoke()
    print("[ppt-browser-smoke] ok")


if __name__ == "__main__":
    main()
